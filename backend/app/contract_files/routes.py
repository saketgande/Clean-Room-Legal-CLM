import hashlib
import logging
import secrets

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Query,
    Request,
    Response,
    UploadFile,
    status,
)
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from slowapi.util import get_remote_address
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.contract_files.models import (
    ContractEdit,
    ContractFile,
    ContractShare,
    ContractTextSnapshot,
    ContractVersion,
    StorageObject,
)
from app.contract_files.schemas import (
    ContractFileResponse,
    ContractEditDecisionRequest,
    ContractEditResponse,
    ContractShareCreate,
    ContractShareCreateResponse,
    ContractShareResponse,
    ContractTextSnapshotResponse,
    ContractVersionResponse,
    ExternalCommentCreate,
    ExternalCommentResponse,
    ExternalShareResponse,
)
from app.contract_files.service import (
    add_version_from_upload,
    next_version_number,
    requeue_contract_ai_jobs,
)
from app.contracts.models import Contract
from app.contracts.service import get_contract_for_user
from app.contracts.comments_service import add_counterparty_comment, list_shared_comments
from app.core.audit import write_audit_log, write_timeline_event
from app.core.config import settings
from app.core.database import utcnow
from app.core.deps import get_db, require_permission
from app.core.enums import ContractVersionSource, ShareAccessMode, StorageBackend
from app.core.rate_limit import limiter
from app.integrations.storage import storage_service

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/contracts/{contract_id}", tags=["contract-files"])
external_share_router = APIRouter(prefix="/external-shares", tags=["external-shares"])

EXTERNAL_TEXT_EXCERPT_CHARS = 12_000

# External-share passcode brute-force lockout. After this many consecutive
# failed passcode attempts (keyed per share-token + client IP in Redis), the
# share locks for an exponentially growing window. Counters are best-effort:
# if Redis is unreachable we fail OPEN (no lockout) so a Redis outage never
# breaks legitimate share access — the per-IP rate limit still applies.
SHARE_PASSCODE_MAX_ATTEMPTS = 5
SHARE_PASSCODE_LOCKOUT_BASE_SECONDS = 30
SHARE_PASSCODE_LOCKOUT_MAX_SECONDS = 3600
SHARE_PASSCODE_ATTEMPT_TTL_SECONDS = 3600


def _share_rate_limit_key(request: Request) -> str:
    """Rate-limit key for external-share endpoints: share token + client IP.

    Pinning on the token (from the path) as well as the IP means a single
    leaked link can't be hammered from one host, while distinct shares get
    independent buckets.
    """
    token = request.path_params.get("token", "")
    return f"share:{token}:{get_remote_address(request)}"


@router.get("/files", response_model=list[ContractFileResponse])
def list_contract_files(
    contract_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:read")),
):
    get_contract_for_user(db, contract_id=contract_id, user=current_user)
    return db.scalars(
        select(ContractFile).where(
            ContractFile.org_id == current_user.org_id,
            ContractFile.contract_id == contract_id,
            ContractFile.deleted_at.is_(None),
        )
    ).all()


@router.get("/versions", response_model=list[ContractVersionResponse])
def list_contract_versions(
    contract_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:read")),
):
    get_contract_for_user(db, contract_id=contract_id, user=current_user)
    return db.scalars(
        select(ContractVersion)
        .where(
            ContractVersion.org_id == current_user.org_id,
            ContractVersion.contract_id == contract_id,
            ContractVersion.deleted_at.is_(None),
        )
        .order_by(ContractVersion.version_number.asc())
    ).all()


@router.post(
    "/versions",
    response_model=ContractVersionResponse,
    status_code=status.HTTP_201_CREATED,
)
@limiter.limit(settings.rate_limit_contract_upload)
async def upload_contract_version(
    contract_id: str,
    request: Request,
    response: Response,
    file: UploadFile = File(...),
    change_summary: str | None = Form(default=None),
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:update")),
):
    """Create a new authoritative version of an existing contract from an
    uploaded .docx — used by the Word add-in to push the edited draft back."""
    _ = response  # present for slowapi's rate-limit header injection
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    return await add_version_from_upload(
        db,
        contract=contract,
        upload=file,
        user=current_user,
        change_summary=change_summary,
        request_id=getattr(request.state, "request_id", None),
    )


@router.post(
    "/counterparty-revision",
    response_model=ContractVersionResponse,
    status_code=status.HTTP_201_CREATED,
)
@limiter.limit(settings.rate_limit_contract_upload)
async def log_counterparty_revision(
    contract_id: str,
    request: Request,
    response: Response,
    file: UploadFile = File(...),
    change_summary: str | None = Form(default=None),
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:update")),
):
    """Log a counterparty's returned redlined document as a COUNTERPARTY_REVISION
    version. Re-opens REVIEW if the contract had moved past it and flags the
    contract for an auto playbook deviation review, which fires once the new
    version's clause extraction completes."""
    _ = response
    from app.contracts.lifecycle import transition_contract_stage
    from app.core.enums import ContractLifecycleStage

    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    req_id = getattr(request.state, "request_id", None)
    version = await add_version_from_upload(
        db,
        contract=contract,
        upload=file,
        user=current_user,
        change_summary=change_summary or "Counterparty revision received",
        source=ContractVersionSource.COUNTERPARTY_REVISION,
        request_id=req_id,
    )
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    if contract.lifecycle_stage in (ContractLifecycleStage.APPROVAL, ContractLifecycleStage.SIGNATURE):
        try:
            transition_contract_stage(
                db,
                contract=contract,
                to_stage=ContractLifecycleStage.REVIEW,
                actor_user_id=current_user.id,
                reason="Counterparty revision received — re-opening review",
                override=True,
                override_authorized=True,
                request_id=req_id,
            )
        except Exception:  # noqa: BLE001 - re-open is best-effort
            pass
    meta = dict(contract.metadata_json or {})
    meta["auto_review_pending"] = True
    contract.metadata_json = meta
    write_timeline_event(
        db,
        org_id=contract.org_id,
        resource_type="contract",
        resource_id=contract.id,
        event_type="contract.counterparty_revision",
        title="Counterparty revision received",
        actor_user_id=current_user.id,
        request_id=req_id,
        details={"contract_version_id": version.id},
    )
    db.commit()
    db.refresh(version)
    return version


@router.get("/versions/{version_id}/text", response_model=ContractTextSnapshotResponse)
def get_version_text_snapshot(
    contract_id: str,
    version_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:read")),
):
    get_contract_for_user(db, contract_id=contract_id, user=current_user)
    version = db.get(ContractVersion, version_id)
    if version is None or version.contract_id != contract_id or version.org_id != current_user.org_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract version not found")
    if version.text_snapshot_id is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Text snapshot not found")
    snapshot = db.get(ContractTextSnapshot, version.text_snapshot_id)
    if snapshot is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Text snapshot not found")
    return snapshot


@router.get("/versions/{version_id}/download")
def download_contract_version(
    contract_id: str,
    version_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:read")),
):
    get_contract_for_user(db, contract_id=contract_id, user=current_user)
    version = db.get(ContractVersion, version_id)
    if version is None or version.contract_id != contract_id or version.org_id != current_user.org_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract version not found")
    storage_object = db.get(StorageObject, version.storage_object_id)
    if storage_object is None or storage_object.org_id != current_user.org_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Stored file not found")
    try:
        path = storage_service.path_for_read(storage_object.storage_key)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Stored file bytes not found")
    # Force a download rather than letting the browser render the bytes inline:
    # an inline HTML/SVG masquerading as an allowed type would otherwise execute
    # in our origin (stored-XSS / phishing surface).
    return FileResponse(
        path,
        media_type=storage_object.mime_type,
        filename=storage_object.filename,
        content_disposition_type="attachment",
    )


@router.get("/edits", response_model=list[ContractEditResponse])
def list_contract_edits(
    contract_id: str,
    status_filter: str | None = None,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract:redline")),
):
    get_contract_for_user(db, contract_id=contract_id, user=current_user)
    query = select(ContractEdit).where(
        ContractEdit.org_id == current_user.org_id,
        ContractEdit.contract_id == contract_id,
    )
    if status_filter:
        query = query.where(ContractEdit.status == status_filter)
    return db.scalars(query.order_by(ContractEdit.created_at.desc())).all()


class ManualEditProposal(BaseModel):
    original_text: str = Field(min_length=3)
    replacement_text: str = ""  # empty string = propose deleting the passage
    rationale: str | None = None
    start_hint: int | None = None  # viewer's char offset, disambiguates repeats


def _split_paragraphs(text: str) -> list[str]:
    return [p for p in text.split("\n") if p.strip()]


def _build_manual_redline_docx(
    *,
    contract_title: str,
    base_version_number: int,
    author_name: str,
    pre_text: str,
    original_text: str,
    replacement_text: str,
    post_text: str,
) -> bytes:
    """Full-document .docx with the proposed change as a NATIVE Word tracked
    change (w:del + w:ins), so a lawyer opening it in Word sees a real redline."""
    from io import BytesIO

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from app.ai.tool_runtime import (
        _append_deleted_text,
        _append_inserted_text,
        _enable_word_track_revisions,
    )

    document = Document()
    _enable_word_track_revisions(document, OxmlElement=OxmlElement, qn=qn)
    document.add_heading(contract_title, level=1)
    document.add_paragraph(f"Redline proposal on V{base_version_number}")
    for para in _split_paragraphs(pre_text):
        document.add_paragraph(para)
    revision_paragraph = document.add_paragraph()
    _append_deleted_text(
        revision_paragraph,
        original_text,
        author=author_name,
        revision_id="1",
        OxmlElement=OxmlElement,
        qn=qn,
    )
    if replacement_text:
        _append_inserted_text(
            revision_paragraph,
            replacement_text,
            author=author_name,
            revision_id="2",
            OxmlElement=OxmlElement,
            qn=qn,
        )
    for para in _split_paragraphs(post_text):
        document.add_paragraph(para)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_plain_docx(*, title: str, subtitle: str, text: str) -> bytes:
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    if subtitle:
        document.add_paragraph(subtitle)
    for para in _split_paragraphs(text):
        document.add_paragraph(para)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _store_generated_docx(
    db: Session, *, org_id: str, user_id: str, filename: str, content: bytes
) -> StorageObject:
    stored = storage_service.save_bytes(
        org_id=org_id,
        filename=filename,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=content,
    )
    storage_object = StorageObject(
        org_id=org_id,
        storage_key=stored.storage_key,
        filename=stored.filename,
        mime_type=stored.mime_type,
        size_bytes=stored.size_bytes,
        sha256_hash=stored.sha256_hash,
        storage_backend=StorageBackend.LOCAL_VOLUME,
        created_by_user_id=user_id,
        updated_by_user_id=user_id,
    )
    db.add(storage_object)
    db.flush()
    return storage_object


@router.post(
    "/edits/propose", response_model=ContractEditResponse, status_code=status.HTTP_201_CREATED
)
def propose_contract_edit(
    contract_id: str,
    payload: ManualEditProposal,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract:redline")),
):
    """User-authored redline: select text in the document, propose a change.
    Creates the same structure as AI redlines — a non-authoritative proposal
    version whose snapshot has the change applied, plus a proposed
    ContractEdit — so accept/reject rides the existing pipeline."""
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    base_version = (
        db.get(ContractVersion, contract.current_authoritative_version_id)
        if contract.current_authoritative_version_id
        else None
    )
    if base_version is None or not base_version.text_snapshot_id:
        raise HTTPException(
            status.HTTP_409_CONFLICT, "Contract has no text to redline yet"
        )
    snapshot = db.get(ContractTextSnapshot, base_version.text_snapshot_id)
    if snapshot is None or not (snapshot.text or "").strip():
        raise HTTPException(
            status.HTTP_409_CONFLICT, "Contract has no extracted text to redline"
        )
    text = snapshot.text

    # Locate the selection: trust the viewer's offset when it matches, else
    # fall back to the first occurrence of the quoted passage.
    start = -1
    hint = payload.start_hint
    if (
        hint is not None
        and 0 <= hint <= len(text) - len(payload.original_text)
        and text[hint : hint + len(payload.original_text)] == payload.original_text
    ):
        start = hint
    else:
        start = text.find(payload.original_text)
    if start < 0:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            "Selected text was not found in the current document text",
        )
    end = start + len(payload.original_text)
    proposed_text = text[:start] + payload.replacement_text + text[end:]

    author_name = current_user.full_name or "Reviewer"
    docx_bytes = _build_manual_redline_docx(
        contract_title=contract.title,
        base_version_number=base_version.version_number,
        author_name=author_name,
        pre_text=text[:start],
        original_text=payload.original_text,
        replacement_text=payload.replacement_text,
        post_text=text[end:],
    )
    contract_file = db.get(ContractFile, base_version.contract_file_id)
    if contract_file is None:
        raise HTTPException(status.HTTP_409_CONFLICT, "Contract file record missing")
    storage_object = _store_generated_docx(
        db,
        org_id=current_user.org_id,
        user_id=current_user.id,
        filename=f"{contract.title[:60]}-redline-v{base_version.version_number}.docx",
        content=docx_bytes,
    )
    proposal_version = ContractVersion(
        org_id=current_user.org_id,
        contract_id=contract.id,
        contract_file_id=contract_file.id,
        version_number=next_version_number(db, contract_file.id),
        storage_object_id=storage_object.id,
        source=ContractVersionSource.USER_REDLINE,
        change_summary=(payload.rationale or f"Manual redline by {author_name}")[:240],
        is_authoritative=False,
        created_by_user_id=current_user.id,
        updated_by_user_id=current_user.id,
    )
    db.add(proposal_version)
    db.flush()
    proposal_snapshot = ContractTextSnapshot(
        org_id=current_user.org_id,
        contract_id=contract.id,
        contract_version_id=proposal_version.id,
        extraction_method="manual_edit_text",
        extraction_quality_score=1.0,
        text=proposed_text,
        page_map=snapshot.page_map,
        ocr_provider=snapshot.ocr_provider,
        validation_status="complete",
        created_by_user_id=current_user.id,
        updated_by_user_id=current_user.id,
    )
    db.add(proposal_snapshot)
    db.flush()
    proposal_version.text_snapshot_id = proposal_snapshot.id

    edit = ContractEdit(
        org_id=current_user.org_id,
        contract_id=contract.id,
        contract_version_id=base_version.id,
        edit_type="manual",
        status="proposed",
        original_text=payload.original_text,
        replacement_text=payload.replacement_text or None,
        rationale=payload.rationale,
        citation=[
            {
                "type": "assistant_edit_version",
                "contract_version_id": proposal_version.id,
            },
            {
                "type": "anchor",
                "start": start,
                "end": end,
                "matched": True,
                "applied": True,
                "risk_level": "medium",
            },
            {"type": "author", "name": author_name, "kind": "user"},
        ],
        created_by_user_id=current_user.id,
        updated_by_user_id=current_user.id,
    )
    db.add(edit)
    db.flush()
    write_audit_log(
        db,
        org_id=current_user.org_id,
        actor_user_id=current_user.id,
        action="contract.manual_edit_proposed",
        resource_type="contract_edit",
        resource_id=edit.id,
        after={
            "contract_id": contract.id,
            "proposal_version_id": proposal_version.id,
        },
    )
    write_timeline_event(
        db,
        org_id=current_user.org_id,
        resource_type="contract",
        resource_id=contract.id,
        event_type="contract.manual_edit_proposed",
        title="Manual redline proposed",
        actor_user_id=current_user.id,
        details={"contract_edit_id": edit.id, "contract_version_id": proposal_version.id},
    )
    db.commit()
    db.refresh(edit)
    return edit


class ManualTextUpdate(BaseModel):
    text: str = Field(min_length=1)
    change_summary: str | None = None


@router.put("/text", response_model=ContractVersionResponse)
def update_contract_text(
    contract_id: str,
    payload: ManualTextUpdate,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract:redline")),
):
    """Direct in-document editing (Word-style): replace the contract's current
    text. Versions stay immutable — this creates a NEW authoritative version
    (source=manual_edit) with a fresh snapshot and a regenerated .docx, so the
    edit is auditable and reversible via version history."""
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    base_version = (
        db.get(ContractVersion, contract.current_authoritative_version_id)
        if contract.current_authoritative_version_id
        else None
    )
    if base_version is None:
        raise HTTPException(status.HTTP_409_CONFLICT, "Contract has no version to edit")
    base_snapshot = (
        db.get(ContractTextSnapshot, base_version.text_snapshot_id)
        if base_version.text_snapshot_id
        else None
    )
    if base_snapshot is not None and base_snapshot.text == payload.text:
        raise HTTPException(status.HTTP_409_CONFLICT, "No changes to save")
    contract_file = db.get(ContractFile, base_version.contract_file_id)
    if contract_file is None:
        raise HTTPException(status.HTTP_409_CONFLICT, "Contract file record missing")

    author_name = current_user.full_name or "Reviewer"
    summary = (payload.change_summary or f"Manual edit by {author_name}")[:240]
    docx_bytes = _build_plain_docx(
        title=contract.title,
        subtitle=summary,
        text=payload.text,
    )
    storage_object = _store_generated_docx(
        db,
        org_id=current_user.org_id,
        user_id=current_user.id,
        filename=f"{contract.title[:60]}-manual-edit.docx",
        content=docx_bytes,
    )
    new_version = ContractVersion(
        org_id=current_user.org_id,
        contract_id=contract.id,
        contract_file_id=contract_file.id,
        version_number=next_version_number(db, contract_file.id),
        storage_object_id=storage_object.id,
        source=ContractVersionSource.MANUAL_EDIT,
        change_summary=summary,
        is_authoritative=False,
        created_by_user_id=current_user.id,
        updated_by_user_id=current_user.id,
    )
    db.add(new_version)
    db.flush()
    snapshot = ContractTextSnapshot(
        org_id=current_user.org_id,
        contract_id=contract.id,
        contract_version_id=new_version.id,
        extraction_method="manual_edit_text",
        extraction_quality_score=1.0,
        text=payload.text,
        page_map=base_snapshot.page_map if base_snapshot else None,
        ocr_provider=base_snapshot.ocr_provider if base_snapshot else None,
        validation_status="complete",
        created_by_user_id=current_user.id,
        updated_by_user_id=current_user.id,
    )
    db.add(snapshot)
    db.flush()
    new_version.text_snapshot_id = snapshot.id

    # Exclusive authoritative flip — same pattern as accepting a redline.
    versions = db.scalars(
        select(ContractVersion).where(
            ContractVersion.org_id == current_user.org_id,
            ContractVersion.contract_id == contract_id,
            ContractVersion.deleted_at.is_(None),
        )
    ).all()
    for version in versions:
        version.is_authoritative = version.id == new_version.id
        version.updated_by_user_id = current_user.id
    contract.current_authoritative_version_id = new_version.id
    contract.current_contract_file_id = contract_file.id
    contract_file.current_version_id = new_version.id
    contract.updated_by_user_id = current_user.id
    contract_file.updated_by_user_id = current_user.id

    write_audit_log(
        db,
        action="contract.text_manually_edited",
        resource_type="contract_version",
        resource_id=new_version.id,
        org_id=current_user.org_id,
        actor_user_id=current_user.id,
        metadata={"contract_id": contract.id, "summary": summary},
    )
    write_timeline_event(
        db,
        org_id=current_user.org_id,
        resource_type="contract",
        resource_id=contract.id,
        event_type="contract.text_manually_edited",
        title="Document edited manually",
        actor_user_id=current_user.id,
        details={"contract_version_id": new_version.id, "summary": summary},
    )
    db.commit()
    db.refresh(new_version)
    return new_version


@router.get("/export-docx")
def export_contract_docx(
    contract_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract:read")),
):
    """Export the CURRENT authoritative text as a .docx — so what you download
    always reflects accepted redlines, not the originally uploaded binary."""
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    version = (
        db.get(ContractVersion, contract.current_authoritative_version_id)
        if contract.current_authoritative_version_id
        else None
    )
    snapshot = (
        db.get(ContractTextSnapshot, version.text_snapshot_id)
        if version is not None and version.text_snapshot_id
        else None
    )
    if snapshot is None or not (snapshot.text or "").strip():
        raise HTTPException(status.HTTP_409_CONFLICT, "No extracted text to export")
    content = _build_plain_docx(
        title=contract.title,
        subtitle=f"Current text · V{version.version_number}",
        text=snapshot.text,
    )
    filename = f"{contract.title[:60]}-v{version.version_number}.docx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/edits/{edit_id}/accept", response_model=ContractEditResponse)
def accept_contract_edit(
    contract_id: str,
    edit_id: str,
    payload: ContractEditDecisionRequest,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract:redline")),
):
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    edit = _get_contract_edit(db, contract_id=contract_id, edit_id=edit_id, org_id=current_user.org_id)
    if edit.status != "proposed":
        raise HTTPException(status.HTTP_409_CONFLICT, "Only proposed edits can be accepted")
    proposal_version = _proposal_version_for_edit(db, edit=edit, org_id=current_user.org_id)
    contract_file = db.get(ContractFile, proposal_version.contract_file_id)
    if contract_file is None or contract_file.org_id != current_user.org_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract file not found")
    versions = db.scalars(
        select(ContractVersion).where(
            ContractVersion.org_id == current_user.org_id,
            ContractVersion.contract_id == contract_id,
            ContractVersion.deleted_at.is_(None),
        )
    ).all()
    for version in versions:
        version.is_authoritative = version.id == proposal_version.id
        version.updated_by_user_id = current_user.id
    proposal_version.change_summary = _decision_summary(
        proposal_version.change_summary,
        decision="accepted",
        comment=payload.comment,
    )
    contract.current_authoritative_version_id = proposal_version.id
    contract.current_contract_file_id = contract_file.id
    contract_file.current_version_id = proposal_version.id
    edit.status = "accepted"
    edit.updated_by_user_id = current_user.id
    contract.updated_by_user_id = current_user.id
    contract_file.updated_by_user_id = current_user.id
    write_audit_log(
        db,
        action="contract.edit_accepted",
        resource_type="contract_edit",
        resource_id=edit.id,
        org_id=current_user.org_id,
        actor_user_id=current_user.id,
        metadata={
            "contract_id": contract_id,
            "contract_version_id": proposal_version.id,
            "comment": payload.comment,
        },
    )
    write_timeline_event(
        db,
        org_id=current_user.org_id,
        resource_type="contract",
        resource_id=contract.id,
        event_type="contract.edit_accepted",
        title="Assistant tracked change accepted",
        actor_user_id=current_user.id,
        details={"contract_edit_id": edit.id, "contract_version_id": proposal_version.id},
    )
    requeue_contract_ai_jobs(
        db, user=current_user, contract=contract, version=proposal_version
    )
    db.commit()
    db.refresh(edit)
    return edit


@router.post("/edits/{edit_id}/reject", response_model=ContractEditResponse)
def reject_contract_edit(
    contract_id: str,
    edit_id: str,
    payload: ContractEditDecisionRequest,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract:redline")),
):
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    edit = _get_contract_edit(db, contract_id=contract_id, edit_id=edit_id, org_id=current_user.org_id)
    if edit.status != "proposed":
        raise HTTPException(status.HTTP_409_CONFLICT, "Only proposed edits can be rejected")
    proposal_version = _proposal_version_for_edit(db, edit=edit, org_id=current_user.org_id)
    proposal_version.change_summary = _decision_summary(
        proposal_version.change_summary,
        decision="rejected",
        comment=payload.comment,
    )
    edit.status = "rejected"
    edit.updated_by_user_id = current_user.id
    proposal_version.updated_by_user_id = current_user.id
    write_audit_log(
        db,
        action="contract.edit_rejected",
        resource_type="contract_edit",
        resource_id=edit.id,
        org_id=current_user.org_id,
        actor_user_id=current_user.id,
        metadata={
            "contract_id": contract_id,
            "contract_version_id": proposal_version.id,
            "comment": payload.comment,
        },
    )
    write_timeline_event(
        db,
        org_id=current_user.org_id,
        resource_type="contract",
        resource_id=contract.id,
        event_type="contract.edit_rejected",
        title="Assistant tracked change rejected",
        actor_user_id=current_user.id,
        details={"contract_edit_id": edit.id, "contract_version_id": proposal_version.id},
    )
    db.commit()
    db.refresh(edit)
    return edit


@router.post("/versions/{version_id}/restore", response_model=ContractVersionResponse)
def restore_contract_version(
    contract_id: str,
    version_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:update")),
):
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    version = db.get(ContractVersion, version_id)
    if version is None or version.contract_id != contract_id or version.org_id != current_user.org_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract version not found")
    contract_file = db.get(ContractFile, version.contract_file_id)
    if contract_file is None or contract_file.org_id != current_user.org_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract file not found")
    restored_version = ContractVersion(
        org_id=current_user.org_id,
        contract_id=contract_id,
        contract_file_id=contract_file.id,
        version_number=next_version_number(db, contract_file.id),
        storage_object_id=version.storage_object_id,
        source=ContractVersionSource.RESTORED,
        change_summary=f"Restored from version {version.version_number}",
        is_authoritative=True,
        created_by_user_id=current_user.id,
        updated_by_user_id=current_user.id,
    )
    db.add(restored_version)
    db.flush()

    source_snapshot = (
        db.get(ContractTextSnapshot, version.text_snapshot_id) if version.text_snapshot_id else None
    )
    if source_snapshot is not None:
        restored_snapshot = ContractTextSnapshot(
            org_id=current_user.org_id,
            contract_id=contract_id,
            contract_version_id=restored_version.id,
            extraction_method=source_snapshot.extraction_method,
            extraction_quality_score=source_snapshot.extraction_quality_score,
            text=source_snapshot.text,
            page_map=source_snapshot.page_map,
            ocr_provider=source_snapshot.ocr_provider,
            validation_status=source_snapshot.validation_status,
            created_by_user_id=current_user.id,
            updated_by_user_id=current_user.id,
        )
        db.add(restored_snapshot)
        db.flush()
        restored_version.text_snapshot_id = restored_snapshot.id

    existing_versions = db.scalars(
        select(ContractVersion).where(
            ContractVersion.org_id == current_user.org_id,
            ContractVersion.contract_id == contract_id,
            ContractVersion.deleted_at.is_(None),
        )
    ).all()
    for row in existing_versions:
        row.is_authoritative = row.id == restored_version.id
        row.updated_by_user_id = current_user.id
    contract_file.current_version_id = restored_version.id
    contract_file.updated_by_user_id = current_user.id
    contract.current_contract_file_id = contract_file.id
    contract.current_authoritative_version_id = restored_version.id
    contract.updated_by_user_id = current_user.id
    write_audit_log(
        db,
        action="contract.version_restored",
        resource_type="contract_version",
        resource_id=restored_version.id,
        org_id=current_user.org_id,
        actor_user_id=current_user.id,
        metadata={
            "contract_id": contract.id,
            "contract_file_id": contract_file.id,
            "restored_from_version_id": version.id,
            "restored_from_version_number": version.version_number,
        },
    )
    write_timeline_event(
        db,
        org_id=current_user.org_id,
        resource_type="contract",
        resource_id=contract.id,
        event_type="contract.version_restored",
        title="Contract version restored",
        actor_user_id=current_user.id,
        details={
            "contract_version_id": restored_version.id,
            "version_number": restored_version.version_number,
            "restored_from_version_id": version.id,
            "restored_from_version_number": version.version_number,
        },
    )
    requeue_contract_ai_jobs(
        db, user=current_user, contract=contract, version=restored_version
    )
    db.commit()
    db.refresh(restored_version)
    return restored_version


@router.delete("/versions/{version_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_contract_version(
    contract_id: str,
    version_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:delete")),
):
    get_contract_for_user(db, contract_id=contract_id, user=current_user)
    version = db.get(ContractVersion, version_id)
    if version is None or version.contract_id != contract_id or version.org_id != current_user.org_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract version not found")
    if version.is_authoritative:
        raise HTTPException(status.HTTP_409_CONFLICT, "Authoritative version cannot be deleted")
    contract_file = db.get(ContractFile, version.contract_file_id)
    if contract_file is not None and contract_file.current_version_id == version_id:
        raise HTTPException(status.HTTP_409_CONFLICT, "Current file version cannot be deleted")

    version.deleted_at = utcnow()
    version.deleted_by_user_id = current_user.id
    version.updated_by_user_id = current_user.id
    if version.text_snapshot_id:
        snapshot = db.get(ContractTextSnapshot, version.text_snapshot_id)
        if snapshot is not None and snapshot.deleted_at is None:
            snapshot.deleted_at = version.deleted_at
            snapshot.deleted_by_user_id = current_user.id
            snapshot.updated_by_user_id = current_user.id

    active_storage_refs = db.scalar(
        select(func.count(ContractVersion.id)).where(
            ContractVersion.org_id == current_user.org_id,
            ContractVersion.storage_object_id == version.storage_object_id,
            ContractVersion.deleted_at.is_(None),
            ContractVersion.id != version.id,
        )
    )
    if not active_storage_refs:
        storage_object = db.get(StorageObject, version.storage_object_id)
        if storage_object is not None and storage_object.deleted_at is None:
            storage_object.deleted_at = version.deleted_at
            storage_object.deleted_by_user_id = current_user.id
            storage_object.updated_by_user_id = current_user.id

    write_audit_log(
        db,
        action="contract.version_deleted",
        resource_type="contract_version",
        resource_id=version.id,
        org_id=current_user.org_id,
        actor_user_id=current_user.id,
        metadata={"contract_id": contract_id, "storage_object_id": version.storage_object_id},
    )
    write_timeline_event(
        db,
        org_id=current_user.org_id,
        resource_type="contract",
        resource_id=contract_id,
        event_type="contract.version_deleted",
        title="Contract version deleted",
        actor_user_id=current_user.id,
        details={"contract_version_id": version.id, "version_number": version.version_number},
    )
    db.commit()


@router.get("/shares", response_model=list[ContractShareResponse])
def list_contract_shares(
    contract_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:share")),
):
    get_contract_for_user(db, contract_id=contract_id, user=current_user)
    return db.scalars(
        select(ContractShare)
        .where(
            ContractShare.org_id == current_user.org_id,
            ContractShare.contract_id == contract_id,
            ContractShare.deleted_at.is_(None),
        )
        .order_by(ContractShare.created_at.desc())
    ).all()


@router.post(
    "/shares",
    response_model=ContractShareCreateResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_contract_share(
    contract_id: str,
    payload: ContractShareCreate,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:share")),
):
    contract = get_contract_for_user(db, contract_id=contract_id, user=current_user)
    if payload.contract_version_id:
        version = db.get(ContractVersion, payload.contract_version_id)
        if version is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract version not found")
    elif contract.current_authoritative_version_id:
        version = db.get(ContractVersion, contract.current_authoritative_version_id)
    else:
        version = None
    if version is not None and (version.org_id != current_user.org_id or version.contract_id != contract_id):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract version not found")
    if payload.access_mode == ShareAccessMode.VIEW_ONLY and payload.download_allowed:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "View-only shares cannot allow downloads")
    token = secrets.token_urlsafe(32)
    share = ContractShare(
        org_id=current_user.org_id,
        contract_id=contract_id,
        contract_version_id=version.id if version else None,
        token_hash=_hash_secret(token),
        passcode_hash=_hash_secret(payload.passcode) if payload.passcode else None,
        access_mode=payload.access_mode,
        expires_at=payload.expires_at,
        download_allowed=payload.download_allowed or payload.access_mode == ShareAccessMode.DOWNLOAD_ALLOWED,
        created_by_user_id=current_user.id,
        updated_by_user_id=current_user.id,
    )
    db.add(share)
    db.flush()
    write_audit_log(
        db,
        action="contract.share_created",
        resource_type="contract_share",
        resource_id=share.id,
        org_id=current_user.org_id,
        actor_user_id=current_user.id,
        metadata={
            "contract_id": contract_id,
            "contract_version_id": share.contract_version_id,
            "download_allowed": share.download_allowed,
            "expires_at": share.expires_at.isoformat() if share.expires_at else None,
        },
    )
    write_timeline_event(
        db,
        org_id=current_user.org_id,
        resource_type="contract",
        resource_id=contract_id,
        event_type="contract.share_created",
        title="Contract share created",
        actor_user_id=current_user.id,
        details={"share_id": share.id, "download_allowed": share.download_allowed},
    )
    db.commit()
    db.refresh(share)
    return {"share": share, "token": token}


@router.post("/shares/{share_id}/revoke", response_model=ContractShareResponse)
def revoke_contract_share(
    contract_id: str,
    share_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(require_permission("contract_file:share")),
):
    get_contract_for_user(db, contract_id=contract_id, user=current_user)
    share = db.get(ContractShare, share_id)
    if (
        share is None
        or share.org_id != current_user.org_id
        or share.contract_id != contract_id
        or share.deleted_at is not None
    ):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract share not found")
    share.revoked_at = utcnow()
    share.updated_by_user_id = current_user.id
    write_audit_log(
        db,
        action="contract.share_revoked",
        resource_type="contract_share",
        resource_id=share.id,
        org_id=current_user.org_id,
        actor_user_id=current_user.id,
        metadata={"contract_id": contract_id},
    )
    db.commit()
    db.refresh(share)
    return share


@external_share_router.get("/{token}", response_model=ExternalShareResponse)
@limiter.limit("20/minute", key_func=_share_rate_limit_key)
def view_external_share(
    request: Request,
    response: Response,
    token: str,
    passcode: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    share = _get_active_share(db, token=token, passcode=passcode, request=request)
    contract = db.get(Contract, share.contract_id)
    if contract is None or contract.deleted_at is not None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Shared contract not found")
    version = _share_version(db, share=share, contract=contract)
    storage_object = db.get(StorageObject, version.storage_object_id) if version else None
    snapshot = db.get(ContractTextSnapshot, version.text_snapshot_id) if version and version.text_snapshot_id else None
    text = snapshot.text if snapshot else ""
    write_audit_log(
        db,
        action="contract.external_share_viewed",
        resource_type="contract_share",
        resource_id=share.id,
        org_id=share.org_id,
        metadata={"contract_id": contract.id, "contract_version_id": version.id if version else None},
    )
    db.commit()
    return ExternalShareResponse(
        contract_id=contract.id,
        contract_version_id=version.id if version else None,
        title=contract.title,
        filename=storage_object.filename if storage_object else None,
        access_mode=share.access_mode,
        download_allowed=share.download_allowed,
        text_excerpt=text[:EXTERNAL_TEXT_EXCERPT_CHARS] if text else None,
        text_truncated=len(text) > EXTERNAL_TEXT_EXCERPT_CHARS,
    )


def _share_contract(db: Session, share) -> Contract:
    contract = db.get(Contract, share.contract_id)
    if contract is None or contract.deleted_at is not None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Shared contract not found")
    return contract


@external_share_router.get("/{token}/comments", response_model=list[ExternalCommentResponse])
@limiter.limit("30/minute", key_func=_share_rate_limit_key)
def list_external_share_comments(
    request: Request,
    response: Response,
    token: str,
    passcode: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    """Shared (counterparty-visible) comments on the shared contract. Internal
    comments are never exposed here."""
    share = _get_active_share(db, token=token, passcode=passcode, request=request)
    contract = _share_contract(db, share)
    return list_shared_comments(db, contract=contract)


@external_share_router.post(
    "/{token}/comments",
    response_model=ExternalCommentResponse,
    status_code=status.HTTP_201_CREATED,
)
@limiter.limit("10/minute", key_func=_share_rate_limit_key)
def add_external_share_comment(
    request: Request,
    response: Response,
    token: str,
    payload: ExternalCommentCreate,
    passcode: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    """Let the counterparty leave a (shared) comment via the share link."""
    share = _get_active_share(db, token=token, passcode=passcode, request=request)
    contract = _share_contract(db, share)
    return add_counterparty_comment(
        db,
        contract=contract,
        author_name=payload.author_name,
        body=payload.body,
        request_id=getattr(request.state, "request_id", None),
    )


@external_share_router.get("/{token}/download")
@limiter.limit("20/minute", key_func=_share_rate_limit_key)
def download_external_share(
    request: Request,
    token: str,
    passcode: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    share = _get_active_share(db, token=token, passcode=passcode, request=request)
    if not share.download_allowed:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Download is disabled for this share")
    contract = db.get(Contract, share.contract_id)
    if contract is None or contract.deleted_at is not None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Shared contract not found")
    version = _share_version(db, share=share, contract=contract)
    if version is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Shared contract version not found")
    storage_object = db.get(StorageObject, version.storage_object_id)
    if storage_object is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Stored file not found")
    try:
        path = storage_service.path_for_read(storage_object.storage_key)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Stored file bytes not found")
    write_audit_log(
        db,
        action="contract.external_share_downloaded",
        resource_type="contract_share",
        resource_id=share.id,
        org_id=share.org_id,
        metadata={"contract_id": contract.id, "contract_version_id": version.id},
    )
    db.commit()
    # attachment, never inline — see download_contract_version for rationale.
    # This matters most on the unauthenticated external-share path.
    return FileResponse(
        path,
        media_type=storage_object.mime_type,
        filename=storage_object.filename,
        content_disposition_type="attachment",
    )


_share_lockout_redis = None
_share_lockout_redis_init = False


def _get_lockout_redis():
    """Lazily build a Redis client for passcode-lockout counters.

    Returns ``None`` (and disables lockout) if the ``redis`` client can't be
    constructed, so a missing/broken Redis never blocks share access. The
    client is cached after the first attempt.
    """
    global _share_lockout_redis, _share_lockout_redis_init
    if _share_lockout_redis_init:
        return _share_lockout_redis
    _share_lockout_redis_init = True
    try:
        import redis  # imported lazily; only needed for the lockout counter

        _share_lockout_redis = redis.Redis.from_url(
            settings.redis_url, socket_timeout=0.25, socket_connect_timeout=0.25
        )
    except Exception as exc:  # pragma: no cover - defensive: redis missing/misconfigured
        logger.warning("Share passcode lockout disabled — Redis unavailable: %s", exc)
        _share_lockout_redis = None
    return _share_lockout_redis


def _share_lockout_keys(token: str, ip: str) -> tuple[str, str]:
    # Hash the token so the raw share secret never lands in a Redis key.
    digest = _hash_secret(f"{token}:{ip}")
    return f"share_pc_fail:{digest}", f"share_pc_lock:{digest}"


def _check_share_lockout(token: str, ip: str) -> None:
    """Raise 429 if this share+IP is currently locked out. Fails open."""
    client = _get_lockout_redis()
    if client is None:
        return
    _, lock_key = _share_lockout_keys(token, ip)
    try:
        if client.get(lock_key) is not None:
            raise HTTPException(
                status.HTTP_429_TOO_MANY_REQUESTS,
                "Too many incorrect passcode attempts — try again later.",
            )
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover - Redis hiccup: fail open
        logger.warning("Share lockout check failed, allowing request: %s", exc)


def _register_share_passcode_failure(token: str, ip: str) -> None:
    """Count a failed passcode attempt; arm an exponential lockout past the cap.

    Backoff doubles each excess failure (BASE, 2*BASE, 4*BASE, …) capped at
    SHARE_PASSCODE_LOCKOUT_MAX_SECONDS. Best-effort: Redis errors are swallowed.
    """
    client = _get_lockout_redis()
    if client is None:
        return
    fail_key, lock_key = _share_lockout_keys(token, ip)
    try:
        attempts = client.incr(fail_key)
        if attempts == 1:
            client.expire(fail_key, SHARE_PASSCODE_ATTEMPT_TTL_SECONDS)
        if attempts >= SHARE_PASSCODE_MAX_ATTEMPTS:
            overflow = attempts - SHARE_PASSCODE_MAX_ATTEMPTS
            lock_seconds = min(
                SHARE_PASSCODE_LOCKOUT_BASE_SECONDS * (2**overflow),
                SHARE_PASSCODE_LOCKOUT_MAX_SECONDS,
            )
            client.set(lock_key, "1", ex=lock_seconds)
    except Exception as exc:  # pragma: no cover - Redis hiccup: don't block auth flow
        logger.warning("Failed to record share passcode failure: %s", exc)


def _clear_share_passcode_failures(token: str, ip: str) -> None:
    """Reset counters after a successful passcode. Best-effort."""
    client = _get_lockout_redis()
    if client is None:
        return
    fail_key, lock_key = _share_lockout_keys(token, ip)
    try:
        client.delete(fail_key, lock_key)
    except Exception:  # pragma: no cover - non-critical cleanup
        pass


def _get_active_share(
    db: Session, *, token: str, passcode: str | None, request: Request | None = None
) -> ContractShare:
    # Client IP for brute-force keying; falls back to a constant when no request
    # is threaded (e.g. internal callers) so the helpers still key consistently.
    ip = get_remote_address(request) if request is not None else "internal"
    share = db.scalar(
        select(ContractShare).where(
            ContractShare.token_hash == _hash_secret(token),
            ContractShare.deleted_at.is_(None),
        )
    )
    if share is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Share not found")
    now = utcnow()
    if share.revoked_at is not None or (share.expires_at is not None and share.expires_at < now):
        raise HTTPException(status.HTTP_410_GONE, "Share is no longer active")
    if share.passcode_hash:
        # Reject early while a lockout window is active, before doing the compare.
        _check_share_lockout(token, ip)
        if not secrets.compare_digest(_hash_secret(passcode or ""), share.passcode_hash):
            _register_share_passcode_failure(token, ip)
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Passcode required")
        _clear_share_passcode_failures(token, ip)
    elif passcode:
        # No passcode is required on this share but the caller supplied one
        # anyway. Log it so abuse review can spot probing behaviour, but do
        # not block — legitimate users sometimes resubmit forms with stale
        # values from a passcode-gated link they used previously.
        write_audit_log(
            db,
            action="contract.external_share_passcode_unexpected",
            resource_type="contract_share",
            resource_id=share.id,
            org_id=share.org_id,
            metadata={"contract_id": share.contract_id},
        )
        db.commit()
    return share


def _share_version(db: Session, *, share: ContractShare, contract: Contract) -> ContractVersion | None:
    version_id = share.contract_version_id or contract.current_authoritative_version_id
    if not version_id:
        return None
    version = db.get(ContractVersion, version_id)
    if version is None or version.org_id != share.org_id or version.contract_id != contract.id:
        return None
    return version


def _get_contract_edit(db: Session, *, contract_id: str, edit_id: str, org_id: str) -> ContractEdit:
    edit = db.get(ContractEdit, edit_id)
    if edit is None or edit.org_id != org_id or edit.contract_id != contract_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Contract edit not found")
    return edit


def _proposal_version_for_edit(db: Session, *, edit: ContractEdit, org_id: str) -> ContractVersion:
    proposal_version_id = None
    for citation in edit.citation or []:
        if isinstance(citation, dict) and citation.get("type") == "assistant_edit_version":
            proposal_version_id = citation.get("contract_version_id")
            break
    if not proposal_version_id:
        # Resilience for edits whose citation lost the version link: fall back
        # to the most recent non-authoritative assistant-edit/redline version
        # for this contract (the proposal version a redline batch produced).
        fallback = db.scalar(
            select(ContractVersion)
            .where(
                ContractVersion.org_id == org_id,
                ContractVersion.contract_id == edit.contract_id,
                ContractVersion.is_authoritative.is_(False),
                ContractVersion.deleted_at.is_(None),
                ContractVersion.source.in_(
                    [
                        ContractVersionSource.ASSISTANT_EDIT,
                        ContractVersionSource.PLAYBOOK_REDLINE,
                        ContractVersionSource.ASSISTANT_GENERATED,
                        ContractVersionSource.USER_REDLINE,
                    ]
                ),
            )
            .order_by(ContractVersion.created_at.desc())
        )
        if fallback is None:
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                "Edit has no assistant version to apply",
            )
        return fallback
    version = db.get(ContractVersion, proposal_version_id)
    if version is None or version.org_id != org_id or version.contract_id != edit.contract_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assistant edit version not found")
    return version


def _decision_summary(summary: str | None, *, decision: str, comment: str | None) -> str:
    suffix = f"Assistant edit {decision}."
    if comment:
        suffix = f"{suffix} Comment: {comment}"
    if not summary:
        return suffix
    return f"{summary}\n{suffix}"


def _hash_secret(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()
