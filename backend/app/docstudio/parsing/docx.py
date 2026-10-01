"""Word parsing.

Two things about `.docx` that a naive reader gets wrong, and both are silent:

**Numbering is not in the text.** A numbered paragraph stores only "item of list
N at level L"; Word counts and draws "2.1" while laying out the page. So
`paragraph.text` for `2.1 Confidential Information...` is just `Confidential
Information...`. Clause boundaries are usually *found* by that leading number, so
an auto-numbered contract arrives as one undifferentiated block and no clause is
separable. Worse, the numbering can live on the paragraph OR on its style (Word's
own `List Number` does the latter) and styles inherit through `basedOn`.

**`Document.paragraphs` silently omits every table.** In a contract that is the
fee schedule, the SLA tiers, the liability cap and the signature block.

Both are handled here by walking the XML body directly.
"""

from dataclasses import replace
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from docx.oxml.ns import qn

from .base import ParsedBlock, ParsedDocument, UnsupportedFormat
from .labels import label_only, level_for, split_label

# A zip bomb guard: a .docx is a zip, and an attacker controls it.
_MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
_MAX_COMPRESSION_RATIO = 100

_ROMAN = [
    (1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
    (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"),
]


def _roman(n: int) -> str:
    out = ""
    for value, sign in _ROMAN:
        while n >= value:
            out += sign
            n -= value
    return out


def _letter(n: int) -> str:
    # Word repeats the letter past z (aa, bb, cc) rather than carrying like a
    # base-26 number.
    return chr(ord("a") + (n - 1) % 26) * ((n - 1) // 26 + 1)


def _format_counter(n: int, fmt: str) -> str:
    if fmt == "lowerLetter":
        return _letter(n)
    if fmt == "upperLetter":
        return _letter(n).upper()
    if fmt == "lowerRoman":
        return _roman(n)
    if fmt == "upperRoman":
        return _roman(n).upper()
    if fmt == "decimalZero":
        return f"{n:02d}"
    return str(n)


def _child_val(element, tag: str) -> str | None:
    child = element.find(qn(tag))
    return None if child is None else child.get(qn("w:val"))


def _num_pr(element) -> tuple[str, int] | None:
    """(numId, ilvl) from a paragraph's or a style's pPr, if it carries one."""
    if element is None:
        return None
    p_pr = element.find(qn("w:pPr"))
    num_pr = None if p_pr is None else p_pr.find(qn("w:numPr"))
    if num_pr is None:
        return None
    num_id = _child_val(num_pr, "w:numId")
    if num_id is None:
        return None
    return num_id, int(_child_val(num_pr, "w:ilvl") or 0)


class _Numbering:
    """Recomputes the clause numbers Word never writes down.

    Counting is stateful — a level increments and every deeper level restarts —
    which is why no library hands this over and why `label()` must be called
    once per paragraph in document order.
    """

    def __init__(self, document):
        self._levels: dict[str, dict[int, dict]] = {}
        self._abstract_of: dict[str, str] = {}
        self._overrides: dict[str, dict[int, int]] = {}
        self._counters: dict[str, dict[int, int]] = {}
        self._style_cache: dict[str, tuple[str, int] | None] = {}
        try:
            root = document.part.numbering_part.element
        except Exception:
            return  # nothing in this document is numbered
        for abstract in root.findall(qn("w:abstractNum")):
            abstract_id = abstract.get(qn("w:abstractNumId"))
            if abstract_id is None:
                continue
            levels: dict[int, dict] = {}
            for lvl in abstract.findall(qn("w:lvl")):
                levels[int(lvl.get(qn("w:ilvl")) or 0)] = {
                    "fmt": _child_val(lvl, "w:numFmt") or "decimal",
                    "text": _child_val(lvl, "w:lvlText") or "",
                    "start": int(_child_val(lvl, "w:start") or 1),
                    # "Legal numbering": render every parent level as an arabic
                    # numeral whatever its own format. Common in contracts that
                    # mix "Article IV" with "4.1".
                    "is_lgl": lvl.find(qn("w:isLgl")) is not None,
                }
            self._levels[abstract_id] = levels
        for num in root.findall(qn("w:num")):
            num_id = num.get(qn("w:numId"))
            abstract = num.find(qn("w:abstractNumId"))
            if num_id is None or abstract is None:
                continue
            self._abstract_of[num_id] = abstract.get(qn("w:val"))
            for override in num.findall(qn("w:lvlOverride")):
                start = override.find(qn("w:startOverride"))
                if start is not None:
                    ilvl = int(override.get(qn("w:ilvl")) or 0)
                    self._overrides.setdefault(num_id, {})[ilvl] = int(
                        start.get(qn("w:val")) or 1
                    )

    def label(self, paragraph) -> tuple[str, int, str | None]:
        """(visible number, level, list). ("", 1, None) when not numbered.

        The list is Word's abstract numbering definition, because a level only
        means something inside its own list: an (a)/(b) list under clause 1.1
        is usually a separate list that starts again at level 1.
        """
        if not self._levels:
            return "", 1, None
        found = _num_pr(paragraph._p) or self._from_style(paragraph)
        if found is None:
            return "", 1, None
        num_id, ilvl = found
        if num_id == "0":
            return "", 1, None  # Word's explicit "not numbered", used to lift a style
        abstract_id = self._abstract_of.get(num_id)
        levels = self._levels.get(abstract_id or "")
        if not levels:
            return "", 1, None
        level = levels.get(ilvl)
        if level is None or level["fmt"] in {"bullet", "none"}:
            # Bullets carry no clause identity and the marker would land inside
            # every quoted citation. Their nesting level still counts.
            return "", ilvl + 1, abstract_id
        # Counters key on the abstract definition so the several numIds Word
        # emits for one visual list share a sequence.
        counters = self._counters.setdefault(abstract_id, {})
        start = self._overrides.get(num_id, {}).get(ilvl, level["start"])
        counters[ilvl] = counters.get(ilvl, start - 1) + 1
        for deeper in [k for k in counters if k > ilvl]:
            del counters[deeper]
        label = level["text"]
        for position in range(ilvl + 1):
            parent = levels.get(position) or level
            fmt = "decimal" if (level["is_lgl"] and position < ilvl) else parent["fmt"]
            value = counters.get(position, parent["start"])
            label = label.replace(f"%{position + 1}", _format_counter(value, fmt))
        return label.strip(), ilvl + 1, abstract_id

    def _from_style(self, paragraph) -> tuple[str, int] | None:
        """Word's built-in `List Number` defines its numbering in styles.xml,
        and a house style inherits it through `basedOn`. Reading only the
        paragraph finds nothing for any of them."""
        style = paragraph.style
        key = getattr(style, "style_id", None)
        if key in self._style_cache:
            return self._style_cache[key]
        found = None
        seen = 0
        while style is not None and seen < 20:  # guard a basedOn cycle
            found = _num_pr(getattr(style, "element", None))
            if found is not None:
                break
            style = style.base_style
            seen += 1
        if key is not None:
            self._style_cache[key] = found
        return found


# Indentation is measured in twips — twentieths of a point, 1440 to the inch.
# Values are bucketed to an eighth of an inch so a stray 722 does not become a
# nesting level of its own.
_TWIPS_PER_EIGHTH_INCH = 180
# A document could indent a dozen different amounts; beyond a handful the depth
# stops meaning anything a reader would recognise as structure.
_MAX_INDENT_LEVELS = 5


def _effective_indent(paragraph) -> int:
    """How far this paragraph is indented, in twips.

    Checked on the paragraph and then up its style chain, for the same reason
    numbering is: Word templates put indentation on a named style ("Quote") and
    house styles inherit it through `basedOn`.
    """
    for element in _style_chain(paragraph):
        p_pr = element.find(qn("w:pPr"))
        ind = None if p_pr is None else p_pr.find(qn("w:ind"))
        if ind is None:
            continue
        # w:start is the newer spelling of w:left.
        raw = ind.get(qn("w:left")) or ind.get(qn("w:start"))
        if raw is not None:
            try:
                return max(0, int(raw))
            except ValueError:
                return 0
    return 0


def _style_chain(paragraph):
    """The paragraph element, then each style it inherits from."""
    yield paragraph._p
    style = paragraph.style
    seen = 0
    while style is not None and seen < 20:  # guard a basedOn cycle
        element = getattr(style, "element", None)
        if element is not None:
            yield element
        style = style.base_style
        seen += 1


def _indent_levels(indents: list[int]) -> dict[int, int]:
    """Map each distinct indent to a depth, shallowest first.

    Computed across the whole document rather than per paragraph: what makes an
    indent mean "one level in" is that other paragraphs sit further out, and
    that is only knowable once every paragraph has been seen.
    """
    buckets = sorted({i // _TWIPS_PER_EIGHTH_INCH for i in indents})
    return {bucket: depth for depth, bucket in enumerate(buckets[:_MAX_INDENT_LEVELS])}


def _guard_zip(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = [i for i in archive.infolist() if not i.is_dir()]
            uncompressed = sum(i.file_size for i in infos)
            compressed = sum(max(i.compress_size, 1) for i in infos)
    except BadZipFile as exc:
        raise UnsupportedFormat("Not a valid DOCX container") from exc
    if uncompressed > _MAX_UNCOMPRESSED_BYTES:
        raise UnsupportedFormat("DOCX expands beyond the safe extraction limit")
    if compressed and uncompressed / compressed > _MAX_COMPRESSION_RATIO:
        raise UnsupportedFormat("DOCX compression ratio exceeds the safe limit")


def _iter_body(parent_elm, parent):
    """Paragraphs and tables in document order.

    python-docx keeps them in separate collections with no ordering between
    them, so the only way to read the document as written is to walk the body.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _paragraph_text(paragraph) -> str:
    """Every character Word would show, including tracked insertions.

    `Paragraph.text` reads only the runs sitting directly in the paragraph, so
    anything Word wrapped loses its text entirely. A counterparty redline is
    exactly that shape:

        Liability is capped at <w:ins>twelve (12) months</w:ins>
                               <w:del>six (6) months</w:del> of fees.

    which came out as "Liability is capped at  of fees." — a sentence that
    reads as though there is no cap at all. The same blindness drops hyperlink
    text, because Word wraps that too.

    Insertions are kept and deletions dropped, i.e. the document as it reads
    with the changes accepted. That is the text under negotiation: the wording
    the other side is proposing, which is the thing a reviewer needs to see.
    Deleted wording is history and belongs in the redline, not in the clause.
    """
    parts: list[str] = []
    for node in paragraph._p.iter():
        tag = node.tag
        if tag == qn("w:t"):
            # w:delText is the spec's tag for deleted text and is skipped by
            # not being matched here. Some producers put a plain w:t inside
            # w:del anyway, so check the ancestry rather than trusting the tag.
            if not _inside_deletion(node):
                parts.append(node.text or "")
        elif tag == qn("w:tab"):
            parts.append("\t")
        elif tag in {qn("w:br"), qn("w:cr")}:
            parts.append("\n")
    return "".join(parts)


def _inside_deletion(node) -> bool:
    parent = node.getparent()
    while parent is not None:
        if parent.tag == qn("w:del"):
            return True
        parent = parent.getparent()
    return False


def _render_table(table) -> str:
    """One row per line, cells tab-separated — what pasting into a text editor
    gives. No markdown pipes: this text is quoted back in citations, where added
    syntax reads as damage."""
    lines: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            # A merged cell is returned once per grid position it spans, as the
            # same element; without this a header merged across three columns
            # is emitted three times.
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            # A cell holds paragraphs AND further tables. Reading only its
            # paragraphs drops a nested table whole: a fee schedule laid out
            # inside a layout table loses every figure in it, silently.
            inner = " ".join(_render_cell(cell)).replace("\n", " ")
            cells.append(" ".join(inner.split()))
        line = "\t".join(cells)
        if line.strip():
            lines.append(line)
    return "\n".join(lines)


def _paragraph_block(paragraph, numbering, *, forced_label: str | None = None) -> ParsedBlock | None:
    """One paragraph, with its clause number resolved.

    Shared by the body and by table cells so a clause inside a layout table is
    read exactly like one outside it. `numbering.label` counts, so this must be
    called once per paragraph in document order.
    """
    text = _paragraph_text(paragraph).strip()
    label, level, word_list = numbering.label(paragraph)
    if forced_label:
        # The number lived in its own cell, so it never reached the text.
        label, level, word_list = forced_label, level_for(forced_label), None
    # Where Word's own numbering set the level, it is the author's statement
    # of nesting rather than something inferred from how the number looks.
    declared = level if word_list else None
    if not text:
        return None
    if label:
        kind = "list_item"
    else:
        label, remainder = split_label(text)
        if label:
            text, level = remainder, level_for(label)
        kind = "heading" if _looks_like_heading(paragraph, text) else "paragraph"
    return ParsedBlock(
        text=text,
        kind=kind,
        number_label=label or None,
        level=level,
        declared_level=declared,
        declared_list=word_list,
    )


def _layout_table_blocks(table, numbering) -> list[ParsedBlock] | None:
    """Clauses from a table used for layout, or None if it holds data.

    Contracts put clauses in tables for two reasons: a single-column table to
    draw a box around a clause, and a two-column table with the number in the
    first column. Both are prose, and rendering them as a grid makes every
    clause inside them unaddressable — a liability cap in a bordered box is not
    a clause, cannot be cited and cannot be commented on.

    A fee schedule is a table and must stay one, so the tests are deliberately
    narrow: exactly one cell per row, or exactly two with the first holding a
    clause number and nothing else. Anything else returns None.
    """
    widths = {len(row.cells) for row in table.rows}
    if widths == {1}:
        blocks = []
        for row in table.rows:
            for item in _iter_body(row.cells[0]._tc, row.cells[0]):
                if hasattr(item, "rows"):
                    return None  # a table inside it: treat the whole thing as data
                block = _paragraph_block(item, numbering)
                if block is not None:
                    blocks.append(block)
        return blocks or None

    if widths != {2}:
        return None
    labels = [label_only(row.cells[0].text) for row in table.rows]
    if not all(labels):
        return None
    blocks = []
    for row, label in zip(table.rows, labels):
        for item in _iter_body(row.cells[1]._tc, row.cells[1]):
            if hasattr(item, "rows"):
                return None
            block = _paragraph_block(item, numbering, forced_label=label)
            if block is not None:
                blocks.append(block)
                label = None  # only the first paragraph of the cell carries it
    return blocks or None


def _render_cell(cell) -> list[str]:
    """A cell's paragraphs and nested tables, in the order they appear."""
    parts: list[str] = []
    for item in _iter_body(cell._tc, cell):
        if hasattr(item, "rows"):  # a table inside this cell
            rendered = _render_table(item)
            if rendered:
                parts.append(rendered)
        else:
            text = _paragraph_text(item).strip()
            if text:
                parts.append(text)
    return parts


def _looks_like_heading(paragraph, text: str) -> bool:
    style = (getattr(paragraph.style, "name", "") or "").lower()
    if style.startswith("heading") or style == "title":
        return True
    # Short, fully upper-case lines are how most contract templates mark a
    # section when they do not use a heading style.
    stripped = text.strip()
    return bool(stripped) and len(stripped) < 120 and stripped == stripped.upper()


class DocxParser:
    name = "docx"
    # 2: Word's numbering levels reach the tree (`tree.py`), so parents differ.
    # 3: "Item 5" read as a number; a numbered run after "WHEREAS:" is asked about.
    # 4: numbering that starts again at 1 is asked about, not a new part.
    version = "4"

    def parse(self, content: bytes, *, filename: str) -> ParsedDocument:
        _guard_zip(content)
        from docx import Document

        try:
            document = Document(BytesIO(content))
        except KeyError as exc:  # a zip, but not a Word package (e.g. a renamed .xlsx)
            raise UnsupportedFormat("The file is a zip archive but not a Word document") from exc
        numbering = _Numbering(document)
        blocks: list[ParsedBlock] = []
        indents: list[int] = []
        warnings: list[str] = []

        for item in _iter_body(document.element.body, document):
            if hasattr(item, "rows"):  # Table
                laid_out = _layout_table_blocks(item, numbering)
                if laid_out is not None:
                    blocks.extend(laid_out)
                    indents.extend([0] * len(laid_out))
                    continue
                rendered = _render_table(item)
                if rendered:
                    blocks.append(ParsedBlock(text=rendered, kind="table"))
                    indents.append(0)
                continue
            block = _paragraph_block(item, numbering)
            if block is None:
                continue
            blocks.append(block)
            indents.append(_effective_indent(item))

        # Indentation is how a Word document shows nesting that its numbering
        # does not — an indented block quote of a statute belongs under the
        # paragraph quoting it, and nothing in its text says so. Applied as a
        # floor, never a ceiling: where a number already implies a deeper
        # level, the number wins, because that is what a reader cites.
        depths = _indent_levels(indents)
        blocks = [
            block
            if (depth := depths.get(indent // _TWIPS_PER_EIGHTH_INCH, 0) + 1) <= block.level
            else replace(block, level=depth)
            for block, indent in zip(blocks, indents)
        ]

        if not blocks:
            warnings.append("No readable text found in this Word document.")
        return ParsedDocument(blocks=blocks, page_count=None, warnings=warnings)
