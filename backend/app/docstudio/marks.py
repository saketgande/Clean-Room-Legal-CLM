"""Changes to a PDF, kept beside it and never in it.

A PDF is a printout: each word is painted at a fixed place, so typing into it
cannot push the rest of the page along, and its fonts usually hold only the
letters the document used. A suggested change to a PDF is therefore a *mark*
on the page — a proposal annotation, anchored to its words the way a comment
is — and the PDF's bytes do not change. What becomes of an agreed mark is
chosen per document, each way checked:

* `marked_up` — the original PDF plus standard PDF annotations (strike-outs,
  carets, notes), written as an incremental update: the original bytes are
  kept, so a signature over them still verifies.
* `carry` — a Word version of a typed PDF with the agreed marks as tracked
  changes, as `redline.make_editable` makes one.
* `patch` — the agreed words swapped on the page itself, only where the new
  words fit their line and every letter is in the embedded font, and only if
  the page then reads exactly as agreed and nothing else on it moved.
* `amendment` — a separate document listing the agreed changes clause by
  clause: the way a signed contract is changed.
"""

from __future__ import annotations

import io
import os
import re
import tempfile
from collections import defaultdict

from sqlalchemy import select

from .anchoring import ORPHANED
from .annotations import _record, annotate, locate, remove
from .models import DsAnnotation, DsClause, DsVersion
from .ocr import _plain
from .redline import (
    SCANNED,
    Edit,
    RedlineError,
    _file,
    _next_version,
    is_scan,
    propose,
    word_content,
)

PDF_MIME = "application/pdf"
STATUSES = ("open", "agreed", "rejected")


def _require_pdf(version) -> None:
    if version.mime_type != PDF_MIME:
        raise RedlineError("Marks are for PDFs. A Word file takes tracked changes in the file itself.")


# --- placing and deciding marks -------------------------------------------------------


def add(db, version, edits: list[Edit], *, author: str, kind: str = "user") -> tuple[list, list]:
    """Each edit as a proposal on its words. An edit whose words are not found,
    or found more than once with nothing to tell them apart, is refused with
    the reason — never placed on a guess."""
    _require_pdf(version)
    text = version.flat_text or ""
    placed, refused = [], []
    for edit in edits:
        try:
            start, end = locate(text, edit.find, prefix=edit.before[-400:], suffix=edit.after[:400])
        except LookupError as exc:
            refused.append((edit, str(exc)))
            continue
        if " ".join(_plain(text[start:end]).split()) == " ".join(edit.replace.split()):
            refused.append((edit, "The new wording is the same as the old."))
            continue
        mark = annotate(db, version_id=version.id, start=start, end=end, body=edit.why or "",
                        kind="proposal", author_name=author, proposed_text=edit.replace)
        mark.author_kind = kind
        placed.append(mark)
    return placed, refused


def highlight(db, version, start: int, end: int, *, author: str):
    _require_pdf(version)
    return annotate(db, version_id=version.id, start=start, end=end, body="", kind="highlight", author_name=author)


def set_status(db, annotation_id: str, status: str) -> DsAnnotation:
    """Agree to a proposal, reject it, or reopen it — recorded, as every decision is."""
    if status not in STATUSES:
        raise ValueError(f"status must be one of {', '.join(STATUSES)}")
    mark = db.get(DsAnnotation, annotation_id)
    if mark is None or mark.kind != "proposal":
        raise LookupError("No such suggested change.")
    if mark.status != status:
        mark.status = status
        version = db.get(DsVersion, mark.version_id)
        _record(db, version, f"proposal.{'reopened' if status == 'open' else status}", {"annotation_id": mark.id})
        db.flush()
    return mark


def _marks(db, version, *, statuses=None, kinds=("proposal",)) -> list[DsAnnotation]:
    query = select(DsAnnotation).where(
        DsAnnotation.version_id == version.id,
        DsAnnotation.kind.in_(kinds),
        DsAnnotation.anchor_state != ORPHANED,
        DsAnnotation.parent_annotation_id.is_(None),
    )
    if statuses:
        query = query.where(DsAnnotation.status.in_(statuses))
    return list(db.scalars(query.order_by(DsAnnotation.anchor_start)))


def _edit(mark: DsAnnotation) -> Edit:
    return Edit(find=mark.anchor_quote_exact or "", replace=mark.proposed_text or "",
                before=mark.anchor_quote_prefix or "", after=mark.anchor_quote_suffix or "", why=mark.body or "")


def _agreed(db, version) -> list[DsAnnotation]:
    found = _marks(db, version, statuses=("agreed",))
    if not found:
        raise LookupError("Agree to at least one suggested change first.")
    return found


# --- the agreed marks as a Word version -------------------------------------------------


def carry(db, version) -> dict:
    """A Word version of this typed PDF with every agreed mark as a tracked change
    by the person who made it — one new version, so the marks are placed on the
    same words they were made on. A mark that cannot be placed stays agreed on
    the PDF, with the reason; one that is placed is now a change in the file,
    and leaves the margin."""
    _require_pdf(version)
    if is_scan(version):
        # The scan can be rebuilt as Word for reading and editing (redline's
        # word_content does it from what OCR read), but the changes to a signed
        # contract are not tracked changes on a machine's reading of it.
        raise RedlineError(SCANNED)
    agreed = _agreed(db, version)
    content, made = word_content(db, version)
    groups: dict[str, list[DsAnnotation]] = defaultdict(list)
    for mark in agreed:
        groups[mark.author_name or "You"].append(mark)
    carried, refused, suggested = [], [], []
    for author, marks in groups.items():
        edits = [_edit(m) for m in marks]
        owner = {id(edit): mark for edit, mark in zip(edits, marks, strict=True)}
        content, applied = propose(content, edits, author=author)  # hands back the same edits
        for edit, ids in applied.edits:
            carried.append(owner[id(edit)])
            suggested.append((author, applied.when, edit, ids))
        refused += [{"find": edit.find[:200], "why": why} for edit, why in applied.refused]
    filename = made.pop("filename")
    for mark in carried:  # now tracked changes in the file; their reasons go with them
        remove(db, mark.id)
    made["carried"] = len(carried)
    made["refused"] = refused
    result = _next_version(db, version, content, filename=filename, event="redline.made_editable", details=made)
    new_version = db.get(DsVersion, result.version_id)
    for author, when, edit, ids in suggested:
        _record(db, new_version, "redline.suggested", {
            "author": author, "date": when, "refused": [],
            "changes": [{"ids": ids, "find": edit.find[:300], "replace": edit.replace[:300], "why": edit.why[:600]}],
        })
    return {"carried": len(carried), "refused": refused, "version_id": result.version_id}


# --- where a mark's words are on the page ------------------------------------------------


def _tokens(text: str) -> list[str]:
    return [t for t in (re.sub(r"\W+", "", w.lower()) for w in _plain(text).split()) if t]


def _regions(db, mark: DsAnnotation) -> list[dict]:
    clause = db.scalars(
        select(DsClause).where(DsClause.version_id == mark.version_id, DsClause.clause_id == mark.anchor_clause_id)
    ).first()
    return list(clause.source_regions or []) if clause is not None else []


def _boxes(doc, regions: list[dict], quote: str, prefix: str) -> list[tuple[int, object]]:
    """The mark's words on the page, one box per printed line: (page index, rect).
    Empty when they cannot be found for certain — on a scan, or across a hyphen."""
    import fitz

    words = []  # (token, page index, rect, line key)
    for region in regions:
        index = int(region.get("page") or 0) - 1
        if not 0 <= index < len(doc) or not region.get("bbox"):
            continue
        page, box = doc[index], region["bbox"]
        clip = fitz.Rect(box["x0"] * page.rect.width - 2, box["y0"] * page.rect.height - 2,
                         box["x1"] * page.rect.width + 2, box["y1"] * page.rect.height + 2)
        for x0, y0, x1, y1, word, block, line, _ in page.get_text("words", clip=clip, sort=True):
            token = re.sub(r"\W+", "", word.lower())
            if token:
                words.append((token, index, fitz.Rect(x0, y0, x1, y1), (index, block, line)))
    wanted, before = _tokens(quote), _tokens(prefix)[-6:]
    if not wanted:
        return []
    tokens = [w[0] for w in words]
    hits = [i for i in range(len(tokens) - len(wanted) + 1) if tokens[i : i + len(wanted)] == wanted]
    if len(hits) > 1 and before:  # the words before it tell identical passages apart
        def agree(i):
            return sum(1 for a, b in zip(reversed(tokens[:i]), reversed(before), strict=False) if a == b)
        best = max(agree(i) for i in hits)
        hits = [i for i in hits if agree(i) == best]
    if len(hits) != 1:
        return []
    lines: dict[tuple, object] = {}
    for _, index, rect, key in words[hits[0] : hits[0] + len(wanted)]:
        lines[key] = lines[key] | rect if key in lines else rect
    return [(key[0], rect) for key, rect in lines.items()]


def _kind(mark: DsAnnotation) -> tuple[str, str, str]:
    """delete | insert-after | insert-before | replace, with the old and new words."""
    old, new = " ".join((mark.anchor_quote_exact or "").split()), " ".join((mark.proposed_text or "").split())
    if not new:
        return "delete", old, new
    if new.startswith(old):
        return "insert-after", old, new[len(old):].strip()
    if new.endswith(old):
        return "insert-before", old, new[: -len(old)].strip()
    return "replace", old, new


def _describe(mark: DsAnnotation) -> str:
    kind, old, new = _kind(mark)
    return {
        "delete": f"Delete “{old}”",
        "insert-after": f"Insert “{new}” after “{old}”",
        "insert-before": f"Insert “{new}” before “{old}”",
        "replace": f"Replace “{old}” with “{new}”",
    }[kind]


# --- output 1: the marked-up PDF --------------------------------------------------------


def marked_up(db, version) -> tuple[bytes, dict]:
    """The original PDF with every open and agreed mark, highlight and comment as
    a standard PDF annotation any reader shows — added as an incremental update,
    so the file begins with the original's exact bytes. Returns the file and
    how many marks sat on their words and how many became notes on their clause."""
    import fitz

    _require_pdf(version)
    original = _file(version)
    marks = _marks(db, version, statuses=("open", "agreed", "resolved"),
                   kinds=("proposal", "highlight", "comment"))
    on_words = as_notes = 0
    with tempfile.TemporaryDirectory() as folder:
        path = os.path.join(folder, "original.pdf")
        with open(path, "wb") as handle:
            handle.write(original)
        doc = fitz.open(path)
        for mark in marks:
            regions = _regions(db, mark)
            boxes = _boxes(doc, regions, mark.anchor_quote_exact or "", mark.anchor_quote_prefix or "")
            who = mark.author_name or "You"
            if mark.kind == "proposal":
                text = _describe(mark) + (f"\n{mark.body}" if mark.body else "")
            elif mark.kind == "comment":
                replies = db.scalars(select(DsAnnotation).where(DsAnnotation.parent_annotation_id == mark.id))
                text = "\n".join([mark.body or ""] + [f"— {r.body}" for r in replies])
            else:
                text = ""
            if boxes:
                on_words += 1
                for index, rect in boxes:
                    page = doc[index]
                    if mark.kind == "proposal" and _kind(mark)[0] in ("delete", "replace"):
                        annot = page.add_strikeout_annot(rect)
                    else:
                        annot = page.add_highlight_annot(rect)
                    annot.set_info(title=who, content=text)
                    annot.update()
                if mark.kind == "proposal" and (mark.proposed_text or "").strip():
                    index, last = boxes[-1]
                    caret = doc[index].add_caret_annot(last.tr)
                    caret.set_info(title=who, content=text)
                    caret.update()
            elif regions and text:
                as_notes += 1
                region = regions[0]
                page = doc[int(region["page"]) - 1]
                point = fitz.Point(region["bbox"]["x0"] * page.rect.width, region["bbox"]["y0"] * page.rect.height)
                note = page.add_text_annot(point, text)
                note.set_info(title=who)
                note.update()
        doc.save(path, incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
        doc.close()
        with open(path, "rb") as handle:
            marked = handle.read()
    if not marked.startswith(original):
        raise RedlineError("The marked-up file did not keep the original's bytes, so it was not made.")
    return marked, {"on_words": on_words, "as_notes": as_notes}


# --- output 2: the agreed words written on the page itself -------------------------------


def _read_words(content: bytes) -> tuple[int, list[str]]:
    """How our own reader sees a PDF: its number of blocks and its words."""
    from .parsing.registry import parser_for

    blocks = parser_for(PDF_MIME).parse(content, filename="x.pdf").blocks
    return len(blocks), [w for b in blocks for w in f"{b.number_label or ''} {_plain(b.text)}".split()]


def patch(db, version) -> dict:
    """The agreed changes that fit, written on the page — inside the drawing
    instruction of the words they replace (pdftext.py), so the page's reading
    order is kept and everything after them stays where it was.

    Checked, or nothing is written: each change must fit its line; afterwards
    every changed page must be pixel-identical outside the changed words and
    every other page identical, and our own reader must read the new file as
    the old with exactly the agreed words changed and nothing else. Kept as an
    incremental update and made the next version; the original stays.
    """
    import difflib

    import fitz

    from .pdftext import PdfTextError, replace

    _require_pdf(version)
    original = _file(version)
    agreed = _agreed(db, version)
    before = fitz.open(stream=original, filetype="pdf")
    done, refused, changed = [], [], defaultdict(list)
    with tempfile.TemporaryDirectory() as folder:
        path = os.path.join(folder, "original.pdf")
        with open(path, "wb") as handle:
            handle.write(original)
        doc = fitz.open(path)
        for mark in agreed:
            new = " ".join((mark.proposed_text or "").split())
            old = " ".join((mark.anchor_quote_exact or "").split())
            try:
                if not new:
                    raise RedlineError("A deletion would leave a gap in the line. Use the Word version or an amendment.")
                boxes = _boxes(doc, _regions(db, mark), old, mark.anchor_quote_prefix or "")
                if len(boxes) != 1:
                    raise RedlineError("Only words on one printed line can be changed on the page itself."
                                       if boxes else "Those words could not be found exactly on the page.")
                index, rect = boxes[0]
                content, old_width, new_width = replace(doc.tobytes(), index, old, new,
                                                        mark.anchor_quote_prefix or "")
                drawn = new_width * (rect.width / old_width) if old_width else new_width
                page = doc[index]
                words = page.get_text("words")
                after = [w[0] for w in words if abs(w[3] - rect.y1) < 1.5 and w[0] >= rect.x1 - 0.5]
                room = (min(after) - rect.x0 - 0.5) if after else max(w[2] for w in words) - rect.x0
                if drawn > room:
                    raise RedlineError("The new words are longer than the room on their line.")
                if rect.width - drawn > 0.3 * rect.height:  # about one space: more shows as a gap
                    raise RedlineError("The new words are shorter than the old by more than a space, and would "
                                       "leave a visible gap in the line. Use the Word version or an amendment.")
            except (RedlineError, PdfTextError) as exc:
                refused.append({"find": old[:200], "why": str(exc)})
                continue
            xrefs = page.get_contents()
            doc.update_stream(xrefs[0], content)
            for xref in xrefs[1:]:  # the page's instructions were read as one list; they go back as one
                doc.update_stream(xref, b"")
            changed[index].append(fitz.Rect(rect.x0 - 1, rect.y0 - 2, rect.x0 + max(rect.width, drawn) + 1,
                                            rect.y1 + 2))
            done.append((mark, old, new))
        if not done:
            doc.close()
            return {"patched": 0, "refused": refused, "version_id": None}
        # Check the pages: nothing moved outside the changed words.
        for index in range(len(doc)):
            a = before[index].get_pixmap(matrix=fitz.Matrix(1, 1), alpha=False)
            b = doc[index].get_pixmap(matrix=fitz.Matrix(1, 1), alpha=False)
            if (a.width, a.height) != (b.width, b.height) or (a.samples != b.samples and index not in changed):
                raise RedlineError(f"Page {index + 1} changed although no change was on it, so nothing was written.")
            if a.samples == b.samples:
                continue
            stride, moved = a.width * a.n, 0
            for y in range(a.height):
                row_a, row_b = a.samples[y * stride : (y + 1) * stride], b.samples[y * stride : (y + 1) * stride]
                if row_a != row_b:
                    moved += sum(1 for x in range(a.width)
                                 if row_a[x * a.n : (x + 1) * a.n] != row_b[x * a.n : (x + 1) * a.n]
                                 and not any(r.contains(fitz.Point(x + 0.5, y + 0.5)) for r in changed[index]))
            if moved:
                raise RedlineError(f"{moved} pixels changed on page {index + 1} outside the changed words, "
                                   "so nothing was written.")
        doc.save(path, incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
        doc.close()
        with open(path, "rb") as handle:
            patched = handle.read()
    if not patched.startswith(original):
        raise RedlineError("The patched file did not keep the original's bytes, so it was not made.")
    # Check the reading: our reader must see the old document with exactly the agreed words changed.
    (blocks_before, words_before), (blocks_after, words_after) = _read_words(original), _read_words(patched)
    differences = [(words_before[i1:i2], words_after[j1:j2]) for tag, i1, i2, j1, j2
                   in difflib.SequenceMatcher(None, words_before, words_after, autojunk=False).get_opcodes()
                   if tag != "equal"]
    agreed_words = [(old.split(), new.split()) for _, old, new in done]
    if blocks_before != blocks_after or sorted(map(str, differences)) != sorted(map(str, agreed_words)):
        raise RedlineError("Read back, the changed PDF did not say exactly the agreed words in their places, "
                           "so nothing was written.")
    from .service import ingest, record_event

    details = {"changes": [{"find": old[:300], "replace": new[:300], "by": m.author_name or "You"}
                           for m, old, new in done], "refused": refused,
               "checks": "read back word for word; pixel-identical outside the changed words; original bytes kept"}
    for mark, _, _ in done:  # now in the file's own words
        remove(db, mark.id)
    result = ingest(db, org_id=version.org_id, content=patched, filename=version.filename or "document.pdf",
                    mime_type=PDF_MIME, document_id=version.document_id)
    record_event(db, org_id=version.org_id, document_id=version.document_id, version_id=result.version_id,
                 event_type="pdf.patched", details=details)
    return {"patched": len(done), "refused": refused, "version_id": result.version_id}


# --- output 3: an amendment --------------------------------------------------------------


def amendment(db, version, title: str) -> bytes:
    """The agreed changes as an amendment: a separate document that says, clause
    by clause, what is deleted and what is inserted — how a signed contract is
    changed. The contract itself is not touched."""
    from docx import Document

    from .ask import label_of
    from .service import clauses_for

    agreed = _agreed(db, version)
    clauses = clauses_for(db, version.id)
    by_id = {c.clause_id: c for c in clauses}
    document = Document()
    document.add_heading("Amendment", level=1)
    document.add_paragraph(f"to {title} (the “Agreement”)")
    document.add_paragraph("The parties agree to amend the Agreement as follows:")
    for n, mark in enumerate(agreed, start=1):
        label = label_of(by_id.get(mark.anchor_clause_id), by_id)
        where = f"clause {label}" if re.match(r"^[\w.()-]*\d[\w.()-]*$", label) else f"the paragraph beginning “{label}”"
        text = _describe(mark)
        document.add_paragraph(f"{n}. In {where}: {text[0].lower()}{text[1:]}.")
    document.add_paragraph("Except as amended above, the Agreement remains unchanged and in full force and effect.")
    table = document.add_table(rows=4, cols=2)
    for r, field in enumerate(("Signed for and on behalf of", "Name", "Title", "Date")):
        for c in range(2):
            table.cell(r, c).text = f"{field}:"
    out = io.BytesIO()
    document.save(out)
    _record(db, version, "amendment.drafted", {"changes": [_describe(m) for m in agreed]})
    return out.getvalue()
