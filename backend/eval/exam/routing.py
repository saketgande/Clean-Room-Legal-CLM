"""The routing set: free-text requests drafted by an AI, and the owner's picks (FR-3, FR-21, FR-25).

The exam asks the app's intake triage to route about 50 email / chat / notice style requests and
scores how many go where the owner says they should. This module builds that set:

* ``snapshot_workflows`` freezes today's enabled workflows (read-only on the live database); the
  owner's picks are made against that list, so ``routing_fresh`` can say when it has changed.
* ``draft_routing`` (the paid command) has an AI write the requests, at least one per workflow,
  and writes the sheet the owner marks. A drafted request is an INPUT, never an answer key
  (FR-3): the suggested workflow in the sheet is only a suggestion until the owner marks it.
* ``load_routing`` turns the returned sheet into the keys, stored with the fingerprint of the
  workflow list they were picked against.

Private files (git- and docker-ignored, ``eval/private/``): ``routing/requests.json``,
``routing/workflows.json`` (the snapshot), ``keys/routing.json``, the sheet in ``sheets/``.
They are the format ``eval.exam.content.load_content`` reads. No client text is involved here,
so unlike pre-fill no ``eval.real_contract_sent`` row is written.

Import this module only after ``eval.exam.env.isolate()``: it reaches the app's models and the
Claude client, which import the database engine.
"""

import hashlib
import json
import os
import re
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.ai.gateway.features import AIFeature, FeatureKind, PromptLayout
from app.ai.gateway.ledger import AICallContext, record_call
from app.ai.prompt_versions import PromptBundle, hash_text, model_config_hash
from app.analytics.routes import token_cost
from app.core.config import settings
from app.core.enums import AICallStatus, AIValidationStatus
from app.integrations import claude
from app.workflows.models import Workflow
from eval.exam import runner
from eval.exam.content import (
    PRIVATE_DIR,
    ROUTING_DRAFT_MODEL,
    RoutingRequest,
    routing_requests_sha256,
)
from eval.exam.env import ExamRefused, LiveDB, Owner, audit
from eval.exam.sheets import SHEETS_DIR, SheetRejected, load_sheet, write_routing_sheet

ROUTING_DIR = PRIVATE_DIR / "routing"
REQUESTS_PATH = ROUTING_DIR / "requests.json"
SNAPSHOT_PATH = ROUTING_DIR / "workflows.json"
KEYS_PATH = PRIVATE_DIR / "keys" / "routing.json"

# FR-21: "about 50 requests spread across the 32 workflows" = 2 each for the first 18 (by the
# order the app tries workflows in), 1 for the rest. Fewer workflows than 18 get 2 each.
_DOUBLED = 18
_CHANNELS = ("email", "chat", "notice")

# Drafting is one short call per workflow. 1,200 tokens is several times the longest answer
# asked for (2 messages of at most 120 words in a small JSON object), so a cut-off answer means
# something went wrong; the worst case in the estimate uses it too.
_MAX_TOKENS = 1_200
_EXPECTED_OUTPUT_TOKENS_PER_REQUEST = 200  # ~100 words + subject + JSON, with some margin

_CALL_TYPE = "exam_routing_draft"
_PROMPT_VERSION = "2"

_SYSTEM = (
    "You write realistic messages that employees of an Indian company send to its in-house legal "
    "team. Each message is an input to a test of how the legal team's intake system routes "
    "work, so it must read exactly like something a real person would send: ordinary business "
    "language, concrete but invented company and people names, amounts in INR, dates, places. "
    "Never name the workflow, the intake system or any routing; never say the message is a test. "
    "Write 40 to 120 words per message."
)
_USER = """The legal team has this workflow:

Name: {name}
What it is for: {description}
When it is chosen: {criteria}

Write {count} different message(s) that a person would send when they need exactly this kind of \
work done, without using the workflow's name. One message per channel, in this order: {channels}.
- email: a subject line and a short body.
- chat: a few words as the subject and the chat message as the body.
- notice: a notice, demand or formal letter received from a third party, forwarded to legal \
(subject = its heading, body = what it says and what the sender wants).

Answer as JSON with one key per channel ({channels}), each \
{{"subject": "...", "body": "..."}}."""
_MESSAGE = {
    "type": "object",
    "properties": {"subject": {"type": "string"}, "body": {"type": "string"}},
    "required": ["subject", "body"],
    "additionalProperties": False,
}


def _schema(channels: list[str]) -> dict:
    """One required key per channel, so the count is part of the answer's shape. An array cannot
    be held to a length: a run stopped at request 10 of 32 when the model wrote one message
    where two were asked."""
    return {
        "type": "object",
        "properties": {channel: _MESSAGE for channel in channels},
        "required": list(channels),
        "additionalProperties": False,
    }


# Local objects, never registered in the app's feature registry (the exam's calls are not an
# app feature): they exist so the call leaves the same ledger row every other AI call does.
_FEATURE = AIFeature(
    key=_CALL_TYPE,
    kind=FeatureKind.TASK,
    label="Exam: routing request drafts",
    category="AI exam",
    prompt_key=_CALL_TYPE,
    default_prompt_version=_PROMPT_VERSION,
    layout=PromptLayout.FEATURE_SYSTEM,
    guard=False,
    max_tokens=_MAX_TOKENS,
    temperature=1.0,  # unused: complete_json sends no temperature (sonnet-5-5 rejects one)
    tool_name="routing_requests",
)
_PROMPT = PromptBundle(
    prompt_key=_CALL_TYPE,
    version=_PROMPT_VERSION,
    prompt_hash=hash_text(f"{_SYSTEM}\n\n{_USER}"),
    shared_system_prompt="",
    skill_prompt=_SYSTEM,
    model_name=ROUTING_DRAFT_MODEL,
    model_config_hash=model_config_hash({"model": ROUTING_DRAFT_MODEL, "max_tokens": _MAX_TOKENS}),
)


# --- the workflow snapshot ---------------------------------------------------------------


def workflow_fingerprint(workflows: list[dict]) -> str:
    """sha256 of the enabled workflows' (id, name, enabled, version, criteria, steps), by id.

    Anything that can change where a request should go changes this; a description or an
    ``eval_order`` edit does not (``version`` moves when the builder saves a real change).
    """
    rows = sorted(
        (
            [w["id"], w["name"], w["enabled"], w["version"], w["criteria"], w["steps"]]
            for w in workflows
            if w["enabled"]
        ),
        key=lambda row: row[0],
    )
    canonical = json.dumps(rows, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def snapshot_workflows(live: LiveDB) -> dict:
    """Today's enabled workflows of the live organisation: ``{fingerprint, taken_at, workflows}``.

    Read-only. Ordered the way the app tries workflows (``eval_order``), with name and id as
    tie-breaks so the same list always gives the same snapshot and the same request ids.
    """
    with live.sessions() as db:
        rows = db.scalars(
            select(Workflow)
            .where(Workflow.org_id == live.org_id, Workflow.enabled.is_(True))
            .order_by(Workflow.eval_order, Workflow.name, Workflow.id)
        ).all()
        workflows = [
            {
                "id": w.id,
                "name": w.name,
                "description": w.description,
                "enabled": bool(w.enabled),
                "is_builtin": bool(w.is_builtin),
                "eval_order": w.eval_order,
                "version": w.version,
                "criteria": w.criteria or {},
                "steps": w.steps or [],
            }
            for w in rows
        ]
    return {
        "fingerprint": workflow_fingerprint(workflows),
        "taken_at": datetime.now(UTC).isoformat(),
        "workflows": workflows,
    }


def routing_fresh(live: LiveDB, snapshot: dict | None) -> bool:
    """True while the live workflow list is the one the picks were made against (AC-21).

    Picks against a list that has since gained, lost or changed a workflow no longer say where a
    request should go today, so the runner leaves routing out instead of scoring them.
    """
    return snapshot is not None and snapshot_workflows(live)["fingerprint"] == snapshot.get(
        "fingerprint"
    )


def coverage(snapshot: dict, requests: list[RoutingRequest]) -> dict:
    """How the requests spread over the workflows (FR-21: how many are covered, how many each)."""
    names = {w["id"]: w["name"] for w in snapshot["workflows"]}
    per_workflow = {name: 0 for name in names.values()}
    for request in requests:
        name = names.get(request.drafted_for_workflow_id)
        if name is not None:
            per_workflow[name] += 1
    return {
        "workflows": len(names),
        "workflows_covered": sum(1 for n in per_workflow.values() if n),
        "requests": len(requests),
        "requests_per_workflow": per_workflow,
    }


# --- drafting ----------------------------------------------------------------------------


def plan_calls(snapshot: dict) -> list[tuple[dict, list[str]]]:
    """One ``(workflow, channels)`` per call, in snapshot order: its channels are the requests
    to write. The channel rotates with the workflow's place, and a workflow's second request
    uses the next channel, so no workflow gets two of the same kind."""
    return [
        (
            wf,
            [_CHANNELS[k % 3], *([_CHANNELS[(k + 1) % 3]] if k < _DOUBLED else [])],
        )
        for k, wf in enumerate(snapshot["workflows"])
    ]


def _user_prompt(workflow: dict, channels: list[str]) -> str:
    return _USER.format(
        name=workflow["name"],
        description=workflow["description"] or "(no description)",
        # Only what is set: the empty keys of criteria are noise to the model.
        criteria=json.dumps({k: v for k, v in workflow["criteria"].items() if v}),
        count=len(channels),
        channels=", ".join(channels),
    )


def _payload(workflow: dict, channels: list[str]) -> dict:
    """The request ``complete_json`` will send, for the estimate (within the few dozen
    characters of cache markup it adds to the system prompt)."""
    return {
        "model": ROUTING_DRAFT_MODEL,
        "max_tokens": _MAX_TOKENS,
        "system": _SYSTEM,
        "messages": [{"role": "user", "content": _user_prompt(workflow, channels)}],
        "output_config": {"format": {"type": "json_schema", "schema": _schema(channels)}},
    }


def estimate_draft(snapshot: dict) -> runner.Estimate:
    """What drafting will cost: every call is synchronous (not batched) at the drafting model."""
    expected = worst = 0.0
    calls = plan_calls(snapshot)
    for workflow, channels in calls:
        payload = _payload(workflow, channels)
        chars = len(_SYSTEM) + len(payload["messages"][0]["content"])
        expected += token_cost(
            chars // 4, _EXPECTED_OUTPUT_TOKENS_PER_REQUEST * len(channels), ROUTING_DRAFT_MODEL
        )
        worst += runner.worst_cost(payload, batched=False)
    return runner.Estimate(
        delivery="sync",
        calls=len(calls),
        batch_calls=0,
        sync_calls=len(calls),
        usd=expected,
        list_usd=expected,
        worst_case_usd=worst,
        by_call_type={_CALL_TYPE: {"calls": len(calls), "batch": 0, "usd": round(expected, 4)}},
    )


def _drafts_from(response, channels: list[str]) -> list[dict]:
    """One ``{subject, body}`` per channel, in channel order, or ``ValueError`` saying why not."""
    if response.stop_reason == "max_tokens":
        raise ValueError(f"answer cut off at max_tokens ({_MAX_TOKENS})")
    text = next((b.get("text") for b in response.content_blocks if b.get("type") == "text"), None)
    try:
        answer = json.loads(text or "")
        drafts = [answer.get(channel) for channel in channels]
    except (ValueError, AttributeError):
        raise ValueError("the answer is not the requested JSON") from None
    if not all(
        isinstance(d, dict)
        and isinstance(d.get("subject"), str)
        and isinstance(d.get("body"), str)
        and d["subject"].strip()
        and d["body"].strip()
        for d in drafts
    ):
        raise ValueError(f"expected a subject and a body for each of: {', '.join(channels)}")
    return drafts


def _same_model(a: str, b: str) -> bool:
    """Equal once a ``-YYYYMMDD`` snapshot suffix is ignored."""
    return re.sub(r"-\d{8}$", "", a) == re.sub(r"-\d{8}$", "", b)


async def _draft_one(
    live: LiveDB, owner: Owner, run_id: str, workflow: dict, channels: list[str]
) -> list[dict]:
    """One drafting call and its ledger row (live ledger, stamped with the run)."""
    ctx = AICallContext(
        org_id=live.org_id,
        user_id=owner.user_id,
        resource=("workflow", workflow["id"]),
        eval_run_id=run_id,
    )
    # Names the workflow and channels only: the drafted text is not stored in the ledger.
    log_input = {"workflow_id": workflow["id"], "channels": channels}

    def log(**fields) -> None:
        record_call(
            sessions=live.sessions,
            feature=_FEATURE,
            ctx=ctx,
            prompt=_PROMPT,
            model=ROUTING_DRAFT_MODEL,
            log_input=log_input,
            **fields,
        )

    try:
        response = await claude.claude_client.complete_json(
            org_id=live.org_id,
            system_prompt=_SYSTEM,
            user_prompt=_user_prompt(workflow, channels),
            json_schema=_schema(channels),
            max_tokens=_MAX_TOKENS,
            model=ROUTING_DRAFT_MODEL,
        )
    except runner.CeilingReached:
        raise  # nothing was sent, so there is nothing to record
    except Exception as exc:
        log(
            status=AICallStatus.FAILED,
            validation_status=AIValidationStatus.NOT_VALIDATED,
            error_class=type(exc).__name__,
        )
        raise
    try:
        drafts = _drafts_from(response, channels)
    except ValueError as exc:
        log(
            status=AICallStatus.VALIDATION_FAILED,
            validation_status=AIValidationStatus.INVALID,
            response=response,
            validation_error=str(exc),
        )
        raise
    log(
        status=AICallStatus.SUCCEEDED,
        validation_status=AIValidationStatus.VALID,
        response=response,
    )
    return drafts


def _assemble(drafted: list[tuple[dict, list[str], list[dict]]]) -> list[RoutingRequest]:
    """Number the requests: every workflow's first request, then every second one. Any prefix of
    the ids (the sample takes the first ten) therefore spreads over as many workflows as it can,
    instead of using up its slots on two requests per workflow."""
    firsts: list = []
    seconds: list = []
    for workflow, channels, texts in drafted:
        for n, (channel, text) in enumerate(zip(channels, texts, strict=True)):
            (firsts if n == 0 else seconds).append((workflow, channel, text))
    return [
        RoutingRequest(
            item_id=f"rt-{i:02d}",
            channel=channel,
            subject=text["subject"].strip(),
            body=text["body"].strip(),
            drafted_for_workflow_id=workflow["id"],
            drafting_model=ROUTING_DRAFT_MODEL,
        )
        for i, (workflow, channel, text) in enumerate(firsts + seconds, start=1)
    ]


def _write_json(path: Path, data) -> None:
    """Whole file or nothing: a half-written key file would load as someone's picks."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, path)


async def draft_routing(*, live: LiveDB, owner: Owner, ask: Callable[[str], str]) -> dict | None:
    """The ``draft-routing`` command: draft the requests, write the sheet for the owner to mark.

    Returns ``{run_id, requests, sheet_id, sheet, coverage, usd}``, or ``None`` when the owner
    does not give the go-ahead (nothing is spent). Raises ``ExamRefused`` before asking for
    anything if there is nothing to draft for or a routing set already exists, and
    ``RuntimeError`` when a call fails or the ceiling is reached part-way: the run is recorded
    ``incomplete`` and NOTHING is written, because a set that misses a workflow would break
    the "at least one per workflow" promise (FR-21).

    ponytail: one call at a time (32 short calls, minutes). Gather them under a small semaphore
    if the workflow list ever grows enough to matter. A set that exists is never replaced from
    here, because the owner's picks belong to those exact requests: delete the three files
    named in the refusal to start over.
    """
    existing = [p for p in (REQUESTS_PATH, SNAPSHOT_PATH, KEYS_PATH) if p.exists()]
    if existing:
        raise ExamRefused(
            "a routing set already exists and the owner's picks belong to those exact requests; "
            "to draft a new one on purpose, delete "
            + ", ".join(f"eval/private/{p.parent.name}/{p.name}" for p in existing)
            + " first (a new set means a new baseline)"
        )
    snapshot = snapshot_workflows(live)
    if not snapshot["workflows"]:
        raise ExamRefused("there are no enabled workflows to draft requests for")
    estimate = estimate_draft(snapshot)

    same_model = _same_model(ROUTING_DRAFT_MODEL, settings.claude_model)
    print(
        f"Drafting {estimate.calls} calls for {len(snapshot['workflows'])} workflows with "
        f"{ROUTING_DRAFT_MODEL}; no contract text is sent."
    )
    if same_model:
        print(
            f"Warning: {ROUTING_DRAFT_MODEL} is also the model under test. Requests drafted by "
            "the model that will route them tend to be easier for it; the record will say so."
        )
    run = runner.approve_paid_run(
        kind="draft_routing",
        live=live,
        owner=owner,
        estimate=estimate,
        ask=ask,
        extra={
            "drafting_model": ROUTING_DRAFT_MODEL,
            "model_under_test": settings.claude_model,
            "drafted_by_model_under_test": same_model,
            "workflows": len(snapshot["workflows"]),
            "workflow_fingerprint": snapshot["fingerprint"],
        },
    )
    if run is None:
        return None

    claude.exam_meter = run.meter
    outcome, reason = "incomplete", "interrupted"
    results_sha256 = None
    summary: dict = {"versions": {"models": [ROUTING_DRAFT_MODEL]}}
    result: dict | None = None
    try:
        drafted = []
        for done, (workflow, channels) in enumerate(plan_calls(snapshot), start=1):
            try:
                texts = await _draft_one(live, owner, run.run_id, workflow, channels)
            except runner.CeilingReached as exc:
                reason = "ceiling_reached"
                raise RuntimeError(f"{exc}; nothing was written") from exc
            except Exception as exc:
                # The call's failure is on the ledger ("contains errors"); the run itself
                # stopped short, which the plan's reasons call "interrupted".
                reason = "interrupted"
                raise RuntimeError(
                    f"drafting stopped at '{workflow['name']}' ({exc}); nothing was written, "
                    "and the calls made so far were paid for"
                ) from exc
            drafted.append((workflow, channels, texts))
            print(f"  drafted {done}/{estimate.calls}: {workflow['name']}")
        requests = _assemble(drafted)
        rows = [r.model_dump() for r in requests]
        # The sheet first: it checks every row against the snapshot, so a problem there leaves
        # no requests file behind for a later run to trip over.
        sheet_id = write_routing_sheet(rows, snapshot=snapshot, path=SHEETS_DIR)
        _write_json(SNAPSHOT_PATH, snapshot)
        _write_json(REQUESTS_PATH, [r.model_dump(mode="json") for r in requests])
        results_sha256 = hashlib.sha256(REQUESTS_PATH.read_bytes()).hexdigest()
        outcome, reason = "complete", None
        result = {
            "run_id": run.run_id,
            "requests": len(requests),
            "sheet_id": sheet_id,
            "sheet": SHEETS_DIR / f"routing-{sheet_id}.xlsx",
            "coverage": coverage(snapshot, requests),
        }
        summary["headline"] = {
            "requests": len(requests),
            "workflows_covered": result["coverage"]["workflows_covered"],
        }
    finally:
        actual = runner.ledger_actual(live.sessions, run.run_id, org_id=live.org_id)
        summary["contains_errors"] = actual["failed_calls"] > 0
        summary["failed_call_rate"] = (
            actual["failed_calls"] / actual["calls"] if actual["calls"] else 0.0
        )
        runner.finish_paid_run(
            run,
            outcome=outcome,
            reason=reason,
            actual=actual,
            results_sha256=results_sha256,
            summary=summary,
        )
    result["usd"] = actual["usd"]
    return result


# --- the owner's picks -------------------------------------------------------------------


def _sheet_id(path: Path) -> str:
    """The id a sheet was made with (its hidden ``meta`` tab); the audit row is filed under it."""
    wb = load_workbook(path, read_only=True)
    try:
        for name, value, *_ in wb["meta"].iter_rows(values_only=True):
            if name == "sheet_id":
                return str(value)
    finally:
        wb.close()
    raise SheetRejected(["this spreadsheet has no sheet id"])


def load_routing(path: Path, *, live: LiveDB, owner: Owner) -> dict:
    """The ``load-routing`` command: the owner's marked sheet becomes the routing keys.

    All or nothing: ``SheetRejected`` lists every problem and no key is written (FR-19). The
    keys are stored with the fingerprint of the workflow list they were picked against and a
    digest of the request texts they were picked on; an unmarked request is stored as
    ``unmarked`` so the report can count it and the scorer leaves it out. Loading again
    replaces the earlier keys. Returns the counts by status.

    A sheet from an EARLIER draft (the owner deleted the routing files and drafted again) is
    valid in itself and has the same ids and workflows, so only the request texts tell it apart:
    its picks are about other messages, and loading them would key today's requests wrongly.
    """
    kind, keys = load_sheet(path)
    if kind != "routing":
        raise SheetRejected(["this is a facts sheet; routing picks are loaded from a routing sheet"])
    try:
        snapshot = json.loads(SNAPSHOT_PATH.read_text(encoding="utf-8"))
        requests = [
            RoutingRequest.model_validate(r)
            for r in json.loads(REQUESTS_PATH.read_text(encoding="utf-8"))
        ]
    except OSError:
        raise SheetRejected(["no routing requests have been drafted yet"]) from None
    problems = []
    if {k.workflow_fingerprint for k in keys} != {snapshot["fingerprint"]}:
        problems.append("the sheet was made against a different workflow list than the saved one")
    if {k.item_id for k in keys} != {r.item_id for r in requests}:
        problems.append("the sheet's requests are not the drafted requests")
    elif {k.requests_sha256 for k in keys} != {
        routing_requests_sha256([r.model_dump() for r in requests])
    }:
        problems.append(
            "the sheet was made from other request texts than the saved ones (an earlier "
            "draft?): load the sheet the last draft-routing wrote"
        )
    if problems:
        raise SheetRejected(problems)

    sheet_id = _sheet_id(Path(path))
    _write_json(KEYS_PATH, [k.model_dump(mode="json") for k in keys])
    drafted_for = {r.item_id: r.drafted_for_workflow_id for r in requests}
    mark = {"confirmed": "correct", "corrected": "wrong", "unmarked": None}
    counts = {s: sum(k.status == s for k in keys) for s in mark}
    audit(
        live,
        action="eval.labels_loaded",
        actor=owner,
        resource_type="eval_labels",
        resource_id=sheet_id,
        after={
            "kind": "routing",
            "sheet_sha256": keys[0].sheet_sha256 if keys else None,
            "counts": counts,
            "rows": [
                {
                    "item_id": k.item_id,
                    "suggested_value": drafted_for[k.item_id],
                    "mark": mark[k.status],
                    "final_value": k.workflow_id,
                }
                for k in keys
            ],
        },
    )
    return counts
