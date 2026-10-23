"""The spreadsheets the owner fills in: the facts sheet and the routing sheet (FR-18, FR-21).

``write_*`` makes an XLSX the owner can mark in Excel; ``load_sheet`` turns the returned file
into answer keys. The loader is the real check: worksheet protection on the question columns is
only an accident guard, so every locked cell is compared with what was written, and a file that
does not match the questions it was made from is rejected WHOLE with every problem listed
(FR-19), never loaded in part.

Rows are marked ``correct`` (the suggestion becomes the key), ``wrong`` (the owner's own value
becomes the key) or left blank (``unmarked``: kept in the list so the report can count it, but
never scored). A suggestion is never a key until the owner marks it (FR-3).

Everything written here quotes real contracts, so a sheet is only ever written to or read from
``SHEETS_DIR`` (git- and docker-ignored, FR-36). Stdlib, openpyxl and ``eval.exam.content``
only: like ``content.py`` this must import without pulling the app in.

Row dicts the writers take (the exam modules that build sheets pass these)::

    facts   {doc_id, contract (file name), field, value (None | str | list[str]), quote
             (already checked against the frozen text, or None), char_start, char_end, page
             (int | None), suggested_by, file_sha256, frozen_text_sha256, [question]}
    routing a ``RoutingRequest`` dump: {item_id, channel, subject, body,
             drafted_for_workflow_id, ...}

``path`` may be a ``.xlsx`` file or a folder; for a folder the file is named
``<kind>-<sheet_id>.xlsx`` because the id only exists once the sheet is written. The rows the
sheet was made from are saved beside it as ``<sheet_id>.expected.json``.
"""

import hashlib
import io
import json
import re
import uuid
from collections import Counter
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from eval.exam.content import PRIVATE_DIR, OwnerFactKey, RoutingKey, routing_requests_sha256

SHEETS_DIR = PRIVATE_DIR / "sheets"
TEXTS_DIR = PRIVATE_DIR / "owner" / "texts"  # <file_sha256>.txt, written by owner-import

NOT_STATED = "not stated in the document"
NO_SUPPORT = "no supporting sentence found"
NOTHING_TO_QUOTE = "(not stated, so there is no sentence to show)"
NO_WORKFLOW = "(no workflow — a person should route it)"

FACT_QUESTIONS = {
    "parties": "Who are the parties to this contract?",
    "effective_date": "On what date does this contract start (take effect)?",
    "end_date_or_term": "On what date does this contract end, or how long does it run?",
    "governing_law": "Which law governs this contract?",
    "notice_period": "How much notice is needed to end this contract or stop it renewing?",
    "renewal_term": "Does this contract renew, and on what terms?",
}

FACTS_HEADERS = [
    "row_key",
    "contract",
    "question",
    "suggested_answer",
    "supporting_sentence",
    "where",
    "your_mark",
    "correct_value_if_wrong",
    "notes",
]
ROUTING_HEADERS = [
    "row_key",
    "channel",
    "request",
    "suggested_workflow",
    "your_mark",
    "correct_workflow_if_wrong",
    "notes",
]
_TABS = {"facts": "Facts", "routing": "Routing"}
_HEADERS = {"facts": FACTS_HEADERS, "routing": ROUTING_HEADERS}
_WIDTHS = {
    "facts": [18, 30, 40, 40, 60, 24, 12, 40, 30],
    "routing": [18, 10, 70, 34, 12, 40, 30],
}
_MARKS = ("correct", "wrong")
_SHEET_ID = re.compile(r"[0-9a-f]{12}")

_FACTS_README = [
    "How to fill in this sheet",
    "",
    (
        "1. Each row is one fact about one contract: the question, the answer suggested by an "
        "AI, and the sentence from the contract that the suggestion rests on."
    ),
    (
        "2. In your_mark choose correct if the suggested answer is right, or wrong if it is "
        "not. Leave it blank if you cannot tell: the row is then left out of the exam and "
        "counted."
    ),
    (
        "3. If you choose wrong, write the right answer in correct_value_if_wrong. Put several "
        f"parties in one cell separated by | . If the contract does not say, write: {NOT_STATED}"
    ),
    (
        f"4. A row that reads '{NO_SUPPORT}' can only be marked wrong: the sentence the "
        "suggestion relied on is not in the contract."
    ),
    (
        "5. Only your_mark, correct_value_if_wrong and notes can be edited. Do not add, delete "
        "or reorder rows: a sheet that no longer matches the questions it was made from is "
        "rejected as a whole."
    ),
]
_ROUTING_README = [
    "How to fill in this sheet",
    "",
    (
        "1. Each row is a request someone might send in by email, chat or notice. An AI drafted "
        "it for the workflow shown in suggested_workflow; that is only a suggestion."
    ),
    (
        "2. In your_mark choose correct if that workflow is where the request should go, or "
        "wrong if it should go somewhere else. Leave it blank if you cannot tell: the row is "
        "then left out of the exam and counted."
    ),
    (
        "3. If you choose wrong, pick the right workflow in correct_workflow_if_wrong. Pick "
        f"'{NO_WORKFLOW}' if no workflow fits."
    ),
    (
        "4. Only your_mark, correct_workflow_if_wrong and notes can be edited. Do not add, "
        "delete or reorder rows: a sheet that no longer matches the questions it was made from "
        "is rejected as a whole."
    ),
]


class SheetRejected(Exception):
    """The returned sheet cannot be loaded; ``problems`` lists every reason (nothing loaded)."""

    def __init__(self, problems: list[str]):
        self.problems = list(problems)
        super().__init__("sheet rejected, nothing loaded: " + "; ".join(self.problems))


# --- small helpers ----------------------------------------------------------------------


def _in_sheets_dir(path: Path) -> bool:
    return Path(path).resolve().is_relative_to(SHEETS_DIR.resolve())


def _cell(text: str) -> str:
    """Text as it will sit in a cell: openpyxl refuses control characters (OCR output has
    form feeds) and Excel treats a cell over 32,767 characters as a corrupt file."""
    return ILLEGAL_CHARACTERS_RE.sub("", text)[:32_000]


def _same(a: str, b: str) -> bool:
    # Whitespace-insensitive: Excel and XML normalise line endings, and a changed space is not
    # a changed question.
    return a.split() == b.split()


def _norm(label: str) -> str:
    """A workflow label as the loader reads it back: spacing and case are not told apart (the
    owner types or pastes the pick, and Excel's dropdown match is case-insensitive)."""
    return " ".join(label.split()).casefold()


def _text(value) -> str:
    """A cell as the owner left it. Typing 2024-04-01 into an Excel cell stores a date, and a
    bare number is a float, so neither may reach the key as something other than the text the
    owner meant."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _clean_value(field: str, value) -> str | list[str] | None:
    """One spelling of a value for the key: "not stated" and empty are None; parties a list."""
    if isinstance(value, list):
        items = [v.strip() for v in value if v and v.strip()]
        return items or None
    text = (value or "").strip()
    if not text or text.casefold() == NOT_STATED:
        return None
    if field == "parties":
        return [p.strip() for p in text.split("|") if p.strip()] or None
    return text


def _show(value: str | list[str] | None) -> str:
    if value is None:
        return NOT_STATED
    return " | ".join(value) if isinstance(value, list) else value


def _locate(needle: str, text: str) -> tuple[int, int] | None:
    """Where ``needle`` appears word for word in ``text``, ignoring only spacing and line
    breaks (the same test the pre-fill quote check applies). First occurrence."""
    words = needle.split()
    if not words:
        return None
    match = re.search(r"\s+".join(re.escape(w) for w in words), text)
    return match.span() if match else None


def _row_key(sheet_id: str, row_id: str, field: str = "") -> str:
    return hashlib.sha256((sheet_id + row_id + field).encode("utf-8")).hexdigest()[:16]


def _put(ws, row: int, col: int, text: str, *, locked: bool = True, bold: bool = False):
    cell = ws.cell(row=row, column=col, value=text)
    # A contract sentence (or a request) that starts with "=" must stay text, never a formula.
    cell.data_type = "s"
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if bold:
        cell.font = Font(bold=True)
    if not locked:
        cell.protection = Protection(locked=False)
    return cell


# --- writing ----------------------------------------------------------------------------


def _write_workbook(
    target: Path,
    *,
    kind: str,
    rows: list[list[str]],
    mark_options: list[str],
    readme: list[str],
    meta: dict[str, str],
    workflow_labels: list[str] | None = None,
) -> None:
    """The workbook both sheets share: header, locked question columns, the last three
    columns (mark, correction, notes) editable, a mark dropdown per row."""
    headers, widths = _HEADERS[kind], _WIDTHS[kind]
    mark_col = len(headers) - 3  # 0-based: your_mark, then the correction, then notes
    wb = Workbook()
    ws = wb.active
    ws.title = _TABS[kind]
    guide = wb.create_sheet("README")
    info = wb.create_sheet("meta")
    info.sheet_state = "hidden"

    for c, name in enumerate(headers, start=1):
        _put(ws, 1, c, name, bold=True)
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = widths[c - 1]
    ws.column_dimensions["A"].hidden = True
    ws.freeze_panes = "A2"

    dropdowns: dict[str, DataValidation] = {}
    for opts in set(mark_options):
        dv = DataValidation(
            type="list",
            formula1=f'"{opts}"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="your_mark",
            error=f"Choose one of: {opts}. Leave it blank if you cannot tell.",
        )
        ws.add_data_validation(dv)
        dropdowns[opts] = dv
    if workflow_labels is not None:
        # Excel caps an inline list at 255 characters; 32 workflow names do not fit.
        choices = wb.create_sheet("workflows")
        choices.sheet_state = "hidden"
        for r, label in enumerate(workflow_labels, start=1):
            _put(choices, r, 1, label)
        pick = DataValidation(
            type="list",
            formula1=f"=workflows!$A$1:$A${len(workflow_labels)}",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="correct_workflow_if_wrong",
            error="Pick a workflow from the list.",
        )
        ws.add_data_validation(pick)

    for r, (cells, opts) in enumerate(zip(rows, mark_options, strict=True), start=2):
        for c, text in enumerate(cells, start=1):
            editable = c - 1 >= mark_col
            cell = _put(ws, r, c, text, locked=not editable)
            if c - 1 == mark_col:
                dropdowns[opts].add(cell)
            elif c - 1 == mark_col + 1:
                cell.number_format = "@"  # keep a typed 2024-04-01 as text, not an Excel date
                if workflow_labels is not None:
                    pick.add(cell)

    ws.protection.sheet = True
    ws.protection.formatColumns = False  # protection locks cells, not the owner's column widths
    ws.protection.formatRows = False
    for r, line in enumerate(readme, start=1):
        _put(guide, r, 1, line, bold=(r == 1))
    guide.column_dimensions["A"].width = 110
    for r, (key, value) in enumerate(meta.items(), start=1):
        _put(info, r, 1, key)
        _put(info, r, 2, value)
    wb.active = wb.sheetnames.index(_TABS[kind])
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)


def _target(path: Path, kind: str, sheet_id: str) -> Path:
    path = Path(path)
    target = path if path.suffix.lower() == ".xlsx" else path / f"{kind}-{sheet_id}.xlsx"
    if not _in_sheets_dir(target):
        raise ValueError(f"sheets are written to {SHEETS_DIR} only")
    return target


def _save_expected(target: Path, sheet_id: str, expected: dict) -> None:
    (target.parent / f"{sheet_id}.expected.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=1), encoding="utf-8"
    )


def write_facts_sheet(rows: list[dict], *, path: Path) -> str:
    """Write the facts sheet (one row per owner document x field); returns the sheet id.

    A stated suggestion whose quote could not be found in the frozen text reads
    "no supporting sentence found" and the dropdown offers only ``wrong`` (FR-17). A "not
    stated" suggestion has no sentence to check, so it can be confirmed: the owner is the
    check.
    """
    keys = [(r["doc_id"], r["field"]) for r in rows]
    if len(set(keys)) != len(keys):
        raise ValueError("duplicate (doc_id, field) rows")
    sheet_id = uuid.uuid4().hex[:12]
    target = _target(path, "facts", sheet_id)

    table: list[list[str]] = []
    options: list[str] = []
    expected_rows: list[dict] = []
    for row in rows:
        field = row["field"]
        value = _clean_value(field, row["value"])
        quote = row["quote"]
        unsupported = value is not None and quote is None
        if value is None:
            sentence = NOTHING_TO_QUOTE
        else:
            sentence = NO_SUPPORT if unsupported else quote
        where = ""
        if not unsupported and row["char_start"] is not None:
            span = f"characters {row['char_start']}–{row['char_end']}"
            page = row.get("page")
            where = f"page {page}, {span}" if page else span
        locked = [
            _cell(row["contract"]),
            _cell(row.get("question") or FACT_QUESTIONS[field]),
            _cell(_show(value)),
            _cell(sentence),
            where,
        ]
        row_key = _row_key(sheet_id, row["doc_id"], field)
        table.append([row_key, *locked, "", "", ""])
        options.append("wrong" if unsupported else "correct,wrong")
        expected_rows.append(
            {
                "row_key": row_key,
                "label": f"{row['doc_id']} / {field}",
                "locked": locked,
                "unsupported": unsupported,
                "doc_id": row["doc_id"],
                "field": field,
                "value": value,
                "quote": quote,
                "char_start": row["char_start"],
                "char_end": row["char_end"],
                "suggested_by": row["suggested_by"],
                "file_sha256": row["file_sha256"],
                "frozen_text_sha256": row["frozen_text_sha256"],
            }
        )
    frozen = sorted({(r["doc_id"], r["file_sha256"], r["frozen_text_sha256"]) for r in rows})
    manifest_sha = hashlib.sha256(json.dumps(frozen).encode("utf-8")).hexdigest()
    now = datetime.now(UTC).isoformat()
    meta = {
        "sheet_id": sheet_id,
        "kind": "facts",
        "generated_at": now,
        "row_count": str(len(rows)),
        "frozen_manifest_sha256": manifest_sha,
    }
    _write_workbook(
        target,
        kind="facts",
        rows=table,
        mark_options=options,
        readme=_FACTS_README,
        meta=meta,
    )
    _save_expected(
        target,
        sheet_id,
        {"sheet_id": sheet_id, "kind": "facts", "generated_at": now, "rows": expected_rows},
    )
    return sheet_id


def write_routing_sheet(rows: list[dict], *, snapshot: dict, path: Path) -> str:
    """Write the routing sheet (one row per drafted request); returns the sheet id.

    ``snapshot`` is the workflow snapshot the owner picks against; its fingerprint is carried
    into every key so a later change to the workflow list is noticed (AC-21).
    """
    live = [w for w in snapshot["workflows"] if w.get("enabled", True)]
    # Two workflows may share a name, or differ only by case or spacing ("MSA review" /
    # "msa  Review"), which the loader cannot tell apart; the owner must still be able to name
    # exactly one, or a pick would key the wrong workflow.
    names = Counter(_norm(w["name"]) for w in live)
    label_of = {
        w["id"]: w["name"] if names[_norm(w["name"])] == 1 else f"{w['name']} ({w['id'][:8]})"
        for w in live
    }
    taken = Counter(_norm(label) for label in [*label_of.values(), NO_WORKFLOW])
    if any(n > 1 for n in taken.values()):
        raise ValueError("two workflows cannot be told apart by name, even with their ids")
    if len({r["item_id"] for r in rows}) != len(rows):
        raise ValueError("duplicate item_id rows")
    sheet_id = uuid.uuid4().hex[:12]
    target = _target(path, "routing", sheet_id)

    table: list[list[str]] = []
    expected_rows: list[dict] = []
    for row in rows:
        suggested = row["drafted_for_workflow_id"]
        if suggested not in label_of:
            raise ValueError(f"{row['item_id']} was drafted for a workflow not in the snapshot")
        locked = [
            _cell(row["channel"]),
            _cell(f"{row['subject']}\n\n{row['body']}"),
            _cell(label_of[suggested]),
        ]
        row_key = _row_key(sheet_id, row["item_id"])
        table.append([row_key, *locked, "", "", ""])
        expected_rows.append(
            {
                "row_key": row_key,
                "label": row["item_id"],
                "locked": locked,
                "item_id": row["item_id"],
                "drafted_for_workflow_id": suggested,
            }
        )
    labels = [*label_of.values(), NO_WORKFLOW]
    now = datetime.now(UTC).isoformat()
    # The request texts the picks are made on: a sheet from an earlier draft (same ids, same
    # workflows, other texts) must not be loadable against today's requests.
    requests_sha256 = routing_requests_sha256(rows)
    meta = {
        "sheet_id": sheet_id,
        "kind": "routing",
        "generated_at": now,
        "row_count": str(len(rows)),
        "workflow_fingerprint": snapshot["fingerprint"],
        "requests_sha256": requests_sha256,
    }
    _write_workbook(
        target,
        kind="routing",
        rows=table,
        mark_options=["correct,wrong"] * len(table),
        readme=_ROUTING_README,
        meta=meta,
        workflow_labels=labels,
    )
    _save_expected(
        target,
        sheet_id,
        {
            "sheet_id": sheet_id,
            "kind": "routing",
            "generated_at": now,
            "workflow_fingerprint": snapshot["fingerprint"],
            "requests_sha256": requests_sha256,
            "workflows": {**{label: wid for wid, label in label_of.items()}, NO_WORKFLOW: None},
            "rows": expected_rows,
        },
    )
    return sheet_id


# --- loading ----------------------------------------------------------------------------


def _meta(wb) -> dict[str, str]:
    if "meta" not in wb.sheetnames:
        return {}
    return {
        _text(a): _text(b) for a, b, *_ in wb["meta"].iter_rows(values_only=True) if _text(a)
    }


def _frozen_text(file_sha256: str) -> str | None:
    path = TEXTS_DIR / f"{file_sha256}.txt"
    return path.read_text(encoding="utf-8") if path.is_file() else None


def load_sheet(path: Path) -> tuple[str, list[OwnerFactKey] | list[RoutingKey]]:
    """Validate a returned sheet and turn it into keys: ``("facts" | "routing", keys)``.

    All or nothing: any problem raises ``SheetRejected`` listing every one and no key is
    returned. Unmarked rows come back with status ``unmarked`` so the report can count them;
    the scorer leaves them out. A mark that disagrees with the correction cell (a correction
    typed against ``correct`` or against no mark) is a problem, not a guess: the system never
    decides what the owner meant.
    """
    path = Path(path)
    if not _in_sheets_dir(path):
        raise SheetRejected([f"sheets are read from {SHEETS_DIR} only"])
    try:
        data = path.read_bytes()
    except OSError:
        raise SheetRejected(["the spreadsheet could not be read"]) from None
    try:
        wb = load_workbook(io.BytesIO(data))
    except Exception:
        raise SheetRejected(["not a readable .xlsx spreadsheet"]) from None

    meta = _meta(wb)
    sheet_id = meta.get("sheet_id", "")
    saved = path.parent / f"{sheet_id}.expected.json"
    if not _SHEET_ID.fullmatch(sheet_id) or not saved.is_file():
        raise SheetRejected(["this spreadsheet was not made by the exam (unknown sheet id)"])
    expected = json.loads(saved.read_text(encoding="utf-8"))
    kind = expected["kind"]
    if meta.get("kind") != kind or _TABS[kind] not in wb.sheetnames:
        raise SheetRejected([f"this is not a {kind} sheet as it was made (tabs were changed)"])

    headers = _HEADERS[kind]
    ncols = len(headers)
    problems: list[str] = []
    found: dict[str, list[list]] = {}
    unkeyed: list[int] = []
    sheet_rows = list(wb[_TABS[kind]].iter_rows(values_only=True))
    if [_text(v) for v in (sheet_rows[0] if sheet_rows else ())[:ncols]] != headers:
        problems.append("the column headings were changed")
    for number, values in enumerate(sheet_rows[1:], start=2):
        cells = [*values[:ncols], *([None] * (ncols - len(values)))]
        if not any(_text(v) for v in cells):
            continue
        key = _text(cells[0])
        if key:
            found.setdefault(key, []).append(cells)
        else:
            unkeyed.append(number)
    problems += [f"row {n} has no row_key" for n in unkeyed]

    wanted = {r["row_key"]: r for r in expected["rows"]}
    problems += [
        f"row_key {k} appears {len(v)} times" for k, v in found.items() if len(v) > 1
    ]
    problems += [
        f"a row that is not on the sheet that was made ({k})"
        for k in sorted(found.keys() - wanted.keys())
    ]
    problems += [f"{r['label']}: its row is missing" for k, r in wanted.items() if k not in found]

    now = datetime.now(UTC)
    sheet_sha256 = hashlib.sha256(data).hexdigest()
    by_label: dict[str, set[str | None]] = {}
    for label, wid in expected.get("workflows", {}).items():
        by_label.setdefault(_norm(label), set()).add(wid)
    keys: list = []
    for row_key, spec in wanted.items():
        if len(found.get(row_key, [])) != 1:
            continue
        cells = found[row_key][0]
        label = spec["label"]
        n_locked = len(spec["locked"])
        for shown, wrote in zip(cells[1 : 1 + n_locked], spec["locked"], strict=True):
            if not _same(_text(shown), wrote):
                problems.append(f"{label}: a question column was changed")
                break
        mark = _text(cells[1 + n_locked]).casefold()
        correction = _text(cells[2 + n_locked])
        if mark not in ("", *_MARKS):
            problems.append(f"{label}: your_mark must be correct, wrong or blank")
            continue
        if mark == "correct" and spec.get("unsupported"):
            problems.append(f"{label}: marked correct but it has no supporting sentence")
            continue
        if mark == "wrong" and not correction:
            problems.append(f"{label}: marked wrong but no correct value was written")
            continue
        if mark != "wrong" and correction:
            problems.append(f"{label}: a correct value was written but the mark is not wrong")
            continue
        status = {"correct": "confirmed", "wrong": "corrected", "": "unmarked"}[mark]

        if kind == "routing":
            workflow_id = None
            if mark == "correct":
                workflow_id = spec["drafted_for_workflow_id"]
            elif mark == "wrong":
                picked = by_label.get(_norm(correction), set())
                if not picked:
                    problems.append(f"{label}: the corrected workflow is not in the list")
                    continue
                if len(picked) > 1:  # a sheet made before labels were kept apart
                    problems.append(f"{label}: the corrected workflow name fits more than one")
                    continue
                [workflow_id] = picked
            keys.append(
                RoutingKey(
                    item_id=spec["item_id"],
                    status=status,
                    workflow_id=workflow_id,
                    workflow_fingerprint=expected["workflow_fingerprint"],
                    sheet_sha256=sheet_sha256,
                    loaded_at=now,
                    requests_sha256=expected.get("requests_sha256"),
                )
            )
            continue

        value, quote, start, end = None, None, None, None
        if mark == "correct":
            value, quote = spec["value"], spec["quote"]
            start, end = spec["char_start"], spec["char_end"]
        elif mark == "wrong":
            value = _clean_value(spec["field"], correction)
            text = _frozen_text(spec["file_sha256"]) if isinstance(value, str) else None
            span = _locate(value, text) if text else None
            if span:
                start, end = span
                quote = text[start:end]
        keys.append(
            OwnerFactKey(
                doc_id=spec["doc_id"],
                field=spec["field"],
                status=status,
                value=value,
                quote=quote,
                char_start=start,
                char_end=end,
                suggested_value=spec["value"],
                suggested_quote=spec["quote"],
                suggested_by=spec["suggested_by"],
                file_sha256=spec["file_sha256"],
                frozen_text_sha256=spec["frozen_text_sha256"],
                sheet_sha256=sheet_sha256,
                loaded_at=now,
            )
        )
    if problems:
        raise SheetRejected(problems)
    return kind, keys
