"""Comparing a counterparty's returned version with the one we sent.

Guards the failures that make a negotiation round unsafe: a change the reviewer
never sees (an edit they didn't track, a clause they quietly removed), a
renumbered clause shown as rewritten, our own agreed wording reported as a
change of theirs, and a reply version that loses what we decided.
"""

import io
from types import SimpleNamespace

from docx import Document
from lxml import etree

from app.contract_files import revisions
from app.contract_files.text_extraction import extract_docx_text

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

BEFORE = """1. Definitions. Confidential Information means information marked confidential.

2. Term. This Agreement lasts one (1) year.

3. Survival. Obligations survive for two (2) years.

4. Governing law. This Agreement is governed by the laws of New York."""

# Ours: we changed survival (2 -> 5 years) and governing law (New York -> India).
OURS = """1. Definitions. Confidential Information means information marked confidential.

2. Term. This Agreement lasts one (1) year.

3. Survival. Obligations survive for five (5) years.

4. Governing law. This Agreement is governed by the laws of India."""


def _by_label(rows):
    return {r["label"].split(".")[0]: r for r in rows}


def test_a_clause_we_changed_and_they_kept_is_ours_not_theirs():
    """Our wording coming back untouched is agreement, not a change to review."""
    rows = revisions.compare(OURS, OURS, before_text=BEFORE)
    assert {r["kind"] for r in rows} == {"ours"}


def test_their_counter_revert_and_plain_change_are_told_apart():
    theirs = OURS.replace("lasts one (1) year", "lasts two (2) years") \
        .replace("five (5) years", "three (3) years") \
        .replace("laws of India", "laws of New York")
    rows = _by_label(revisions.compare(OURS, theirs, before_text=BEFORE))
    assert rows["2"]["kind"] == "changed"      # a clause we never touched
    assert rows["3"]["kind"] == "countered"    # changed what we proposed
    assert rows["4"]["kind"] == "reverted"     # put their original back
    assert ["-", "five (5)"] in rows["3"]["parts"] and ["+", "three (3)"] in rows["3"]["parts"]


def test_renumbering_is_not_a_change_but_a_removed_clause_is():
    """Deleting clause 2 renumbers everything after it; only the deletion is news."""
    theirs = """1. Definitions. Confidential Information means information marked confidential.

2. Survival. Obligations survive for five (5) years.

3. Governing law. This Agreement is governed by the laws of India."""
    rows = revisions.compare(OURS, theirs)
    assert [r["kind"] for r in rows] == ["removed"]
    assert "one (1) year" in rows[0]["our_text"]


def test_an_added_clause_is_found_in_place():
    theirs = OURS.replace("3. Survival.", "3. Residuals. Either party may use Residuals.\n\n4. Survival.")
    rows = revisions.compare(OURS, theirs)
    added = [r for r in rows if r["kind"] == "added"]
    assert len(added) == 1 and "Residuals" in added[0]["their_text"]


def _docx(paragraphs) -> bytes:
    """A Word file whose paragraphs may carry tracked changes: each paragraph is
    a list of (kind, text) with kind '' (plain), 'ins' or 'del'."""
    doc = Document()
    for runs in paragraphs:
        p = doc.add_paragraph()
        for kind, text in runs:
            r = etree.SubElement(p._p, f"{{{W}}}r")
            t = etree.SubElement(r, f"{{{W}}}{'delText' if kind == 'del' else 't'}")
            t.text = text
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            if kind:
                wrap = etree.Element(f"{{{W}}}{kind}", {f"{{{W}}}id": "1", f"{{{W}}}author": "Globex"})
                p._p.remove(r)
                wrap.append(r)
                p._p.append(wrap)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_tracked_insertions_are_read_and_an_unmarked_edit_is_flagged():
    """Their redline marks the survival change but edits the term silently.
    The extractor must read the marked insertion (it used to drop it), and the
    silent edit must be the one flagged."""
    ours = "1. Term. This Agreement lasts one (1) year.\n\n2. Survival. Obligations survive for five (5) years."
    data = _docx([
        [("", "1. Term. This Agreement lasts two (2) years.")],  # edited, not tracked
        [("", "2. Survival. Obligations survive for "), ("del", "five (5)"), ("ins", "three (3)"), ("", " years.")],
    ])
    theirs = extract_docx_text(Document(io.BytesIO(data)))
    assert "survive for three (3) years" in theirs  # the tracked insertion is kept

    rows = _by_label(revisions.compare(ours, theirs, unmarked=revisions.unmarked_view(data)))
    assert rows["1"]["unmarked"] is True
    assert rows["2"]["unmarked"] is False


def test_a_clean_copy_marks_nothing_so_nothing_is_flagged():
    """With no tracked changes at all, every difference is unmarked — flagging
    each one would be noise, so the round says so once instead."""
    assert revisions.unmarked_view(_docx([[("", "1. Term. Two years.")]])) is None


def test_our_reply_puts_our_wording_back_only_where_we_pushed_back():
    theirs = OURS.replace("lasts one (1) year", "lasts two (2) years") \
        .replace("five (5) years", "three (3) years") \
        .replace("laws of India", "laws of New York")
    rows = revisions.compare(OURS, theirs, before_text=BEFORE)
    decisions = {"2": "accepted", "3": "kept", "4": "countered"}
    changes = []
    for r in rows:
        key = r["label"].split(".")[0]
        changes.append(SimpleNamespace(**r, decision=decisions[key],
                                       counter_text="4. Governing law. The laws of England and Wales."
                                       if key == "4" else None))
    reply = revisions.counter_text(theirs, changes)
    assert "lasts two (2) years" in reply          # accepted theirs
    assert "five (5) years" in reply               # kept ours
    assert "England and Wales" in reply and "New York" not in reply  # our counter
    assert "Definitions" in reply                  # untouched clauses survive


def test_a_round_through_the_api(monkeypatch):
    """Their file logged -> the round lists their changes, flags the unmarked
    one -> decisions -> finishing writes our reply version with our wording
    back. Runs the real routes and database; AI jobs are switched off."""
    import pytest
    from fastapi.testclient import TestClient

    import app.contract_files.service as files_service
    import app.models
    from app.contracts.models import Contract
    from app.core.database import SessionLocal, utcnow
    from app.main import app

    monkeypatch.setattr(files_service, "dispatch_job", lambda db, *, job: job)
    c = TestClient(app)
    login = c.post("/api/v1/auth/login", json={"email": "admin@example.com", "password": "local-dev-password"})
    if login.status_code != 200:
        pytest.skip("no seeded admin in this database")
    h = {"authorization": "Bearer " + login.json()["access_token"]}
    docx_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    up = c.post("/api/v1/contracts/upload", headers=h, data={"title": "Revision round test NDA"},
                files={"file": ("nda.docx", _docx([[("", "1. Definitions.")]]), docx_type)})
    assert up.status_code == 202, up.text
    cid = up.json()["contract"]["id"]
    try:
        for text in (BEFORE, OURS):  # their original, then the version we sent
            assert c.put(f"/api/v1/contracts/{cid}/text", headers=h, json={"text": text}).status_code == 200
        theirs = _docx([
            [("", "1. Definitions. Confidential Information means information marked confidential.")],
            [("", "")],
            [("", "2. Term. This Agreement lasts two (2) years.")],  # not tracked
            [("", "")],
            [("", "3. Survival. Obligations survive for "), ("del", "five (5)"), ("ins", "three (3)"), ("", " years.")],
            [("", "")],
            [("", "4. Governing law. This Agreement is governed by the laws of India.")],
        ])
        r = c.post(f"/api/v1/contracts/{cid}/counterparty-revision", headers=h,
                   files={"file": ("globex-v3.docx", theirs, docx_type)})
        assert r.status_code == 201, r.text

        rnd = c.get(f"/api/v1/contracts/{cid}/revisions/current", headers=h).json()
        by = {ch["label"].split(".")[0]: ch for ch in rnd["changes"]}
        assert rnd["tracked"] is True and rnd["status"] == "open"
        assert by["2"]["kind"] == "changed" and by["2"]["unmarked"] is True
        assert by["3"]["kind"] == "countered" and by["3"]["unmarked"] is False
        assert by["4"]["kind"] == "ours" and by["4"]["decision"] == "agreed"

        finish = f"/api/v1/contracts/{cid}/revisions/{rnd['id']}/finish"
        assert c.post(finish, headers=h).status_code == 409  # undecided changes block it
        base = f"/api/v1/contracts/{cid}/revisions/{rnd['id']}/changes/"
        assert c.post(base + by["2"]["id"], headers=h, json={"decision": "accepted"}).status_code == 200
        assert c.post(base + by["3"]["id"], headers=h, json={"decision": "kept"}).status_code == 200
        done = c.post(finish, headers=h).json()
        assert done["status"] == "closed" and done["outcome"] == "counter" and done["outcome_version_id"]

        text = c.get(f"/api/v1/contracts/{cid}/versions/{done['outcome_version_id']}/text", headers=h).json()["text"]
        assert "two (2) years" in text and "five (5) years" in text and "laws of India" in text
    finally:
        # The upload queued AI jobs that were never sent; the stale-job sweeper
        # sends such jobs later, which would spend real calls on this contract.
        from app.jobs.models import JobRun

        db = SessionLocal()
        for job in db.query(JobRun).filter(JobRun.resource_type == "contract", JobRun.resource_id == cid,
                                           JobRun.status == "queued"):
            job.status = "cancelled"
        db.get(Contract, cid).deleted_at = utcnow()
        db.commit()
        db.close()
