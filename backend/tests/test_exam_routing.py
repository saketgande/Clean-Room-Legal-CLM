"""The routing set: today's workflows, the drafted requests, the owner's picks.

specs/005-ai-eval-baseline T016 + T022 (FR-3, FR-19, FR-21, FR-25, FR-39; AC-21). Nothing here can
spend money or touch real data: the Anthropic client is the fake from
``tests/exam_fake_anthropic.py``, workflows are invented rows, private files live under
``tmp_path``, and the "live database" is a scratch Postgres database holding only the five
tables the command writes or reads (audit_log, a_i_call_log, workflow, user, role), created per
module and dropped afterwards, so the shared dev database never receives an ``eval.*`` audit row.
"""

from __future__ import annotations

import asyncio
import json
import re
import uuid
from contextlib import contextmanager
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import NullPool

import app.models  # noqa: F401  (registers every mapper, so metadata has the tables below)
from app.core.config import settings
from app.core.database import Base, new_uuid
from app.core.models import AICallLog, AuditLog
from app.integrations import claude
from app.workflows.models import Workflow
from eval.exam import content as content_mod
from eval.exam import routing, sheets
from eval.exam.content import ROUTING_DRAFT_MODEL, RoutingRequest, routing_requests_sha256
from eval.exam.env import ExamRefused, LiveDB, Owner
from tests import exam_fake_anthropic as fake_mod
from tests.exam_fake_anthropic import install

SCRATCH_TABLES = ("role", "user", "audit_log", "a_i_call_log", "workflow")
OWNER = Owner(user_id="owner-user-id", email="owner@example.com")


@contextmanager
def scratch_database():
    """A throwaway Postgres database on the test server with ``SCRATCH_TABLES`` in it."""
    base = make_url(settings.database_url)
    name = f"legal_clm_scratch_{uuid.uuid4().hex[:10]}"
    admin = create_engine(
        base.set(database="postgres"), isolation_level="AUTOCOMMIT", poolclass=NullPool
    )
    try:
        with admin.connect() as conn:
            conn.execute(text(f'CREATE DATABASE "{name}"'))
    except SQLAlchemyError as exc:
        admin.dispose()
        pytest.skip(f"this database user cannot CREATE DATABASE ({type(exc).__name__})")
    engine = create_engine(base.set(database=name), poolclass=NullPool)
    try:
        Base.metadata.create_all(engine, tables=[Base.metadata.tables[t] for t in SCRATCH_TABLES])
        yield engine
    finally:
        engine.dispose()
        with admin.connect() as conn:
            conn.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
        admin.dispose()


@pytest.fixture(scope="module")
def engine():
    with scratch_database() as e:
        yield e


@pytest.fixture
def live(engine) -> LiveDB:
    """A ``LiveDB`` with its own org id, so tests never see each other's rows."""
    return LiveDB(
        url=engine.url.render_as_string(hide_password=False),
        sessions=sessionmaker(bind=engine, autoflush=False, expire_on_commit=False),
        org_id=new_uuid(),
    )


@pytest.fixture
def private(tmp_path, monkeypatch):
    """Every private path the module uses, under tmp_path: no test reads or writes eval/private."""
    folder = tmp_path / "routing"
    monkeypatch.setattr(routing, "ROUTING_DIR", folder)
    monkeypatch.setattr(routing, "REQUESTS_PATH", folder / "requests.json")
    monkeypatch.setattr(routing, "SNAPSHOT_PATH", folder / "workflows.json")
    monkeypatch.setattr(routing, "KEYS_PATH", tmp_path / "keys" / "routing.json")
    monkeypatch.setattr(routing, "SHEETS_DIR", tmp_path / "sheets")
    monkeypatch.setattr(sheets, "SHEETS_DIR", tmp_path / "sheets")
    return tmp_path


def add_workflows(live, n, *, disabled=0, org=None, names=None):
    with live.sessions() as db:
        for i in range(n + disabled):
            db.add(
                Workflow(
                    org_id=org or live.org_id,
                    name=names[i] if names else f"Flow {i:02d}",
                    description=f"does thing {i}",
                    enabled=i < n,
                    eval_order=10 + i,
                    criteria={"match_keyword": f"kw{i}", "conditions": []},
                    steps=[{"id": "s1", "type": "human_task", "name": "Review", "config": {}}],
                )
            )
        db.commit()


def reply(salt: str = ""):
    """The fake model answering the drafting prompt: as many requests as asked for. ``salt``
    changes the wording, so a second draft of the same workflows reads differently."""

    def message_for(payload, usage=None):
        channels = payload["output_config"]["format"]["schema"]["required"]
        body = {
            channel: {"subject": f"s{i}{salt}", "body": f"body {i}{salt} " + "w " * 50}
            for i, channel in enumerate(channels)
        }
        return {
            "id": "m",
            "type": "message",
            "role": "assistant",
            "model": payload["model"],
            "content": [{"type": "text", "text": json.dumps(body)}],
            "stop_reason": "end_turn",
            "usage": dict(usage or fake_mod.USAGE),
        }

    return message_for


def go(live, *, ceiling="25", approve=True):
    """The ``draft-routing`` command with the owner typing the go-ahead and the ceiling."""
    prompts: list[str] = []

    def ask(prompt):
        prompts.append(prompt)
        return ceiling if approve else ""

    result = asyncio.run(routing.draft_routing(live=live, owner=OWNER, ask=ask))
    return result, prompts


def audits(live, action):
    with live.sessions() as db:
        return list(
            db.scalars(
                select(AuditLog).where(AuditLog.org_id == live.org_id, AuditLog.action == action)
            )
        )


def calls(live):
    with live.sessions() as db:
        return list(db.scalars(select(AICallLog).where(AICallLog.org_id == live.org_id)))


def draft_and_fill(live, monkeypatch, marks, *, workflows=3, salt=""):
    """Draft a small set on the fake model, then mark its sheet the way the owner would:
    ``marks`` = one ``(mark, correction)`` per row, from the top."""
    add_workflows(live, workflows)
    install(monkeypatch)
    monkeypatch.setattr(fake_mod, "message_for", reply(salt))
    result, _ = go(live)
    fill_sheet(result["sheet"], marks)
    return result


def fill_sheet(path, marks):
    wb = load_workbook(path)
    ws = wb["Routing"]
    for row, (mark, correction) in zip(range(2, ws.max_row + 1), marks, strict=False):
        ws.cell(row=row, column=5, value=mark)
        ws.cell(row=row, column=6, value=correction)
    wb.save(path)


# --- the workflow snapshot (FR-21, AC-21) -----------------------------------------------------


def test_snapshot_is_the_enabled_workflows_of_this_org_and_its_fingerprint_moves(live):
    """Another org's and disabled workflows must not enter the snapshot (the picks would then be
    about workflows the app never routes to); a changed workflow must move the fingerprint, and a
    description or order edit must not (it cannot change where a request goes)."""
    add_workflows(live, 3, disabled=2)
    add_workflows(live, 2, org=new_uuid())
    snapshot = routing.snapshot_workflows(live)
    assert [w["name"] for w in snapshot["workflows"]] == ["Flow 00", "Flow 01", "Flow 02"]
    assert routing.routing_fresh(live, snapshot)
    assert not routing.routing_fresh(live, None)

    def flow_01(db):
        return db.scalars(
            select(Workflow).where(Workflow.org_id == live.org_id, Workflow.name == "Flow 01")
        ).one()

    with live.sessions() as db:
        flow_01(db).version = 2
        db.commit()
    assert not routing.routing_fresh(live, snapshot)
    again = routing.snapshot_workflows(live)
    assert again["fingerprint"] != snapshot["fingerprint"]

    with live.sessions() as db:
        flow = flow_01(db)
        flow.description, flow.eval_order = "x", 99
        db.commit()
    assert routing.snapshot_workflows(live)["fingerprint"] == again["fingerprint"]


def test_the_plan_spreads_about_50_requests_over_the_32_workflows():
    """FR-21: two requests for each of the first 18 workflows and one for each of the other 14
    is 50; fewer workflows than that still get two each, and no workflow gets the same channel
    twice."""
    snapshot = {"workflows": [{"id": f"w{i}", "name": f"n{i}"} for i in range(32)]}
    plan = routing.plan_calls(snapshot)
    assert sum(len(channels) for _, channels in plan) == 50
    assert all(len(set(channels)) == len(channels) for _, channels in plan)
    assert [len(channels) for _, channels in plan] == [2] * 18 + [1] * 14
    small = routing.plan_calls({"workflows": snapshot["workflows"][:3]})
    assert sum(len(channels) for _, channels in small) == 6


def test_the_estimate_prices_every_call_synchronously_at_the_drafting_model(live):
    """FR-31: the owner sees the drafting cost before approving; drafting is never batched (a
    one-time content step), and the worst case bounds the expected cost."""
    add_workflows(live, 32)
    estimate = routing.estimate_draft(routing.snapshot_workflows(live))
    assert (estimate.calls, estimate.batch_calls, estimate.sync_calls) == (32, 0, 32)
    assert 0 < estimate.usd < estimate.worst_case_usd < 1.5


# --- drafting (FR-21, FR-39) ------------------------------------------------------------------


def test_drafting_writes_the_requests_the_sheet_and_the_ledger_and_audit_rows(
    live, private, monkeypatch, capsys
):
    """FR-21/FR-39: one call per workflow on the drafting model in structured-output form (no
    temperature, no tools); the requests are numbered so any prefix spreads over many workflows;
    the snapshot is saved beside them; every call is on the ledger stamped with the run and holds
    no drafted text; the approval and the immutable run record are on the audit trail."""
    add_workflows(live, 5)
    fake = install(monkeypatch)
    monkeypatch.setattr(fake_mod, "message_for", reply())
    result, _ = go(live)
    assert result["requests"] == 10
    assert result["coverage"]["workflows_covered"] == 5
    assert set(result["coverage"]["requests_per_workflow"].values()) == {2}
    assert len(fake.streamed) == 5
    for payload in fake.streamed:
        assert payload["model"] == ROUTING_DRAFT_MODEL
        assert not {"temperature", "tools", "tool_choice"} & payload.keys()
        assert payload["output_config"]["format"]["type"] == "json_schema"

    requests = [
        RoutingRequest.model_validate(r) for r in json.loads(routing.REQUESTS_PATH.read_text())
    ]
    assert [r.item_id for r in requests] == [f"rt-{i:02d}" for i in range(1, 11)]
    assert {r.drafting_model for r in requests} == {ROUTING_DRAFT_MODEL}
    assert len({r.drafted_for_workflow_id for r in requests[:5]}) == 5  # a prefix spreads out
    snapshot = json.loads(routing.SNAPSHOT_PATH.read_text())
    assert snapshot["fingerprint"] == routing.snapshot_workflows(live)["fingerprint"]
    assert result["sheet"].is_file()

    rows = calls(live)
    assert len(rows) == 5 and {r.eval_run_id for r in rows} == {result["run_id"]}
    assert {r.model for r in rows} == {ROUTING_DRAFT_MODEL}
    assert {r.prompt_key for r in rows} == {"exam_routing_draft"}
    assert all(r.status == "succeeded" and r.batched is False for r in rows)
    assert all("body" not in json.dumps(r.input_payload) for r in rows)

    [approved] = audits(live, "eval.spend_approved")
    assert approved.after["kind"] == "draft_routing"
    assert approved.after["drafting_model"] == ROUTING_DRAFT_MODEL
    assert approved.after["drafted_by_model_under_test"] is False
    [done] = audits(live, "eval.run_finished")
    assert done.after["outcome"] == "complete" and done.after["actual"]["calls"] == 5
    assert done.after["results_sha256"]
    assert claude.exam_meter is None
    out = capsys.readouterr().out
    assert "Estimate for this draft routing run" in out and "Delivery: sync" in out


def test_without_a_go_ahead_nothing_is_spent_and_nothing_is_written(live, private, monkeypatch):
    """FR-31: a refusal makes no call, writes no file, and is on the record."""
    add_workflows(live, 3)
    fake = install(monkeypatch)
    result, _ = go(live, approve=False)
    assert result is None and fake.streamed == []
    assert not routing.REQUESTS_PATH.exists()
    assert len(audits(live, "eval.spend_refused")) == 1
    assert claude.exam_meter is None


def test_an_existing_routing_set_is_never_replaced_and_nothing_is_asked(live, private, monkeypatch):
    """FR-3: the owner's picks belong to those exact requests; drafting again must refuse (naming
    what to delete) before it asks for money, not overwrite them."""
    add_workflows(live, 2)
    install(monkeypatch)
    routing.KEYS_PATH.parent.mkdir(parents=True)
    routing.KEYS_PATH.write_text("[]")
    with pytest.raises(ExamRefused, match="keys/routing.json"):
        go(live)


def test_with_no_enabled_workflow_there_is_nothing_to_draft_for(live, private, monkeypatch):
    install(monkeypatch)
    with pytest.raises(ExamRefused, match="no enabled workflows"):
        go(live)


def test_the_answer_shape_requires_one_message_per_channel():
    """A run stopped at request 10 of 32 because the drafting answer was a free-length list and
    the model wrote one message where two were asked. Each channel is now a required key, so the
    count is enforced by the answer's shape, and an answer missing one is still refused."""
    schema = routing._schema(["email", "notice"])
    assert schema["required"] == ["email", "notice"]
    assert schema["additionalProperties"] is False

    def response(answer: dict) -> SimpleNamespace:
        text = {"type": "text", "text": json.dumps(answer)}
        return SimpleNamespace(stop_reason="end_turn", content_blocks=[text])

    short = response({"email": {"subject": "s", "body": "b"}})
    with pytest.raises(ValueError, match="email, notice"):
        routing._drafts_from(short, ["email", "notice"])
    both = response({c: {"subject": c, "body": "b"} for c in ("notice", "email")})
    drafts = routing._drafts_from(both, ["email", "notice"])
    assert [d["subject"] for d in drafts] == ["email", "notice"]  # channel order, not key order


def test_a_bad_answer_stops_the_run_and_writes_nothing(live, private, monkeypatch):
    """FR-21: a set that misses a workflow would break 'at least one request per workflow', so an
    answer that is not what was asked stops the run: nothing is written, the call is on the
    ledger as failed, and the run record says it stopped short ('interrupted', one of the plan's
    reasons) and 'contains errors'."""
    add_workflows(live, 4)
    install(monkeypatch)
    good = reply()

    def flaky(payload, usage=None):
        message = good(payload, usage)
        if "Flow 02" in payload["messages"][0]["content"]:
            message["content"][0]["text"] = '{"requests": []}'
        return message

    monkeypatch.setattr(fake_mod, "message_for", flaky)
    with pytest.raises(RuntimeError, match="Flow 02"):
        go(live)
    assert not routing.REQUESTS_PATH.exists() and not routing.SNAPSHOT_PATH.exists()
    assert not (private / "sheets").exists()
    rows = sorted(calls(live), key=lambda r: r.created_at)
    assert [r.status for r in rows] == ["succeeded", "succeeded", "validation_failed"]
    [done] = audits(live, "eval.run_finished")
    assert done.after["outcome"] == "incomplete"
    assert done.after["incomplete_reason"] == "interrupted"
    assert done.after["contains_errors"] is True
    assert claude.exam_meter is None


def test_the_ceiling_stops_drafting_before_the_call_that_could_pass_it(live, private, monkeypatch):
    """FR-32: spend never passes the ceiling; the meter refuses before any request is made."""
    add_workflows(live, 3)
    fake = install(monkeypatch)
    monkeypatch.setattr(fake_mod, "message_for", reply())
    with pytest.raises(RuntimeError, match="ceiling"):
        go(live, ceiling="0.001")
    assert fake.streamed == [] and not routing.REQUESTS_PATH.exists()
    [done] = audits(live, "eval.run_finished")
    assert done.after["incomplete_reason"] == "ceiling_reached"


def test_a_provider_error_is_a_failed_call_on_the_ledger(live, private, monkeypatch):
    """FR-40: a failed drafting call is counted and listed, never silently dropped."""
    add_workflows(live, 2)
    fake = install(monkeypatch)
    fake.stream_error = ValueError("boom")
    with pytest.raises(RuntimeError, match="boom"):
        go(live)
    [row] = calls(live)
    assert row.status == "failed" and row.error_class == "ValueError"
    [done] = audits(live, "eval.run_finished")
    assert done.after["incomplete_reason"] == "interrupted" and done.after["contains_errors"] is True


def test_drafting_with_the_model_under_test_is_warned_and_recorded(
    live, private, monkeypatch, capsys
):
    """FR-21: requests written by the model that will route them tend to be easier for it; that is
    said on screen and recorded, not hidden (a date suffix on the model name is ignored)."""
    add_workflows(live, 1)
    install(monkeypatch)
    monkeypatch.setattr(fake_mod, "message_for", reply())
    monkeypatch.setattr(settings, "claude_model", f"{ROUTING_DRAFT_MODEL}-20260101")
    go(live)
    assert "also the model under test" in capsys.readouterr().out
    [approved] = audits(live, "eval.spend_approved")
    assert approved.after["drafted_by_model_under_test"] is True


# --- the owner's picks (FR-19, FR-21, AC-21) --------------------------------------------------


def test_picks_become_keys_against_the_snapshot_and_are_audited(live, private, monkeypatch):
    """FR-3/FR-19/FR-21: correct -> the drafted-for workflow, wrong -> the owner's pick (or 'no
    workflow'), blank -> unmarked and counted; every key carries the workflow fingerprint and
    the request digest it was picked against; the audit row lists each row without the request
    text; and a workflow added later makes the routing stale."""
    result = draft_and_fill(
        live,
        monkeypatch,
        [
            ("correct", None),
            ("wrong", "Flow 02"),
            ("wrong", "(no workflow — a person should route it)"),
            (None, None),
        ],
    )
    counts = routing.load_routing(result["sheet"], live=live, owner=OWNER)
    assert counts == {"confirmed": 1, "corrected": 2, "unmarked": 3}
    keys = json.loads(routing.KEYS_PATH.read_text())
    assert [k["item_id"] for k in keys] == [f"rt-0{i}" for i in range(1, 7)]
    snapshot = json.loads(routing.SNAPSHOT_PATH.read_text())
    assert {k["workflow_fingerprint"] for k in keys} == {snapshot["fingerprint"]}
    requests = json.loads(routing.REQUESTS_PATH.read_text())
    assert {k["requests_sha256"] for k in keys} == {routing_requests_sha256(requests)}
    assert keys[2]["workflow_id"] is None and keys[2]["status"] == "corrected"
    assert keys[3]["workflow_id"] is None and keys[3]["status"] == "unmarked"

    [loaded] = audits(live, "eval.labels_loaded")
    assert loaded.resource_id == result["sheet_id"]
    assert loaded.after["counts"] == counts and len(loaded.after["rows"]) == 6
    assert "body" not in json.dumps(loaded.after)

    # the content loader's models accept everything written
    assert len(content_mod._models(content_mod.RoutingRequest, routing.REQUESTS_PATH)) == 6
    assert len(content_mod._models(content_mod.RoutingKey, routing.KEYS_PATH)) == 6
    add_workflows(live, 1)  # another enabled workflow: the list the picks were made on is gone
    assert not routing.routing_fresh(live, snapshot)


def test_a_rejected_sheet_loads_nothing(live, private, monkeypatch):
    """FR-19: a mark with no correction rejects the whole sheet; no key file, no audit row."""
    result = draft_and_fill(live, monkeypatch, [("wrong", None)])
    with pytest.raises(sheets.SheetRejected, match="no correct value"):
        routing.load_routing(result["sheet"], live=live, owner=OWNER)
    assert not routing.KEYS_PATH.exists() and audits(live, "eval.labels_loaded") == []


def test_a_sheet_made_against_another_workflow_list_is_rejected(live, private, monkeypatch):
    """AC-21: picks made against a workflow list other than the saved snapshot say nothing about
    today's routing and must not become keys."""
    result = draft_and_fill(live, monkeypatch, [("correct", None)])
    snapshot = json.loads(routing.SNAPSHOT_PATH.read_text())
    snapshot["fingerprint"] = "0" * 64
    routing.SNAPSHOT_PATH.write_text(json.dumps(snapshot))
    with pytest.raises(sheets.SheetRejected, match="different workflow list"):
        routing.load_routing(result["sheet"], live=live, owner=OWNER)
    assert not routing.KEYS_PATH.exists()


def test_a_facts_sheet_is_not_loaded_as_routing(live, private):
    rows = [
        {
            "doc_id": "owner-01", "contract": "x.pdf", "field": "parties", "value": ["A"],
            "quote": "A", "char_start": 0, "char_end": 1, "page": 1, "suggested_by": "m",
            "file_sha256": "a" * 64, "frozen_text_sha256": "b" * 64,
        }
    ]
    sheet_id = sheets.write_facts_sheet(rows, path=private / "sheets")
    with pytest.raises(sheets.SheetRejected, match="facts sheet"):
        routing.load_routing(private / "sheets" / f"facts-{sheet_id}.xlsx", live=live, owner=OWNER)


def test_loading_before_anything_was_drafted_is_rejected(live, private):
    """A routing sheet with no saved requests or snapshot cannot be matched to anything."""
    snapshot = {"fingerprint": "f" * 64, "workflows": [{"id": "wf1", "name": "One", "enabled": True}]}
    rows = [
        {
            "item_id": "rt-01", "channel": "email", "subject": "s", "body": "b",
            "drafted_for_workflow_id": "wf1", "drafting_model": "m",
        }
    ]
    sheet_id = sheets.write_routing_sheet(rows, snapshot=snapshot, path=private / "sheets")
    with pytest.raises(sheets.SheetRejected, match="no routing requests"):
        routing.load_routing(
            private / "sheets" / f"routing-{sheet_id}.xlsx", live=live, owner=OWNER
        )


# --- a sheet is bound to the exact requests it was made from (T022, FR-21) --------------------


def redraft(live, monkeypatch, salt):
    """The owner deletes the routing files and drafts again (the only way to start over)."""
    routing.REQUESTS_PATH.unlink()
    routing.SNAPSHOT_PATH.unlink()
    monkeypatch.setattr(fake_mod, "message_for", reply(salt))
    result, _ = go(live)
    return result


def test_a_sheet_from_an_earlier_draft_is_rejected_against_a_redrafted_set(
    live, private, monkeypatch
):
    """FR-21: after a re-draft the ids (rt-01..), the workflows and so the fingerprint are all the
    same, and the old sheet is valid in itself -- but its picks are about OTHER messages. Loading it
    would key today's requests with answers to yesterday's, so only the request texts can tell
    them apart. Nothing is loaded; the new draft's own sheet still loads."""
    first = draft_and_fill(live, monkeypatch, [("correct", None)] * 3, salt=" (first draft)")
    second = redraft(live, monkeypatch, " (second draft)")
    assert first["sheet_id"] != second["sheet_id"]
    with pytest.raises(sheets.SheetRejected, match="other request texts") as raised:
        routing.load_routing(first["sheet"], live=live, owner=OWNER)
    assert "different workflow list" not in "; ".join(raised.value.problems)
    assert not routing.KEYS_PATH.exists() and audits(live, "eval.labels_loaded") == []

    fill_sheet(second["sheet"], [("correct", None)] * 3)
    assert routing.load_routing(second["sheet"], live=live, owner=OWNER)["confirmed"] == 3


def test_a_sheet_is_still_good_when_a_redraft_reads_exactly_the_same(live, private, monkeypatch):
    """The binding is to the texts, not to the draft: the same messages are the same questions, so
    picks made on them are not thrown away just because the files were drafted again."""
    first = draft_and_fill(live, monkeypatch, [("correct", None)] * 3, salt=" same")
    redraft(live, monkeypatch, " same")
    assert routing.load_routing(first["sheet"], live=live, owner=OWNER)["confirmed"] == 3


def test_requests_edited_after_the_sheet_was_made_reject_the_picks(live, private, monkeypatch):
    """The saved requests are an input to every run: a hand edit of one after the owner marked
    the sheet means the owner never saw what the run asks."""
    result = draft_and_fill(live, monkeypatch, [("correct", None)] * 3)
    requests = json.loads(routing.REQUESTS_PATH.read_text())
    requests[0]["body"] += " Also, please call me."
    routing.REQUESTS_PATH.write_text(json.dumps(requests))
    with pytest.raises(sheets.SheetRejected, match="other request texts"):
        routing.load_routing(result["sheet"], live=live, owner=OWNER)
    assert not routing.KEYS_PATH.exists()


def test_a_sheet_made_before_requests_were_bound_is_rejected(live, private, monkeypatch):
    """A sheet whose saved record has no request digest cannot be shown to match: refusing it is
    the safe answer (the owner re-drafts or regenerates) rather than trusting the ids."""
    result = draft_and_fill(live, monkeypatch, [("correct", None)] * 3)
    saved = result["sheet"].parent / f"{result['sheet_id']}.expected.json"
    record = json.loads(saved.read_text())
    del record["requests_sha256"]
    saved.write_text(json.dumps(record))
    with pytest.raises(sheets.SheetRejected, match="other request texts"):
        routing.load_routing(result["sheet"], live=live, owner=OWNER)


def test_workflows_whose_names_differ_only_by_case_are_drafted_and_picked_correctly(
    live, private, monkeypatch
):
    """FR-21: 'MSA review' and 'msa  Review' read as one name to the sheet loader; they must get
    distinct labels (an id suffix), or a pick of one would key the other."""
    add_workflows(live, 2, names=["MSA review", "msa  Review"])
    install(monkeypatch)
    monkeypatch.setattr(fake_mod, "message_for", reply())
    result, _ = go(live)
    snapshot = json.loads(routing.SNAPSHOT_PATH.read_text())
    second = snapshot["workflows"][1]
    ws = load_workbook(result["sheet"])["Routing"]
    labels = {ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)}
    assert len(labels) == 2 and all(re.search(r" \([0-9a-f]{8}\)$", label) for label in labels)
    fill_sheet(result["sheet"], [("wrong", f"msa  Review ({second['id'][:8]})")])
    routing.load_routing(result["sheet"], live=live, owner=OWNER)
    first_key = json.loads(routing.KEYS_PATH.read_text())[0]
    assert first_key["workflow_id"] == second["id"]


# --- the content loader refuses inconsistent routing files (T022) -------------------------------


def write_private(root, requests=None, keys=None):
    for rel, value in (("routing/requests.json", requests), ("keys/routing.json", keys)):
        if value is not None:
            path = root / "private" / rel
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(value), encoding="utf-8")


def request(item_id="rt-01", body="Need an NDA."):
    return {
        "item_id": item_id, "channel": "email", "subject": "NDA", "body": body,
        "drafted_for_workflow_id": "wf-1", "drafting_model": ROUTING_DRAFT_MODEL,
    }


def routing_key(item_id="rt-01", requests_sha256=None):
    return {
        "item_id": item_id, "status": "confirmed", "workflow_id": "wf-1",
        "workflow_fingerprint": "f" * 64, "sheet_sha256": "1" * 64,
        "loaded_at": "2026-10-07T00:00:00+00:00", "requests_sha256": requests_sha256,
    }


def load_from(tmp_path, monkeypatch):
    monkeypatch.setattr(content_mod, "CONTENT_DIR", tmp_path / "content")
    monkeypatch.setattr(content_mod, "PRIVATE_DIR", tmp_path / "private")
    return content_mod.load_content()


def test_consistent_routing_files_load(tmp_path, monkeypatch):
    """The checks below must not refuse the files `draft-routing` and `load-routing` write."""
    requests = [request("rt-01"), request("rt-02", body="Renew our MSA.")]
    digest = routing_requests_sha256(requests)
    write_private(tmp_path, requests, [routing_key("rt-01", digest), routing_key("rt-02", digest)])
    content = load_from(tmp_path, monkeypatch)
    assert len(content.routing) == 2 and len(content.routing_keys) == 2


@pytest.mark.parametrize(
    ("requests", "keys", "problem"),
    [
        ([request("rt-01"), request("rt-01", body="another")], [], "duplicate routing request rt-01"),
        ([request("rt-01")], [routing_key("rt-09")], "routing key for unknown request rt-09"),
        ([request("rt-01")], [routing_key("rt-01"), routing_key("rt-01")], "duplicate routing key"),
        ([], [routing_key("rt-01")], "routing key for unknown request rt-01"),
    ],
)
def test_routing_files_whose_ids_disagree_are_refused(tmp_path, monkeypatch, requests, keys, problem):
    """A duplicate request id would be sent twice and scored once; a key for a request that is not
    there scores nothing and reads as 'unpicked' for ever. Neither fails later, so the loader
    refuses them before money is estimated."""
    write_private(tmp_path, requests, keys)
    with pytest.raises(ValueError, match="inconsistent") as raised:
        load_from(tmp_path, monkeypatch)
    assert problem in str(raised.value)


def test_keys_picked_on_other_request_texts_are_refused(tmp_path, monkeypatch):
    """The same guard at load time as at `load-routing`: keys stored against one set of texts and
    requests that now read differently (a hand edit, a restored backup) are not a key."""
    old = routing_requests_sha256([request("rt-01", body="the text the owner saw")])
    write_private(tmp_path, [request("rt-01", body="a different text")], [routing_key("rt-01", old)])
    with pytest.raises(ValueError, match="different request texts"):
        load_from(tmp_path, monkeypatch)


def test_the_request_digest_follows_the_text_and_not_the_order_or_the_drafting_model():
    """The digest binds a pick to what it was about: any change to the message, its channel, its id
    or the workflow it was drafted for changes it; the order of the list and the drafting model
    (provenance) do not."""
    a, b = request("rt-01"), request("rt-02", body="Renew our MSA.")
    base = routing_requests_sha256([a, b])
    assert routing_requests_sha256([b, a]) == base
    assert routing_requests_sha256([a, {**b, "drafting_model": "another-model"}]) == base
    for change in ({"body": "x"}, {"subject": "x"}, {"channel": "chat"},
                   {"drafted_for_workflow_id": "wf-2"}, {"item_id": "rt-03"}):
        assert routing_requests_sha256([a, {**b, **change}]) != base
