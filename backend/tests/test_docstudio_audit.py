"""Extraction cannot check itself, so something else has to count.

When a run of text is skipped, what remains is still a grammatical sentence and
still looks exactly like a document that parsed correctly. A counterparty's
redline came out as "Liability is capped at  of fees." — a clause with its only
number missing, reading perfectly, passing every quality measure the pipeline
had.

This counts the text the file holds, by a different route from the parser that
read it, and reports the gap. Crude on purpose: characters, not structure,
because it is looking for the one thing a structure-aware check cannot see —
content that is simply absent.
"""

from io import BytesIO

from docx import Document
from docx.oxml import parse_xml

from app.docstudio.audit import UNEXPLAINED_THRESHOLD, audit, describe
from app.docstudio.parsing.registry import DOCX_MIME, parser_for
from app.docstudio.structure import build

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _bytes(document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _audit(content, extracted, dropped=0):
    return audit(
        content, mime_type=DOCX_MIME, extracted_text=extracted, deliberately_dropped=dropped
    )


def _redlined_docx() -> bytes:
    document = Document()
    document.add_paragraph("1. TERM. This Agreement runs for three years.")
    document.element.body.insert(1, parse_xml(f'''<w:p {W}>
      <w:r><w:t xml:space="preserve">Liability is capped at </w:t></w:r>
      <w:ins w:id="1" w:author="Counsel" w:date="2026-01-01T00:00:00Z">
        <w:r><w:t>twelve (12) months of fees as agreed between the parties</w:t></w:r></w:ins>
      <w:r><w:t xml:space="preserve">.</w:t></w:r></w:p>'''))
    return _bytes(document)


# --- it catches the loss ----------------------------------------------------


def test_a_dropped_tracked_insertion_is_detected():
    """The defect this file exists for, simulated by extracting the text a
    reader blind to `w:ins` would produce."""
    content = _redlined_docx()
    blind = "1. TERM. This Agreement runs for three years.\n\nLiability is capped at ."

    result = _audit(content, blind)

    assert result.is_suspicious
    assert result.unexplained > 40


def test_the_finding_names_what_the_document_uses():
    """"5% of the text is missing" sends an engineer reading XML. "…and the
    document uses tracked insertions" sends them to the cause."""
    result = _audit(_redlined_docx(), "Liability is capped at .")

    assert "tracked insertions" in describe(result)


def test_a_correct_extraction_raises_nothing():
    """The guard that keeps this usable. A check that fires on every document
    is a check everybody turns off."""
    content = _redlined_docx()
    extracted = build(parser_for(DOCX_MIME).parse(content, filename="c.docx")).flat_text

    assert _audit(content, extracted).is_suspicious is False


# --- what must not be mistaken for a loss -----------------------------------


def test_whitespace_differences_are_not_a_loss():
    """Extraction collapses runs of spaces, drops empty paragraphs and re-joins
    wrapped lines. Counting raw length would flag every document ever parsed."""
    document = Document()
    document.add_paragraph("1.    TERM.     Three     years.")
    document.add_paragraph("")
    document.add_paragraph("2. FEES.")

    result = _audit(_bytes(document), "1. TERM. Three years.\n\n2. FEES.")

    assert result.unexplained == 0


def test_deliberately_removed_furniture_is_not_a_loss():
    """Page headers are stripped on purpose. Without subtracting them, every
    document with a running header looks like a parser failure and the real
    signal is lost in the noise."""
    document = Document()
    document.add_paragraph("CONFIDENTIAL - DRAFT")
    document.add_paragraph("1. TERM. Three years.")

    result = _audit(_bytes(document), "1. TERM. Three years.", dropped=len("CONFIDENTIAL-DRAFT"))

    assert result.unexplained == 0
    assert result.is_suspicious is False


def test_struck_wording_is_not_expected_in_the_extraction():
    """Deleted text is not part of the document as it reads, so the source
    count must exclude it — otherwise every redline reports a false loss the
    size of everything the other side removed."""
    document = Document()
    document.element.body.insert(0, parse_xml(f'''<w:p {W}>
      <w:r><w:t xml:space="preserve">The term is </w:t></w:r>
      <w:del w:id="2" w:author="X" w:date="2026-01-01T00:00:00Z">
        <w:r><w:delText>three years and no longer than that</w:delText></w:r></w:del>
      <w:r><w:t>five years.</w:t></w:r></w:p>'''))

    result = _audit(_bytes(document), "The term is five years.")

    assert result.unexplained == 0


# --- boundaries -------------------------------------------------------------


def test_an_unreadable_file_is_not_reported_as_a_loss():
    """A file the auditor cannot open says nothing rather than claiming 100%
    of it went missing."""
    assert _audit(b"not a docx at all", "some text") is None


def test_a_format_with_no_independent_reader_is_skipped():
    """Silence is honest. Inventing a comparison for a format there is no
    second way to read would produce a number nobody could act on."""
    assert audit(b"anything", mime_type="text/plain", extracted_text="x") is None


def test_the_threshold_leaves_room_for_normal_variation():
    """Set too tight this fires constantly and gets ignored; too loose and a
    missing clause hides inside it."""
    assert 0 < UNEXPLAINED_THRESHOLD <= 0.05


def test_an_empty_file_does_not_divide_by_zero():
    document = Document()
    result = _audit(_bytes(document), "")

    assert result.ratio == 0.0
    assert result.is_suspicious is False
