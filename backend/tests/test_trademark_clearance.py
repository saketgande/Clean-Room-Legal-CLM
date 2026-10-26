"""Trademark clearance case: the workflow definition, saved searches, the
decision and filing records and the step-owner rule. DB-backed (Postgres)."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import delete

import app.models  # noqa: F401  (registers every mapper)
from app.auth.models import User
from app.core.config import settings
from app.core.enums import SourceStatus, TrademarkDecision
from app.core.models import ResourceTimelineEvent
from app.intake.models import IntakeRequest
from app.notifications.models import Notification
from app.trademarks.clearance_service import CLEARANCE_MARKER_FIELD, TrademarkClearanceService
from app.trademarks.models import Trademark
from app.trademarks.providers import similarity
from app.trademarks.schemas import (
    DecisionRequest,
    FilingRequest,
    SearchRunUpdate,
    SearchSimilarRequest,
    SearchSimilarResponse,
    SearchSummary,
    SourceResult,
)
from app.workflows.builtin import BUILTIN_FLOWS
from app.workflows.models import WorkflowRun, WorkflowStepRun
from app.workflows.service import _matches, _skip_reason


def _clearance_flow() -> dict:
    return next(f for f in BUILTIN_FLOWS if f["name"] == "Trademark Clearance & Filing")


# --- the flow definition (pure) ---------------------------------------------


def _request(**overrides):
    base = {
        "type_label": "Trademark",
        "description": "",
        "priority": "Medium",
        "department": None,
        "field_values": {},
    }
    return SimpleNamespace(**{**base, **overrides})


def test_clearance_flow_matches_only_requests_that_carry_the_marker():
    criteria = _clearance_flow()["criteria"]
    assert _matches(criteria, _request(field_values={CLEARANCE_MARKER_FIELD: True}))
    # Free text about trademarks - e.g. a cease-and-desist - must never select it.
    assert not _matches(criteria, _request(description="Cease and desist over trademark infringement"))
    assert not _matches(criteria, _request(field_values={"mark": "ACME"}))


def test_clearance_flow_covers_the_raci_steps_in_order():
    steps = _clearance_flow()["steps"]
    assert [s["name"] for s in steps] == [
        "Requisition & preliminary search",
        "Detailed search & risk mapping",
        "TM Legal review",
        "Final decision",
        "Escalation (TM Head / GC)",
        "Filing preparation",
        "Filing & docketing",
    ]
    by_name = {s["name"]: s for s in steps}
    assert by_name["TM Legal review"]["config"]["return_to"] == 1
    # Escalation only runs when flagged; filing is skipped for a rejected mark.
    assert by_name["Escalation (TM Head / GC)"]["cond"] == {"field": "tm_escalate", "op": "eq", "value": True}
    assert by_name["Filing & docketing"]["cond"] == {"field": "tm_decision", "op": "ne", "value": "reject"}
    rejected = _request(field_values={"tm_decision": "reject"})
    filing = by_name["Filing & docketing"]
    # The engine skips a step only when its rule is false; one it cannot evaluate yet (no
    # decision recorded) still runs, so a case can never skip filing by accident.
    assert _skip_reason(filing, rejected) == "Condition not met"
    assert _skip_reason(filing, _request(field_values={})) is None
    assert _skip_reason(filing, _request(field_values={"tm_decision": "proceed"})) is None


# --- a clearance case in the database ---------------------------------------


@pytest.fixture
def case(db_session):
    org_id = str(uuid.uuid4())
    suffix = uuid.uuid4().hex[:8]

    def make_user(label: str) -> User:
        user = User(
            id=str(uuid.uuid4()),
            org_id=org_id,
            email=f"{label}-{suffix}@clearance.test",
            full_name=label.title(),
            hashed_password="x",
            status="active",
        )
        db_session.add(user)
        return user

    owner, other = make_user("owner"), make_user("other")
    db_session.flush()

    request = IntakeRequest(
        id=str(uuid.uuid4()),
        org_id=org_id,
        ref=f"REQ-{suffix}",
        requester_user_id=owner.id,
        type_label="Trademark",
        description="Trademark clearance",
        field_values={"mark": "ZQ", CLEARANCE_MARKER_FIELD: True},
        submitted_at=datetime.now(UTC),
    )
    db_session.add(request)
    db_session.flush()
    run = WorkflowRun(
        id=str(uuid.uuid4()),
        org_id=org_id,
        request_id=request.id,
        flow_name="Trademark Clearance & Filing",
        flow_version=1,
        steps=[{"name": "Final decision"}],
        status="waiting",
        current_index=0,
        context={},
    )
    db_session.add(run)
    db_session.flush()
    step = WorkflowStepRun(
        id=str(uuid.uuid4()),
        org_id=org_id,
        flow_run_id=run.id,
        idx=0,
        step_type="human_task",
        step_name="Final decision",
        status="waiting_human",
        assignee_user_id=owner.id,
    )
    trademark = Trademark(
        id=str(uuid.uuid4()),
        org_id=org_id,
        name="ZQ",
        jurisdiction="IN",
        status="draft",
        source_intake_request_id=request.id,
    )
    db_session.add_all([step, trademark])
    db_session.commit()

    yield SimpleNamespace(
        org_id=org_id, owner=owner, other=other, request=request, run=run, step=step, trademark=trademark
    )

    db_session.rollback()
    db_session.execute(delete(Notification).where(Notification.org_id == org_id))
    db_session.execute(delete(ResourceTimelineEvent).where(ResourceTimelineEvent.org_id == org_id))
    db_session.execute(delete(Trademark).where(Trademark.org_id == org_id))
    db_session.execute(delete(WorkflowStepRun).where(WorkflowStepRun.org_id == org_id))
    db_session.execute(delete(WorkflowRun).where(WorkflowRun.org_id == org_id))
    db_session.execute(delete(IntakeRequest).where(IntakeRequest.org_id == org_id))
    db_session.execute(delete(User).where(User.org_id == org_id))
    db_session.commit()


def test_decision_is_limited_to_the_step_owner(db_session, case):
    service = TrademarkClearanceService(db_session)
    with pytest.raises(HTTPException) as denied:
        service.record_decision(
            user=case.other, trademark_id=case.trademark.id, payload=DecisionRequest(decision=TrademarkDecision.PROCEED)
        )
    assert denied.value.status_code == 403

    case.step.status = "done"
    db_session.commit()
    with pytest.raises(HTTPException) as idle:
        service.record_decision(
            user=case.owner, trademark_id=case.trademark.id, payload=DecisionRequest(decision=TrademarkDecision.PROCEED)
        )
    assert idle.value.status_code == 409


def test_conditional_decision_needs_conditions_and_feeds_the_workflow(db_session, case):
    service = TrademarkClearanceService(db_session)
    with pytest.raises(HTTPException) as missing:
        service.record_decision(
            user=case.owner,
            trademark_id=case.trademark.id,
            payload=DecisionRequest(decision=TrademarkDecision.CONDITIONAL),
        )
    assert missing.value.status_code == 422

    trademark = service.record_decision(
        user=case.owner,
        trademark_id=case.trademark.id,
        payload=DecisionRequest(decision=TrademarkDecision.CONDITIONAL, conditions="Device logo only", escalate=True),
    )
    assert trademark.decision == "conditional"
    assert trademark.decision_conditions == "Device logo only"
    assert trademark.decided_by_user_id == case.owner.id
    db_session.refresh(case.request)
    assert case.request.field_values["tm_decision"] == "conditional"
    assert case.request.field_values["tm_escalate"] is True


def test_rejection_closes_the_mark_and_an_override_reopens_it(db_session, case):
    service = TrademarkClearanceService(db_session)
    rejected = service.record_decision(
        user=case.owner, trademark_id=case.trademark.id, payload=DecisionRequest(decision=TrademarkDecision.REJECT)
    )
    assert rejected.status == "abandoned"
    overridden = service.record_decision(
        user=case.owner, trademark_id=case.trademark.id, payload=DecisionRequest(decision=TrademarkDecision.PROCEED)
    )
    assert overridden.decision == "proceed"
    assert overridden.status == "draft"


def test_filing_records_the_docket_and_marks_the_mark_filed(db_session, case):
    service = TrademarkClearanceService(db_session)
    with pytest.raises(HTTPException) as empty:
        service.record_filing(user=case.owner, trademark_id=case.trademark.id, payload=FilingRequest())
    assert empty.value.status_code == 422

    trademark = service.record_filing(
        user=case.owner,
        trademark_id=case.trademark.id,
        payload=FilingRequest(application_number="IN-1", filing_form="TM-A", docket_number="DK-9"),
    )
    assert (trademark.application_number, trademark.filing_form, trademark.docket_number) == ("IN-1", "TM-A", "DK-9")
    assert trademark.status == "filed"
    assert trademark.filed_on is not None
    assert trademark.renewal_due_on is not None


def test_case_summary_and_awaiting_action_follow_the_current_step(db_session, case):
    service = TrademarkClearanceService(db_session)
    summary = service.case_summaries(org_id=case.org_id, request_ids=[case.request.id])[case.request.id]
    assert summary.step_name == "Final decision"
    assert summary.step_status == "waiting_human"
    assert summary.assignee_user_id == case.owner.id

    assert [a.trademark_id for a in service.awaiting_action(user=case.owner)] == [case.trademark.id]
    assert service.awaiting_action(user=case.other) == []
    assert service.step_counts(org_id=case.org_id, request_ids=[case.request.id]) == {"Final decision": 1}


def test_saved_search_attaches_to_the_mark_is_rated_and_exports_to_excel(db_session, case):
    service = TrademarkClearanceService(db_session)
    request = SearchSimilarRequest(trademark_name="ZQ", description="insect spray")
    response = SearchSimilarResponse(
        query_id=f"srch_{uuid.uuid4().hex[:8]}",
        sources={
            "internal_portfolio": SourceResult(
                status=SourceStatus.COMPLETE,
                results=[{"name": "ZQX", "similarity_score": 0.7, "risk_level": "medium", "nice_class": "5"}],
            )
        },
        summary=SearchSummary(
            high_risk_count=0, medium_risk_count=1, low_risk_count=0, recommendation="review_recommended"
        ),
    )
    run = service.save_search_run(user=case.owner, request=request, response=response)
    assert run.trademark_id is None

    service.attach_run(org_id=case.org_id, query_id=response.query_id, trademark_id=case.trademark.id)
    db_session.commit()
    runs = service.list_runs(org_id=case.org_id, trademark_id=case.trademark.id)
    assert [r.id for r in runs] == [run.id]
    assert runs[0].purpose == "preliminary"

    rated = service.update_run(user=case.owner, run_id=run.id, payload=SearchRunUpdate(risk_level="high", notes="Close"))
    assert (rated.risk_level, rated.notes) == ("high", "Close")

    filename, content = service.export_run_xlsx(user=case.owner, run_id=run.id)
    assert filename == f"trademark-search-{response.query_id}.xlsx"
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Summary", "internal_portfolio"]
    summary = {row[0].value: row[1].value for row in workbook["Summary"].iter_rows()}
    assert summary["Overall risk (LTE)"] == "high"
    assert workbook["internal_portfolio"]["A1"].value == "name"
    assert workbook["internal_portfolio"]["A2"].value == "ZQX"

    # Another org can never read the run.
    with pytest.raises(HTTPException) as foreign:
        service.get_run(org_id=str(uuid.uuid4()), run_id=run.id)
    assert foreign.value.status_code == 404


def test_a_case_search_never_matches_the_mark_under_review(db_session, case, monkeypatch):
    vector = [1.0] + [0.0] * 383
    monkeypatch.setattr(similarity, "embed_texts", lambda texts: [vector for _ in texts])
    case.trademark.embedding = vector
    db_session.commit()

    from_case = SearchSimilarRequest(trademark_name="ZQ", trademark_id=case.trademark.id)
    plain = SearchSimilarRequest(trademark_name="ZQ")
    own = similarity._search_internal_portfolio(db_session, org_id=case.org_id, request=from_case, settings=settings)
    everyone = similarity._search_internal_portfolio(db_session, org_id=case.org_id, request=plain, settings=settings)

    assert own.results == []
    assert [r["trademark_id"] for r in everyone.results] == [case.trademark.id]


def test_audit_log_records_searches_exports_and_documents(db_session, case):
    from app.core.audit import write_audit_log

    service = TrademarkClearanceService(db_session)

    request = SearchSimilarRequest(trademark_name="ZQ", trademark_id=case.trademark.id)
    response = SearchSimilarResponse(
        query_id=f"srch_{uuid.uuid4().hex[:8]}",
        sources={},
        summary=SearchSummary(high_risk_count=0, medium_risk_count=0, low_risk_count=0, recommendation="proceed"),
    )
    run = service.save_search_run(user=case.owner, request=request, response=response)
    service.export_run_xlsx(user=case.owner, run_id=run.id)
    write_audit_log(
        db_session,
        action="intake.document.added",
        resource_type="intake_request",
        resource_id=case.request.id,
        org_id=case.org_id,
        actor_user_id=case.owner.id,
        after={"filename": "report.pdf"},
    )
    db_session.commit()

    entries = service.activity(org_id=case.org_id, trademark_id=case.trademark.id)
    titles = {entry.event_type: entry.title for entry in entries}
    assert titles["trademark.search_run"] == "Detailed search run saved"
    assert titles["trademark.search_exported"].startswith("Search report exported")
    assert titles["document.added"] == "Document attached: report.pdf"
    assert {entry.actor_name for entry in entries} == {"Owner"}
    assert [e.created_at for e in entries] == sorted((e.created_at for e in entries), reverse=True)
