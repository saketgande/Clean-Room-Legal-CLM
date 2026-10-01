"""The Word editor's save path (app/contract_files/editor.py).

Also guards the 2026-09-30 bug: closing the editor on a generated file (a draft's
Word export) came back "changed", made a version that differed only by the
export's title line, and stranded every pending suggested change on the version
before it (Accept then answered 409). A save whose words match what was opened
is not a version, and the editor is read-only while suggestions are pending.

The editor saves through an unauthenticated callback, so this guards what
stands in for a session: a save must be signed by the editor, carry our own
signed link naming who opened it (not a file link reused), fetch the edited
file only from the editor itself whatever address the callback names, and
turn into exactly one new version — never a duplicate when the editor sends
the same file again on close. Runs the real routes and database; AI jobs off.
"""

import io

import jwt
import pytest
from docx import Document
from fastapi.testclient import TestClient

import app.contract_files.service as files_service
import app.models  # noqa: F401
from app.contract_files import editor
from app.core.config import settings

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _word(text: str) -> bytes:
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_a_signed_save_becomes_the_next_version(monkeypatch):
    from app.contracts.models import Contract
    from app.core.database import SessionLocal, utcnow
    from app.jobs.models import JobRun
    from app.main import app

    # Switched on here so the check runs in CI too, where no editor is configured.
    monkeypatch.setattr(settings, "onlyoffice_url", "http://localhost:8082")
    monkeypatch.setattr(settings, "onlyoffice_jwt_secret", "test-editor-secret-" + "x" * 24)
    monkeypatch.setattr(files_service, "dispatch_job", lambda db, *, job: job)
    fetched: list[str] = []
    reply = {"file": _word("1. Term. This Agreement lasts five (5) years.")}

    class FakeClient:
        def __init__(self, *a, **k): ...
        async def __aenter__(self): return self
        async def __aexit__(self, *a): return False

        async def get(self, url):
            fetched.append(url)
            import httpx
            return httpx.Response(200, content=reply["file"], request=httpx.Request("GET", url))

    monkeypatch.setattr(editor.httpx, "AsyncClient", FakeClient)
    c = TestClient(app)
    login = c.post("/api/v1/auth/login", json={"email": "admin@example.com", "password": "local-dev-password"})
    if login.status_code != 200:
        pytest.skip("no seeded admin in this database")
    h = {"authorization": "Bearer " + login.json()["access_token"]}
    up = c.post("/api/v1/contracts/upload", headers=h, data={"title": "Word editor test"},
                files={"file": ("t.docx", _word("1. Term."), DOCX)})
    cid = up.json()["contract"]["id"]
    try:
        assert c.put(f"/api/v1/contracts/{cid}/text", headers=h,
                     json={"text": "1. Term. This Agreement lasts one (1) year."}).status_code == 200
        cfg = c.get(f"/api/v1/contracts/{cid}/editor/config", headers=h).json()
        assert cfg["enabled"] and cfg["config"]["editorConfig"]["mode"] == "edit"
        save_url = cfg["config"]["editorConfig"]["callbackUrl"].split("/api/v1", 1)[1]
        file_url = cfg["config"]["document"]["url"].split("/api/v1", 1)[1]

        def signed(status: int, secret=None) -> dict:
            body = {"key": "k", "status": status, "url": "http://evil.example/steal?x=1"}
            return {**body, "token": jwt.encode(body, secret or settings.onlyoffice_jwt_secret, algorithm="HS256")}

        versions = lambda: len(c.get(f"/api/v1/contracts/{cid}/versions", headers=h).json())
        before = versions()

        assert c.post("/api/v1" + save_url, json=signed(6, secret="not-the-editor")).status_code == 403
        file_token_as_save = "/api/v1/editor/contract-saved?" + file_url.split("?", 1)[1]
        assert c.post(file_token_as_save, json=signed(6)).status_code == 403
        assert c.post("/api/v1" + save_url, json=signed(1)).json() == {"error": 0}  # someone opened it
        assert versions() == before and not fetched

        # Closed without edits: the same words in a freshly built file (other bytes).
        served = Document(io.BytesIO(c.get("/api/v1" + file_url).content))
        same_words = reply["file"]
        reply["file"] = _word("\n".join(p.text for p in served.paragraphs))
        assert c.post("/api/v1" + save_url, json=signed(2)).json() == {"error": 0}
        assert versions() == before
        reply["file"] = same_words

        assert c.post("/api/v1" + save_url, json=signed(6)).json() == {"error": 0}
        assert fetched and fetched[0].startswith(settings.onlyoffice_internal_url) and "evil" not in fetched[0]
        assert versions() == before + 1
        contract = c.get(f"/api/v1/contracts/{cid}", headers=h).json()
        text = c.get(f"/api/v1/contracts/{cid}/versions/{contract['current_authoritative_version_id']}/text",
                     headers=h).json()["text"]
        assert "five (5) years" in text

        # The editor sends the file again when it closes: not a second version.
        assert c.post("/api/v1" + save_url, json=signed(2)).json() == {"error": 0}
        assert versions() == before + 1

        # A suggested change waiting on the current version: Word opens read-only.
        from app.contract_files.models import ContractEdit

        db = SessionLocal()
        cur = db.get(Contract, cid).current_authoritative_version_id
        db.add(ContractEdit(org_id=db.get(Contract, cid).org_id, contract_id=cid, contract_version_id=cur,
                            edit_type="playbook_redline", status="proposed", original_text="five (5) years",
                            replacement_text="three (3) years"))
        db.commit()
        db.close()
        cfg = c.get(f"/api/v1/contracts/{cid}/editor/config", headers=h).json()
        assert cfg["config"]["editorConfig"]["mode"] == "view" and "waiting" in cfg["read_only_reason"]
    finally:
        db = SessionLocal()
        for job in db.query(JobRun).filter(JobRun.resource_type == "contract", JobRun.resource_id == cid,
                                           JobRun.status == "queued"):
            job.status = "cancelled"
        db.get(Contract, cid).deleted_at = utcnow()
        db.commit()
        db.close()
