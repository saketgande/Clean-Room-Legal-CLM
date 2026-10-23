"""The exam run: batch delivery, capture and replay, the ceiling, resume and abandon.

specs/005-ai-eval-baseline T018 (FR-1, FR-10, FR-11, FR-24, FR-25, FR-30, FR-32, FR-40, FR-41;
AC-1, AC-24, AC-30, AC-32, AC-40, AC-41). The run is driven end to end through
``runner.execute_exam_run`` against a small exam built in the test, with the app's REAL code on
the other side of the Claude boundary: the upload pipeline's read and persist steps, the job
handlers' skills, playbook review, Contract Brain ``ask`` and intake triage. Only the provider is
fake (``tests/exam_fake_anthropic.py``), so nothing here can spend money, and no contract text is
real: every document is generated below.

Two scratch Postgres databases stand in for the two real ones, so the live development data is
never touched: a "live" one holding only the tables the gate writes (audit_log, a_i_call_log,
user, role) and an "exam" one named ``legal_clm_eval_...`` with the whole schema. The two things a
real run does that a test cannot repeat are swapped at the seams ``_rebuild_exam_db`` (drop and
migrate a database) and ``_exam_sessions`` (the process's default engine); both are covered by the
isolation test, which runs a whole exam in a fresh interpreter.
"""

from __future__ import annotations

import asyncio
import functools
import hashlib
import io
import json
import uuid
from contextlib import contextmanager
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from sqlalchemy import create_engine, func, select, text, update
from sqlalchemy.engine import make_url
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import NullPool

import app.models  # noqa: F401  (registers every table, so create_all builds the whole schema)
from app.ai import embeddings
from app.ai.gateway import ledger
from app.ai.models import AISkillRun
from app.analytics.routes import token_cost
from app.contract_brain.models import BrainQuery, ClauseExtraction
from app.contract_brain.retrieval import clause_language_matches
from app.contracts.models import Contract
from app.core.config import settings
from app.core.database import Base
from app.core.models import AICallLog
from app.intake.models import IntakeRequest
from app.integrations import claude
from app.integrations.storage import storage_service
from app.jobs.models import JobRun
from eval.exam import content as content_mod
from eval.exam import evaldb, report, runner, score
from eval.exam import owner as owner_mod
from eval.exam import routing as routing_mod
from eval.exam.content import (
    OWNER_FILES,
    ExamContent,
    ExamDocument,
    ExtractionItem,
    FieldKey,
    OwnerFactKey,
    OwnerManifestEntry,
    PlantedChange,
    PortfolioQuestion,
    QAItem,
    ReviewItem,
    ReviewRuleKey,
    RoutingKey,
    RoutingRequest,
    SourceCategory,
    app_code_sha256,
    content_fingerprint,
)
from eval.exam.env import EVAL_DB_NAME, ExamRefused, LiveDB
from tests import exam_fake_anthropic
from tests.exam_fake_anthropic import batch_result, install
from tests.test_exam_gate import OWNER, SCRATCH_TABLES, Typist, audit_rows, make_live

# --- two scratch databases --------------------------------------------------------------------


@contextmanager
def scratch(name_prefix: str, tables=None):
    base = make_url(settings.database_url)
    name = f"{name_prefix}_{uuid.uuid4().hex[:10]}"
    admin = create_engine(
        base.set(database="postgres"), isolation_level="AUTOCOMMIT", poolclass=NullPool
    )
    try:
        with admin.connect() as conn:
            conn.execute(text(f'CREATE DATABASE "{name}"'))
    except SQLAlchemyError as exc:
        admin.dispose()
        pytest.skip(f"this database user cannot CREATE DATABASE ({type(exc).__name__})")
    engine = create_engine(base.set(database=name), poolclass=NullPool)
    try:
        with engine.begin() as conn:
            conn.execute(text("CREATE EXTENSION IF NOT EXISTS vector"))
            conn.execute(text("CREATE EXTENSION IF NOT EXISTS pg_trgm"))
        Base.metadata.create_all(engine, tables=tables)
        yield engine
    finally:
        engine.dispose()
        with admin.connect() as conn:
            conn.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
        admin.dispose()


@pytest.fixture(scope="module")
def exam_engine():
    with scratch(f"{EVAL_DB_NAME}_batch") as engine:
        yield engine


@pytest.fixture(scope="module")
def live_engine():
    with scratch("legal_clm_scratch", [Base.metadata.tables[t] for t in SCRATCH_TABLES]) as engine:
        yield engine


@pytest.fixture
def live(live_engine) -> LiveDB:
    return make_live(live_engine)


def wipe(engine) -> None:
    """Every table empty again (what ``reset`` + ``migrate`` give a real run)."""
    names = ", ".join(f'"{t.name}"' for t in Base.metadata.sorted_tables)
    with engine.begin() as conn:
        conn.execute(text(f"TRUNCATE TABLE {names} RESTART IDENTITY CASCADE"))


# --- a small exam: documents that are generated here, answers that are made up here ------------

TERMS = (
    "SERVICES AGREEMENT\n\n"
    "This Services Agreement is made on 1 April 2024 between {a} Private Limited, a company "
    "registered in Mumbai, India ({a}) and {b} Pvt. Ltd. ({b}).\n\n"
    "1. Definitions\n"
    "In this Agreement the Services means the services described in Schedule 1.\n\n"
    "2. Term and renewal\n"
    "This Agreement starts on 1 April 2024 and continues for three years. It renews automatically "
    "for successive one year terms unless either party gives ninety days written notice.\n\n"
    "3. Limitation of liability\n"
    "Each party's total liability under this Agreement is limited to the fees paid in the twelve "
    "months before the claim.\n\n"
    "4. Governing law\n"
    "This Agreement is governed by the laws of India and the courts at Mumbai have exclusive "
    "jurisdiction.\n\n"
    "5. Payment\n"
    "The Customer shall pay each invoice within thirty days. Total fees are INR 2,500,000.\n"
)
# Different wording throughout, so a clause of the first contract is "the same wording" as a clause of
# the owner's contract (which is built from TERMS) and of no other.
TERMS_B = (
    "SUPPLY AGREEMENT\n\n"
    "This Supply Agreement is made on 1 June 2025 between {a} Limited ({a}) and {b} Limited ({b}).\n\n"
    "1. Scope\n"
    "The Supplier shall deliver the goods listed in Schedule A to the Buyer's warehouse.\n\n"
    "2. Delivery\n"
    "Goods are delivered within ten working days of an order and risk passes on delivery.\n\n"
    "3. Price\n"
    "The Buyer pays the price in Schedule A within forty five days of each invoice.\n\n"
    "4. Law\n"
    "This Agreement is governed by the laws of Singapore.\n"
)
GOVERNING_LAW = "This Agreement is governed by the laws of India and the courts at Mumbai have exclusive"
PLANTED = "Each party's total liability under this Agreement is unlimited."


def text_of(a: str, b: str, terms: str = TERMS) -> str:
    return terms.format(a=a, b=b)


@functools.cache  # the same words give the same bytes in every world, so content fingerprints agree
def docx_of(body: str) -> bytes:
    document = docx.Document()
    for paragraph in body.split("\n\n"):
        document.add_paragraph(paragraph)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def doc(doc_id, file, data: bytes, **changes) -> ExamDocument:
    values = dict(
        doc_id=doc_id,
        category=SourceCategory.CONSTRUCTED,
        dataset=None,
        file=file,
        mime_type="text/plain",
        file_sha256=sha(data),
        title=Path(file).stem,
        ai_steps=["metadata", "clauses", "brain"],
    )
    values.update(changes)
    return ExamDocument(**values)


PLAYBOOK = {
    "name": "Exam MSA playbook",
    "description": "Rules the exam reviews against.",
    "rules": [
        {
            "clause_type": "limitation_of_liability",
            "rule_type": "standard_position",
            "preferred_position": "Liability is capped at twelve months of fees.",
            "risk_level": "high",
        },
        {
            "clause_type": "governing_law",
            "rule_type": "standard_position",
            "preferred_position": "Indian law.",
            "risk_level": "medium",
        },
    ],
}


def rule_keys(planted: bool) -> list[ReviewRuleKey]:
    change = PlantedChange(
        rule="limitation_of_liability",
        status="deviation",
        original_text=None,
        changed_text=PLANTED,
        anchor_text="3. Limitation of liability\n" + "x" * 40,
        char_start=640,
        severity="high",
        severity_source="Exam MSA playbook / limitation_of_liability",
    )
    return [
        ReviewRuleKey(
            rule="limitation_of_liability",
            status="deviation" if planted else "compliant",
            severity="high" if planted else None,
            planted=change if planted else None,
        ),
        ReviewRuleKey(rule="governing_law", status="compliant"),
    ]


class World(SimpleNamespace):
    """What a test needs: the two databases, the fake provider, the exam content and its files."""


def build_world(
    tmp_path, monkeypatch, exam_engine, live_engine, *, owner_doc=True, routing=False,
    owner_frozen_sha=None,
):
    wipe(exam_engine)
    docs_dir = tmp_path / "contracts"
    (docs_dir / "public").mkdir(parents=True)
    owner_dir = tmp_path / "owner_files"
    owner_dir.mkdir()
    runs_dir = tmp_path / "runs"

    files: dict[str, bytes] = {
        "public/cuad-t1.txt": text_of("Acme", "Beta").encode(),
        "public/cp-t1.txt": text_of("Cora", "Delta").encode(),
        "public/scan-cp-t1.txt": text_of("Cora", "Delta").encode(),
        "public/ae-t1.txt": text_of("Eagle", "Falcon", TERMS_B).encode(),
        "public/pf-t1.txt": text_of("Gamma", "Harbor", TERMS_B).encode(),
    }
    # The scanned copy has its own bytes: the same words with a stamp line in front.
    files["public/scan-cp-t1.txt"] = b"DocuSign Envelope ID: 0123456789ABCDEF\n" + files[
        "public/scan-cp-t1.txt"
    ]
    for name, data in files.items():
        (docs_dir / name).write_bytes(data)

    documents = [
        doc("cuad-t1", "public/cuad-t1.txt", files["public/cuad-t1.txt"], category=SourceCategory.PUBLIC, dataset="cuad"),
        doc("cp-t1", "public/cp-t1.txt", files["public/cp-t1.txt"], ai_steps=["metadata"]),
        doc("scan-cp-t1", "public/scan-cp-t1.txt", files["public/scan-cp-t1.txt"], ai_steps=["metadata"], twin_of="cp-t1"),
        doc("ae-t1", "public/ae-t1.txt", files["public/ae-t1.txt"], uploader="portfolio"),
        doc("pf-t1", "public/pf-t1.txt", files["public/pf-t1.txt"], uploader="portfolio"),
    ]
    owner_keys: list[OwnerFactKey] = []
    manifest: list[OwnerManifestEntry] = []
    if owner_doc:
        name = OWNER_FILES[11]  # "Mutual NDA — Northwind Robotics Inc..docx"
        data = docx_of(text_of("Northwind", "Orchid"))
        (owner_dir / name).write_bytes(data)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        frozen = asyncio.run(
            owner_mod.read_as_app(data, mime_type=mime, filename=name, ocr=owner_mod.FrozenOCR())
        )
        entry = OwnerManifestEntry(
            doc_id="owner-12",
            name=name,
            file_sha256=sha(data),
            mime_type=mime,
            pages=0,
            kind="text",
            frozen_text_sha256=owner_frozen_sha or sha(frozen.text.encode()),
            frozen_method=frozen.method,
            imported_at=datetime.now(UTC),
        )
        manifest.append(entry)
        documents.append(
            ExamDocument(
                doc_id="owner-12",
                category=SourceCategory.OWNER,
                dataset=None,
                file=name,
                mime_type=mime,
                file_sha256=entry.file_sha256,
                title=Path(name).stem,
                ai_steps=["metadata", "clauses"],
            )
        )
        owner_keys = [
            OwnerFactKey(
                doc_id="owner-12",
                field="governing_law",
                status="confirmed",
                value="India",
                quote=GOVERNING_LAW,
                char_start=None,
                char_end=None,
                suggested_value="India",
                suggested_quote=GOVERNING_LAW,
                suggested_by="claude-sonnet-5-5",
                file_sha256=entry.file_sha256,
                frozen_text_sha256=entry.frozen_text_sha256,
                sheet_sha256="s" * 64,
                loaded_at=datetime.now(UTC),
            )
        ]
        monkeypatch.setattr(owner_mod, "read_manifest", lambda: manifest)
    monkeypatch.setattr(owner_mod, "stale_keys", lambda keys, manifest=None: [])

    review = [
        ReviewItem(
            item_id="rv-cp-t1", doc_id="cp-t1", category=SourceCategory.CONSTRUCTED,
            source="constructed", playbook="cp_msa", rules=rule_keys(planted=True),
        ),
        ReviewItem(
            item_id="rv-ae-t1", doc_id="ae-t1", category=SourceCategory.CONSTRUCTED,
            source="constructed", playbook="cp_msa", control=True, rules=rule_keys(planted=False),
        ),
    ]
    extraction = [
        ExtractionItem(
            item_id="ex-cuad-t1", doc_id="cuad-t1", category=SourceCategory.PUBLIC, source="CUAD",
            fields=[
                FieldKey(field="governing_law", kind="jurisdiction", value="India", quote=GOVERNING_LAW),
                FieldKey(field="effective_date", kind="date", value="2024-04-01", quote=None),
            ],
        ),
        ExtractionItem(
            item_id="ex-cp-t1", doc_id="cp-t1", category=SourceCategory.CONSTRUCTED, source="constructed",
            fields=[FieldKey(field="governing_law", kind="jurisdiction", value="India", quote=GOVERNING_LAW)],
        ),
    ]
    qa = [
        QAItem(
            item_id="qa-001", doc_id="cuad-t1", category=SourceCategory.PUBLIC, source="LegalBench-RAG",
            question="Which law governs the Agreement and which courts have jurisdiction?",
            answer="Indian law; the Mumbai courts.", passages=[GOVERNING_LAW],
        ),
        QAItem(
            item_id="qa-002", doc_id="cuad-t1", category=SourceCategory.PUBLIC, source="LegalBench-RAG",
            question="How many contracts are in my portfolio?",
            answer="three", passages=["unused"],
        ),
    ]
    portfolio = [
        PortfolioQuestion(
            item_id="pq-01", question="Which contracts renew automatically for one year terms?",
            fact="auto_renew", expected_doc_ids=["ae-t1", "pf-t1"],
        )
    ]
    requests: list[RoutingRequest] = []
    routing_keys: list[RoutingKey] = []
    snapshot = None
    if routing:
        workflow = {
            "id": str(uuid.uuid4()), "name": "NDA review", "description": "Reviews NDAs",
            "enabled": True, "is_builtin": False, "eval_order": 1, "version": 1,
            "criteria": {}, "steps": [],
        }
        fingerprint = routing_mod.workflow_fingerprint([workflow])
        snapshot = {"fingerprint": fingerprint, "taken_at": "2026-10-07T00:00:00Z", "workflows": [workflow]}
        requests = [
            RoutingRequest(
                item_id="rt-01", channel="email", subject="Need an NDA", body="Please send us an NDA.",
                drafted_for_workflow_id=workflow["id"], drafting_model="claude-sonnet-5-5",
            )
        ]
        routing_keys = [
            RoutingKey(
                item_id="rt-01", status="confirmed", workflow_id=workflow["id"],
                workflow_fingerprint=fingerprint, sheet_sha256="r" * 64, loaded_at=datetime.now(UTC),
            )
        ]
        monkeypatch.setattr(routing_mod, "routing_fresh", lambda live, snap: True)

    content = ExamContent(
        documents=documents, review=review, extraction=extraction, qa=qa, portfolio=portfolio,
        routing=requests, owner_keys=owner_keys, routing_keys=routing_keys,
        playbooks={"cp_msa": PLAYBOOK},
        sample={
            "review": ["rv-cp-t1", "rv-scan-cp-t1"], "extraction": ["ex-cp-t1"], "qa": ["qa-001"],
            "portfolio": ["pq-01"], "routing": ["rt-01"], "portfolio_docs": ["ae-t1"],
        },
        workflow_snapshot=snapshot, owner_manifest=manifest,
    )

    exam_sessions = sessionmaker(bind=exam_engine, autoflush=False, expire_on_commit=False)
    monkeypatch.setattr(runner, "_exam_sessions", lambda: exam_sessions)
    monkeypatch.setattr(runner, "_rebuild_exam_db", lambda live, name: wipe(exam_engine))
    monkeypatch.setattr(runner, "RUNS_DIR", runs_dir)
    monkeypatch.setattr(report, "RUNS_DIR", runs_dir)
    monkeypatch.setattr(runner, "POLL_SECONDS", 0)
    for module in (evaldb, score):
        monkeypatch.setattr(module, "DOCS_DIR", docs_dir)
    monkeypatch.setattr(evaldb, "WORK_DIR", tmp_path)
    monkeypatch.setattr(evaldb, "_OWNER_FILES_DIR", owner_dir)
    monkeypatch.setattr(storage_service, "root", (tmp_path / "storage").resolve())
    monkeypatch.setattr(ledger, "_default_sessions", lambda: exam_sessions())
    # The job handler opens its own session on the process's default engine, which a real exam
    # run has already moved onto the exam database (env.isolate); here it is the scratch one.
    monkeypatch.setattr(runner.job_tasks, "SessionLocal", exam_sessions)
    # Embeddings are not what is under test here (the isolation test uses the real model): fixed
    # vectors keep every run fast and independent of the model being installed.
    monkeypatch.setattr(
        embeddings, "_embed_local", lambda texts: [embeddings._deterministic_mock_vector(t) for t in texts])
    fake = install(monkeypatch)
    status = fake.messages.batches.retrieve

    async def retrieve_with_a_datetime(batch_id):
        # The SDK's to_dict() leaves timestamps as datetimes, which a state file cannot hold.
        reply = await status(batch_id)
        body = {**reply.to_dict(), "ended_at": datetime.now(UTC)}
        return SimpleNamespace(to_dict=lambda: body)

    fake.messages.batches.retrieve = retrieve_with_a_datetime
    live = make_live(live_engine)
    return World(
        content=content, fake=fake, live=live, exam=exam_sessions, runs=runs_dir, tmp=tmp_path,
        engine=exam_engine, docs=docs_dir, manifest=manifest,
    )


@pytest.fixture
def world(tmp_path, monkeypatch, exam_engine, live_engine):
    return build_world(tmp_path, monkeypatch, exam_engine, live_engine)


def run_exam(world, *, mode="full", delivery="batch", ceiling="25", change=None, target=None, typist=None):
    typist = typist or Typist(ceiling)
    return asyncio.run(
        runner.execute_exam_run(
            live=world.live, owner=OWNER, mode=mode, delivery=delivery, change=change,
            target=target, ask=typist, content=world.content,
        )
    )


def finished(world):
    [row] = audit_rows(world.live, "eval.run_finished")
    return row


def live_ledger(world, **filters):
    with world.live.sessions() as db:
        return list(db.scalars(select(AICallLog).where(AICallLog.org_id == world.live.org_id)))




class Crash(BaseException):
    """A process that dies (power cut, kill -9): not an Exception, so nothing in the app catches it."""


def run_id_of(world) -> str:
    [row] = audit_rows(world.live, "eval.spend_approved")
    return row.resource_id


def state_of(world, run_id=None) -> dict:
    return json.loads((world.runs / (run_id or run_id_of(world)) / "state.json").read_text())


def resume(world, *, answers="y", **kwargs):
    return asyncio.run(
        runner.resume_exam_run(
            live=world.live, owner=OWNER, run_id=kwargs.pop("run_id", None) or run_id_of(world),
            ask=lambda prompt: answers, content=kwargs.pop("content", world.content),
        )
    )


def exam_rows(world, model, *where):
    with world.exam() as db:
        return list(db.scalars(select(model).where(*where)))


# =============================================================================================
# The run, end to end -- done once, read many times
# =============================================================================================


class Spy(Typist):
    """Records how much had reached the provider at the moment each question was asked."""

    def __init__(self, world, **kwargs):
        super().__init__(**kwargs)
        self.world, self.sent = world, []

    def __call__(self, prompt):
        self.sent.append(len(self.world.fake.streamed) + len(self.world.fake.batch_creates))
        return super().__call__(prompt)


@pytest.fixture(scope="module")
def happy(tmp_path_factory, exam_engine, live_engine):
    """One complete batch run of the small exam, and everything the tests below read back from
    it: the record, what reached the provider, both ledgers and the audit trail. Computed inside
    a patch context that is closed before the tests run, so nothing leaks between tests."""
    with pytest.MonkeyPatch.context() as mp:
        world = build_world(tmp_path_factory.mktemp("happy"), mp, exam_engine, live_engine)
        spy = Spy(world)
        results = run_exam(world, typist=spy)
        run_id = results["run_id"]
        with world.exam() as db:
            exam_ledger = db.scalar(select(func.count()).select_from(AICallLog))
            left_running = db.scalar(
                select(func.count()).select_from(AISkillRun).where(AISkillRun.status == "running")
            )
            contracts = {c.title: c.id for c in db.scalars(select(Contract))}
            queries = list(db.scalars(select(BrainQuery)))
        return SimpleNamespace(
            world=world,
            spy=spy,
            results=results,
            scores=score.score_run(results, world.content),
            batch_creates=world.fake.batch_creates,
            streamed=world.fake.streamed,
            live_rows=live_ledger(world),
            approved=audit_rows(world.live, "eval.spend_approved")[0],
            finished=finished(world),
            exam_ledger=exam_ledger,
            left_running=left_running,
            contracts=contracts,
            queries=queries,
            results_path=world.runs / run_id / "results.json",
            run_id=run_id,
        )


def test_a_full_batch_run_completes_and_every_item_is_scored(happy):
    """AC-1: the whole small exam runs, every planned item is answered and scored, and the items
    come in the order the content fixes (review, extraction, answers, portfolio), the same every run."""
    r = happy.results
    assert r["outcome"] == "complete" and r["incomplete_reason"] is None
    assert r["mode"] == "full" and r["delivery"] == "batch" and r["kind"] == "exam"
    assert r["contains_errors"] is False and r["failed_call_rate"] == 0.0
    assert [(i["item_id"], i["status"]) for i in r["items"]] == [
        ("rv-cp-t1", "scored"), ("rv-ae-t1", "scored"), ("rv-scan-cp-t1", "scored"),
        ("ex-cuad-t1", "scored"), ("ex-cp-t1", "scored"), ("ex-scan-cp-t1", "scored"),
        ("ex-owner-12", "scored"), ("qa-001", "scored"), ("qa-002", "scored"), ("pq-01", "scored"),
    ]
    assert [d["doc_id"] for d in r["documents"]] == [
        "cuad-t1", "cp-t1", "scan-cp-t1", "ae-t1", "pf-t1", "owner-12"]
    assert all(d["status"] == "ok" for d in r["documents"])


def test_public_and_constructed_calls_go_out_as_one_batch_per_phase_and_the_owners_do_not(happy):
    """FR-41/AC-41: two phases, one batch each, carrying exactly the requests the app built (the
    custom id is the hash of the params); the owner's contract is called normally and its text
    appears in no batch."""
    assert [len(b) for b in happy.batch_creates] == [8, 5]
    for batch in happy.batch_creates:
        for request in batch:
            assert request["custom_id"] == runner.custom_id(request["params"])
            assert "stream" not in request["params"]  # a valid non-streaming batch request as is
    assert [b["phase"] for b in happy.results["batches"]] == [1, 2]
    assert all(b["succeeded"] == b["requests"] for b in happy.results["batches"])
    assert all(isinstance(b["ended_at"], str) for b in happy.results["batches"])
    # the owner's labelling and metadata, and the portfolio question's filter and search answer
    assert len(happy.streamed) == 4
    assert "Northwind" in json.dumps(happy.streamed)
    assert "Northwind" not in json.dumps(happy.batch_creates)
    by_doc = {c["doc_id"]: c["batched"] for c in happy.results["calls"] if c["doc_id"]}
    assert by_doc["owner-12"] is False
    assert all(batched for doc_id, batched in by_doc.items() if doc_id != "owner-12")


def test_the_second_phase_is_sent_after_the_first_is_stored(happy):
    """Review and answers read what phase 1 stored, so they are captured and sent only afterwards:
    phase 2 holds the three reviews and the two contract questions. The portfolio question is not
    in it -- a "which contracts" question makes a filter call and then a search (or table reads),
    more calls than a batch can replay for one unit -- so its calls go out normally. And since
    Step 3e only a PORTFOLIO "how many contracts" is answered from the database (the filter, no
    call): the same words asked about one contract are a search like any other."""
    kinds = [r["params"]["tool_choice"]["name"] for r in happy.batch_creates[1]]
    assert sorted(kinds) == ["return_contract_brain_answer"] * 2 + ["return_playbook_review"] * 3
    qa2 = next(i for i in happy.results["items"] if i["item_id"] == "qa-002")
    assert qa2["prediction"]["answer_mode"] == "rag"
    assert any(c["call_type"] == "contract_brain_answer" and c["item_ids"] == ["qa-002"]
               for c in happy.results["calls"])
    asked = calls_of(happy.results, "pq-01")
    assert sorted(c["call_type"] for c in asked) == ["contract_brain_answer", "portfolio_filter"]
    assert not any(c["batched"] for c in asked)


def test_capture_leaves_no_trace_in_the_ledger_and_every_call_is_priced_at_what_was_paid(happy):
    """Capturing a request aborts the app's call before it records anything, so the exam ledger
    holds exactly the calls that were answered; batch rows cost half what the same tokens cost
    in a normal call, and the batch's recorded spend is the sum of its calls."""
    calls = happy.results["calls"]
    assert happy.exam_ledger == len(calls) == 17
    batched = [c for c in calls if c["batched"]]
    owner_call = next(c for c in calls if not c["batched"])
    assert owner_call["cost_usd"] == pytest.approx(
        token_cost(100, 20, "claude-sonnet-4-6"), abs=1e-6)
    assert all(c["cost_usd"] == pytest.approx(owner_call["cost_usd"] / 2, abs=1e-6) for c in batched)
    phase_1 = sum(c["cost_usd"] for c in batched if c["call_type"] != "playbook_review"
                  and c["call_type"] != "contract_brain_answer")
    assert happy.results["batches"][0]["usd"] == pytest.approx(phase_1, abs=1e-5)
    assert happy.finished.after["actual"]["usd"] == pytest.approx(
        sum(c["cost_usd"] for c in calls), abs=1e-5)
    assert happy.left_running > 0  # aborted skill runs, harmless and wiped with the database


def test_the_live_ledger_holds_exam_traffic_only_marked_and_without_text(happy):
    """FR-34/FR-38/FR-41: every call is copied to the live ledger stamped with the run and
    flagged batched or not; text and exam-database ids stay behind."""
    assert len(happy.live_rows) == 17
    assert {r.eval_run_id for r in happy.live_rows} == {happy.run_id}
    assert sum(r.batched for r in happy.live_rows) == 13  # 17 calls; the owner's 2 and pq-01's 2 went normally
    assert all(r.provider_request_id.startswith("msgbatch_fake") for r in happy.live_rows if r.batched)
    assert not any(r.provider_request_id.startswith("msgbatch") for r in happy.live_rows if not r.batched)
    assert all(r.input_payload is None and r.validated_output is None and r.job_id is None
               for r in happy.live_rows)


def test_the_run_record_is_immutable_pinned_to_its_results_and_the_results_are_read_only(happy):
    """FR-39/AC-39: one audit row says who, when, mode, versions, estimate, ceiling, outcome and
    cost; its results hash is the file's; the file cannot be written."""
    after = happy.finished.after
    assert after["kind"] == "exam" and after["mode"] == "full" and after["delivery"] == "batch"
    assert after["outcome"] == "complete" and after["ceiling_usd"] == 25.0
    assert after["approved_by"]["email"] == OWNER.email and after["started_by"]["email"] == OWNER.email
    assert after["versions"]["content_fingerprint"] == content_fingerprint(happy.world.content)
    assert after["versions"]["app_code_sha256"] == app_code_sha256()
    assert after["versions"]["model_under_test"] == "claude-sonnet-4-6"
    assert after["versions"]["models"] == ["claude-sonnet-4-6"]
    assert {p["prompt_key"] for p in after["versions"]["prompts"]} == {
        "clause_labeling", "contract_metadata_extraction", "playbook_review", "contract_brain_answer",
        "portfolio_filter"}
    assert set(after["headline"]) >= {"review_f1", "extraction_accuracy", "qa_evidence_found"}
    assert after["results_sha256"] == hashlib.sha256(happy.results_path.read_bytes()).hexdigest()
    assert happy.results_path.stat().st_mode & 0o222 == 0
    assert json.loads(happy.results_path.read_text())["run_id"] == happy.run_id
    assert (happy.results_path.parent / "detail.md").is_file()
    assert [b["phase"] for b in after["batches"]] == [1, 2]


def test_the_estimate_shown_before_the_go_ahead_bounds_what_was_spent(happy):
    """FR-31/FR-33: the estimate was computed from the real requests before any money moved; its
    worst case is a true upper bound on the actual cost. The portfolio question is priced as a
    filter call plus one table read per portfolio contract (its fact, renewal, is wording no
    contract field answers), all at full price. This run made no reads -- the fake provider says
    the question is not a list question, so the app searched -- which is why the estimate has
    two reads more than the ledger: it prices what the question is expected to do."""
    assert happy.spy.sent == [0]  # the ceiling (the go-ahead) was asked with nothing sent
    assert happy.spy.prompts[0].startswith("Spending ceiling for run")
    estimate = happy.approved.after["estimate"]
    assert estimate["worst_case_usd"] >= happy.finished.after["actual"]["usd"] > 0
    assert estimate["usd"] > 0 and estimate["list_usd"] > estimate["usd"]
    assert (estimate["calls"], estimate["batch_calls"], estimate["sync_calls"]) == (18, 13, 5)
    assert set(estimate["by_call_type"]) == {
        "clause_labeling", "contract_metadata_extraction", "playbook_review", "contract_brain_answer",
        "portfolio_filter", "tabular_cell_extraction"}
    assert estimate["by_call_type"]["tabular_cell_extraction"]["calls"] == 2  # ae-t1 and pf-t1
    assert happy.approved.after["ceiling_usd"] == 25.0
    assert happy.approved.after["versions_at_start"]["content_fingerprint"] == (
        happy.finished.after["versions"]["content_fingerprint"])


def test_answers_come_from_the_apps_own_search_over_the_exam_database_only(happy):
    """FR-24/FR-37/AC-24: the question is answered through the app's own ask, with the context its
    search returned (never empty) from the exam contracts, and nothing else."""
    qa1 = next(i for i in happy.results["items"] if i["item_id"] == "qa-001")["prediction"]
    assert qa1["context_chars"] > 0 and qa1["answer_mode"] == "rag"
    assert qa1["sources"] and GOVERNING_LAW in " ".join(s["text"] for s in qa1["sources"])
    cuad = happy.contracts["cuad-t1"]
    assert {s["contract_id"] for s in qa1["sources"] if s["contract_id"]} == {cuad}
    portfolio = next(i for i in happy.results["items"] if i["item_id"] == "pq-01")["prediction"]
    assert portfolio["context_chars"] > 0
    assert len(happy.queries) == 3  # the three questions, all in the exam database
    seen = {c for q in happy.queries for c in (q.retrieval_metadata or {}).get("contract_ids", [])}
    assert seen <= set(happy.contracts.values())
    portfolio_users = {happy.contracts["ae-t1"], happy.contracts["pf-t1"]}
    ids = next(q for q in happy.queries if q.query_scope == "portfolio").retrieval_metadata["contract_ids"]
    assert set(ids) == portfolio_users  # the portfolio user sees what it uploaded, nothing else


def test_each_call_reports_how_much_of_the_document_it_saw(happy):
    """FR-10/AC-10: review, extraction and labelling calls carry characters seen out of total, and
    the scorer turns them into the average; labelling sends excerpts, so it sees less than all."""
    seen = [c for c in happy.results["calls"] if c["chars_total"]]
    assert {c["call_type"] for c in seen} == {
        "clause_labeling", "contract_metadata_extraction", "playbook_review"}
    assert all(0 < c["chars_seen"] <= c["chars_total"] for c in seen)
    labelling = next(c for c in seen if c["call_type"] == "clause_labeling")
    assert labelling["chars_seen"] < labelling["chars_total"]
    coverage = happy.scores["measures"]["coverage_avg"]
    assert coverage["n"] == len(seen) and 0 < coverage["value"] <= 1


def test_the_record_is_what_the_scorer_reads(happy):
    """The results file satisfies the scorer's contract: every measure it computes has an n, the
    run is current, and nothing was left unscored."""
    run = happy.scores["run"]
    assert run["content_current"] is True and run["delivery"] == "batch"
    assert happy.scores["failed_items"] == [] and happy.scores["failed_calls"] == []
    assert happy.scores["measures"]["truncation_rate"]["value"] == 0.0
    assert happy.scores["measures"]["qa_search_found"]["n"] == 2
    # a question about one contract is always a search (only a portfolio-scope count is answered
    # from the database), so both of these had a search context to show
    assert happy.scores["counts"]["qa"]["empty_context"] == 0


# =============================================================================================
# Before the go-ahead: nothing is sent, and what cannot run is refused
# =============================================================================================


def test_nothing_is_sent_before_the_go_ahead_and_a_refusal_spends_nothing(world):
    """FR-31/AC-31: by the time the owner is asked, the documents are read, the requests recorded
    and priced -- and not one has left the machine. A refusal ends the run with an audit row and
    no paid call, no run record and no meter left behind."""
    spy = Spy(world, approve=False)
    with pytest.raises(ExamRefused, match="go-ahead was not given"):
        run_exam(world, typist=spy)
    assert spy.sent == [0]  # asked exactly once, with nothing sent
    assert world.fake.streamed == [] and world.fake.batch_creates == []
    assert audit_rows(world.live, "eval.spend_approved") == []
    assert len(audit_rows(world.live, "eval.spend_refused")) == 1
    assert live_ledger(world) == [] and claude.exam_meter is None
    assert not world.runs.exists() or not list(world.runs.iterdir())


def test_a_planned_scan_that_ocr_has_not_read_refuses_before_any_money_is_asked(world, monkeypatch):
    """FR-20: a scan with no frozen OCR cannot be read at all and the app would not say so, so the
    run is refused up front -- before an estimate, a question, or a database rebuild."""
    item = owner_mod.OCRItem(
        sha="x" * 64, doc_id="owner-12", name="scan.pdf", path=Path("scan.pdf"),
        mime_type="application/pdf", pages=3, owner=True, forced=False,
    )

    async def missing(content, manifest):
        return [item]

    monkeypatch.setattr(owner_mod, "plan_ocr", missing)
    typist = Typist()
    with pytest.raises(ExamRefused, match=r"owner-12.*Run `ocr` first"):
        run_exam(world, typist=typist)
    assert typist.prompts == [] and world.fake.streamed == []


def test_reading_the_documents_cannot_make_a_paid_call_even_if_the_code_that_reads_them_changes(
    world, monkeypatch
):
    """FR-31: "nothing is spent before the go-ahead" must not rest on today's ingestion code
    happening to make no AI call. Reading runs under a capture-mode meter, so a call made there
    stops the run before it leaves the machine -- unmetered and unapproved -- and no meter or
    approval is left behind."""

    async def reads_then_asks(db, document, *, actors, ocr):
        await claude.claude_client._post_messages(
            org_id=None,
            json_payload={
                "model": settings.claude_model, "max_tokens": 10,
                "messages": [{"role": "user", "content": "hello"}],
            },
        )

    monkeypatch.setattr(evaldb, "ingest_document", reads_then_asks)
    typist = Typist()
    with pytest.raises(ExamRefused, match="tried to make a paid AI call"):
        run_exam(world, typist=typist)
    assert world.fake.streamed == [] and world.fake.batch_creates == [] and typist.prompts == []
    assert claude.exam_meter is None and audit_rows(world.live, "eval.spend_approved") == []


@pytest.mark.parametrize(
    "change, target, message",
    [("x", None, "go together"), (None, "review", "go together"), ("x", "weather", "--target is one of")],
)
def test_a_change_and_its_target_area_go_together_and_are_checked_before_anything_runs(
    world, change, target, message
):
    """FR-28: the targeted area is named before the change's runs; half a pair or an unknown area
    is refused before any work is done."""
    typist = Typist()
    with pytest.raises(ExamRefused, match=message):
        run_exam(world, change=change, target=target, typist=typist)
    assert typist.prompts == [] and audit_rows(world.live) == []


def test_a_sample_run_asks_a_fixed_subset_records_its_target_and_is_not_a_baseline(world, happy):
    """FR-30/AC-30: sample mode runs the fixed subset (and the owner's ordinary contract), costs
    less than the full run it is a sample of, says so in its report, and can never form a
    baseline. The change and its target are on record before any call was made (FR-28)."""
    results = run_exam(world, mode="sample", change="whole-document review", target="review")
    assert [i["item_id"] for i in results["items"]] == [
        "rv-cp-t1", "rv-scan-cp-t1", "ex-cp-t1", "ex-scan-cp-t1", "ex-owner-12", "qa-001", "pq-01"]
    assert {d["doc_id"] for d in results["documents"]} == {
        "cp-t1", "scan-cp-t1", "cuad-t1", "ae-t1", "owner-12"}
    assert results["mode"] == "sample" and results["outcome"] == "complete"
    [approved] = audit_rows(world.live, "eval.spend_approved")
    assert approved.after["change"] == "whole-document review"
    assert approved.after["target_area"] == "review" and approved.after["main_measure"] == "review_f1"
    assert results["change"] == "whole-document review" and results["target_area"] == "review"
    assert approved.after["estimate"]["usd"] < happy.approved.after["estimate"]["usd"]
    assert "sample" in report.run_flags(results)
    detail = (world.runs / results["run_id"] / "detail.md").read_text()
    assert report.SAMPLE_LABEL in detail
    problems = report.check_pair({**finished(world).after, "run_id": results["run_id"]}, None)
    assert any("only full runs count" in p for p in problems)


def test_two_runs_on_unchanged_content_ask_the_same_questions_about_the_same_documents(world, happy):
    """FR-1/AC-1: the item list and the document list are fixed by the content, in the same
    order, every run -- and the content fingerprint, which a baseline pair must share, is the same."""
    again = run_exam(world)
    assert [(i["item_id"], i["task"], i["doc_id"]) for i in again["items"]] == [
        (i["item_id"], i["task"], i["doc_id"]) for i in happy.results["items"]]
    assert [d["doc_id"] for d in again["documents"]] == [d["doc_id"] for d in happy.results["documents"]]
    assert again["versions"]["content_fingerprint"] == happy.results["versions"]["content_fingerprint"]
    assert again["estimate"]["calls"] == happy.results["estimate"]["calls"]
    assert again["run_id"] != happy.results["run_id"] and again["outcome"] == "complete"


# =============================================================================================
# Delivery modes
# =============================================================================================


def test_sync_delivery_sends_every_call_normally(world):
    """FR-41: under sync delivery nothing is batched -- every call is a normal request, so the
    estimate, the ledger and the record all say so."""
    results = run_exam(world, delivery="sync")
    assert results["outcome"] == "complete" and results["delivery"] == "sync"
    assert world.fake.batch_creates == [] and results["batches"] == []
    assert len(world.fake.streamed) == len(results["calls"]) == 17
    assert not any(c["batched"] for c in results["calls"])
    assert all(c["latency_ms"] is not None for c in results["calls"])
    [approved] = audit_rows(world.live, "eval.spend_approved")
    assert approved.after["delivery"] == "sync" and approved.after["estimate"]["batch_calls"] == 0
    assert not any(r.batched for r in live_ledger(world))


# =============================================================================================
# Failures: unreturned, errored, changed -- all counted, never scored as answers
# =============================================================================================


def calls_of(results, unit_id):
    """The calls a portfolio unit made itself -- not the phase-1 calls on the documents it reads,
    which carry its item id too."""
    return [c for c in results["calls"] if c["item_ids"] == [unit_id] and c["doc_id"] is None]


def failed_calls_of(world):
    return exam_rows(world, AICallLog, AICallLog.status == "failed")


def test_requests_that_never_come_back_are_failed_calls_and_their_items_are_not_scored(world):
    """FR-40/FR-41/AC-40/AC-41: when the provider's window ends and requests expired, each is a
    failed call (listed, flagged "contains errors"), its item is failed -- never scored as if it
    were an answer -- and the owner's synchronous contract is unaffected."""
    world.fake.default_outcome = "expired"
    results = run_exam(world)
    items = {i["item_id"]: i for i in results["items"]}
    assert results["contains_errors"] is True and results["outcome"] == "complete"
    assert results["failed_call_rate"] == pytest.approx(13 / 17)
    assert items["ex-owner-12"]["status"] == "scored"
    for item_id in ("rv-cp-t1", "rv-ae-t1", "qa-001"):
        assert items[item_id]["status"] == "failed" and items[item_id]["prediction"] is None
        assert items[item_id]["failure"]["reason"].startswith("call failed")
    # the portfolio question's own calls were ordinary ones and went through; what it read did not
    assert items["pq-01"]["status"] == "failed" and items["pq-01"]["prediction"] is None
    assert items["pq-01"]["failure"]["reason"].startswith("upstream: ")
    assert items["ex-cuad-t1"]["failure"]["reason"].startswith("upstream: ")
    assert items["qa-002"]["status"] == "failed"  # it reads a document whose calls failed
    rows = failed_calls_of(world)
    assert len(rows) == 13 and {r.error_class for r in rows} == {"BatchCallFailed"}
    assert all(c["cost_usd"] == 0 for c in results["calls"] if c["status"] != "succeeded")
    assert finished(world).after["failed_call_rate"] == pytest.approx(13 / 17)
    assert finished(world).after["contains_errors"] is True
    scores = score.score_run(results, world.content)
    assert scores["measures"]["review_f1"]["value"] is None  # nothing was scored as an answer
    assert len(scores["failed_calls"]) == 13


def poison(world, error_type):
    """Make the first request of the first batch end in ``error_type``; the retry batch (the
    second create) answers normally."""
    fake = world.fake
    create = fake.messages.batches.create

    async def create_and_poison(*, requests):
        if len(fake.batch_creates) == 1:
            fake.outcomes.clear()  # the second batch: whatever it is, it answers
        batch = await create(requests=requests)
        if len(fake.batch_creates) == 1:
            fake.outcomes[requests[0]["custom_id"]] = batch_result("errored", error_type=error_type)
        return batch

    fake.messages.batches.create = create_and_poison


def test_a_transient_error_is_retried_once_in_its_own_batch(world):
    """The synchronous path retries overloads, so a batch request that ended in one is resent
    ONCE, alone, and its answer replayed like any other (FR-40: no failed call if the retry
    succeeded)."""
    poison(world, "overloaded_error")
    results = run_exam(world)
    first, retry = world.fake.batch_creates[0], world.fake.batch_creates[1]
    assert len(world.fake.batch_creates) == 3  # phase 1, its retry, phase 2
    assert [r["custom_id"] for r in retry] == [first[0]["custom_id"]]
    assert results["outcome"] == "complete" and results["failed_call_rate"] == 0.0
    phase_1 = [b for b in results["batches"] if b["phase"] == 1]
    assert [b["retry_of"] for b in phase_1] == [None, phase_1[0]["batch_id"]]
    assert phase_1[0]["errored"] == 1 and phase_1[1]["succeeded"] == 1
    assert any(r.provider_request_id.startswith(f"{phase_1[1]['batch_id']}:") for r in live_ledger(world))


def test_an_error_that_a_second_try_cannot_cure_is_not_retried(world):
    """A rejected request would be rejected again (and billed for the trouble): only transient
    errors get the retry batch; this one is a failed call and its items fail."""
    poison(world, "invalid_request_error")
    results = run_exam(world)
    assert len(world.fake.batch_creates) == 2
    assert results["failed_call_rate"] == pytest.approx(1 / 17)
    items = {i["item_id"]: i for i in results["items"]}
    assert items["ex-cuad-t1"]["status"] == "failed" and items["qa-001"]["status"] == "failed"
    assert items["rv-cp-t1"]["status"] == "scored"


def test_an_answer_cut_off_at_max_tokens_is_a_truncated_call_and_it_is_not_scored(world):
    """FR-11/AC-11: a batch answer that stopped at max_tokens reaches the record with that stop
    reason, so the report can count it as truncated. The app refuses an unfinished answer (a
    validation failure), so its item is failed -- never scored on half an answer -- and, as the app
    records it, it is also one of the failed calls the 1% baseline limit (D-13) counts.
    OPEN QUESTION FOR THE OWNER: should a truncated answer count toward that limit? The exam is
    expected to find truncation, and today's system may truncate more than 1% of its calls."""
    fake = world.fake
    create = fake.messages.batches.create

    async def create_and_cut_off(*, requests):
        batch = await create(requests=requests)
        if len(fake.batch_creates) == 1:
            first = requests[0]
            cut = {**exam_fake_anthropic.message_for(first["params"], fake.usage), "stop_reason": "max_tokens"}
            fake.outcomes[first["custom_id"]] = {"type": "succeeded", "message": cut}
        return batch

    fake.messages.batches.create = create_and_cut_off
    results = run_exam(world)
    [cut] = [c for c in results["calls"] if c["stop_reason"] == "max_tokens"]
    assert cut["status"] == "validation_failed" and cut["batched"] is True
    items = {i["item_id"]: i for i in results["items"]}
    assert cut["item_ids"] and all(
        items[i]["status"] == "failed" and items[i]["prediction"] is None for i in cut["item_ids"])
    scores = score.score_run(results, world.content)
    assert scores["measures"]["truncation_rate"]["value"] == pytest.approx(1 / 17)
    assert results["failed_call_rate"] == pytest.approx(1 / 17) and results["contains_errors"] is True
    assert results["outcome"] == "complete"


def test_a_request_whose_prompt_changed_since_capture_is_never_answered_with_another_result(world):
    """FR-41: the batch answer belongs to the exact request that was sent. If the prompt the app
    builds at replay differs (here: the contract's title, which every prompt carries), there is no
    answer for it -- a failed call, not somebody else's result."""
    fake = world.fake
    create = fake.messages.batches.create

    async def create_then_rename(*, requests):
        batch = await create(requests=requests)
        if len(fake.batch_creates) == 2:  # phase 2 is out: change a title phase 1 stored
            with world.exam() as db:
                db.execute(update(Contract).where(Contract.title == "cp-t1").values(title="renamed"))
                db.commit()
        return batch

    fake.messages.batches.create = create_then_rename
    results = run_exam(world)
    items = {i["item_id"]: i for i in results["items"]}
    assert items["rv-cp-t1"]["status"] == "failed"
    assert items["rv-cp-t1"]["failure"]["reason"].startswith("call failed")
    assert items["rv-ae-t1"]["status"] == "scored" and items["rv-scan-cp-t1"]["status"] == "scored"
    [row] = failed_calls_of(world)
    assert row.error_class == "BatchCallFailed" and row.request_id == "rv-cp-t1"
    assert results["failed_call_rate"] == pytest.approx(1 / 17)


def test_a_unit_that_raises_fails_its_item_and_nothing_else(world, monkeypatch):
    """A bug or a bad input in one unit is that item's failure, listed with its reason; the rest of
    the run is unaffected and the run still completes (FR-40)."""
    real = runner.ask_contract_brain

    async def explode(*, payload, request, db, current_user):
        if request.state.request_id == "qa-001":
            raise ValueError("the question could not be put")
        return await real(payload=payload, request=request, db=db, current_user=current_user)

    monkeypatch.setattr(runner, "ask_contract_brain", explode)
    results = run_exam(world)
    items = {i["item_id"]: i for i in results["items"]}
    assert items["qa-001"]["status"] == "failed"
    assert items["qa-001"]["failure"]["reason"] == "unit failed: ValueError: the question could not be put"
    assert items["pq-01"]["status"] == "scored" and results["outcome"] == "complete"
    assert results["contains_errors"] is True and results["failed_call_rate"] == 0.0


def test_a_document_whose_graph_cannot_be_built_fails_the_questions_that_search_it(world, monkeypatch):
    """A search over a knowledge graph that lacks a document is not the exam the app takes, so the
    questions that depend on it fail -- unasked, spending nothing -- while what does not read the
    graph (its extraction, its review) is scored as usual; and the run is flagged."""
    real = runner.ingest_contract_brain

    def broken_for_cuad(db, **kwargs):
        if kwargs["contract"].title == "cuad-t1":
            raise RuntimeError("the graph broke")
        return real(db, **kwargs)

    monkeypatch.setattr(runner, "ingest_contract_brain", broken_for_cuad)
    results = run_exam(world)
    items = {i["item_id"]: i for i in results["items"]}
    for item_id in ("qa-001", "qa-002"):
        assert items[item_id]["status"] == "failed"
        assert items[item_id]["failure"]["reason"] == (
            "upstream: brain ingestion failed: RuntimeError: the graph broke")
    assert items["ex-cuad-t1"]["status"] == "scored" and items["pq-01"]["status"] == "scored"
    assert results["contains_errors"] is True and results["outcome"] == "complete"
    assert not any(c["item_ids"] == ["qa-001"] for c in results["calls"])  # never asked
    assert len(world.fake.batch_creates[1]) == 3  # the three reviews (the portfolio question is direct)


def test_a_document_that_cannot_be_read_fails_its_items_and_sends_nothing_for_it(world):
    """FR-40: a document the exam cannot read (here: its bytes are not the ones the exam recorded)
    keeps its items in the record as failed -- listed, not dropped -- and costs nothing; the rest
    of the exam goes on."""
    (world.docs / "public" / "cp-t1.txt").write_bytes(b"tampered")
    results = run_exam(world)
    docs = {d["doc_id"]: d for d in results["documents"]}
    items = {i["item_id"]: i for i in results["items"]}
    assert docs["cp-t1"]["status"] == "failed" and "sha256" in docs["cp-t1"]["failure"]
    assert items["rv-cp-t1"]["status"] == "failed" and items["ex-cp-t1"]["status"] == "failed"
    assert items["rv-cp-t1"]["failure"]["reason"].startswith("upstream: ")
    assert items["rv-scan-cp-t1"]["status"] == "scored"
    assert not any(c["doc_id"] == "cp-t1" for c in results["calls"])
    assert results["contains_errors"] is True and results["outcome"] == "complete"
    scored = score.score_run(results, world.content)
    assert {f["item_id"] for f in scored["failed_items"]} == {"rv-cp-t1", "ex-cp-t1"}


def test_an_owner_contract_whose_text_no_longer_matches_its_frozen_text_leaves_the_run(
    tmp_path, monkeypatch, exam_engine, live_engine
):
    """FR-20/AC-20: keys confirmed against one text must not score another. A contract whose
    text differs from the frozen text it was confirmed against is left out -- its keys listed as
    out of date, nothing sent for it -- and the rest of the exam goes on."""
    world = build_world(tmp_path, monkeypatch, exam_engine, live_engine, owner_frozen_sha="0" * 64)
    results = run_exam(world)
    assert results["excluded"]["owner_keys_out_of_date"] == ["owner-12:governing_law"]
    assert all(i["item_id"] != "ex-owner-12" for i in results["items"])
    # the owner's contract was never sent anywhere: the only direct calls are the portfolio question's
    assert len(world.fake.streamed) == 2 and "Northwind" not in json.dumps(world.fake.streamed)
    assert results["outcome"] == "complete" and results["failed_call_rate"] == 0.0


# =============================================================================================
# The ceiling is a hard limit (FR-32)
# =============================================================================================


def test_batch_admission_is_by_worst_case_so_spend_cannot_pass_the_ceiling(world, happy):
    """FR-32/AC-32: requests are admitted in planned order only while their WORST case fits under
    the ceiling, because a sent batch cannot be recalled. The rest are not attempted: the run is
    incomplete, nothing it did not attempt counts as failed, and what was spent is under the
    ceiling. The owner's calls that the ceiling refused leave no failed call behind."""
    first_four = [r["params"] for r in happy.batch_creates[0][:4]]
    ceiling = sum(runner.worst_cost(p, batched=True) for p in first_four) + 0.002
    results = run_exam(world, ceiling=f"{ceiling:.4f}")
    assert results["outcome"] == "incomplete" and results["incomplete_reason"] == "ceiling_reached"
    assert [b["requests"] for b in results["batches"]] == [4]  # phase 2 was never even captured
    items = {i["item_id"]: i["status"] for i in results["items"]}
    assert items["ex-cuad-t1"] == items["ex-cp-t1"] == items["ex-scan-cp-t1"] == "scored"
    assert items["ex-owner-12"] == items["rv-cp-t1"] == items["qa-001"] == items["pq-01"] == "not_attempted"
    assert finished(world).after["actual"]["usd"] <= float(f"{ceiling:.4f}")
    assert results["failed_call_rate"] == 0.0 and results["contains_errors"] is False
    assert failed_calls_of(world) == [] and not any(r.status != "succeeded" for r in live_ledger(world))
    assert world.fake.streamed == []  # the owner's contract was refused by the ceiling, not sent
    assert exam_rows(world, AICallLog, AICallLog.error_class == "CeilingReached") == []
    assert finished(world).after["outcome"] == "incomplete"
    problems = report.check_pair({**finished(world).after, "run_id": results["run_id"]}, None)
    assert any("did not complete" in p for p in problems)  # never a baseline


def test_a_ceiling_too_small_for_any_request_sends_nothing_at_all(world):
    """The hard limit also holds at its extreme: with no room for even the first request, nothing
    goes out, nothing is spent, and every item is not attempted."""
    results = run_exam(world, ceiling="0.0001")
    assert world.fake.streamed == [] and world.fake.batch_creates == []
    assert results["outcome"] == "incomplete" and results["batches"] == []
    assert {i["status"] for i in results["items"]} == {"not_attempted"}
    assert finished(world).after["actual"]["usd"] == 0


def test_a_call_the_ceiling_refused_is_never_stamped_or_copied_however_late_its_row_lands(world):
    """FR-32/FR-40: a refusal is not a call. The run deletes the gateway's failure row for it, but
    one that commits between that delete and the stamp must not be tagged as the run's and copied
    to the live ledger either (it would show on the usage page and in the failed-call count for a
    call that never left the machine). The stamp itself leaves such a row alone."""
    results = run_exam(world, mode="sample")
    run_id = results["run_id"]
    with world.exam() as db:
        template = db.scalars(select(AICallLog)).first()
        columns = [c.key for c in AICallLog.__mapper__.column_attrs if c.key != "id"]
        late = AICallLog(**{name: getattr(template, name) for name in columns})
        late.eval_run_id, late.error_class, late.status = None, "CeilingReached", "failed"
        db.add(late)
        db.commit()
        late_id = late.id
    copied = len(live_ledger(world))
    assert runner.stamp_and_mirror(
        run_id=run_id, exam_sessions=world.exam, live=world.live, owner=OWNER) == 0
    with world.exam() as db:
        assert db.get(AICallLog, late_id).eval_run_id is None
    assert len(live_ledger(world)) == copied
    actual = runner.ledger_actual(world.live.sessions, run_id, org_id=world.live.org_id)
    assert actual["failed_calls"] == 0 and actual["calls"] == copied


def die_before_the_owners_calls(monkeypatch):
    """The process is killed right after phase 1's batch is sent, before the owner's synchronous
    contract is called (those calls run while the batch is processing)."""
    real = runner._Run._run_chains

    async def dying(self, units, mode, *, record):
        if mode == "live":
            monkeypatch.setattr(runner._Run, "_run_chains", real)  # the next process is healthy
            raise Crash()
        return await real(self, units, mode, record=record)

    monkeypatch.setattr(runner._Run, "_run_chains", dying)


def test_a_resumed_run_holds_back_what_the_batch_in_flight_could_cost_so_the_ceiling_stays_hard(
    world, happy, monkeypatch
):
    """FR-32: a sent batch cannot be recalled, so while it is out its worst case is held and the
    owner's direct calls are admitted next to that hold. A process that dies in that window and is
    resumed must hold the same amount: the ledger's spend knows nothing of requests not yet
    replayed, so without the re-hold a call the original run would have refused goes out and spend
    can pass the ceiling."""
    batch = [r["params"] for r in happy.batch_creates[0]]
    held = sum(runner.worst_cost(p, batched=True) for p in batch)
    margin = 0.003  # room for the batch's own rounding, not for a call as big as an owner's
    assert min(runner.worst_cost(p, batched=False) for p in happy.streamed) > margin
    ceiling = held + margin
    die_before_the_owners_calls(monkeypatch)
    with pytest.raises(Crash):
        run_exam(world, ceiling=f"{ceiling:.4f}")
    state = state_of(world)
    assert state["step"] == "submitted" and state["units"]["cl-owner-12"] == "pending"
    assert world.fake.streamed == [] and len(world.fake.batch_creates) == 1

    reserved = []
    real = runner._Run._run_chains

    async def watching(self, units, mode, *, record):
        if mode == "live":
            reserved.append(self.meter.reserved_usd)
        return await real(self, units, mode, record=record)

    monkeypatch.setattr(runner._Run, "_run_chains", watching)
    results = resume(world)

    assert reserved[0] == pytest.approx(held, rel=1e-3)  # held again, as the first process held it
    assert world.fake.streamed == []  # the owner's contract was refused by the ceiling, not sent
    items = {i["item_id"]: i["status"] for i in results["items"]}
    assert items["ex-owner-12"] == "not_attempted"
    assert results["outcome"] == "incomplete" and results["incomplete_reason"] == "ceiling_reached"
    assert finished(world).after["actual"]["usd"] <= ceiling
    assert len(world.fake.batch_creates) == 1  # and phase 1's batch was not sent a second time


# =============================================================================================
# A run that waits on a batch: resume, abandon, and the create call that may have succeeded
# =============================================================================================


def die_while_waiting(monkeypatch):
    """The process is killed the first time it asks how far the batch has got. By then the owner's
    synchronous contract has been called and the phase-1 batch is out."""
    real = runner._Run._finish_batch

    async def dying(self, batch):
        monkeypatch.setattr(runner._Run, "_finish_batch", real)  # the next process is healthy
        raise Crash()

    monkeypatch.setattr(runner._Run, "_finish_batch", dying)


def test_a_run_that_dies_while_waiting_resumes_without_repeating_anything(world, monkeypatch):
    """FR-32/FR-39: the state on disk is all that --resume needs. It polls the batch that is
    already out (no second batch, no second charge), does not call the owner's contract again,
    and finishes the same run under the ceiling the owner approved -- and the process that died
    left the run waiting, not crashed."""
    die_while_waiting(monkeypatch)
    with pytest.raises(Crash):
        run_exam(world)
    run_id = run_id_of(world)
    state = state_of(world)
    assert (state["phase"], state["step"]) == (1, "submitted")
    assert state["units"]["cl-owner-12"] == "done" and state["units"]["md-owner-12"] == "done"
    assert claude.exam_meter is None and audit_rows(world.live, "eval.run_finished") == []
    assert len(world.fake.streamed) == 2 and len(world.fake.batch_creates) == 1
    assert runner.reconcile_crashed_runs(world.live, exam_sessions=world.exam) == []  # waiting

    results = resume(world)

    assert results["run_id"] == run_id and results["outcome"] == "complete"
    assert len(world.fake.batch_creates) == 2  # phase 2's only: phase 1's was never sent again
    # the owner's contract was not called a second time: 2 more are the portfolio question's
    assert len(world.fake.streamed) == 4
    assert len(results["calls"]) == 17 and {i["status"] for i in results["items"]} == {"scored"}
    [resumed] = audit_rows(world.live, "eval.run_resumed")
    assert resumed.resource_id == run_id and resumed.after["phase"] == 1
    assert resumed.after["step"] == "submitted" and resumed.after["resumed_by"]["email"] == OWNER.email
    [row] = audit_rows(world.live, "eval.run_finished")
    assert row.after["outcome"] == "complete" and row.after["approved_by"]["email"] == OWNER.email
    assert row.after["ceiling_usd"] == 25.0
    ids = [r.id for r in live_ledger(world)]
    assert len(ids) == len(set(ids)) == 17  # every call copied once, none twice
    assert results["started_at"] < results["finished_at"]


def test_a_new_run_is_refused_while_another_waits_on_a_batch(world, monkeypatch):
    """A new run rebuilds the exam database from nothing; the waiting run's answers have nowhere
    to go without it, so the second run is refused and told what to do."""
    die_while_waiting(monkeypatch)
    with pytest.raises(Crash):
        run_exam(world)
    typist = Typist()
    with pytest.raises(ExamRefused, match=r"is waiting on a batch.*--resume .*--abandon"):
        run_exam(world, typist=typist)
    assert typist.prompts == [] and len(world.fake.batch_creates) == 1


def test_abandoning_cancels_the_pending_batch_and_finishes_the_run_incomplete(world, monkeypatch):
    """FR-39/FR-40: an abandoned run cancels its batch, copies what was spent to the live ledger,
    and is finished as incomplete -- so it can never become part of a baseline -- and the exam
    database is free to be rebuilt."""
    die_while_waiting(monkeypatch)
    with pytest.raises(Crash):
        run_exam(world)
    run_id = run_id_of(world)
    asyncio.run(runner.abandon_exam_run(live=world.live, owner=OWNER, run_id=run_id))
    assert world.fake.cancelled == ["msgbatch_fake1"]
    [row] = audit_rows(world.live, "eval.run_finished")
    assert row.after["outcome"] == "incomplete" and row.after["incomplete_reason"] == "abandoned"
    assert row.after["results_sha256"] is None
    assert len(live_ledger(world)) == 2  # the owner's two calls were paid for
    # The cancelled batch is billed for the requests it had already answered, though nobody
    # replayed them: they are in the record's cost, not in a ledger row.
    answered = len(world.fake.batch_creates[0])
    model = world.fake.batch_creates[0][0]["params"]["model"]
    each = token_cost(100, 20, model, batched=True)  # the fake's usage for one answer
    unreplayed = row.after["actual"]["by_call_type"]["unreplayed_batch_answers"]
    assert unreplayed == {"calls": answered, "batch": answered, "usd": pytest.approx(answered * each)}
    assert row.after["actual"]["calls"] == 2 + answered
    ledger_usd = sum(
        token_cost(r.prompt_tokens, r.completion_tokens, r.model, batched=bool(r.batched))
        for r in live_ledger(world)
    )
    assert row.after["actual"]["usd"] == pytest.approx(ledger_usd + answered * each)
    assert any("did not complete" in p for p in report.check_pair({**row.after, "run_id": run_id}, None))
    runner._refuse_if_waiting(world.live)  # no longer blocks a new run
    with pytest.raises(ExamRefused, match="already finished"):
        resume(world)
    with pytest.raises(ExamRefused, match="already finished"):
        asyncio.run(runner.abandon_exam_run(live=world.live, owner=OWNER, run_id=run_id))


def submit_dies(world, monkeypatch, how):
    """The first ``create_batch`` either answers with an error (the reply may have been lost after
    the provider made the batch) or the process dies in the middle of it."""
    real = claude.claude_client.create_batch
    state = {"first": True}

    async def create(requests):
        if state["first"]:
            state["first"] = False
            world.fake.batch_creates.append(requests)  # the provider may well have made it
            if how == "error":
                raise RuntimeError("the connection dropped before the reply")
            raise Crash()
        return await real(requests)

    monkeypatch.setattr(claude.claude_client, "create_batch", create)


@pytest.mark.parametrize(
    "how, raises, step", [("error", runner.RunPaused, "submit_failed"), ("died", Crash, "submitting")]
)
def test_a_batch_create_that_may_have_succeeded_is_only_sent_again_when_the_owner_says_so(
    world, monkeypatch, how, raises, step
):
    """A create call with no reply may have made -- and billed -- a batch nobody has the id of.
    Both ways it can end (an error, or the process dying) leave the run WAITING (never crashed:
    reconcile leaves it alone); resuming asks the owner; "no" changes nothing; "yes" sends again
    with the first batch's worst case still held under the ceiling, so the limit stays a limit."""
    submit_dies(world, monkeypatch, how)
    with pytest.raises(raises):
        run_exam(world)
    assert state_of(world)["step"] == step
    assert runner.reconcile_crashed_runs(world.live, exam_sessions=world.exam) == []
    assert len(world.fake.batch_creates) == 1
    lost_worst = sum(runner.worst_cost(r["params"], batched=True) for r in world.fake.batch_creates[0])

    asked = []
    with pytest.raises(runner.RunPaused, match="still waiting"):
        asyncio.run(runner.resume_exam_run(
            live=world.live, owner=OWNER, run_id=run_id_of(world), content=world.content,
            ask=lambda prompt: asked.append(prompt) or "n"))
    assert "may have been created without a reply" in asked[0]
    assert len(world.fake.batch_creates) == 1  # "no" sent nothing

    held = []
    admit = runner.Meter.admit

    def watching(self, payloads):
        held.append(self.reserved_usd)
        return admit(self, payloads)

    monkeypatch.setattr(runner.Meter, "admit", watching)
    results = resume(world, answers="y")
    assert held[0] == pytest.approx(lost_worst)  # the lost batch's worst case was still held
    assert results["outcome"] == "complete" and results["failed_call_rate"] == 0.0
    assert len(world.fake.batch_creates) == 3  # the lost one, its replacement, phase 2


def test_abandoning_a_run_whose_batch_create_never_answered_tells_the_owner_to_cancel_it_by_hand(
    world, monkeypatch, capsys
):
    """FR-40: a create call that errored may have made a batch nobody has the id of. Abandon cannot
    cancel what it cannot name: it says so (the owner can in the console), cancels nothing, and
    still ends the run as incomplete -- never leaves it waiting."""
    submit_dies(world, monkeypatch, "error")
    with pytest.raises(runner.RunPaused):
        run_exam(world)
    run_id = run_id_of(world)
    asyncio.run(runner.abandon_exam_run(live=world.live, owner=OWNER, run_id=run_id))
    assert "check the Anthropic console and cancel it there" in capsys.readouterr().out
    assert world.fake.cancelled == []
    [row] = audit_rows(world.live, "eval.run_finished")
    assert row.after["incomplete_reason"] == "abandoned" and row.after["actual"]["usd"] == 0
    runner._refuse_if_waiting(world.live)  # no longer blocks a new run


def test_abandoning_still_ends_the_run_when_the_cancelled_batch_cannot_be_read(
    world, monkeypatch, capsys
):
    """What the cancelled batch had already answered is read to price it. If that read fails the
    abandon must still finish (an unfinished run blocks every new one), and say that the cost on
    record may be too low rather than present it as complete."""
    die_while_waiting(monkeypatch)
    with pytest.raises(Crash):
        run_exam(world)

    async def unreadable(batch_id):
        raise RuntimeError("the results could not be fetched")

    monkeypatch.setattr(claude.claude_client, "batch_results", unreadable)
    asyncio.run(runner.abandon_exam_run(live=world.live, owner=OWNER, run_id=run_id_of(world)))
    assert "cost recorded for this run may be too low" in capsys.readouterr().out
    assert world.fake.cancelled == ["msgbatch_fake1"]
    [row] = audit_rows(world.live, "eval.run_finished")
    assert row.after["incomplete_reason"] == "abandoned" and row.after["actual"]["calls"] == 2


def test_a_batch_settled_again_after_a_resume_is_not_counted_twice(world, monkeypatch):
    """The run record's per-batch cost: the process dies while the RETRY batch is out, so the
    first batch (already settled once, its cost saved) is settled again on resume. Its figure is
    what its own answers cost, not that plus the same again."""
    fake = world.fake
    answers = fake.messages.batches.results

    async def first_batch_overloaded(batch_id):
        items = await answers(batch_id)
        first = fake.batches["msgbatch_fake1"][0]["custom_id"]
        failure = batch_result("errored", error_type="overloaded_error")
        overloaded = SimpleNamespace(custom_id=first, result=SimpleNamespace(to_dict=lambda: failure))

        async def changed():
            async for item in items:
                yield overloaded if batch_id == "msgbatch_fake1" and item.custom_id == first else item

        return changed()

    fake.messages.batches.results = first_batch_overloaded
    real = runner._Run._finish_batch
    asked = []

    async def dies_waiting_for_the_retry(self, batch):
        asked.append(batch["batch_id"])
        if len(asked) == 2:
            monkeypatch.setattr(runner._Run, "_finish_batch", real)
            raise Crash()
        return await real(self, batch)

    monkeypatch.setattr(runner._Run, "_finish_batch", dies_waiting_for_the_retry)
    with pytest.raises(Crash):
        run_exam(world)
    assert state_of(world)["step"] == "retry_submitted"

    results = resume(world)

    each = token_cost(100, 20, fake.batch_creates[0][0]["params"]["model"], batched=True)
    first, retry = [b for b in results["batches"] if b["phase"] == 1]
    assert first["errored"] == 1 and retry["succeeded"] == 1
    assert first["usd"] == pytest.approx((len(fake.batch_creates[0]) - 1) * each)
    assert retry["usd"] == pytest.approx(each)
    assert results["outcome"] == "complete" and results["failed_call_rate"] == 0.0


def test_resume_refuses_what_it_cannot_continue_safely(world, monkeypatch, happy):
    """--resume refuses (and says --abandon) when the run is unknown or finished, when the app code
    or the exam content is not what the run started with, when its exam database is gone, and when
    its state file is missing -- never guessing."""
    with pytest.raises(ExamRefused, match="no approved exam run"):
        resume(world, run_id="no-such-run")
    with pytest.raises(ExamRefused, match="already finished"):
        asyncio.run(runner.resume_exam_run(
            live=happy.world.live, owner=OWNER, run_id=happy.run_id, ask=lambda p: "y",
            content=happy.world.content))
    die_while_waiting(monkeypatch)
    with pytest.raises(Crash):
        run_exam(world)
    other = world.content.model_copy(update={"sample": {"review": ["rv-cp-t1"]}})
    with pytest.raises(ExamRefused, match="changed since this run started.*--abandon"):
        resume(world, content=other)
    gone = create_engine(
        make_url(settings.database_url).set(database="legal_clm_eval_never_made"), poolclass=NullPool)
    monkeypatch.setattr(runner, "_exam_sessions", lambda: sessionmaker(bind=gone))
    with pytest.raises(ExamRefused, match="is gone.*--abandon"):
        resume(world)
    monkeypatch.setattr(runner, "_exam_sessions", lambda: world.exam)
    (world.runs / run_id_of(world) / "state.json").unlink()
    with pytest.raises(ExamRefused, match="no saved state.*--abandon"):
        resume(world)
    gone.dispose()


def test_a_run_that_crashed_is_closed_as_incomplete_by_the_next_run(world, monkeypatch):
    """FR-40/AC-40: a process that dies with no batch outstanding leaves an approved run with no
    finish record; the next run closes it -- incomplete, reason crashed, what it spent copied to the
    live ledger -- so it can never become part of a baseline; and then runs normally."""
    real = runner._Run._brain

    async def dies(self):
        monkeypatch.setattr(runner._Run, "_brain", real)
        raise Crash()

    monkeypatch.setattr(runner._Run, "_brain", dies)
    with pytest.raises(Crash):
        run_exam(world)
    crashed = run_id_of(world)
    assert audit_rows(world.live, "eval.run_finished") == []
    results = run_exam(world)
    first, second = audit_rows(world.live, "eval.run_finished")
    assert first.resource_id == crashed and first.after["outcome"] == "incomplete"
    assert first.after["incomplete_reason"] == "crashed" and first.after["actual"]["calls"] == 10
    assert second.resource_id == results["run_id"] and second.after["outcome"] == "complete"
    assert any("did not complete" in p for p in report.check_pair({**first.after, "run_id": crashed}, None))


# =============================================================================================
# What the app's own code makes of the answers: flags, facts, quotes, routing
# =============================================================================================

REVIEW_ANSWER = {
    "deviations": [
        {   # numbered by the AI -> the rule it numbered; has the original text -> a deviation
            "rule_index": 1, "clause_type": "limitation_of_liability", "severity": "high",
            "issue": "liability is unlimited", "original_text": "Each party's total liability",
            "citations": [{"quote": PLANTED}],
        },
        {   # no number, a clause type that canonicalises to a rule's -> that rule; no original -> missing
            "clause_type": "Governing Law", "severity": "medium", "issue": "no governing law",
            "citations": [{"quote": GOVERNING_LAW}],
        },
        {   # no number, no matching type -> found by the text it quotes (the planted change)
            "clause_type": "indemnity", "severity": "low", "issue": "changed wording",
            "citations": [{"quote": PLANTED}],
        },
        {   # nothing ties it to a rule -> unattributed, which the scorer counts as a false alarm
            "clause_type": "something odd", "severity": "low", "issue": "odd",
            "citations": [{"quote": "Nothing in the contract says anything like this at all"}],
        },
    ],
    "summary": "four findings",
    "citations": [],
}
METADATA_ANSWER = {
    "counterparty_name": "Acme Private Limited and Beta Pvt. Ltd.", "jurisdiction": "India",
    "value_amount": 2500000, "currency": "INR", "effective_date": "2024-04-01",
    "expiration_date": "2027-04-01", "confidence": "high", "citations": [{"quote": GOVERNING_LAW}],
}
LABELS_ANSWER = {
    "labels": [
        {"segment_id": f"S{n}", "clause_type": kind, "confidence": "high"}
        for n, kind in enumerate(
            ["renewal", "term_and_termination", "limitation_of_liability", "governing_law",
             "payment_terms", "notices"], 1)
    ],
}


# Every segment is the same type, so wherever two contracts share a clause's wording the app's
# "same wording elsewhere" search will find the pair.
ONE_TYPE_LABELS = {
    "labels": [
        {"segment_id": f"S{n}", "clause_type": "governing_law", "confidence": "high"}
        for n in range(1, 10)
    ],
}


def test_flags_facts_and_quotes_are_read_the_way_the_app_reads_them(world, monkeypatch):
    """FR-10/FR-12/FR-22/FR-23: a deviation is attributed to a rule by its number, then its clause
    type, then the text it quotes -- never by peeking at whether it is right; extraction is read
    from the contract row and stored clauses the app's skills filled; each quote is checked word
    for word against the source."""
    base = exam_fake_anthropic.structured_payload_by_tool
    monkeypatch.setattr(exam_fake_anthropic, "structured_payload_by_tool", lambda: {
        **base(), "return_playbook_review": REVIEW_ANSWER,
        "return_contract_metadata_extraction": METADATA_ANSWER,
        "return_clause_labeling": LABELS_ANSWER})
    results = run_exam(world)
    items = {i["item_id"]: i for i in results["items"]}
    flags = items["rv-cp-t1"]["prediction"]["flags"]
    assert [(f["rule"], f["status"], f["severity"], f["attributed_by"]) for f in flags] == [
        ("limitation_of_liability", "deviation", "high", "rule_index"),
        ("governing_law", "missing", "medium", "clause_type"),
        ("limitation_of_liability", "missing", "low", "text"),
        (None, "missing", "low", "none"),
    ]
    assert flags[0]["quote"] == PLANTED
    scores = score.score_run(results, world.content)
    assert {a["item_id"] for a in scores["false_alarms"]} >= {"rv-ae-t1"}  # flags on the control
    assert scores["missed_high_critical"] == []  # the planted high-severity change WAS found

    fields = items["ex-cuad-t1"]["prediction"]["fields"]
    assert fields["parties"] == ["Acme Private Limited", "Beta Pvt. Ltd."]
    assert (fields["effective_date"], fields["governing_law"]) == ("2024-04-01", "India")
    assert fields["value"] == {"amount": "2500000.00", "currency": "INR"}
    clauses = {c.clause_type: c.text for c in exam_rows(world, ClauseExtraction)
               if c.contract_id == next(d["contract_id"] for d in results["documents"] if d["doc_id"] == "cuad-t1")}
    assert fields["liability_cap"] == [clauses["limitation_of_liability"]]
    assert fields["renewal_term"] and fields["renewal_term"][0] in (
        clauses.get("renewal"), clauses.get("term_and_termination"))
    assert scores["measures"]["extraction_accuracy"]["value"] == 1.0  # every key fact is right

    review = next(c for c in results["calls"] if c["item_ids"] == ["rv-cp-t1"])
    by_text = {q["text"]: q for q in review["quotes"]}
    assert by_text[GOVERNING_LAW]["word_for_word"] is True and by_text[GOVERNING_LAW]["align_score"] == 100.0
    assert by_text[PLANTED]["word_for_word"] is False  # the source says something else
    assert by_text["Nothing in the contract says anything like this at all"]["word_for_word"] is False
    validity = scores["measures"]["quote_validity"]
    assert 0 < validity["value"] < 1 and validity["n"] > 0
    assert any(q["text"] == PLANTED for q in scores["private"]["invalid_quotes"])


def test_the_owners_contracts_cannot_turn_up_in_a_batched_request(world, monkeypatch):
    """FR-35/FR-41: the app's search names other contracts whose clauses read the same ("also
    appears in <title>"), and the runner may open every contract in the exam database -- the
    owner's included. Their contracts are taken out of the app's reach before phase 2, so no client
    title can reach a batched request, while the public and constructed titles still can (the
    lens is working) and the owner's contract is back, intact, in the record."""
    base = exam_fake_anthropic.structured_payload_by_tool
    monkeypatch.setattr(exam_fake_anthropic, "structured_payload_by_tool",
                        lambda: {**base(), "return_clause_labeling": ONE_TYPE_LABELS})
    results = run_exam(world)
    sent = json.dumps(world.fake.batch_creates)
    assert "Mutual NDA" not in sent and "Northwind" not in sent
    owner = next(i for i in results["items"] if i["item_id"] == "ex-owner-12")
    assert owner["status"] == "scored"
    [row] = [c for c in exam_rows(world, Contract) if c.title.startswith("Mutual NDA")]
    assert row.deleted_at is None  # back, for the record
    # and the path was real: with the owner's contract in reach, the app's search DOES name it
    contracts = exam_rows(world, Contract)
    cuad = next(c for c in contracts if c.title == "cuad-t1")
    with world.exam() as db:
        facts = clause_language_matches(
            db, org_id=cuad.org_id, contract_ids=[cuad.id], accessible_ids={c.id for c in contracts})
    assert any(row.title in f["fact"] for f in facts)


def test_a_portfolio_answer_names_the_contracts_it_names_by_whole_title(world, monkeypatch):
    """FR-9: an answer that came from SEARCH (here the fake provider says the question is not one
    the filter answers) names contracts only in its text, so they are read from it by title (whole
    words only, so "ae-t10" is not "ae-t1") and only among the portfolio's own documents. The
    answer is a search answer, and says so."""
    sentence = "The Buyer pays the price in Schedule A within forty five days of each invoice."
    answer = {
        "answer": "Only ae-t1 qualifies, and cuad-t1; the draft ae-t10 is not one of ours.",
        "citations": [{"quote": sentence}], "confidence": "high",
    }
    base = exam_fake_anthropic.structured_payload_by_tool
    monkeypatch.setattr(exam_fake_anthropic, "structured_payload_by_tool",
                        lambda: {**base(), "return_contract_brain_answer": answer})
    results = run_exam(world)
    item = next(i for i in results["items"] if i["item_id"] == "pq-01")
    assert item["prediction"]["named_doc_ids"] == ["ae-t1"]  # not pf-t1, ae-t10 or cuad-t1 (not ours)
    assert item["prediction"]["citations"] and item["prediction"]["context_chars"] > 0
    assert item["prediction"]["answer_mode"] == "rag"
    scored = score.score_run(results, world.content)
    assert scored["portfolio"]["questions"][0]["wrongly_included"] == 0
    assert scored["portfolio"]["questions"][0]["missed_doc_ids"] == ["pf-t1"]


# --- "which contracts ..." questions: the filter decides, and the AI reads what no field holds --

SUPPLY = "The Supplier shall deliver the goods listed in Schedule A to the Buyer's warehouse."
RENEWAL_WORDING = {"is_list_question": True, "wording": ["Does this contract renew automatically?"]}


def answers_by_contract(monkeypatch, yes_for: str):
    """The fake provider answers a table read "yes" (quoting the contract) for the contract whose
    text carries ``yes_for`` and "no" for any other."""
    real = exam_fake_anthropic.message_for

    def message_for(payload, usage=None):
        message = real(payload, usage)
        if (payload.get("tool_choice") or {}).get("name") == "return_tabular_cell_extraction":
            yes = yes_for in json.dumps(payload["messages"])
            message["content"][0]["input"] = {
                "answer": "Yes" if yes else "No", "verdict": "yes" if yes else "no",
                "confidence": "high", "not_found": not yes,
                "citations": [{"quote": SUPPLY}] if yes else [],
            }
        return message

    monkeypatch.setattr(exam_fake_anthropic, "message_for", message_for)


def the_filter_says(monkeypatch, answer: dict, **others):
    base = exam_fake_anthropic.structured_payload_by_tool
    monkeypatch.setattr(exam_fake_anthropic, "structured_payload_by_tool",
                        lambda: {**base(), "return_portfolio_filter": answer, **others})


def test_a_question_about_wording_has_every_contract_read_to_the_end_inside_its_unit(world, monkeypatch):
    """Step 3e: for wording no contract field holds ("renews automatically") the app's filter hands
    each candidate contract to the AI as a table read, queued as jobs a worker would run. The exam
    has no worker, so the unit runs those jobs itself with the job handler's own entry point and
    only then asks for the list the way the page does. Failure guarded: the unit returning while the
    reads were still pending (a list of nothing, scored as a total miss), the reads' calls
    belonging to no unit (an unreported cost), or the answer being read for names instead of taken
    from the filter."""
    the_filter_says(monkeypatch, RENEWAL_WORDING)
    answers_by_contract(monkeypatch, yes_for="Eagle")  # ae-t1 says yes, pf-t1 says no
    results = run_exam(world)
    item = next(i for i in results["items"] if i["item_id"] == "pq-01")
    assert item["status"] == "scored"
    prediction = item["prediction"]
    assert prediction["answer_mode"] == "database_filter"
    assert prediction["named_doc_ids"] == ["ae-t1"]
    assert (prediction["checked"], prediction["total"], prediction["unknown"]) == (2, 2, 0)
    assert prediction["reading_failed"] == 0 and "ae-t1" in prediction["answer"]
    # the filter call and one read per contract: all this unit's, all ordinary calls
    mine = calls_of(results, "pq-01")
    assert sorted(c["call_type"] for c in mine) == [
        "portfolio_filter", "tabular_cell_extraction", "tabular_cell_extraction"]
    assert not any(c["batched"] for c in mine) and all(c["status"] == "succeeded" for c in mine)
    reads = [c for c in mine if c["call_type"] == "tabular_cell_extraction"]
    assert [q["word_for_word"] for c in reads for q in c["quotes"]] == [True]  # the "yes" quote
    # none of it went into the batch, and the jobs ran to the end
    sent = [r["params"]["tool_choice"]["name"] for b in world.fake.batch_creates for r in b]
    assert not {"return_portfolio_filter", "return_tabular_cell_extraction"} & set(sent)
    jobs = exam_rows(world, JobRun, JobRun.job_type == "tabular_cell_extraction")
    assert len(jobs) == 2 and {j.status for j in jobs} == {"succeeded"}
    assert {j.idempotency_key.split("#")[0] for j in jobs} == {"exam:pq-01"}
    # the saved answer is the final list, as the page would have it after polling
    saved = next(q for q in exam_rows(world, BrainQuery) if q.query_scope == "portfolio")
    assert saved.retrieval_metadata["reading"]["done"] == 2 and len(saved.retrieval_metadata["matches"]) == 1
    scored = score.score_run(results, world.content)["portfolio"]["questions"][0]
    assert (scored["found"], scored["missed_doc_ids"], scored["wrongly_included"]) == (1, ["pf-t1"], 0)


def test_a_question_a_contract_field_answers_makes_one_call_and_reads_nothing(world, monkeypatch):
    """pq-09 ("governed by Indian law") is decided from the contract record the metadata step filled,
    over every contract: one filter call, no table read. Failure guarded: the exam reading 33
    contracts for a fact the app already holds, and so paying for what the filter exists to avoid."""
    the_filter_says(monkeypatch, {"is_list_question": True, "governing_law": "India"},
                    return_contract_metadata_extraction=METADATA_ANSWER)
    results = run_exam(world)
    item = next(i for i in results["items"] if i["item_id"] == "pq-01")
    assert item["status"] == "scored" and item["prediction"]["answer_mode"] == "database_filter"
    assert item["prediction"]["named_doc_ids"] == ["ae-t1", "pf-t1"]
    assert (item["prediction"]["checked"], item["prediction"]["total"]) == (2, 2)
    assert [c["call_type"] for c in calls_of(results, "pq-01")] == ["portfolio_filter"]
    assert exam_rows(world, JobRun, JobRun.job_type.like("tabular%")) == []


def test_a_bare_how_many_question_makes_no_ai_call_at_all(world):
    """"How many contracts do we have?" has no condition to propose, so the app answers it from the
    database without the model: the unit makes no call, names every contract the portfolio user can
    open, and costs nothing. Failure guarded: a question that should be free sent to the provider."""
    question = PortfolioQuestion(
        item_id="pq-01", question="How many contracts do we have?", fact="count",
        expected_doc_ids=["ae-t1", "pf-t1"])
    world.content = world.content.model_copy(update={"portfolio": [question]})
    results = run_exam(world)
    item = next(i for i in results["items"] if i["item_id"] == "pq-01")
    assert item["status"] == "scored"
    assert item["prediction"]["answer_mode"] == "database_filter"
    assert item["prediction"]["named_doc_ids"] == ["ae-t1", "pf-t1"]
    assert (item["prediction"]["checked"], item["prediction"]["total"]) == (2, 2)
    assert len(world.fake.streamed) == 2  # the owner's two, and nothing for this question
    assert not [c for c in results["calls"] if c["call_type"] in ("portfolio_filter", "tabular_cell_extraction")]


def test_the_ceiling_stops_a_question_part_way_through_its_reads_and_it_is_not_attempted(world, monkeypatch):
    """FR-32: the reads are metered calls like any other. When the ceiling refuses one, the unit
    stops reading (the rest would be refused too), its item is not attempted -- never scored on a
    half-read list -- the run is incomplete, and the refused call leaves no failed call behind.
    The refusal is tied to the unit through the renamed job, which is what this proves."""
    the_filter_says(monkeypatch, RENEWAL_WORDING)
    answers_by_contract(monkeypatch, yes_for="Eagle")
    real = runner.Meter.before
    reads = []

    def refuse_the_second_read(self, payload):
        if (payload.get("tool_choice") or {}).get("name") == "return_tabular_cell_extraction":
            reads.append(1)
            if len(reads) == 2:
                self.stopped = True
                raise runner.CeilingReached("the ceiling, in this test")
        return real(self, payload)

    monkeypatch.setattr(runner.Meter, "before", refuse_the_second_read)
    results = run_exam(world)
    items = {i["item_id"]: i["status"] for i in results["items"]}
    assert items["pq-01"] == "not_attempted"
    assert results["outcome"] == "incomplete" and results["incomplete_reason"] == "ceiling_reached"
    assert exam_rows(world, AICallLog, AICallLog.error_class == "CeilingReached") == []
    assert results["failed_call_rate"] == 0.0 and failed_calls_of(world) == []


def test_a_read_that_never_ran_fails_the_question_instead_of_scoring_a_half_read_list(world, monkeypatch):
    """A job that could not even start leaves no failed call behind (nothing was asked of the
    provider), and its contracts stay unread. Scoring that list would score the exam's own bug as a
    missed answer, so the unit fails and the report says why. (A read whose AI CALL fails is
    different: the cell and the ledger record it, and the item fails on the call.)"""
    the_filter_says(monkeypatch, RENEWAL_WORDING)

    async def never_runs(job_id):
        raise RuntimeError("the worker could not take the job")

    monkeypatch.setattr(runner.job_tasks, "_run_ai_job", never_runs)
    results = run_exam(world)
    item = next(i for i in results["items"] if i["item_id"] == "pq-01")
    assert item["status"] == "failed"
    assert item["failure"]["reason"] == "unit failed: RuntimeError: the AI read 0 of 2 contracts"
    assert results["outcome"] == "complete" and results["contains_errors"] is True


def test_capturing_a_call_records_the_request_and_leaves_nothing_behind(world):
    """Plan decision 22: capture aborts the app's call at the Claude boundary with a
    BaseException, so the app writes no fallback output and no failure row for a call that was only
    being recorded; the request is kept, keyed by its own hash."""
    plan = runner.plan_run(world.content, "full", "batch")
    run = runner._Run(live=world.live, owner=OWNER, content=world.content, plan=plan,
                      sessions=world.exam, ask=Typist(), change=None, target=None)
    asyncio.run(run.prepare())
    claude.exam_meter = run.meter = runner.Meter(1.0)
    run.meter.mode = "capture"
    try:
        status = asyncio.run(run._run_unit(run.units_by_id["md-cp-t1"]))
    finally:
        claude.exam_meter = None
    payload = run.payloads["md-cp-t1"]
    assert status == "captured" and payload["tool_choice"]["name"] == "return_contract_metadata_extraction"
    assert run.cid_of["md-cp-t1"] == runner.custom_id(payload)
    assert world.fake.streamed == [] and exam_rows(world, AICallLog) == []
    [skill_run] = [s for s in exam_rows(world, AISkillRun) if s.skill_name == "contract_metadata_extraction"]
    assert skill_run.status == "running" and skill_run.output_payload is None
    [contract] = [c for c in exam_rows(world, Contract) if c.title == "cp-t1"]
    assert "ai_suggestions" not in (contract.metadata_json or {})  # no fallback metadata was stored


def test_routing_requests_are_triaged_by_the_app_in_phase_1_and_scored_by_workflow(
    tmp_path, monkeypatch, exam_engine, live_engine
):
    """FR-25: a free-text request is triaged by the app's own triage against the workflows loaded
    into the exam database (never saved as a request), as one batched call in phase 1; its answer
    is the workflow the app picked. Triage runs on a worker thread, so this also proves a capture
    can cross it."""
    world = build_world(tmp_path, monkeypatch, exam_engine, live_engine, routing=True)
    results = run_exam(world)
    item = next(i for i in results["items"] if i["item_id"] == "rt-01")
    assert item["status"] == "scored"
    assert set(item["prediction"]) == {"workflow_id", "needs_human", "source", "degraded"}
    assert item["prediction"]["degraded"] is False
    call = next(c for c in results["calls"] if c["item_ids"] == ["rt-01"])
    assert call["call_type"] == "intake_triage" and call["batched"] is True and call["doc_id"] is None
    assert any(r["unit_id"] == "rt-01" for b in [state_of(world)] for r in [{"unit_id": u} for u in b["units"]])
    assert results["excluded"]["routing_unpicked"] == 0 and results["excluded"]["routing_out_of_date"] is False
    assert score.score_run(results, world.content)["counts"]["routing"]["requests"] == 1
    assert exam_rows(world, IntakeRequest) == []  # never saved


# =============================================================================================
# The plan and the estimate (pure)
# =============================================================================================


@pytest.fixture
def committed(monkeypatch, tmp_path):
    """The exam's committed content (public and constructed), with no private files."""
    monkeypatch.setattr(content_mod, "PRIVATE_DIR", tmp_path)
    return content_mod.load_content()


def count_by_task(plan):
    out: dict[str, int] = {}
    for item in plan.items:
        out[item.task] = out.get(item.task, 0) + 1
    return out


def test_a_full_plan_asks_every_item_and_a_scanned_copy_shares_its_originals_record(committed):
    """FR-1/FR-8: nothing in the content is silently dropped, and each scanned copy is reviewed
    against its original's answer record and read for its facts."""
    plan = runner.plan_run(committed, "full", "batch")
    twins = [d for d in committed.documents if d.twin_of]
    assert twins
    by_task = count_by_task(plan)
    assert by_task["review"] == len(committed.review) + len(twins)
    assert by_task["qa"] == len(committed.qa) == 100 and by_task["portfolio"] == len(committed.portfolio)
    assert by_task["extraction"] >= len(committed.extraction)
    assert {u.unit_id for u in plan.units if u.kind == "review"} >= {f"rv-{d.doc_id}" for d in twins}
    assert plan == runner.plan_run(committed, "full", "batch")  # the same, every time
    assert [u.phase for u in plan.units] == sorted(u.phase for u in plan.units)
    # no owner contracts in this content; and a portfolio question is never batched, because it
    # makes a filter call and then one read per contract where a batch replays one call per unit
    assert all(u.delivery == "batch" for u in plan.units if u.kind != "portfolio")
    assert [u.delivery for u in plan.units if u.kind == "portfolio"] == ["sync"] * 9
    assert all(u.delivery == "sync" for u in runner.plan_run(committed, "full", "sync").units)
    steps = [u.unit_id for u in plan.units if u.doc_id == "cuad-001" and u.phase == 1]
    assert steps == ["cl-cuad-001", "md-cuad-001"]  # labelling before metadata, always
    portfolio_docs = {d.doc_id for d in committed.documents if d.uploader == "portfolio"}
    assert set(plan.portfolio_doc_ids) == portfolio_docs and portfolio_docs <= set(plan.doc_ids)


def test_a_sample_plan_is_the_fixed_subset_and_smaller_than_the_full_plan(committed):
    """FR-30/AC-30: the sample is exactly what sample.json names, with the sub-portfolio, and
    asks far less than a full run."""
    plan = runner.plan_run(committed, "sample", "batch")
    full = runner.plan_run(committed, "full", "batch")
    sample = committed.sample
    ids = {t: [i.item_id for i in plan.items if i.task == t] for t in ("review", "qa", "portfolio")}
    assert ids["qa"] == sample["qa"] and ids["portfolio"] == sample["portfolio"]
    assert len(ids["qa"]) == 9  # the fixed sample asks nine questions
    assert set(ids["review"]) >= set(sample["review"])
    assert plan.portfolio_doc_ids == tuple(sample["portfolio_docs"])
    assert len(plan.units) < len(full.units) / 4 and set(plan.doc_ids) < set(full.doc_ids)
    assert {t for t in count_by_task(plan)} >= {"review", "extraction", "qa", "portfolio"}


def test_only_a_wording_question_is_expected_to_read_contracts_and_the_estimate_prices_the_reads(committed):
    """A question about a fact the contract record has a field for (governing law, expiry, value,
    counterparty, type) reads no contract; one about wording (renews automatically, notice period)
    is expected to read all 40, each priced from its real length. Failure guarded: an estimate that
    prices pq-09's 33 matches as 33 reads (and asks the owner for a ceiling several times too
    high), or one that leaves the reads out and lets the run stop early on a ceiling set from it."""
    plan = runner.plan_run(committed, "full", "batch")
    run = runner._Run(live=None, owner=None, content=committed, plan=plan, sessions=None,
                      ask=None, change=None, target=None)
    run.ingested = {d: SimpleNamespace(text_len=1_000 + n) for n, d in enumerate(plan.portfolio_doc_ids)}
    reads = run.expected_reads()
    facts = {q.item_id: q.fact for q in committed.portfolio}
    assert set(reads) == {"pq-01", "pq-06"} and {facts[i] for i in reads} == {"auto_renewal", "renewal_notice"}
    assert all(len(lengths) == 40 for lengths in reads.values())
    assert facts["pq-09"] == "governing_law" and "pq-09" not in reads
    estimate = runner.build_estimate(plan, {}, model="claude-sonnet-4-6", reads=reads)
    assert estimate.by_call_type["portfolio_filter"]["calls"] == 9  # one for each question
    assert estimate.by_call_type["tabular_cell_extraction"]["calls"] == 80  # two questions, 40 contracts
    assert estimate.by_call_type["tabular_cell_extraction"]["batch"] == 0


def owner_content():
    """Two owner contracts -- an ordinary one and a stamp-only scan with its 'as read today' twin
    -- a third with no key, and one key the owner left unmarked."""
    stamp = datetime.now(UTC)
    manifest, documents, keys = [], [], []
    for number, kind in ((1, "text"), (2, "stamp_only_scan"), (3, "text")):
        name, doc_id = OWNER_FILES[number - 1], f"owner-{number:02d}"
        manifest.append(OwnerManifestEntry(
            doc_id=doc_id, name=name, file_sha256=f"{number}" * 64, mime_type="application/pdf",
            pages=3, kind=kind, frozen_text_sha256="f" * 64, imported_at=stamp))
        base = dict(category=SourceCategory.OWNER, dataset=None, file=name, mime_type="application/pdf",
                    file_sha256=f"{number}" * 64, title=name, ai_steps=["metadata", "clauses"])
        documents.append(ExamDocument(
            doc_id=doc_id, text_mode="forced_ocr" if kind == "stamp_only_scan" else "app", **base))
        if kind == "stamp_only_scan":
            documents.append(ExamDocument(doc_id=f"{doc_id}-asread", as_read_today_of=doc_id, **base))
    for doc_id, field, status in (("owner-01", "governing_law", "confirmed"),
                                  ("owner-01", "parties", "unmarked"),
                                  ("owner-02", "governing_law", "corrected")):
        keys.append(OwnerFactKey(
            doc_id=doc_id, field=field, status=status, value="India", quote=None, char_start=None,
            char_end=None, suggested_value=None, suggested_quote=None, suggested_by="m",
            file_sha256=doc_id[-1] * 64, frozen_text_sha256="f" * 64, sheet_sha256="s" * 64,
            loaded_at=stamp))
    return ExamContent(documents=documents, owner_manifest=manifest, owner_keys=keys)


def test_owner_contracts_are_always_direct_and_what_cannot_score_is_left_out_and_counted(committed):
    """FR-19/FR-20/FR-41: the owner's contracts never go in a batch -- not even in a batch run --
    each stamp-only scan brings its 'as read today' twin, and keys that are unmarked, out of date or
    missing leave their contracts out of the plan while the record counts them."""
    content = owner_content()
    plan = runner.plan_run(content, "full", "batch")
    assert [i.item_id for i in plan.items] == ["ex-owner-01", "ex-owner-02", "ex-owner-02-asread"]
    assert {u.delivery for u in plan.units} == {"sync"}
    assert plan.excluded["owner_docs_without_keys"] == ["owner-03"]
    assert plan.excluded["owner_rows_unmarked"] == 1 and plan.excluded["owner_keys_out_of_date"] == []
    stale = [k for k in content.owner_keys if k.doc_id == "owner-02"]
    cut = runner.plan_run(content, "full", "batch", stale_owner_keys=stale)
    assert [i.item_id for i in cut.items] == ["ex-owner-01"]
    assert cut.excluded["owner_keys_out_of_date"] == ["owner-02:governing_law"]
    assert cut.excluded["owner_docs_without_keys"] == ["owner-02", "owner-03"]
    sample = runner.plan_run(content, "sample", "batch")  # one ordinary, one stamp-only + its twin
    assert [i.item_id for i in sample.items] == ["ex-owner-01", "ex-owner-02", "ex-owner-02-asread"]


def test_routing_leaves_the_plan_when_the_workflow_list_has_changed(tmp_path, monkeypatch, exam_engine, live_engine):
    """FR-21/AC-21: picks made against a workflow list that has since changed are not scored --
    the routing items are not even asked -- and the record says routing is out of date."""
    world = build_world(tmp_path, monkeypatch, exam_engine, live_engine, routing=True)
    assert [i.task for i in runner.plan_run(world.content, "full", "batch").items].count("routing") == 1
    plan = runner.plan_run(world.content, "full", "batch", routing_current=False)
    assert "routing" not in count_by_task(plan) and plan.excluded["routing_out_of_date"] is True
    assert runner.plan_run(world.content, "full", "batch").excluded["routing_unpicked"] == 0


def test_the_estimate_is_priced_from_each_requests_real_size_not_an_average():
    """The 40 CUAD contracts average 162,000 characters: what the owner sees must come from the
    requests the app builds. Input is one token per 4 characters (expected) or per 3 (worst case),
    output the call type's expected length (worst case: max_tokens); batch calls are half price;
    Contract Brain answers, which depend on search results rather than document size, carry a
    constant. A portfolio question is a filter call (a constant) plus one table read per contract
    it is expected to read, priced from that contract's real length, all at full price; its worst
    case also holds the search answer it falls back to."""
    model = "claude-sonnet-4-6"

    def request(chars, max_tokens):
        return {"model": model, "max_tokens": max_tokens,
                "messages": [{"role": "user", "content": "x" * chars}]}

    units = (runner.Unit("md-a", "metadata", 1, "a", "batch"), runner.Unit("md-b", "metadata", 1, "b", "batch"),
             runner.Unit("md-own", "metadata", 1, "own", "sync"), runner.Unit("qa-1", "qa", 2, "a", "batch"),
             runner.Unit("pq-1", "portfolio", 2, None, "sync"))
    plan = runner.Plan("full", "batch", (), (), units, (), {})
    small, large, own = request(2_000, 2048), request(40_000, 2048), request(10_000, 2048)
    read_chars = [4_000, 40_000]  # the two contracts the portfolio question is expected to read
    estimate = runner.build_estimate(
        plan, {"md-a": small, "md-b": large, "md-own": own}, model=model, reads={"pq-1": read_chars})

    def tokens(payload):
        return -(-len(json.dumps(payload, separators=(",", ":"), ensure_ascii=False)) // 4)

    def paid(payload, batched):
        return token_cost(tokens(payload), 1_000, model, batched=batched)

    ask_paid = token_cost(5_600, 1_300, model, batched=True)
    filter_paid = token_cost(runner._FILTER_EXPECTED_IN_TOKENS, 250, model)
    reads_paid = [token_cost(n // 4 + runner._READ_PROMPT_TOKENS, 400, model) for n in read_chars]
    assert estimate.usd == pytest.approx(
        paid(small, True) + paid(large, True) + paid(own, False) + ask_paid + filter_paid
        + sum(reads_paid), abs=1e-3)
    # a bigger request costs exactly its extra input more -- not the same as an average
    assert paid(large, True) - paid(small, True) == pytest.approx(
        token_cost(tokens(large) - tokens(small), 0, model, batched=True), abs=1e-9)
    assert estimate.list_usd == pytest.approx(
        paid(small, False) + paid(large, False) + paid(own, False)
        + token_cost(5_600, 1_300, model) + filter_paid + sum(reads_paid), abs=1e-3)
    worst = (runner.worst_cost(small, batched=True) + runner.worst_cost(large, batched=True)
             + runner.worst_cost(own, batched=False)
             + runner.worst_cost(runner._ask_stand_in(model), batched=True)
             + runner.worst_cost(runner._stand_in(model, "portfolio_filter", runner._FILTER_WORST_IN_CHARS), batched=False)
             + sum(runner.worst_cost(runner._stand_in(model, "tabular_cell_extraction", n + runner._READ_PROMPT_CHARS), batched=False)
                   for n in read_chars)
             + runner.worst_cost(runner._ask_stand_in(model), batched=False))  # the search fallback
    assert estimate.worst_case_usd == pytest.approx(worst, abs=1e-3)
    assert (estimate.calls, estimate.batch_calls, estimate.sync_calls) == (7, 3, 4)
    assert estimate.by_call_type["contract_metadata_extraction"]["calls"] == 3
    assert estimate.by_call_type["contract_metadata_extraction"]["batch"] == 2
    assert estimate.by_call_type["contract_brain_answer"]["calls"] == 1
    assert estimate.by_call_type["portfolio_filter"]["calls"] == 1
    assert estimate.by_call_type["tabular_cell_extraction"]["calls"] == 2
    assert estimate.by_call_type["tabular_cell_extraction"]["batch"] == 0  # a portfolio unit is never batched
    fact_only = runner.build_estimate(plan, {"md-a": small, "md-b": large, "md-own": own}, model=model)
    assert "tabular_cell_extraction" not in fact_only.by_call_type  # no reads expected, none priced
    calibrated = runner.build_estimate(
        plan, {"md-a": small, "md-b": large, "md-own": own}, model=model, reads={"pq-1": read_chars},
        per_call={"contract_metadata_extraction": 0.5})
    assert calibrated.by_call_type["contract_metadata_extraction"]["usd"] == pytest.approx(1.5)
    assert calibrated.worst_case_usd == estimate.worst_case_usd  # the bound never shrinks


def test_a_previous_complete_run_of_the_same_kind_calibrates_the_next_estimate(live):
    """Per-call cost, by call type, from the newest COMPLETE run of the same mode and delivery;
    an incomplete or different one is not a guide."""
    def finish(mode, delivery, outcome, usd):
        run = runner.approve_paid_run(
            kind="exam", live=live, owner=OWNER, ask=Typist("5"),
            estimate=runner.Estimate(delivery=delivery, calls=1, batch_calls=1, sync_calls=0, usd=1,
                                     list_usd=2, worst_case_usd=3, by_call_type={}))
        runner.finish_paid_run(
            run, outcome=outcome, reason=None, results_sha256=None, summary={"mode": mode},
            actual={"usd": usd, "by_call_type": {"playbook_review": {"calls": 4, "batch": 4, "usd": usd}}})
        return run.run_id

    assert runner._calibration(live, "full", "batch") == (None, {})
    finish("full", "batch", "incomplete", 9.0)
    finish("sample", "batch", "complete", 9.0)
    finish("full", "sync", "complete", 9.0)
    assert runner._calibration(live, "full", "batch") == (None, {})
    older, newer = finish("full", "batch", "complete", 0.2), finish("full", "batch", "complete", 0.4)
    assert runner._calibration(live, "full", "batch") == (newer, {"playbook_review": 0.1})
    assert older != newer
