"""The exam's throwaway database: rebuild it, seed it, and read documents into it.

specs/005-ai-eval-baseline T023 (FR-5, FR-24, FR-37, FR-38). ``eval/exam/evaldb.py`` is what
stands between an exam run and the real development data, so these tests prove its guards and
its isolation against a real Postgres: every test builds its own ``legal_clm_eval_test_<id>``
database on the test server and drops it afterwards. Nothing here touches ``legal_clm``, a real
``legal_clm_eval`` (the owner's), ``eval/.work`` or ``eval/private``: stored files and "contract"
documents live under ``tmp_path`` and are generated in the test, and no model, OCR or embedding
provider is called.

Building a database from the app's migrations is the slow part, so one migrated database is made
per module and each test's own database is cloned from it (``CREATE DATABASE .. TEMPLATE``).
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
import uuid
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import docx
import fitz
import pytest
from alembic.script import ScriptDirectory
from celery import Celery
from fastapi import HTTPException
from sqlalchemy import create_engine, event, func, inspect, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import NullPool

import app.contract_files.service as contract_files_service
import app.workflows.service as workflows_service
from app.ai import embeddings
from app.auth.models import Role, User
from app.contract_files.models import (
    ContractDocumentElement,
    ContractEmbedding,
    ContractTextSnapshot,
    ContractVersion,
    StorageObject,
)
from app.contract_files.service import ContractFilesService
from app.contracts.lifecycle import ContractLifecycleService
from app.contracts.models import Contract, ContractStageHistory
from app.core.config import settings
from app.core.database import SessionLocal
from app.core.database import engine as app_engine
from app.core.enums import ContractLifecycleStage, PlaybookStatus, UserStatus
from app.core.models import AICallLog, AuditLog, ResourceTimelineEvent
from app.integrations.ocr import OCRResult
from app.integrations.resend import ResendClient
from app.integrations.sendgrid import SendGridClient
from app.integrations.storage import StorageService
from app.jobs.models import JobRun
from app.jobs.service import JobsService
from app.notifications.models import Notification
from app.organizations.models import Organization
from app.playbooks.library import PLAYBOOK_LIBRARY
from app.playbooks.models import Playbook, PlaybookRule, PlaybookVersion
from app.workflows.models import Workflow, WorkflowRun
from app.workflows.service import WorkflowService
from eval.exam import evaldb
from eval.exam import owner as owner_mod
from eval.exam.content import (
    CONTENT_DIR,
    OWNER_FILES,
    ExamContent,
    ExamDocument,
    OwnerFileRefused,
    OwnerManifestEntry,
    SourceCategory,
)
from eval.exam.env import EVAL_DB_NAME, ExamRefused, LiveDB
from eval.exam.evaldb import IngestFailed
from eval.exam.owner import FrozenOCR
from eval.exam.routing import workflow_fingerprint

# What the app was configured with, before any test repoints a setting at an exam database.
LIVE_URL = settings.database_url
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF = "application/pdf"
SENTENCE = "The Supplier shall provide the Services described in each order. "


# --- throwaway databases ---------------------------------------------------------------------


def url_of(name: str) -> str:
    return make_url(LIVE_URL).set(database=name).render_as_string(hide_password=False)


def live_handle() -> LiveDB:
    """The handle ``reset`` takes. It only reads the server URL, so no organisation is needed."""
    return LiveDB(url=LIVE_URL, sessions=SessionLocal, org_id="not-used")


def run_on_server(sql: str) -> None:
    server = create_engine(LIVE_URL, poolclass=NullPool, isolation_level="AUTOCOMMIT")
    try:
        with server.connect() as conn:
            conn.execute(text(sql))
    finally:
        server.dispose()


def drop_database(name: str) -> None:
    # The only DROP in this file: it can never name anything but a test database.
    assert name.startswith(f"{EVAL_DB_NAME}_test_"), name
    run_on_server(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)')


def database_exists(name: str) -> bool:
    server = create_engine(LIVE_URL, poolclass=NullPool)
    try:
        with server.connect() as conn:
            return (
                conn.execute(
                    text("SELECT 1 FROM pg_database WHERE datname = :n"), {"n": name}
                ).scalar()
                is not None
            )
    finally:
        server.dispose()


def rebuild(name: str) -> None:
    """``evaldb.reset`` on a test database; skipped where the database user cannot create one."""
    try:
        evaldb.reset(live_handle(), name)
    except ExamRefused as exc:
        if str(exc).startswith("could not rebuild"):
            pytest.skip(f"this database user cannot CREATE DATABASE ({exc})")
        raise


@contextmanager
def pointed_at(monkeypatch, name: str) -> Iterator[None]:
    """``settings.database_url`` on a test database, which is what ``migrate`` reads, for the
    length of the block only."""
    with monkeypatch.context() as patch:
        patch.setattr(settings, "database_url", url_of(name))
        yield


@pytest.fixture
def exam_name() -> Iterator[str]:
    name = f"{EVAL_DB_NAME}_test_{uuid.uuid4().hex[:10]}"
    yield name
    drop_database(name)


@pytest.fixture(scope="module")
def migrated_template() -> Iterator[str]:
    """One database built by ``reset`` + ``migrate``, copied by every test that needs a schema."""
    name = f"{EVAL_DB_NAME}_test_tpl_{uuid.uuid4().hex[:10]}"
    rebuild(name)
    patch = pytest.MonkeyPatch()
    try:
        with pointed_at(patch, name):
            evaldb.migrate()
        yield name
    finally:
        patch.undo()
        drop_database(name)


@dataclass(frozen=True)
class ExamDB:
    name: str
    sessions: sessionmaker


@pytest.fixture
def exam_db(migrated_template: str, exam_name: str) -> Iterator[ExamDB]:
    """This test's own migrated, empty exam database, with sessions configured like the app's."""
    run_on_server(f'CREATE DATABASE "{exam_name}" TEMPLATE "{migrated_template}"')
    engine = create_engine(url_of(exam_name), poolclass=NullPool)
    yield ExamDB(
        exam_name,
        sessionmaker(autocommit=False, autoflush=False, expire_on_commit=False, bind=engine),
    )
    engine.dispose()


@pytest.fixture
def actors(exam_db: ExamDB) -> evaldb.ExamActors:
    with exam_db.sessions() as db:
        return evaldb.bootstrap(db)


def count(exam_db: ExamDB, model, *where) -> int:
    with exam_db.sessions() as db:
        return db.scalar(select(func.count()).select_from(model).where(*where)) or 0


@contextmanager
def watching_live_engine() -> Iterator[list[str]]:
    """Every connection the app's own engine (the live database) hands out inside the block."""
    seen: list[str] = []

    def on_checkout(*_args) -> None:
        seen.append("checkout")

    event.listen(app_engine, "checkout", on_checkout)
    try:
        yield seen
    finally:
        event.remove(app_engine, "checkout", on_checkout)


# --- reset (FR-37) ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "name",
    [
        "legal_clm",
        "postgres",
        "template1",
        "legal_clm_prod",
        "legal_clm_stage_eval",
        "exam",
        "",
        "LEGAL_CLM_EVAL",
        "legal_clm_eval-x",
        "legal_clm_eval x",
        'legal_clm_eval"; DROP DATABASE legal_clm; --',
    ],
)
def test_reset_refuses_a_name_that_is_not_an_exam_database(name, monkeypatch):
    """FR-37: ``reset`` drops a database, so the name is the only thing between a mistake and the
    development data. Refused before any connection to the server is made: a name that fails
    must never reach DROP DATABASE, and a name with quotes must never reach the statement."""
    connected: list[str] = []

    def no_connection(*args, **kwargs):
        connected.append("engine")
        raise AssertionError("a refused name must not connect to the server")

    monkeypatch.setattr(evaldb, "create_engine", no_connection)
    with pytest.raises(ExamRefused, match="not an exam database name"):
        evaldb.reset(live_handle(), name)
    assert connected == []


def test_reset_refuses_the_live_database_even_under_an_exam_name(monkeypatch):
    """If the app's own DATABASE_URL ever named an exam-style database, ``reset`` would drop the
    live data; the check compares with the live URL, not only the prefix."""
    name = f"{EVAL_DB_NAME}_test_{uuid.uuid4().hex[:10]}"
    live = LiveDB(url=url_of(name), sessions=SessionLocal, org_id="not-used")
    monkeypatch.setattr(evaldb, "create_engine", lambda *a, **k: pytest.fail("connected"))
    with pytest.raises(ExamRefused, match="must not be the live database"):
        evaldb.reset(live, name)


def test_reset_builds_the_database_from_nothing_every_time(exam_name):
    """FR-5: a run may never inherit the last run's rows. A second reset must leave an empty
    database even while a session from the previous run is still open on it (the FORCE)."""
    rebuild(exam_name)
    assert database_exists(exam_name)

    engine = create_engine(url_of(exam_name))
    try:
        held = engine.connect()  # the previous run's pooled connection, still open
        held.execute(text("CREATE TABLE leftover (id int)"))
        held.commit()

        rebuild(exam_name)

        with pytest.raises(SQLAlchemyError):
            held.execute(text("SELECT 1"))
        held.close()
    finally:
        engine.dispose()
    check = create_engine(url_of(exam_name), poolclass=NullPool)
    try:
        with check.connect() as conn:
            tables = conn.execute(
                text("SELECT count(*) FROM pg_tables WHERE schemaname = 'public'")
            ).scalar()
    finally:
        check.dispose()
    assert tables == 0


def test_reset_forgets_the_previous_runs_files_but_only_in_the_exam_work_folder(
    exam_name, tmp_path, monkeypatch
):
    """The previous run's contract bytes belong to a database that no longer exists, so they go;
    a storage root anywhere else (the app's real contract files) must never be emptied."""
    work = tmp_path / "work"
    monkeypatch.setattr(evaldb, "WORK_DIR", work)

    exam_storage = work / "storage"
    exam_storage.mkdir(parents=True)
    (exam_storage / "old-run.pdf").write_bytes(b"left by the last run")
    monkeypatch.setattr(settings, "storage_root", exam_storage)
    rebuild(exam_name)
    assert exam_storage.is_dir() and list(exam_storage.iterdir()) == []

    real_storage = tmp_path / "app-storage"
    real_storage.mkdir()
    (real_storage / "customer.pdf").write_bytes(b"a real contract")
    monkeypatch.setattr(settings, "storage_root", real_storage)
    rebuild(exam_name)
    assert (real_storage / "customer.pdf").read_bytes() == b"a real contract"


def test_reset_says_why_the_server_refused_without_printing_the_password():
    """A server that cannot be reached must fail as a plain refusal (exit code 2), and the reason
    comes from the driver's one-line message, never from the URL that carries the password."""
    live = LiveDB(
        url="postgresql+psycopg://legal_clm:pw-placeholder@127.0.0.1:1/legal_clm",
        sessions=SessionLocal,
        org_id="not-used",
    )
    with pytest.raises(ExamRefused, match="could not rebuild the exam database") as caught:
        evaldb.reset(live, f"{EVAL_DB_NAME}_test_unreachable")
    assert "pw-placeholder" not in str(caught.value)


# --- migrate (FR-5) --------------------------------------------------------------------------


@pytest.mark.parametrize("name", ["legal_clm", "legal_clm_other", "postgres"])
def test_migrate_refuses_a_database_that_is_not_an_exam_database(name, monkeypatch):
    """``migrate`` runs whatever ``settings.database_url`` says; without its guard a process that
    never called ``isolate()`` would upgrade the live schema."""
    ran: list[str] = []

    class NoAlembic:
        @staticmethod
        def upgrade(*args, **kwargs):
            ran.append("upgrade")

    monkeypatch.setattr(evaldb, "command", NoAlembic)
    monkeypatch.setattr(settings, "database_url", url_of(name))
    with pytest.raises(ExamRefused, match="not an exam database"):
        evaldb.migrate()
    assert ran == []


def test_migrate_builds_the_schema_to_the_single_head_and_keeps_the_apps_loggers(
    exam_name, monkeypatch
):
    """FR-5: the exam database is built by the app's own migrations, to the one head -- two heads
    would make ``upgrade head`` ambiguous. It must also leave logging alone: alembic's env.py
    calls ``fileConfig`` when an ini file is given, which would silence every logger the app has
    already made and hide the app's own warnings from the run."""
    heads = ScriptDirectory(str(evaldb.APP_DIR.parent / "alembic")).get_heads()
    assert len(heads) == 1

    def live_version() -> str | None:
        with SessionLocal() as db:
            return db.execute(text("SELECT version_num FROM alembic_version")).scalar()

    live_before = live_version()
    probe = logging.getLogger("app.test_exam_evaldb.probe")
    probe.disabled = False
    rebuild(exam_name)

    with pointed_at(monkeypatch, exam_name):
        evaldb.migrate()
        evaldb.migrate()  # already at head: a second run must be a no-op, not an error

    check = create_engine(url_of(exam_name), poolclass=NullPool)
    try:
        with check.connect() as conn:
            version = conn.execute(text("SELECT version_num FROM alembic_version")).scalar()
        assert {"contract", "workflow", "a_i_call_log"} <= set(inspect(check).get_table_names())
    finally:
        check.dispose()
    assert version == heads[0]
    assert probe.disabled is False
    assert live_version() == live_before  # the live schema was not part of this


# --- bootstrap (FR-9, FR-24) -----------------------------------------------------------------


def test_bootstrap_creates_the_exam_organisation_and_its_two_users(exam_db):
    """The runner (an admin) uploads and reviews everything; the portfolio user is a legal
    reviewer, not an admin, because admins see every contract and the portfolio questions are only
    meaningful for someone whose access is scoped. It must hold the AI tools permission, or a
    question that needs the AI to read each contract is refused reading. Both must hold their role's permissions after the session is
    gone: the actors are handed to every unit of the run, each in its own session."""
    with exam_db.sessions() as db:
        actors = evaldb.bootstrap(db)

    assert inspect(actors.runner).detached and inspect(actors.portfolio).detached
    assert actors.runner.email == evaldb.RUNNER_EMAIL
    assert actors.portfolio.email == evaldb.PORTFOLIO_EMAIL
    assert "admin_panel:access" in actors.runner.permission_values  # no session: already loaded
    assert actors.portfolio.permission_values
    assert "admin_panel:access" not in actors.portfolio.permission_values
    assert "assistant:use_ai_tools" in actors.portfolio.permission_values

    with exam_db.sessions() as db:
        orgs = db.scalars(select(Organization)).all()
        assert [(o.id, o.slug, o.setup_complete) for o in orgs] == [
            (actors.org_id, "ai-exam", True)
        ]
        users = db.scalars(select(User)).all()
        assert {u.email for u in users} == {evaldb.RUNNER_EMAIL, evaldb.PORTFOLIO_EMAIL}
        assert {u.org_id for u in users} == {actors.org_id}
        assert all(u.status == UserStatus.ACTIVE for u in users)
        role_of = {u.email: db.get(Role, u.active_role_id).name for u in users}
        assert role_of == {evaldb.RUNNER_EMAIL: "admin", evaldb.PORTFOLIO_EMAIL: "legal_reviewer"}
        # Their passwords are random and thrown away: nobody can log in as either.
        assert users[0].hashed_password != users[1].hashed_password


def test_bootstrap_on_a_database_that_already_has_an_organisation_fails_loudly(exam_db, actors):
    """The exam must start from a reset database. A second bootstrap would otherwise quietly
    reuse the first run's organisation, so it is the app's own 409, with nothing added."""
    with exam_db.sessions() as db, pytest.raises(HTTPException) as caught:
        evaldb.bootstrap(db)
    assert caught.value.status_code == 409
    assert count(exam_db, Organization) == 1
    assert count(exam_db, User) == 2


# --- the write guard (FR-37, FR-38) ------------------------------------------------------------


@pytest.mark.parametrize("function", ["bootstrap", "load_playbooks", "load_workflows", "ingest"])
async def test_every_writer_refuses_a_session_that_is_not_on_an_exam_database(function, tmp_path):
    """FR-37/FR-38: a mistake in a caller (the app's own session instead of the exam's) must
    never put an exam organisation, playbook, workflow or contract into the development data.
    The refusal comes before anything is read or written: no connection is even taken."""
    nobody = evaldb.ExamActors(org_id="org", runner=None, portfolio=None)  # never reached
    document = ExamDocument(
        doc_id="t23-live",
        category=SourceCategory.CONSTRUCTED,
        dataset=None,
        file="never-read.docx",
        mime_type=DOCX,
        file_sha256="0" * 64,
        title="never stored",
        ai_steps=[],
    )
    with SessionLocal() as live_session, watching_live_engine() as checkouts:
        with pytest.raises(ExamRefused, match="refusing to write to database"):
            if function == "bootstrap":
                evaldb.bootstrap(live_session)
            elif function == "load_playbooks":
                evaldb.load_playbooks(live_session, ExamContent(), nobody)
            elif function == "load_workflows":
                evaldb.load_workflows(live_session, {"workflows": [], "fingerprint": ""}, nobody)
            else:
                await evaldb.ingest_document(live_session, document, actors=nobody, ocr=FrozenOCR())
        assert checkouts == []


# --- load_playbooks (FR-5, FR-24) --------------------------------------------------------------


def committed_playbooks() -> dict[str, dict]:
    return {
        path.stem: json.loads(path.read_text(encoding="utf-8"))
        for path in sorted((CONTENT_DIR / "playbooks").glob("*.json"))
    }


def test_load_playbooks_publishes_the_library_and_the_exams_own(exam_db, actors):
    """The exam reviews against the app's playbook library plus its own checklists, all
    published, and names a playbook either by the exam key or by the library name. The committed
    playbooks are loaded as they are, so a rule key the model does not have fails here."""
    content = ExamContent(playbooks=committed_playbooks())
    assert {"cp_msa", "cp_dpa", "nda_checklist", "cuad_checklist"} <= set(content.playbooks)

    with exam_db.sessions() as db:
        ids = evaldb.load_playbooks(db, content, actors)

    library_names = {spec["name"] for spec in PLAYBOOK_LIBRARY}
    assert set(ids) == library_names | set(content.playbooks)
    with exam_db.sessions() as db:
        for key, spec in content.playbooks.items():
            playbook = db.get(Playbook, ids[key])
            assert (playbook.name, playbook.org_id) == (spec["name"], actors.org_id)
            assert playbook.status == PlaybookStatus.PUBLISHED
            versions = db.scalars(
                select(PlaybookVersion).where(PlaybookVersion.playbook_id == playbook.id)
            ).all()
            assert [(v.version_number, v.status) for v in versions] == [
                (1, PlaybookStatus.PUBLISHED)
            ]
            assert playbook.current_version_id == versions[0].id
            assert versions[0].source_metadata["key"] == key
            rules = db.scalars(
                select(PlaybookRule)
                .where(PlaybookRule.playbook_version_id == versions[0].id)
                .order_by(PlaybookRule.created_at)
            ).all()
            # The review prompt lists rules in created_at order: it must be the file's order,
            # strictly, or two runs would send different prompts for the same playbook.
            assert [r.clause_type for r in rules] == [r["clause_type"] for r in spec["rules"]]
            assert len({r.created_at for r in rules}) == len(rules)
            assert [r.risk_level for r in rules] == [r["risk_level"] for r in spec["rules"]]
        for name in library_names:
            assert db.get(Playbook, ids[name]).name == name


def test_load_playbooks_without_exam_playbooks_still_seeds_the_library(exam_db, actors):
    """An exam whose content has no playbook of its own (a sample of library items only) must
    still find every library playbook, by name."""
    with exam_db.sessions() as db:
        ids = evaldb.load_playbooks(db, ExamContent(), actors)
    assert set(ids) == {spec["name"] for spec in PLAYBOOK_LIBRARY}


# --- load_workflows (FR-24) --------------------------------------------------------------------


def make_snapshot(*, enabled_only=False) -> dict:
    def workflow(name: str, order: int, *, enabled: bool = True) -> dict:
        return {
            "id": str(uuid.uuid4()),
            "name": name,
            "description": f"{name} (test)",
            "enabled": enabled,
            "is_builtin": order % 2 == 0,
            "eval_order": order,
            "version": order + 1,
            "criteria": {"used_for": [{"form": "nda", "agreement_type": None}]},
            "steps": [{"id": "s1", "type": "approval", "name": "Approve", "config": {}}],
        }

    workflows = [
        workflow("Vendor contract", 20),
        workflow("NDA", 10),
        workflow("Another NDA route", 10),
        workflow("Retired route", 30, enabled=False),
    ]
    if enabled_only:
        workflows = [w for w in workflows if w["enabled"]]
    return {
        "fingerprint": workflow_fingerprint(workflows),
        "taken_at": datetime.now(UTC).isoformat(),
        "workflows": workflows,
    }


def test_load_workflows_copies_the_snapshots_workflows_ids_included(exam_db, actors):
    """The owner's routing picks name these workflow ids, and the intake triage chooses among
    exactly the enabled workflows it finds in the exam database. They are copied, never started,
    and read back to the same fingerprint the picks were made against."""
    snapshot = make_snapshot()
    with exam_db.sessions() as db:
        evaldb.load_workflows(db, snapshot, actors)

    wanted = {w["id"]: w for w in snapshot["workflows"]}
    with exam_db.sessions() as db:
        rows = db.scalars(select(Workflow).order_by(Workflow.created_at)).all()
        assert {r.id for r in rows} == set(wanted)
        for row in rows:
            source = wanted[row.id]
            assert (row.name, row.enabled, row.is_builtin, row.eval_order, row.version) == (
                source["name"],
                source["enabled"],
                source["is_builtin"],
                source["eval_order"],
                source["version"],
            )
            assert (row.criteria, row.steps) == (source["criteria"], source["steps"])
            assert (row.org_id, row.created_by_user_id) == (actors.org_id, actors.runner.id)
        # The app tries workflows in eval_order, then name, then id: a stable load order.
        by_rule = sorted(snapshot["workflows"], key=lambda w: (w["eval_order"], w["name"], w["id"]))
        assert [r.id for r in rows] == [w["id"] for w in by_rule]
    assert count(exam_db, WorkflowRun) == 0  # copied, not run


def test_load_workflows_refuses_a_list_that_is_not_the_one_the_picks_were_made_on(exam_db, actors):
    """FR-24: routing is scored against the workflows the owner saw. If what loads back does not
    reproduce the snapshot's fingerprint (an edited snapshot, a changed rule), the exam stops and
    leaves no workflows behind rather than routing against a list nobody confirmed."""
    snapshot = make_snapshot()
    snapshot["fingerprint"] = "0" * 64
    with exam_db.sessions() as db, pytest.raises(ExamRefused, match="do not match"):
        evaldb.load_workflows(db, snapshot, actors)
    assert count(exam_db, Workflow) == 0


# --- ingest_document (FR-38) -------------------------------------------------------------------


def docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Supply Agreement", level=1)
    for n in range(1, 7):
        document.add_paragraph(f"{n}. " + SENTENCE * 3)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def blank_pdf() -> bytes:
    """One page with no text layer at all: the app would ask OCR to read it."""
    pdf = fitz.open()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    return data


@dataclass
class Workspace:
    """Where an ingest test's documents and stored files live: all under ``tmp_path``."""

    docs: Path
    work: Path
    storage: Path

    def stored(self) -> list[Path]:
        return [p for p in self.storage.rglob("*") if p.is_file()]

    def document(self, name: str, data: bytes, **overrides) -> ExamDocument:
        (self.docs / name).write_bytes(data)
        fields = dict(
            doc_id=f"t23-{uuid.uuid4().hex[:8]}",
            category=SourceCategory.CONSTRUCTED,
            dataset=None,
            file=name,
            mime_type=PDF if name.endswith(".pdf") else DOCX,
            file_sha256=hashlib.sha256(data).hexdigest(),
            title=f"T023 agreement {uuid.uuid4().hex[:8]}",
            ai_steps=["metadata"],
        )
        fields.update(overrides)
        return ExamDocument(**fields)


@pytest.fixture
def workspace(tmp_path, monkeypatch) -> Workspace:
    """The exam's folders moved under ``tmp_path``, and embeddings made local and instant: the
    document is really chunked and written to pgvector, but no model is loaded or downloaded."""
    docs, work = tmp_path / "contracts", tmp_path / "work"
    docs.mkdir()
    storage = StorageService(root=work / "storage")
    monkeypatch.setattr(evaldb, "DOCS_DIR", docs)
    monkeypatch.setattr(evaldb, "WORK_DIR", work)
    monkeypatch.setattr(contract_files_service, "storage_service", storage)
    monkeypatch.setattr(
        embeddings,
        "_embed",
        lambda texts: [embeddings._deterministic_mock_vector(t) for t in texts],
    )
    return Workspace(docs=docs, work=work, storage=storage.root)


@pytest.fixture
def side_effects(monkeypatch) -> list[str]:
    """Everything the upload pipeline starts after it persists a contract, replaced by a
    recorder: Celery sends, job creation and dispatch, the stage-entry triggers, workflow starts,
    and every e-mail transport. Recorders, not failures: the pipeline's own best-effort blocks
    swallow exceptions, which would hide a call from the test."""
    calls: list[str] = []

    def record(name: str):
        def spy(*args, **kwargs):
            calls.append(name)

        return spy

    def record_async(name: str):
        async def spy(*args, **kwargs):
            calls.append(name)

        return spy

    monkeypatch.setattr(Celery, "send_task", record("celery.send_task"))
    monkeypatch.setattr(JobsService, "create_job", record("create_job"))
    monkeypatch.setattr(JobsService, "dispatch_job", record("dispatch_job"))
    monkeypatch.setattr(
        ContractLifecycleService, "transition_contract_stage", record("transition_contract_stage")
    )
    for name in ("start_flow", "start_flow_for_contract"):
        monkeypatch.setattr(workflows_service, name, record_async(name))
        monkeypatch.setattr(WorkflowService, name, record_async(f"WorkflowService.{name}"))
    monkeypatch.setattr(
        workflows_service, "advance_flow_for_contract", record("advance_flow_for_contract")
    )
    monkeypatch.setattr(ResendClient, "send_email", record_async("resend.send_email"))
    monkeypatch.setattr(SendGridClient, "send_email", record_async("sendgrid.send_email"))
    return calls


async def ingest(exam_db: ExamDB, document, actors, ocr=None):
    with exam_db.sessions() as db:
        return await evaldb.ingest_document(db, document, actors=actors, ocr=ocr or FrozenOCR())


async def test_ingest_document_stores_the_document_the_way_an_upload_does(
    exam_db, actors, workspace, side_effects
):
    """FR-5: the exam must read a document with the app's own reader and write the app's own row
    chain (storage object, contract, version, text snapshot, document elements, embeddings), or
    its scores say nothing about the app. The contract lands at Review, as an upload does."""
    data = docx_bytes()
    document = workspace.document("supply.docx", data)

    result = await ingest(exam_db, document, actors)

    assert result.doc_id == document.doc_id
    assert result.elements > 0 and result.frozen_text_matches is None
    with exam_db.sessions() as db:
        contract = db.get(Contract, result.contract_id)
        assert contract.title == document.title
        assert contract.lifecycle_stage == ContractLifecycleStage.REVIEW
        assert contract.org_id == actors.org_id
        assert contract.owner_user_id == actors.runner.id
        assert contract.current_authoritative_version_id == result.version_id

        version = db.get(ContractVersion, result.version_id)
        assert version.contract_id == contract.id and version.is_authoritative
        assert version.text_snapshot_id == result.snapshot_id
        snapshot = db.get(ContractTextSnapshot, result.snapshot_id)
        assert snapshot.extraction_method == result.method
        assert len(snapshot.text) == result.text_len
        assert hashlib.sha256(snapshot.text.encode()).hexdigest() == result.text_sha256
        assert "The Supplier shall provide the Services" in snapshot.text
        assert snapshot.element_count == result.elements
        assert count(exam_db, ContractDocumentElement) == result.elements

        storage_object = db.get(StorageObject, version.storage_object_id)
        assert storage_object.sha256_hash == hashlib.sha256(data).hexdigest()
        # The bytes sit in the exam's own storage, and only there.
        assert [p.read_bytes() for p in workspace.stored()] == [data]

        chunks = db.scalars(select(ContractEmbedding)).all()
        assert chunks and {c.contract_id for c in chunks} == {contract.id}
        # Clause-aligned chunks, not fixed windows: the document elements were already visible
        # to the embedding step, which is why the flush before it matters.
        assert {c.metadata_json["chunk_source"] for c in chunks} == {"element"}


async def test_ingest_document_files_a_portfolio_document_under_the_portfolio_user(
    exam_db, actors, workspace, side_effects
):
    """Portfolio questions are answered through the app's own access rules, so a document the
    record says the portfolio user uploaded must be owned by that member, not by the admin."""
    document = workspace.document("portfolio.docx", docx_bytes(), uploader="portfolio")
    result = await ingest(exam_db, document, actors)
    with exam_db.sessions() as db:
        owner = db.get(User, db.get(Contract, result.contract_id).owner_user_id)
    assert owner.email == evaldb.PORTFOLIO_EMAIL


async def test_ingest_document_starts_nothing_and_never_touches_the_live_database(
    exam_db, actors, workspace, side_effects
):
    """FR-38: reading a document in must not queue jobs, advance a stage (whose triggers start
    jobs and workflows), send a notification or e-mail, or start a workflow run, and nothing may
    reach the live database. Checked at the lowest points a start passes through (the Celery
    send, job creation and dispatch, the transition, workflow starts, mail transports), by the
    rows such a start leaves in the exam database, and by every connection of the app's own
    engine, which is the live one in this process."""
    document = workspace.document("quiet.docx", docx_bytes())

    with watching_live_engine() as live_connections:
        result = await ingest(exam_db, document, actors)

    assert side_effects == []
    assert live_connections == []
    for model in (JobRun, Notification, WorkflowRun, ContractStageHistory, AICallLog):
        assert count(exam_db, model) == 0, model.__name__
    for model in (AuditLog, ResourceTimelineEvent):
        assert count(exam_db, model, model.resource_id == result.contract_id) == 0
    # And afterwards: the live database has never heard of this contract.
    with SessionLocal() as live:
        assert live.get(Contract, result.contract_id) is None
        assert live.get(ContractVersion, result.version_id) is None
        assert live.get(ContractTextSnapshot, result.snapshot_id) is None
        assert (
            live.scalar(select(func.count(Contract.id)).where(Contract.title == document.title))
            == 0
        )


async def test_the_recorders_see_what_the_upload_pipeline_does_after_persisting(
    exam_db, actors, workspace, side_effects
):
    """Control for the test above: the recorders are only evidence if they would have caught the
    call. The pipeline's own advance step (what ``ingest_document`` replaces by setting the stage
    directly) goes through the stage transition, and the recorder sees it."""
    result = await ingest(exam_db, workspace.document("control.docx", docx_bytes()), actors)
    assert side_effects == []
    with exam_db.sessions() as db:
        contract = db.get(Contract, result.contract_id)
        contract.lifecycle_stage = ContractLifecycleStage.INTAKE
        ContractFilesService(db)._advance_upload_to_review(
            contract=contract, user=actors.runner, request_id=None
        )
    assert side_effects == ["transition_contract_stage"]


# --- ingest_document: refusals and failures (FR-37, FR-38, FR-40) ----------------------------


def nothing_stored(exam_db: ExamDB, workspace: Workspace) -> None:
    """A failed document leaves no bytes on disk and no row of the chain in the database."""
    assert workspace.stored() == []
    for model in (
        Contract,
        ContractVersion,
        ContractTextSnapshot,
        StorageObject,
        ContractEmbedding,
    ):
        assert count(exam_db, model) == 0, model.__name__


async def test_ingest_document_refuses_to_store_bytes_outside_the_exam_work_folder(
    exam_db, actors, workspace, side_effects, tmp_path, monkeypatch
):
    """FR-36/FR-37: exam contract bytes may only be written under eval/.work (git- and
    docker-ignored). A process that never set the storage location (the app's real contract
    volume) is refused before a single byte is written."""
    elsewhere = StorageService(root=tmp_path / "the-apps-contract-volume")
    monkeypatch.setattr(contract_files_service, "storage_service", elsewhere)
    document = workspace.document("elsewhere.docx", docx_bytes())

    with pytest.raises(ExamRefused, match=r"backend/eval/\.work"):
        await ingest(exam_db, document, actors)

    assert [p for p in elsewhere.root.rglob("*") if p.is_file()] == []
    nothing_stored(exam_db, workspace)


@pytest.mark.parametrize(
    ("problem", "reason"),
    [
        ("missing", "file is missing; run `build`"),
        ("edited", "sha256 differs"),
        ("outside", "outside the exam's contracts folder"),
    ],
)
async def test_ingest_document_fails_a_file_that_is_not_the_one_the_exam_recorded(
    problem, reason, exam_db, actors, workspace, side_effects
):
    """FR-1/FR-27: the record's sha256 is part of the content fingerprint, so a regenerated or
    edited file must fail loudly (and only its own items) instead of being scored under a
    fingerprint it does not belong to, and a record can never point outside the exam's folder."""
    document = workspace.document("recorded.docx", docx_bytes())
    if problem == "missing":
        (workspace.docs / document.file).unlink()
    elif problem == "edited":
        (workspace.docs / document.file).write_bytes(docx_bytes() + b"edited")
    else:
        secret = workspace.docs.parent / "elsewhere.docx"
        secret.write_bytes(docx_bytes())
        document = document.model_copy(
            update={
                "file": "../elsewhere.docx",
                "file_sha256": hashlib.sha256(secret.read_bytes()).hexdigest(),
            }
        )

    with pytest.raises(IngestFailed, match=document.doc_id) as caught:
        await ingest(exam_db, document, actors)

    assert reason in caught.value.reason
    nothing_stored(exam_db, workspace)


async def test_ingest_document_says_to_run_owner_import_when_the_owner_file_is_not_there(
    exam_db, actors, workspace, tmp_path, monkeypatch
):
    """FR-15: an owner document is read only through the allow-list, from the private folder.
    Not imported yet is a plain instruction, not a crash; a name outside the 14 is refused
    before any file is touched."""
    monkeypatch.setattr(evaldb, "_OWNER_FILES_DIR", tmp_path / "private-owner-files")
    owner_doc = workspace.document(
        "unused.docx",
        docx_bytes(),
        doc_id="owner-12",
        category=SourceCategory.OWNER,
        file=OWNER_FILES[11],
    )
    with pytest.raises(IngestFailed, match="run `owner-import`"):
        await ingest(exam_db, owner_doc, actors)

    outsider = owner_doc.model_copy(update={"file": "Tax return 2025.pdf"})
    with pytest.raises(OwnerFileRefused):
        await ingest(exam_db, outsider, actors)
    nothing_stored(exam_db, workspace)


async def test_ingest_document_leaves_nothing_behind_when_the_text_cannot_be_read(
    exam_db, actors, workspace, side_effects
):
    """The app swallows an OCR failure and keeps the thin native text; the exam must not score
    that as a read. A scan with no frozen OCR fails its own items (telling the owner to run `ocr`),
    and the upload pipeline's rule holds: once bytes are on disk, any failure deletes them."""
    document = workspace.document("scan.pdf", blank_pdf())

    with pytest.raises(IngestFailed, match="none is frozen; run `ocr` first"):
        await ingest(exam_db, document, actors)

    nothing_stored(exam_db, workspace)
    assert side_effects == []


async def test_ingest_document_leaves_no_contract_behind_when_embedding_fails(
    exam_db, actors, workspace, side_effects, monkeypatch
):
    """The row chain and its embeddings are one transaction: a document that cannot be embedded
    must not exist, or a failed document could turn up in a portfolio answer with no chunks."""

    def broken(texts):
        raise RuntimeError("embedding model unavailable")

    monkeypatch.setattr(embeddings, "_embed", broken)
    document = workspace.document("unembeddable.docx", docx_bytes())

    with pytest.raises(IngestFailed, match="RuntimeError: embedding model unavailable"):
        await ingest(exam_db, document, actors)

    nothing_stored(exam_db, workspace)
    assert side_effects == []


class StubOCR:
    """Stands in for the frozen OCR cache: one canned text, and a count of how often it is asked."""

    provider = "reducto"

    def __init__(self, text: str) -> None:
        self.text, self.calls = text, 0

    async def extract_text(self, *, filename: str, mime_type: str, content: bytes) -> OCRResult:
        self.calls += 1
        return OCRResult(
            text=self.text,
            provider="reducto",
            quality_score=0.9,
            metadata={},
            elements=[{"type": "text", "content": self.text}],
        )


async def test_ingest_document_reads_a_stamp_only_scan_from_the_frozen_ocr_text(
    exam_db, actors, workspace, side_effects
):
    """A stamp-only scan has a one-line e-signature text layer, so the app would not ask for OCR;
    ``text_mode="forced_ocr"`` makes the exam read its frozen OCR text instead, the way the app
    reads an OCR'd upload, and only from the injected cache."""
    ocr_text = "1. " + SENTENCE * 6
    ocr = StubOCR(ocr_text)
    document = workspace.document("stamped.pdf", blank_pdf(), text_mode="forced_ocr")

    result = await ingest(exam_db, document, actors, ocr=ocr)

    assert ocr.calls == 1
    assert result.method == "reducto_ocr" and result.text_len == len(ocr_text)
    with exam_db.sessions() as db:
        snapshot = db.get(ContractTextSnapshot, result.snapshot_id)
        assert (snapshot.text, snapshot.ocr_provider) == (ocr_text, "reducto")
    assert side_effects == []


async def test_ingest_document_says_whether_an_owner_documents_frozen_text_still_matches(
    exam_db, actors, workspace, tmp_path, monkeypatch
):
    """FR-20: the owner's answers were confirmed against one text. If the app now reads the file
    differently, the keys are out of date, and the run must be able to tell: True when the text
    read equals the frozen one, False when it does not, and None for the "as read today" twin,
    whose text differs on purpose."""
    private = tmp_path / "private-owner-files"
    private.mkdir()
    data = docx_bytes()
    (private / OWNER_FILES[11]).write_bytes(data)
    monkeypatch.setattr(evaldb, "_OWNER_FILES_DIR", private)

    def entry(frozen_text_sha256: str | None) -> list[OwnerManifestEntry]:
        return [
            OwnerManifestEntry(
                doc_id="owner-12",
                name=OWNER_FILES[11],
                file_sha256=hashlib.sha256(data).hexdigest(),
                mime_type=DOCX,
                pages=1,
                kind="text",
                frozen_text_sha256=frozen_text_sha256,
                imported_at=datetime.now(UTC),
            )
        ]

    owner_doc = workspace.document(
        "unused.docx",
        data,
        doc_id="owner-12",
        category=SourceCategory.OWNER,
        file=OWNER_FILES[11],
    )
    monkeypatch.setattr(owner_mod, "read_manifest", lambda: entry("0" * 64))
    stale = await ingest(exam_db, owner_doc, actors)
    assert stale.frozen_text_matches is False

    monkeypatch.setattr(owner_mod, "read_manifest", lambda: entry(stale.text_sha256))
    fresh = await ingest(exam_db, owner_doc, actors)
    assert fresh.frozen_text_matches is True

    twin = owner_doc.model_copy(
        update={"doc_id": "owner-12-asread", "as_read_today_of": "owner-12"}
    )
    assert (await ingest(exam_db, twin, actors)).frozen_text_matches is None
