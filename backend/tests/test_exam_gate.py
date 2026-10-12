"""The exam's spending gate: who may start a paid run, the go-ahead, the ceiling, the run record.

specs/005-ai-eval-baseline T013 (FR-31, FR-32, FR-34, FR-39, FR-40, FR-41; AC-31, AC-32, AC-34,
AC-39). T015 adds the owner-command checks (FR-3, FR-15, FR-16, FR-17, FR-20, FR-23, FR-35;
AC-15..AC-20, AC-35): frozen OCR, the app's own read path, import, the one-time OCR, the pre-fill
and the key round trip. Nothing here can spend money or touch real data: the Anthropic client is
the fake from ``tests/exam_fake_anthropic.py``, the OCR provider is a fake function, every
"contract" is generated in the test, private files live under ``tmp_path``, and the "live
database" is a scratch Postgres database holding only the four tables the gate writes (audit_log,
a_i_call_log, user, role), created per module and dropped afterwards, so the shared dev database
never receives an ``eval.*`` audit row.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import math
import os
import re
import sys
import uuid
from contextlib import contextmanager
from datetime import UTC, datetime
from pathlib import Path

import docx
import fitz
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import NullPool

import app.models  # noqa: F401  (registers every mapper, so metadata has the tables below)
from app.auth.models import User
from app.core.audit import verify_audit_hash_chain
from app.core.config import settings
from app.core.database import Base, new_uuid
from app.core.models import AICallLog, AuditLog
from app.core.security import hash_password
from app.integrations import claude
from app.integrations.claude import ClaudeClient
from app.integrations.ocr import OCRResult
from app.integrations.reducto import reducto_client
from eval.exam import content as content_mod
from eval.exam import env, runner, sheets
from eval.exam import owner as owner_mod
from eval.exam.content import (
    OWNER_FILES,
    PREFILL_MODEL,
    ExamContent,
    ExamDocument,
    OwnerManifestEntry,
    app_code_sha256,
    content_fingerprint,
)
from eval.exam.env import ExamRefused, LiveDB, Owner
from tests import exam_fake_anthropic
from tests.exam_fake_anthropic import batch_result, install

# The tables the gate and the ledger copy touch (user -> role is its only foreign key).
SCRATCH_TABLES = ("role", "user", "audit_log", "a_i_call_log")
PASSWORD = "a-test-password"
PASSWORD_HASH = hash_password(PASSWORD)


@contextmanager
def scratch_database():
    """A throwaway Postgres database on the test server with ``SCRATCH_TABLES`` in it.

    Only its own tables, and dropped on the way out: the isolation guarantees under test (what
    reaches the "live" database) are checked against a database that is not the shared dev one.
    """
    base = make_url(settings.database_url)
    name = f"legal_clm_scratch_{uuid.uuid4().hex[:10]}"
    admin = create_engine(
        base.set(database="postgres"), isolation_level="AUTOCOMMIT", poolclass=NullPool
    )
    try:
        with admin.connect() as conn:
            conn.execute(text(f'CREATE DATABASE "{name}"'))
    except SQLAlchemyError as exc:
        admin.dispose()
        pytest.skip(f"this database user cannot CREATE DATABASE ({type(exc).__name__})")
    engine = create_engine(base.set(database=name), poolclass=NullPool)
    try:
        Base.metadata.create_all(engine, tables=[Base.metadata.tables[t] for t in SCRATCH_TABLES])
        yield engine
    finally:
        engine.dispose()
        with admin.connect() as conn:
            conn.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
        admin.dispose()


def make_live(engine) -> LiveDB:
    """A ``LiveDB`` over a scratch engine, with its own org id so tests never see each other."""
    return LiveDB(
        url=engine.url.render_as_string(hide_password=False),
        sessions=sessionmaker(bind=engine, autoflush=False, expire_on_commit=False),
        org_id=new_uuid(),
    )


@pytest.fixture(scope="module")
def scratch_engine():
    with scratch_database() as engine:
        yield engine


@pytest.fixture
def live(scratch_engine) -> LiveDB:
    return make_live(scratch_engine)


def audit_rows(live: LiveDB, action: str | None = None) -> list[AuditLog]:
    with live.sessions() as db:
        query = select(AuditLog).where(AuditLog.org_id == live.org_id)
        if action:
            query = query.where(AuditLog.action == action)
        return list(db.scalars(query.order_by(AuditLog.created_at, AuditLog.id)))


OWNER = Owner(user_id="owner-user-id", email="owner@example.com")


def make_estimate(**changes) -> runner.Estimate:
    values = dict(
        delivery="batch",
        calls=486,
        batch_calls=454,
        sync_calls=32,
        usd=9.75,
        list_usd=18.07,
        worst_case_usd=24.5,
        by_call_type={
            "playbook_review": {"calls": 104, "batch": 104, "usd": 2.06},
            "contract_metadata_extraction": {"calls": 90, "batch": 80, "usd": 1.2},
        },
    )
    values.update(changes)
    return runner.Estimate(**values)


class Typist:
    """Plays the owner at the keyboard: records every prompt and answers the ceiling, which is
    the go-ahead (an empty line refuses)."""

    def __init__(self, ceiling: str = "25", *, approve: bool = True) -> None:
        self.ceiling, self.approve = ceiling, approve
        self.prompts: list[str] = []
        self.run_ids: list[str] = []

    def __call__(self, prompt: str) -> str:
        self.prompts.append(prompt)
        if found := re.search(r"for run (\S+) in", prompt):
            self.run_ids.append(found.group(1))
        return self.ceiling if self.approve else ""


def approve(live, *, ask, kind="exam", estimate=None, extra=None):
    return runner.approve_paid_run(
        kind=kind, live=live, owner=OWNER, estimate=estimate or make_estimate(), ask=ask, extra=extra
    )


# --- custom_id and worst_cost (FR-41) -------------------------------------------------------

PAYLOAD = {
    "model": "claude-sonnet-4-6",
    "max_tokens": 1000,
    "temperature": 0.2,
    "system": [{"type": "text", "text": "You review contracts.", "cache_control": {"type": "ephemeral"}}],
    "messages": [{"role": "user", "content": "x" * 2999 + "é"}],
}


def test_custom_id_is_a_hash_of_the_exact_request():
    """A stored batch answer must only ever go back to the identical request: the id has to be
    the same for the same payload however its keys are ordered, and differ for any change."""
    reordered = dict(reversed(list(PAYLOAD.items())))
    changed = {**PAYLOAD, "messages": [{"role": "user", "content": "x" * 3000}]}
    assert runner.custom_id(PAYLOAD) == runner.custom_id(reordered)
    assert runner.custom_id(PAYLOAD) != runner.custom_id(changed)
    assert re.fullmatch(r"[0-9a-f]{64}", runner.custom_id(PAYLOAD))  # the API's 64-character limit


def test_worst_cost_assumes_a_full_length_answer_and_halves_for_batches():
    """The ceiling is a hard limit only because admission uses the most a call can cost: input
    at one token per 3 characters plus an answer of max_tokens, with batch calls at half."""
    from app.analytics.routes import token_cost

    chars = len(PAYLOAD["messages"][0]["content"])
    full = runner.worst_cost(PAYLOAD, batched=False)
    assert full > token_cost(chars // 3, 1000, "claude-sonnet-4-6")  # JSON overhead counted too
    assert full < token_cost(chars // 3 + 400, 1000, "claude-sonnet-4-6")
    assert runner.worst_cost(PAYLOAD, batched=True) == pytest.approx(full / 2)
    assert runner.worst_cost({**PAYLOAD, "max_tokens": 4000}, batched=False) > full
    # priced at the payload's own model, so a cheaper model has a lower worst case
    assert runner.worst_cost({**PAYLOAD, "model": "claude-sonnet-5-5"}, batched=False) < full


# --- the Meter on its own: ceiling, admission, settlement (FR-32, FR-41) ---------------------


def usage_message(payload, input_tokens, output_tokens):
    return {"model": payload["model"], "usage": {"input_tokens": input_tokens,
                                                  "output_tokens": output_tokens}}


def test_a_live_call_is_reserved_at_worst_case_then_settled_at_actual():
    """FR-32: the reservation must be released and replaced by the real cost, or a long run
    would count every call at its worst case and stop long before the ceiling."""
    from app.analytics.routes import token_cost

    meter = runner.Meter(10.0)
    assert meter.before(PAYLOAD) is None
    assert meter.reserved_usd == pytest.approx(runner.worst_cost(PAYLOAD, batched=False))
    meter.after(PAYLOAD, usage_message(PAYLOAD, 1000, 200))
    assert meter.reserved_usd == 0.0
    assert meter.spent_usd == pytest.approx(token_cost(1000, 200, "claude-sonnet-4-6"))


def test_a_failed_call_releases_its_reservation_and_costs_nothing():
    """A request that failed must not keep holding (or count as) spend: the run would shrink its
    own headroom for calls that never produced an answer."""
    meter = runner.Meter(10.0)
    meter.before(PAYLOAD)
    meter.after(PAYLOAD, None)
    assert (meter.reserved_usd, meter.spent_usd) == (0.0, 0.0)


def test_a_call_that_could_pass_the_ceiling_is_refused_and_stops_the_run():
    """AC-32: spend must never pass the ceiling, so the check counts what is already reserved
    by calls in flight, not just what is spent -- and once refused, nothing more is launched."""
    worst = runner.worst_cost(PAYLOAD, batched=False)
    meter = runner.Meter(worst * 1.5)
    meter.before(PAYLOAD)  # one call in flight, holding its worst case
    with pytest.raises(runner.CeilingReached, match="ceiling"):
        meter.before(PAYLOAD)  # a second at the same time could pass the ceiling
    assert meter.stopped
    meter.after(PAYLOAD, usage_message(PAYLOAD, 100, 20))
    small = {**PAYLOAD, "max_tokens": 1, "messages": [{"role": "user", "content": "hi"}]}
    with pytest.raises(runner.CeilingReached):  # even one that would now fit: the run has stopped
        meter.before(small)


def test_spend_never_passes_the_ceiling_even_when_every_answer_is_the_longest_possible():
    """The hard-limit claim: calls that each really cost their worst case, one after another,
    stop before the ceiling -- not on the call after it."""
    worst = runner.worst_cost(PAYLOAD, batched=False)
    ceiling = worst * 7.5
    meter = runner.Meter(ceiling)
    estimated_input = math.ceil(len(json.dumps(PAYLOAD, separators=(",", ":"), ensure_ascii=False)) / 3)
    full_length = usage_message(PAYLOAD, estimated_input, 1000)  # really costs its worst case
    made = 0
    with pytest.raises(runner.CeilingReached):
        while True:
            meter.before(PAYLOAD)
            meter.after(PAYLOAD, full_length)
            made += 1
    assert made == 7 and meter.spent_usd <= ceiling


def test_ceiling_must_be_positive():
    """A zero or negative ceiling would refuse every call; fail at construction, not mid-run."""
    for bad in (0, -1.0, float("nan")):
        with pytest.raises(ValueError):
            runner.Meter(bad)


def test_batch_admission_keeps_planned_order_and_stops_at_the_first_that_does_not_fit():
    """FR-32 + FR-41: a batch cannot be recalled once sent, so requests are admitted in order
    only while their worst case still fits; the rest are not attempted."""
    payloads = [{**PAYLOAD, "messages": [{"role": "user", "content": f"q{i}"}]} for i in range(5)]
    each = runner.worst_cost(payloads[0], batched=True)
    meter = runner.Meter(each * 3.5)
    admitted = meter.admit(payloads)
    assert admitted == payloads[:3]
    assert meter.stopped and meter.reserved_usd == pytest.approx(each * 3)
    assert meter.admit(payloads[3:]) == []  # a stopped meter admits nothing more


def test_settling_a_batch_charges_actual_cost_at_half_price_and_only_for_answered_requests():
    """FR-41: batch rows are billed at 50% and only when answered; an errored or expired request
    costs nothing, and every worst-case hold is released either way."""
    from app.analytics.routes import token_cost

    payloads = [{**PAYLOAD, "messages": [{"role": "user", "content": f"q{i}"}]} for i in range(3)]
    meter = runner.Meter(10.0)
    admitted = meter.admit(payloads)
    ids = [runner.custom_id(p) for p in payloads]
    results = {
        ids[0]: {"type": "succeeded", "message": usage_message(payloads[0], 2000, 500)},
        ids[1]: {"type": "errored", "error": {"type": "error", "error": {"type": "api_error"}}},
        # ids[2]: no result at all (never came back)
    }
    meter.settle_batch(admitted, results)
    assert meter.reserved_usd == 0.0
    assert meter.spent_usd == pytest.approx(token_cost(2000, 500, "claude-sonnet-4-6", batched=True))
    assert meter.spent_usd == pytest.approx(token_cost(2000, 500, "claude-sonnet-4-6") / 2)


def test_capture_records_the_request_and_aborts_with_a_base_exception():
    """Capture must abort the app operation without the app mistaking it for a provider error:
    Deferred has to be a BaseException, or the app's `except Exception` fallbacks would write
    invented output and a failed-call row for a request that was only being recorded."""
    assert issubclass(runner.Deferred, BaseException) and not issubclass(runner.Deferred, Exception)
    meter = runner.Meter(10.0)
    meter.mode = "capture"
    with pytest.raises(runner.Deferred) as caught:
        meter.before(PAYLOAD)
    assert caught.value.custom_id == runner.custom_id(PAYLOAD)
    assert meter.captured == {runner.custom_id(PAYLOAD): PAYLOAD}
    assert meter.reserved_usd == 0.0  # capture spends nothing and holds nothing


def replay_meter(result):
    meter = runner.Meter(10.0)
    meter.mode = "replay"
    cid = runner.custom_id(PAYLOAD)
    meter.results, meter.batch_id_of = {cid: result}, {cid: "msgbatch_abc"}
    return meter, cid


def test_replay_returns_the_stored_answer_with_a_batch_provider_request_id():
    """The ledger marks a row as batched from this id, so its shape is part of the contract."""
    message = usage_message(PAYLOAD, 10, 5)
    meter, cid = replay_meter({"type": "succeeded", "message": message})
    assert meter.before(PAYLOAD) == (message, f"msgbatch_abc:{cid}")


@pytest.mark.parametrize(
    "result, fragment",
    [
        ({"type": "errored", "error": {"type": "error", "error": {"type": "overloaded_error"}}},
         "overloaded_error"),
        ({"type": "expired"}, "expired"),
        ({"type": "canceled"}, "canceled"),
    ],
)
def test_replay_of_an_unanswered_request_is_a_failed_call(result, fragment):
    """FR-40/FR-41: a request the provider did not answer must fail loudly at replay, never
    return a made-up answer; the message says why so the failed-call list is readable."""
    meter, _ = replay_meter(result)
    with pytest.raises(runner.BatchCallFailed, match=fragment):
        meter.before(PAYLOAD)


def test_replay_never_gives_one_requests_answer_to_a_different_request():
    """FR-41: if the app builds a different prompt at replay than at capture (the one known
    cause is a date-dependent fact), the changed request has no stored answer and fails."""
    meter, _ = replay_meter({"type": "succeeded", "message": usage_message(PAYLOAD, 10, 5)})
    changed = {**PAYLOAD, "messages": [{"role": "user", "content": "a different question"}]}
    with pytest.raises(runner.BatchCallFailed, match="not captured"):
        meter.before(changed)


# --- the Meter through the real Claude client hook (FR-32, FR-41) ---------------------------


@pytest.fixture
def client_on_fake(monkeypatch):
    """The real request path on the fake Anthropic client, with the daily-cap Redis calls off."""
    fake = install(monkeypatch)
    monkeypatch.setattr(claude, "reserve_tokens", lambda org_id, estimate: ("key", estimate))
    monkeypatch.setattr(claude, "settle_tokens", lambda reservation, used: None)
    return fake


def ask_once():
    return asyncio.run(
        ClaudeClient().complete_text(
            org_id=None, system_prompt="sys", user_prompt="hi", max_tokens=1000, temperature=0.2
        )
    )


def test_the_ceiling_stops_the_second_call_before_any_request_is_sent(client_on_fake, monkeypatch):
    """AC-32 end to end: the meter installed as claude.exam_meter admits the first call, learns
    its real cost from the response, and refuses the next one BEFORE the provider is asked."""
    from app.analytics.routes import token_cost

    probe = runner.Meter(1000.0)
    monkeypatch.setattr(claude, "exam_meter", probe)
    ask_once()
    [payload] = client_on_fake.streamed
    actual = token_cost(100, 20, payload["model"])
    assert probe.spent_usd == pytest.approx(actual)

    client_on_fake.streamed.clear()
    meter = runner.Meter(runner.worst_cost(payload, batched=False) + actual / 2)
    monkeypatch.setattr(claude, "exam_meter", meter)
    ask_once()
    with pytest.raises(runner.CeilingReached):
        ask_once()
    assert len(client_on_fake.streamed) == 1 and meter.stopped


def test_capture_then_replay_through_the_client_sends_nothing_and_marks_the_row_batched(
    client_on_fake, monkeypatch
):
    """FR-41 end to end: capture aborts the call without a request; replay answers it from the
    stored batch result, again without a request, and the id carries the batch for the stamp."""
    meter = runner.Meter(10.0)
    monkeypatch.setattr(claude, "exam_meter", meter)
    meter.mode = "capture"
    with pytest.raises(runner.Deferred):
        ask_once()
    [(cid, payload)] = meter.captured.items()
    meter.mode = "replay"
    meter.results = {cid: batch_result("succeeded", payload)}
    meter.batch_id_of = {cid: "msgbatch_fake1"}
    response = ask_once()
    assert response.provider_request_id == f"msgbatch_fake1:{cid}"
    assert response.content_blocks[0]["text"] == "A fake answer."
    assert client_on_fake.streamed == []


# --- the go-ahead (FR-31) -------------------------------------------------------------------


def test_a_go_ahead_with_a_ceiling_returns_a_run_and_audits_the_approval_first(live):
    """AC-31/AC-39: the approval row (who, estimate, ceiling, versions) must exist before the
    function returns, so no paid call can ever precede it."""
    extra = {
        "mode": "full",
        "change": None,
        "target_area": None,
        "main_measure": None,
        "versions_at_start": {"model_under_test": "claude-sonnet-4-6", "app_code_sha256": "a" * 64},
    }
    typist = Typist(ceiling="25")
    run = approve(live, ask=typist, extra=extra)
    assert run.run_id == typist.run_ids[0] and run.kind == "exam"
    assert run.ceiling_usd == 25.0 and run.ceiling_credits is None
    assert run.meter.ceiling_usd == 25.0 and claude.exam_meter is None  # installing is the caller's job
    [row] = audit_rows(live, "eval.spend_approved")
    assert row.resource_id == run.run_id and row.resource_type == "eval_run"
    assert row.actor_user_id == OWNER.user_id
    assert row.after["ceiling_usd"] == 25.0 and row.after["delivery"] == "batch"
    assert row.after["approver"] == {"user_id": OWNER.user_id, "email": OWNER.email}
    assert row.after["estimate"]["worst_case_usd"] == 24.5
    assert row.after["mode"] == "full" and row.after["versions_at_start"]["model_under_test"]
    assert audit_rows(live, "eval.spend_refused") == []
    with live.sessions() as db:
        assert verify_audit_hash_chain(db)  # an ordinary link in the immutable chain


def test_the_estimate_is_shown_in_plain_language_before_anything_is_asked(live, capsys):
    """FR-31/FR-32/FR-41: calls split batch/sync, expected cost at the price paid and at list,
    the worst case, the 24-hour wait, and the advice to set about the worst case as a ceiling
    ("charged only for what is used") must all be on screen before the owner types anything."""
    approve(live, ask=Typist(), extra={"versions_at_start": {"model_under_test": "claude-sonnet-4-6"}})
    out = capsys.readouterr().out
    assert "486 (454 batch, 32 sync)" in out
    assert "$9.75" in out and "$18.07" in out and "$24.50" in out
    assert "playbook_review" in out and "contract_metadata_extraction" in out
    assert "claude-sonnet-4-6" in out and "Delivery: batch" in out
    assert "24 hours" in out
    assert "ceiling of about $25.00" in out and "charged only for what is used" in out


def test_enter_is_a_refusal_audited_with_no_provider_call(live, client_on_fake):
    """AC-31: with no go-ahead the estimate is shown and nothing is paid; the refusal is on the
    record and no meter is installed, so no call can be metered into existence."""
    typist = Typist(approve=False)
    assert approve(live, ask=typist) is None
    assert len(typist.prompts) == 1  # the ceiling is the only question
    [row] = audit_rows(live, "eval.spend_refused")
    assert row.resource_id == typist.run_ids[0]
    assert row.after["refused_by"]["email"] == OWNER.email and row.after["reason"] == "no go-ahead"
    assert audit_rows(live, "eval.spend_approved") == []
    assert client_on_fake.streamed == [] and client_on_fake.batch_creates == []
    assert claude.exam_meter is None


@pytest.mark.parametrize("ceiling", ["abc", "0", "-5", "nan", "inf", "12 dollars", "y"])
def test_a_ceiling_that_is_not_a_positive_number_is_a_refusal(live, ceiling):
    """The ceiling is the spending control: an unreadable, zero, negative or infinite one must
    not become 'no limit'. Anything but a positive number is refused and audited."""
    assert approve(live, ask=Typist(ceiling=ceiling)) is None
    [row] = audit_rows(live, "eval.spend_refused")
    assert row.after["reason"] == "no valid spending ceiling"
    assert audit_rows(live, "eval.spend_approved") == []


def test_a_dollar_sign_and_thousands_separator_are_accepted(live):
    """A ceiling typed the way people write money ("$1,000") is the number they meant."""
    assert approve(live, ask=Typist(ceiling="$1,000")).ceiling_usd == 1000.0


def test_a_closed_terminal_is_a_refusal_not_a_crash(live):
    """No terminal to answer on (stdin closed) must refuse, never approve by default."""

    def eof(_prompt):
        raise EOFError

    assert approve(live, ask=eof) is None
    assert len(audit_rows(live, "eval.spend_refused")) == 1


def test_a_go_ahead_cannot_carry_over_to_another_run(live):
    """AC-31: every run asks for its own ceiling, naming its own run id; run A's approval never
    approves run B (pressing Enter on B refuses it)."""
    first = Typist()
    run_a = approve(live, ask=first)
    second = Typist(approve=False)
    assert approve(live, ask=second) is None
    assert first.run_ids[0] == run_a.run_id and second.run_ids[0] != run_a.run_id


def test_extra_cannot_overwrite_who_approved_or_what_was_approved(live):
    """The approval row is evidence: caller-supplied extras are recorded beside the fixed
    fields, never in place of them."""
    forged = {"approver": {"user_id": "someone-else", "email": "x@example.com"}, "ceiling_usd": 1.0}
    run = approve(live, ask=Typist(ceiling="30"), extra=forged)
    [row] = audit_rows(live, "eval.spend_approved")
    assert row.after["approver"]["user_id"] == OWNER.user_id and row.after["ceiling_usd"] == 30.0
    assert run.ceiling_usd == 30.0


def test_ocr_without_a_price_asks_for_the_ceiling_in_credits(live, capsys):
    """FR-20/FR-31: with no REDUCTO_USD_PER_CREDIT the OCR estimate must say so and the ceiling
    is asked, recorded and enforced in credits -- a dollar ceiling would be a guess."""
    estimate = make_estimate(calls=0, batch_calls=0, sync_calls=0, usd=0.0, list_usd=0.0,
                             worst_case_usd=0.0, by_call_type={}, delivery="sync",
                             ocr_files=11, ocr_pages=340, ocr_credits=340, ocr_usd=None)
    typist = Typist(ceiling="400")
    run = approve(live, ask=typist, kind="ocr", estimate=estimate)
    out = capsys.readouterr().out
    assert "11 files, 340 pages, 340 credits" in out
    assert "price not set" in out and "REDUCTO_USD_PER_CREDIT in eval/private/exam.env" in out
    assert typist.prompts[0].endswith("in credits (press Enter to cancel): ")
    assert run.ceiling_credits == 400 and run.ceiling_usd is None and run.meter is None
    [row] = audit_rows(live, "eval.spend_approved")
    assert row.after["ceiling_credits"] == 400 and row.after["ceiling_usd"] is None


def test_ocr_with_a_price_asks_in_dollars_and_a_fractional_credit_ceiling_is_refused(live):
    """With the price known the OCR ceiling is dollars; without it, half a credit is not a
    quantity Reducto sells, so a fractional credit ceiling is a refusal."""
    priced = make_estimate(calls=0, batch_calls=0, sync_calls=0, usd=0.0, list_usd=0.0,
                           worst_case_usd=0.0, by_call_type={}, delivery="sync",
                           ocr_files=11, ocr_pages=340, ocr_credits=340, ocr_usd=5.1)
    run = approve(live, ask=Typist(ceiling="8"), kind="ocr", estimate=priced)
    assert run.ceiling_usd == 8.0 and run.ceiling_credits is None and run.meter is None
    unpriced = make_estimate(calls=0, delivery="sync", ocr_files=2, ocr_pages=9, ocr_credits=9,
                             ocr_usd=None)
    assert approve(live, ask=Typist(ceiling="9.5"), kind="ocr", estimate=unpriced) is None


@pytest.mark.parametrize("kind", ["exam", "prefill", "draft_routing"])
def test_every_ai_run_gets_a_meter_at_its_ceiling(live, kind):
    """FR-32 applies to pre-fill and routing drafting as much as to the exam run."""
    run = approve(live, ask=Typist(ceiling="3"), kind=kind)
    assert run.meter.ceiling_usd == 3.0 and run.meter.mode == "live"


# --- the run record (FR-39, FR-40) ----------------------------------------------------------


def test_finishing_writes_the_one_immutable_run_record_and_removes_the_meter(live, monkeypatch):
    """AC-39: who, when, mode, versions, estimate, ceiling, outcome and actual cost are in one
    audit row that cannot be edited; and the meter is gone so it cannot outlive its run."""
    run = approve(live, ask=Typist(ceiling="25"))
    monkeypatch.setattr(claude, "exam_meter", run.meter)
    actual = {"usd": 9.5, "list_usd": 17.0, "by_call_type": {}, "ocr_usd": None, "ocr_credits": None,
              "calls": 486, "failed_calls": 3}
    runner.finish_paid_run(
        run,
        outcome="complete",
        reason=None,
        actual=actual,
        results_sha256="c" * 64,
        summary={"mode": "full", "versions": {"models": ["claude-sonnet-4-6"]},
                 "contains_errors": True, "failed_call_rate": 0.006, "batches": [{"phase": 1}],
                 "headline": {"review_f1": 0.8}, "outcome": "stopped", "ceiling_usd": 1.0},
    )
    assert claude.exam_meter is None
    [row] = audit_rows(live, "eval.run_finished")
    after = row.after
    assert row.resource_id == run.run_id and row.actor_user_id == OWNER.user_id
    assert after["outcome"] == "complete" and after["ceiling_usd"] == 25.0  # not taken from summary
    assert after["incomplete_reason"] is None and after["results_sha256"] == "c" * 64
    assert after["mode"] == "full" and after["delivery"] == "batch"
    assert after["started_by"]["email"] == OWNER.email and after["approved_by"]["email"] == OWNER.email
    assert datetime.fromisoformat(after["started_at"]) <= datetime.fromisoformat(after["finished_at"])
    assert after["estimate"]["usd"] == 9.75 and after["actual"] == actual
    assert after["contains_errors"] is True and after["failed_call_rate"] == 0.006
    assert after["headline"] == {"review_f1": 0.8} and after["batches"] == [{"phase": 1}]
    # the record cannot be edited or removed afterwards
    with live.sessions() as db:
        stored = db.get(AuditLog, row.id)
        stored.after = {"outcome": "complete", "actual": {"usd": 0.0}}
        with pytest.raises(RuntimeError, match="immutable"):
            db.flush()
        db.rollback()
        db.delete(db.get(AuditLog, row.id))
        with pytest.raises(RuntimeError, match="immutable"):
            db.flush()
        db.rollback()


def test_finishing_leaves_another_runs_meter_alone(live, monkeypatch):
    """Only the run's own meter is removed: finishing an old run must not switch off the meter
    of the run that is metering right now."""
    run = approve(live, ask=Typist())
    current = runner.Meter(5.0)
    monkeypatch.setattr(claude, "exam_meter", current)
    runner.finish_paid_run(run, outcome="complete", reason=None, actual={}, results_sha256=None,
                           summary=None)
    assert claude.exam_meter is current


def test_an_incomplete_run_records_its_reason_and_what_was_spent(live):
    """FR-32/FR-40: a run that hit its ceiling is finished as incomplete with the reason and the
    spend, so it can be read back and can never become part of a baseline."""
    run = approve(live, ask=Typist(ceiling="4"))
    runner.finish_paid_run(run, outcome="incomplete", reason="ceiling_reached",
                           actual={"usd": 3.98, "calls": 200, "failed_calls": 0},
                           results_sha256="d" * 64, summary={"mode": "full"})
    [row] = audit_rows(live, "eval.run_finished")
    assert row.after["outcome"] == "incomplete" and row.after["incomplete_reason"] == "ceiling_reached"
    assert row.after["actual"]["usd"] == 3.98 and row.after["ceiling_usd"] == 4.0


def test_an_unknown_outcome_writes_nothing(live):
    """The run record is permanent, so a typo such as 'completed' must fail before the write."""
    run = approve(live, ask=Typist())
    with pytest.raises(ValueError, match="outcome"):
        runner.finish_paid_run(run, outcome="completed", reason=None, actual={},
                               results_sha256=None, summary=None)
    assert audit_rows(live, "eval.run_finished") == []


# --- who may start a run (FR-31): the owner gate -------------------------------------------

OWNER_EMAIL = "owner@example.com"


@pytest.fixture
def people(live, monkeypatch):
    """An owner, an administrator who is not the owner, and a suspended copy of the owner's
    account, in the scratch 'live' database; the exam is configured with the owner's email."""
    ids = {}
    with live.sessions() as db:
        for key, email, status in (
            ("owner", OWNER_EMAIL, "active"),
            ("admin", "admin@example.com", "active"),
            ("suspended", "suspended-owner@example.com", "suspended"),
        ):
            user = User(org_id=live.org_id, email=f"{live.org_id[:8]}-{email}", full_name=key,
                        hashed_password=PASSWORD_HASH, status=status)
            db.add(user)
            db.flush()
            ids[key] = (user.id, user.email)
        db.commit()
    configured = {"email": ids["owner"][1]}
    monkeypatch.setattr(
        env, "_load_exam_settings",
        lambda: env.ExamSettings(_env_file=None, owner_email=configured["email"]),
    )
    ids["configured"] = configured
    return ids


def authenticate(live, email):
    return env.authenticate_owner(live, email, command="run")


def test_the_owner_is_accepted_and_is_the_default(live, people):
    """The positive case, case-insensitive on the email (people type their address in any case),
    and with no --owner the configured owner is used (the owner dropped the password prompt)."""
    owner_id, email = people["owner"]
    assert authenticate(live, email.upper()) == Owner(user_id=owner_id, email=email)
    assert authenticate(live, None) == Owner(user_id=owner_id, email=email)
    assert audit_rows(live, "eval.access_refused") == []


def test_an_administrator_who_is_not_the_owner_is_refused(live, people):
    """AC-31/FR-31: roles are not consulted, so an admin who names themselves is refused, and the
    refusal is audited."""
    _, admin_email = people["admin"]
    with pytest.raises(ExamRefused, match="only the exam owner"):
        authenticate(live, admin_email)
    [row] = audit_rows(live, "eval.access_refused")
    assert row.after == {"command": "run", "claimed_email": admin_email, "reason": "not_the_owner"}
    assert row.actor_user_id is None


def test_an_owner_who_is_not_active_or_not_in_this_org_is_refused(live, people):
    """A suspended or unknown owner account cannot start a run."""
    people["configured"]["email"] = people["suspended"][1]
    with pytest.raises(ExamRefused, match="no active AEGIS user"):
        authenticate(live, people["suspended"][1])
    people["configured"]["email"] = "nobody@example.com"
    with pytest.raises(ExamRefused, match="no active AEGIS user"):
        authenticate(live, "nobody@example.com")
    reasons = [r.after["reason"] for r in audit_rows(live, "eval.access_refused")]
    assert reasons == ["not_active", "no_such_user"]


def test_with_no_owner_configured_every_command_is_refused(live, people, monkeypatch):
    """FR-31: an exam with no owner set has no one who may run it -- not 'anyone'."""
    monkeypatch.setattr(env, "_load_exam_settings", lambda: env.ExamSettings(_env_file=None))
    with pytest.raises(ExamRefused, match="EVAL_OWNER_EMAIL"):
        authenticate(live, people["owner"][1])
    [row] = audit_rows(live, "eval.access_refused")
    assert row.after["reason"] == "owner_not_configured"


def test_the_exam_settings_come_from_the_private_file_and_default_to_nothing(tmp_path, monkeypatch):
    """With no eval/private/exam.env the exam has no key, no owner and no OCR price (so every
    paid command refuses); with it, the values are read from it and nowhere else."""
    for name in ("EVAL_CLAUDE_API_KEY", "EVAL_OWNER_EMAIL", "REDUCTO_USD_PER_CREDIT"):
        monkeypatch.delenv(name, raising=False)
    path = tmp_path / "exam.env"
    monkeypatch.setitem(env.ExamSettings.model_config, "env_file", path)
    absent = env._load_exam_settings()
    assert (absent.claude_api_key, absent.owner_email, absent.reducto_usd_per_credit) == (
        None, None, None)
    path.write_text("EVAL_CLAUDE_API_KEY=not-a-real-key\nEVAL_OWNER_EMAIL=o@example.com\n"
                    "REDUCTO_USD_PER_CREDIT=0.015\n")
    present = env._load_exam_settings()
    assert (present.claude_api_key, present.owner_email, present.reducto_usd_per_credit) == (
        "not-a-real-key", "o@example.com", 0.015)


# --- isolate() refuses before it changes anything (FR-34, FR-35) ----------------------------


@pytest.fixture
def fresh_process(monkeypatch):
    """Make ``isolate`` see a process whose database engine has not been built yet (as
    ``python -m eval.exam`` is), so its later checks can run. Every refusal under test happens
    before it changes any setting, which the tests assert."""
    monkeypatch.delitem(sys.modules, "app.core.database")
    monkeypatch.setattr(settings, "environment", "test")

    def configure(**exam):
        monkeypatch.setattr(env, "_load_exam_settings", lambda: env.ExamSettings(_env_file=None, **exam))

    return configure


def untouched():
    """Snapshot what isolate would have changed; call the returned function to assert none of it did."""
    before = (settings.database_url, settings.claude_api_key, os.environ.get("ANTHROPIC_API_KEY"),
              settings.storage_root, settings.disable_rbac)

    def check():
        assert (settings.database_url, settings.claude_api_key, os.environ.get("ANTHROPIC_API_KEY"),
                settings.storage_root, settings.disable_rbac) == before

    return check


def test_a_deployed_environment_is_refused(monkeypatch):
    """The exam rebuilds a database and spends money: staging or production is never the place."""
    monkeypatch.setattr(settings, "environment", "production")
    check = untouched()
    with pytest.raises(ExamRefused, match="local/development"):
        env.isolate(paid=False)
    check()


def test_a_process_that_already_built_the_apps_engine_is_refused():
    """The engine is built from settings.database_url when app.core.database is first imported.
    Repointing the settings afterwards would leave it on the live database, so the exam refuses
    to start in such a process (this test process is one: pytest has imported it)."""
    assert "app.core.database" in sys.modules
    with pytest.raises(ExamRefused, match="already exists"):
        env.isolate(paid=False)


def test_a_paid_command_without_its_own_key_is_refused_and_never_falls_back(fresh_process, monkeypatch):
    """AC-34: no exam key means no run, not a run on the key used for real work. Nothing is
    changed before the refusal, so the real key is still the app's and no client was built."""
    fresh_process()  # exam.env has no key
    monkeypatch.setattr(settings, "claude_api_key", "the-apps-real-key")
    check = untouched()
    with pytest.raises(ExamRefused, match="EVAL_CLAUDE_API_KEY"):
        env.isolate(paid=True)
    check()


def test_a_paid_command_with_the_apps_own_key_is_refused(fresh_process, monkeypatch):
    """FR-34: the exam key must differ from the key used for real work, or exam spend could not
    be told apart from (or capped apart from) real spend. Both places the real key can live count."""
    monkeypatch.setattr(settings, "claude_api_key", "shared-key")
    fresh_process(claude_api_key="shared-key")
    with pytest.raises(ExamRefused, match="same as the app's Claude key"):
        env.isolate(paid=True)
    monkeypatch.setattr(settings, "claude_api_key", None)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sdk-fallback-key")  # the SDK's own fallback variable
    fresh_process(claude_api_key="sdk-fallback-key")
    check = untouched()
    with pytest.raises(ExamRefused, match="same as the app's Claude key"):
        env.isolate(paid=True)
    check()


def test_a_paid_command_refuses_a_fake_model_and_third_party_providers(fresh_process, monkeypatch):
    """FR-35: an exam against MOCK_CLAUDE measures nothing, and contract text may go only to
    Anthropic and Reducto, so Voyage/Cohere are refused rather than quietly swapped."""
    fresh_process(claude_api_key="a-distinct-exam-key")
    monkeypatch.setattr(settings, "claude_api_key", "the-apps-real-key")
    monkeypatch.setattr(settings, "mock_claude", True)
    with pytest.raises(ExamRefused, match="MOCK_CLAUDE"):
        env.isolate(paid=True)
    monkeypatch.setattr(settings, "mock_claude", False)
    monkeypatch.setattr(settings, "allow_mock_embeddings", False)
    monkeypatch.setattr(settings, "embedding_provider", "voyage")
    check = untouched()
    with pytest.raises(ExamRefused, match="Voyage and Cohere"):
        env.isolate(paid=True)
    monkeypatch.setattr(settings, "embedding_provider", "local")
    monkeypatch.setattr(settings, "rerank_provider", "cohere")
    with pytest.raises(ExamRefused, match="Voyage and Cohere"):
        env.isolate(paid=True)
    check()


def test_the_exam_database_must_be_the_exams_own_and_never_the_live_one(fresh_process, monkeypatch):
    """A database name that is not the exam's, or that is the live database itself, is refused
    before anything is dropped or rebuilt there."""
    fresh_process(claude_api_key="a-distinct-exam-key")
    check = untouched()
    with pytest.raises(ExamRefused, match="must start with"):
        env.isolate(paid=False, eval_db_name=make_url(settings.database_url).database)
    live_as_exam = make_url(settings.database_url).set(database=env.EVAL_DB_NAME)
    monkeypatch.setattr(settings, "database_url", live_as_exam.render_as_string(hide_password=False))
    with pytest.raises(ExamRefused, match="must not be the live database"):
        env.isolate(paid=False)
    monkeypatch.undo()
    check()



# =============================================================================================
# The owner set (T015): frozen OCR, reading as the app does, import, OCR, pre-fill, the keys
# =============================================================================================

SENTENCE = "The Supplier shall provide the Services described in each order. "
TEXT_FILES = {0, 11, 12}  # Employment Agreement.pdf, the .docx, NDA Form.pdf: a real text layer
STAMP_ONLY = {2, 10}  # INFRAHEDGE-MSA.pdf, MSA_Signed.pdf: an image per page + one stamp line


def text_pdf(tag: int, pages: int = 2) -> bytes:
    """A PDF with a real text layer: what the app reads without OCR."""
    pdf = fitz.open()
    for n in range(pages):
        page = pdf.new_page()
        page.insert_textbox(
            fitz.Rect(50, 50, 550, 780), f"{n + 1}. Agreement {tag}. " + SENTENCE * 8, fontsize=10
        )
    data = pdf.tobytes()
    pdf.close()
    return data


def scan_pdf(tag: int, pages: int = 3, *, stamp: bool = False) -> bytes:
    """Pages that are one image each. With ``stamp`` a one-line e-signature stamp is the only
    text layer (enough of it that the app does not ask for OCR); without it there is none."""
    pdf = fitz.open()
    for n in range(pages):
        page = pdf.new_page()
        pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 280), False)
        pixmap.clear_with(150 + tag)
        page.insert_image(page.rect, pixmap=pixmap)
        if stamp:
            page.insert_text(
                (20, 20), f"DocuSign Envelope ID: 3F2A9C1B-77AA-4D2E-9C11-{tag:04d}{n:08d}", fontsize=7
            )
    data = pdf.tobytes()
    pdf.close()
    return data


def docx_bytes(tag: int) -> bytes:
    document = docx.Document()
    document.add_heading(f"Mutual Non-Disclosure Agreement {tag}", level=1)
    for n in range(1, 9):
        document.add_paragraph(f"{n}. " + SENTENCE * 3)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def owner_folder(root: Path) -> Path:
    """A 'Downloads' folder: the 14 allowed files (3 text, 9 scans, 2 stamp-only scans, the
    layout the real set has) plus a personal file the exam must never touch."""
    src = root / "downloads"
    src.mkdir()
    for i, name in enumerate(OWNER_FILES):
        if name.endswith(".docx"):
            data = docx_bytes(i)
        elif i in TEXT_FILES:
            data = text_pdf(i)
        else:
            data = scan_pdf(i, pages=2 + i % 3, stamp=i in STAMP_ONLY)
        (src / name).write_bytes(data)
    (src / "Tax return 2025.pdf").write_bytes(text_pdf(99))
    return src


def exam_settings(monkeypatch, **values) -> None:
    """The exam's private settings as the test wants them, whatever eval/private/exam.env holds."""
    if "reducto_usd_per_credit" in values:  # that field is read by its environment-variable name
        values["REDUCTO_USD_PER_CREDIT"] = values.pop("reducto_usd_per_credit")
    monkeypatch.setattr(
        env, "_load_exam_settings", lambda: env.ExamSettings(_env_file=None, **values)
    )


@pytest.fixture
def private(tmp_path, monkeypatch):
    """Every private path under tmp_path, so no test reads or writes eval/private."""
    root = tmp_path / "private"
    monkeypatch.setattr(owner_mod, "OWNER_DIR", root / "owner")
    monkeypatch.setattr(owner_mod, "OCR_DIR", root / "ocr")
    monkeypatch.setattr(owner_mod, "KEYS_DIR", root / "keys")
    monkeypatch.setattr(sheets, "SHEETS_DIR", root / "sheets")
    monkeypatch.setattr(sheets, "TEXTS_DIR", root / "owner" / "texts")
    exam_settings(monkeypatch)
    return root


class FakeOCR:
    """Stands in for Reducto: records every file it is asked to read and bills a credit a page."""

    def __init__(self, *, fail=(), empty=(), interrupt=()) -> None:
        self.fail, self.empty, self.interrupt = set(fail), set(empty), set(interrupt)
        self.calls: list[str] = []

    async def __call__(self, *, filename: str, mime_type: str, content: bytes) -> OCRResult:
        self.calls.append(filename)
        if filename in self.interrupt:
            raise KeyboardInterrupt
        if filename in self.fail:
            raise RuntimeError("provider unavailable")
        with fitz.open(stream=content, filetype="pdf") as pdf:
            pages = pdf.page_count
        text = "" if filename in self.empty else f"OCR reading of {filename}. " + SENTENCE * 20
        return OCRResult(
            text=text,
            provider="reducto",
            quality_score=0.9,
            metadata={"num_pages": pages, "credits": pages},
            elements=[{"type": "text", "content": text}],
        )


def ledger_rows(live: LiveDB, run_id: str) -> list[AICallLog]:
    with live.sessions() as db:
        return list(
            db.scalars(
                select(AICallLog)
                .where(AICallLog.org_id == live.org_id, AICallLog.eval_run_id == run_id)
                .order_by(AICallLog.created_at, AICallLog.id)
            )
        )


def sends(live: LiveDB, stage: str | None = None) -> list[AuditLog]:
    """The ``eval.real_contract_sent`` rows (one per stage a send went through), by stage."""
    rows = audit_rows(live, "eval.real_contract_sent")
    return [r for r in rows if stage is None or r.after["stage"] == stage]


def read_as_app_sync(data: bytes, *, ocr: owner_mod.FrozenOCR, name: str = "x.pdf"):
    return asyncio.run(owner_mod.read_as_app(data, mime_type="application/pdf", filename=name, ocr=ocr))


# --- verify_quote (FR-17) -------------------------------------------------------------------

CLAUSE = "This agreement is governed by the laws of India.\nThe courts at Mumbai have jurisdiction."


def test_a_quote_is_found_word_for_word_ignoring_only_spacing_and_line_breaks():
    """FR-17: a PDF wraps lines mid-sentence, so spacing cannot matter -- but a changed word,
    case or punctuation mark is a different sentence, and an empty quote supports nothing."""
    span = owner_mod.verify_quote("governed by the laws of India. The courts at Mumbai", CLAUSE)
    assert span is not None
    assert CLAUSE[span[0] : span[1]] == (
        "governed by the laws of India.\nThe courts at Mumbai"  # the contract's own text
    )
    assert owner_mod.verify_quote("governed  by   the laws of India.", CLAUSE) is not None
    for not_there in (
        "governed by the laws of Mumbai",  # a different word
        "Governed by the laws of India.",  # a different case
        "laws of India The courts at Mumbai",  # a dropped full stop between the sentences
        "the laws of India, and the courts",  # a joined paraphrase
        "",
    ):
        assert owner_mod.verify_quote(not_there, CLAUSE) is None


# --- FrozenOCR: OCR once, then never again (FR-20) ------------------------------------------


def test_a_scan_is_paid_for_once_and_every_later_read_is_free(private):
    """AC-20: the first read goes to the provider and is frozen by the file's sha256; a later
    read -- in another process, with no provider at all -- gets the identical text for free."""
    sent: list[str] = []
    fake = FakeOCR()
    ocr = owner_mod.FrozenOCR(fill=fake, on_send=lambda sha, result: sent.append(sha), run_id="r1")
    data = scan_pdf(1)
    first = asyncio.run(ocr.extract_text(filename="a.pdf", mime_type="application/pdf", content=data))
    sha = hashlib.sha256(data).hexdigest()
    assert fake.calls == ["a.pdf"] and sent == [sha]
    record = json.loads((owner_mod.OCR_DIR / f"{sha}.json").read_text())
    assert record["file_sha256"] == sha and record["run_id"] == "r1" and record["credits"] == 3
    assert record["text"] == first.text and record["provider"] == "reducto"

    again = asyncio.run(
        owner_mod.FrozenOCR().extract_text(filename="a.pdf", mime_type="application/pdf", content=data)
    )
    assert again.text == first.text and again.quality_score == first.quality_score
    assert fake.calls == ["a.pdf"] and sent == [sha]  # nothing was sent the second time


def test_an_unread_scan_never_reaches_a_provider_unless_one_is_given(private):
    """AC-20: a run builds FrozenOCR with no provider, so a scan nobody has OCR-read yet fails
    loudly ('run ocr first') instead of quietly paying for OCR in the middle of a run."""
    with pytest.raises(owner_mod.OCRNotFrozen, match="`ocr` command"):
        asyncio.run(
            owner_mod.FrozenOCR().extract_text(
                filename="a.pdf", mime_type="application/pdf", content=scan_pdf(1)
            )
        )
    assert not owner_mod.OCR_DIR.exists() or not list(owner_mod.OCR_DIR.iterdir())


def test_a_changed_file_is_not_served_the_old_frozen_text(private):
    """FR-20: the cache is keyed by the file's bytes, so a replaced scan is a new file that has
    to be read again; it can never inherit the text frozen for the old one."""
    asyncio.run(
        owner_mod.FrozenOCR(fill=FakeOCR()).extract_text(
            filename="a.pdf", mime_type="application/pdf", content=scan_pdf(1)
        )
    )
    with pytest.raises(owner_mod.OCRNotFrozen):
        asyncio.run(
            owner_mod.FrozenOCR().extract_text(
                filename="a.pdf", mime_type="application/pdf", content=scan_pdf(2)
            )
        )


def test_an_empty_ocr_read_is_reported_but_never_frozen(private):
    """A provider that answers with no text spent credits (so it is reported as sent) but must
    not pin the scan to 'no text' for every later run."""
    sent: list[str] = []
    fake = FakeOCR(empty={"a.pdf"})
    ocr = owner_mod.FrozenOCR(fill=fake, on_send=lambda sha, result: sent.append(sha))
    data = scan_pdf(1)
    for _ in range(2):
        result = asyncio.run(
            ocr.extract_text(filename="a.pdf", mime_type="application/pdf", content=data)
        )
        assert result.text == ""
    assert len(fake.calls) == 2 and len(sent) == 2  # read again, not served from a cache
    assert not owner_mod.OCR_DIR.exists() or not list(owner_mod.OCR_DIR.iterdir())


# --- read_as_app: the app's own reader (FR-23) ----------------------------------------------


def test_a_text_pdf_is_read_natively_and_never_asks_for_ocr(private):
    """The app reads a file with a text layer without OCR; the exam must read it the same way."""
    read = read_as_app_sync(text_pdf(1), ocr=owner_mod.FrozenOCR())
    assert not read.method.endswith("_ocr_failed") and "Agreement 1" in read.text
    assert read.page_map and "1" in read.page_map


def test_a_scan_is_read_through_the_frozen_ocr_and_never_the_real_provider(private, monkeypatch):
    """FR-20/FR-35: the app only OCRs when a Reducto key is set; here the key IS set and the real
    client would explode if touched, yet the frozen OCR is the only provider the read can reach.
    Frozen: the OCR text. Not frozen: the app's own 'wanted OCR and could not' outcome."""
    monkeypatch.setattr(settings, "reducto_api_key", "a-real-looking-key")

    async def real_provider(**_):
        raise AssertionError("the real OCR provider was reached")

    monkeypatch.setattr(reducto_client, "extract_text", real_provider)
    data = scan_pdf(1)
    unread = read_as_app_sync(data, ocr=owner_mod.FrozenOCR())
    assert unread.method.endswith("_ocr_failed")
    asyncio.run(
        owner_mod.FrozenOCR(fill=FakeOCR()).extract_text(
            filename="a.pdf", mime_type="application/pdf", content=data
        )
    )
    frozen = read_as_app_sync(data, ocr=owner_mod.FrozenOCR())
    assert frozen.method == "reducto_ocr" and frozen.text.startswith("OCR reading of")


def test_a_stamp_only_scan_is_now_read_through_ocr_by_the_app(private):
    """FR-23/AC-23: the app used to read a scan whose only text layer is a stamp as the stamp,
    which is what 'as read today' measured. Since 2026-10-07 it sends such a scan to OCR, so
    what the app reads is the frozen OCR text, never the stamp."""
    data = scan_pdf(2, stamp=True)
    asyncio.run(
        owner_mod.FrozenOCR(fill=FakeOCR()).extract_text(
            filename="a.pdf", mime_type="application/pdf", content=data
        )
    )
    today = read_as_app_sync(data, ocr=owner_mod.FrozenOCR())
    assert today.method == "reducto_ocr" and "OCR reading" in today.text


# --- owner-import (FR-15, FR-20) ------------------------------------------------------------


@pytest.fixture
def imported(private, live, tmp_path):
    """The 14 generated files, imported: 3 read as text, 11 image scans (the 2 stamp-only ones
    included, since the app now sends them to OCR too)."""
    src = owner_folder(tmp_path)
    entries = asyncio.run(owner_mod.owner_import(live, OWNER, src_dir=src))
    return src, entries


def test_import_sorts_the_14_files_by_how_the_app_reads_them_and_freezes_only_the_text_ones(
    imported, live
):
    """FR-15/FR-20: kind comes from the app's own read decision (not a guess), only files the app
    reads without OCR are frozen now, and the import is on the audit trail with each file's hash."""
    _, entries = imported
    assert [e.name for e in entries] == list(OWNER_FILES)
    assert [e.doc_id for e in entries] == [f"owner-{n:02d}" for n in range(1, 15)]
    kinds = {e.doc_id: e.kind for e in entries}
    assert {k for k, v in kinds.items() if v == "text"} == {f"owner-{i + 1:02d}" for i in TEXT_FILES}
    assert sum(v == "scan" for v in kinds.values()) == 11  # stamp-only scans are OCR'd now
    for entry in entries:
        frozen = owner_mod._frozen(entry)
        if entry.kind == "text":
            assert entry.frozen_text_sha256 == hashlib.sha256(frozen[0].encode()).hexdigest()
        else:
            assert entry.frozen_text_sha256 is None and frozen is None  # scans wait for `ocr`
        stored = owner_mod.OWNER_DIR / "files" / entry.name
        assert hashlib.sha256(stored.read_bytes()).hexdigest() == entry.file_sha256
    assert entries[12].mime_type == "application/pdf" and entries[11].pages == 0  # the .docx
    [row] = audit_rows(live, "eval.owner_files_imported")
    assert [f["doc_id"] for f in row.after["files"]] == [e.doc_id for e in entries]
    assert row.after["files"][1]["kind"] == "scan" and len(row.after["files"][1]["file_sha256"]) == 64
    assert owner_mod.owner_status()["kinds"] == {"text": 3, "scan": 11}


def test_import_never_opens_or_lists_anything_but_the_14_files(private, live, tmp_path, monkeypatch):
    """AC-15: a personal file in the same folder is neither opened, copied nor even found by a
    directory listing -- the exam asks for its 14 names and nothing else."""
    src = owner_folder(tmp_path)
    opened: list[str] = []
    real_read = Path.read_bytes

    def read_bytes(self):
        opened.append(self.name)
        return real_read(self)

    def refuse_listing(real):
        def inner(path=".", *args, **kwargs):
            if Path(path) == src:
                raise AssertionError("the source folder was listed")
            return real(path, *args, **kwargs)

        return inner

    monkeypatch.setattr(Path, "read_bytes", read_bytes)
    monkeypatch.setattr(os, "listdir", refuse_listing(os.listdir))
    monkeypatch.setattr(os, "scandir", refuse_listing(os.scandir))
    asyncio.run(owner_mod.owner_import(live, OWNER, src_dir=src))
    assert "Tax return 2025.pdf" not in opened
    assert set(OWNER_FILES) <= set(opened)
    assert not (owner_mod.OWNER_DIR / "files" / "Tax return 2025.pdf").exists()


def test_import_with_a_file_missing_copies_nothing(private, live, tmp_path):
    """FR-15: the set is exactly these 14 files; a folder with 13 is refused whole (naming the
    missing one) rather than leaving a half-imported set for keys to be made against."""
    src = owner_folder(tmp_path)
    (src / OWNER_FILES[5]).unlink()
    with pytest.raises(ExamRefused, match=re.escape(OWNER_FILES[5])):
        asyncio.run(owner_mod.owner_import(live, OWNER, src_dir=src))
    assert not owner_mod.OWNER_DIR.exists()
    assert audit_rows(live, "eval.owner_files_imported") == []


# --- the one-time OCR (FR-20, FR-31, FR-35) -------------------------------------------------


def scans(entries):
    return [e for e in entries if e.kind != "text"]


def run_ocr(live, *, ask, fill, content=None):
    return asyncio.run(
        owner_mod.run_ocr(
            live=live, owner=OWNER, ask=ask, fill=fill,
            content=content if content is not None else ExamContent(),
        )
    )


def test_ocr_shows_the_estimate_and_sends_nothing_without_a_go_ahead(imported, live, capsys):
    """AC-20/AC-31: files, pages and credits are on screen before anything is asked; with no
    go-ahead the provider is never called and the refusal is on the record."""
    _, entries = imported
    pages = sum(e.pages for e in scans(entries))
    fake = FakeOCR()
    typist = Typist(approve=False)
    with pytest.raises(ExamRefused, match="nothing was sent"):
        run_ocr(live, ask=typist, fill=fake)
    out = capsys.readouterr().out
    assert f"OCR: 11 files, {pages} pages, {pages} credits" in out
    assert fake.calls == []
    assert len(audit_rows(live, "eval.spend_refused")) == 1
    assert audit_rows(live, "eval.real_contract_sent") == []
    assert all(e.frozen_text_sha256 is None for e in owner_mod.read_manifest() if e.kind != "text")


def test_ocr_without_a_price_warns_and_asks_the_ceiling_in_credits(imported, live, capsys):
    """FR-20: REDUCTO_USD_PER_CREDIT is optional; without it the owner is told so and sets the
    ceiling in credits, which is what is then enforced."""
    typist = Typist(ceiling="500")
    run_ocr(live, ask=typist, fill=FakeOCR())
    out = capsys.readouterr().out
    assert "price not set" in out and "REDUCTO_USD_PER_CREDIT" in out
    assert typist.prompts[-1].endswith(" in credits (press Enter to cancel): ")
    [row] = audit_rows(live, "eval.spend_approved")
    assert row.after["kind"] == "ocr" and row.after["ceiling_credits"] == 500
    assert row.after["ceiling_usd"] is None


def test_ocr_with_a_price_shows_dollars_and_asks_the_ceiling_in_dollars(
    imported, live, capsys, monkeypatch
):
    exam_settings(monkeypatch, reducto_usd_per_credit=0.02)
    _, entries = imported
    pages = sum(e.pages for e in scans(entries))
    typist = Typist(ceiling="5")
    run_ocr(live, ask=typist, fill=FakeOCR())
    out = capsys.readouterr().out
    assert f"about ${pages * 0.02:,.2f}" in out and "price not set" not in out
    assert typist.prompts[-1].endswith(" in USD (press Enter to cancel): ")
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["actual"]["ocr_usd"] == pytest.approx(pages * 0.02)


def test_ocr_reads_every_scan_once_freezes_it_and_audits_each_owner_file(imported, live):
    """FR-20/FR-35/AC-20: the 9 scans AND the 2 stamp-only scans (forced) are read once; each is
    frozen and its text is on disk; every owner file sent is audited with Reducto as the only
    destination; and a second `ocr` finds nothing to do, so no later run pays again."""
    _, entries = imported
    fake = FakeOCR()
    result = run_ocr(live, ask=Typist(ceiling="500"), fill=fake)
    names = [e.name for e in scans(entries)]
    assert fake.calls == names and len(names) == 11 and result["outcome"] == "complete"

    for entry in scans(owner_mod.read_manifest()):
        text, _ = owner_mod._frozen(entry)
        assert text.startswith(f"OCR reading of {entry.name}")
        assert entry.frozen_method == "reducto_ocr"
        assert entry.frozen_text_sha256 == hashlib.sha256(text.encode()).hexdigest()
    # each file is on the record twice: written before it left, then with what it cost
    assert len(sends(live, "sending")) == len(sends(live, "sent")) == 11
    assert len(sends(live)) == 22
    assert {r.after["destination"] for r in sends(live)} == {"reducto"}
    assert {r.after["purpose"] for r in sends(live)} == {"ocr"}
    assert all(r.after["credits"] for r in sends(live, "sent"))
    assert all(r.after["credits"] is None for r in sends(live, "sending"))
    [done] = audit_rows(live, "eval.run_finished")
    pages = sum(e.pages for e in scans(entries))
    assert done.after["kind"] == "ocr" and done.after["outcome"] == "complete"
    assert done.after["actual"]["ocr_credits"] == pages and done.after["actual"]["calls"] == 11

    again = Typist()
    assert run_ocr(live, ask=again, fill=fake)["files"] == 0
    assert again.prompts == [] and fake.calls == names  # no estimate, no ask, no call
    assert len(audit_rows(live, "eval.spend_approved")) == 1


def test_a_scan_read_before_a_crash_or_a_re_import_is_frozen_again_from_its_cache(imported, live):
    """FR-20: the OCR cache is the pinned artifact. A process that died between the cache write
    and the manifest update, or a second `owner-import`, must not leave a read scan looking
    unread (that would make the next run pay for OCR again)."""
    src, entries = imported
    scan = scans(entries)[0]
    asyncio.run(
        owner_mod.FrozenOCR(fill=FakeOCR()).extract_text(
            filename=scan.name, mime_type=scan.mime_type, content=(src / scan.name).read_bytes()
        )
    )
    assert all(e.frozen_text_sha256 is None for e in scans(owner_mod.read_manifest()))
    assert owner_mod.owner_status()["ocr_cached"] == 1
    [recovered] = [e for e in owner_mod.sync_frozen_texts() if e.frozen_text_sha256 and e.kind != "text"]
    assert recovered.doc_id == scan.doc_id and recovered.frozen_method == "reducto_ocr"

    asyncio.run(owner_mod.owner_import(live, OWNER, src_dir=src))
    again = {e.doc_id: e for e in owner_mod.read_manifest()}
    assert again[scan.doc_id].frozen_text_sha256 == recovered.frozen_text_sha256
    assert owner_mod.owner_status()["frozen_texts"] == 4  # 3 text files + the one scan read


def test_a_stamp_only_scan_is_frozen_from_forced_ocr_not_from_its_stamp(
    private, live, tmp_path, monkeypatch
):
    """FR-23: should the app ever again read a stamp-only scan without OCR (as it did before
    2026-10-07), the main figures still need its full text -- so the exam forces OCR for it."""
    monkeypatch.setattr(owner_mod, "_wants_ocr", lambda read: False)  # the old app's decision
    asyncio.run(owner_mod.owner_import(live, OWNER, src_dir=owner_folder(tmp_path)))
    run_ocr(live, ask=Typist(ceiling="500"), fill=FakeOCR())
    stamp_only = next(e for e in owner_mod.read_manifest() if e.kind == "stamp_only_scan")
    frozen_text, _ = owner_mod._frozen(stamp_only)
    assert len(frozen_text) > 1000 and "DocuSign" not in frozen_text


def test_ocr_stops_before_the_file_that_would_pass_the_ceiling(imported, live):
    """FR-32: the ceiling (in credits here) is checked before each file, so spend never passes
    it; the run is incomplete, says why, and the unread scans are simply not frozen."""
    _, entries = imported
    first_two = sum(e.pages for e in scans(entries)[:2])
    fake = FakeOCR()
    result = run_ocr(live, ask=Typist(ceiling=str(first_two)), fill=fake)
    assert len(fake.calls) == 2 and result["outcome"] == "incomplete"
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["outcome"] == "incomplete" and done.after["incomplete_reason"] == "ceiling_reached"
    assert done.after["actual"]["ocr_credits"] == first_two <= done.after["ceiling_credits"]
    frozen = [e for e in scans(owner_mod.read_manifest()) if e.frozen_text_sha256]
    assert [e.name for e in frozen] == fake.calls


def test_a_failed_ocr_read_is_counted_and_listed_never_frozen(imported, live):
    """FR-40: a provider error on one scan is a failed call, shown as such and flagged 'contains
    errors'; the rest are still read, and the failed scan stays unfrozen so a keyed run on it fails.
    Every file was attempted, so the run is complete, not incomplete (and 'failed_calls' is not
    an incomplete_reason: the plan's reasons are ceiling_reached, crashed, interrupted, abandoned)."""
    _, entries = imported
    victim = scans(entries)[2].name
    result = run_ocr(live, ask=Typist(ceiling="500"), fill=FakeOCR(fail={victim}))
    assert result["failed"] == 1 and result["outcome"] == "complete"
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["contains_errors"] is True and done.after["incomplete_reason"] is None
    assert done.after["failed_call_rate"] == pytest.approx(1 / 11)
    errors = [r for r in sends(live) if r.after.get("error")]
    assert [(r.after["stage"], r.after["error"]) for r in errors] == [("failed", "RuntimeError")]
    unfrozen = [e.name for e in scans(owner_mod.read_manifest()) if not e.frozen_text_sha256]
    assert unfrozen == [victim]


def test_a_file_replaced_after_import_is_not_sent_to_ocr(imported, live):
    """A scan swapped on disk after `owner-import` would be paid for and frozen under a hash the
    manifest and the keys do not know; it is never sent. Nothing was sent, so it is NOT a call
    (not counted as one, not a failed one) -- but a planned read was never attempted, so the run
    is incomplete and names the file."""
    _, entries = imported
    victim = scans(entries)[1]
    (owner_mod.OWNER_DIR / "files" / victim.name).write_bytes(scan_pdf(99))
    fake = FakeOCR()
    result = run_ocr(live, ask=Typist(ceiling="500"), fill=fake)
    assert victim.name not in fake.calls and len(fake.calls) == 10
    assert result["failed"] == 0 and result["changed"] == [victim.doc_id]
    assert result["outcome"] == "incomplete"
    assert len(sends(live, "sending")) == len(sends(live, "sent")) == 10
    assert all(r.after["doc_id"] != victim.doc_id for r in sends(live))  # nothing was sent
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["actual"]["calls"] == 10 and done.after["actual"]["failed_calls"] == 0
    assert done.after["contains_errors"] is False and done.after["failed_call_rate"] == 0
    assert done.after["incomplete_reason"] == "interrupted"
    assert done.after["headline"]["files_changed_since_import"] == [victim.doc_id]


def test_each_owner_file_is_on_the_record_before_it_is_sent(imported, live):
    """FR-35: a client file that has left the machine must always be on the audit trail. The row
    is written BEFORE the provider is called -- a crash mid-call would otherwise leave a sent
    contract with no record -- and the provider is not called at all if it cannot be written."""
    real = FakeOCR()
    seen: dict[str, list[str]] = {}

    async def fill(*, filename, mime_type, content):
        seen[filename] = [r.after["stage"] for r in sends(live) if r.after["name"] == filename]
        return await real(filename=filename, mime_type=mime_type, content=content)

    run_ocr(live, ask=Typist(ceiling="500"), fill=fill)
    assert len(seen) == 11 and all(stages == ["sending"] for stages in seen.values())


def test_an_audit_row_that_cannot_be_written_stops_the_run_before_anything_is_sent(
    imported, live, monkeypatch
):
    """FR-35: no record, no send. The refusal is not counted as a failed call: no call was made."""
    fake = FakeOCR()

    def broken(*args, **kwargs):
        raise RuntimeError("audit database unavailable")

    monkeypatch.setattr(env, "audit", broken)
    with pytest.raises(RuntimeError, match="audit database"):
        run_ocr(live, ask=Typist(ceiling="500"), fill=fake)
    assert fake.calls == []


def test_an_interrupted_ocr_leaves_the_send_on_the_record(imported, live):
    """FR-35: Ctrl-C is a KeyboardInterrupt, not an Exception. The file in flight keeps its
    'sending' row and gains an 'interrupted' one, the run is recorded incomplete (interrupted)
    and counts the call that was made, and the scans read before it stay frozen."""
    _, entries = imported
    victim = scans(entries)[3]
    with pytest.raises(KeyboardInterrupt):
        run_ocr(live, ask=Typist(ceiling="500"), fill=FakeOCR(interrupt={victim.name}))
    mine = [r.after["stage"] for r in sends(live) if r.after["doc_id"] == victim.doc_id]
    assert mine == ["sending", "interrupted"]
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["outcome"] == "incomplete" and done.after["incomplete_reason"] == "interrupted"
    assert done.after["actual"]["calls"] == 4 and done.after["actual"]["failed_calls"] == 0
    assert len([e for e in scans(owner_mod.read_manifest()) if e.frozen_text_sha256]) == 3
    assert claude.exam_meter is None


def test_the_ocr_run_record_carries_the_versions_it_ran_on(imported, live, capsys):
    """FR-39: the immutable record says which code, model under test and exam content a run
    started with (taken before the run freezes anything). The estimate shows no 'Model:' line --
    no AI model reads a scan, and the model under test would mislead."""
    content = ExamContent()
    run_ocr(live, ask=Typist(ceiling="500"), fill=FakeOCR(), content=content)
    [done] = audit_rows(live, "eval.run_finished")
    versions = done.after["versions"]
    assert versions["app_code_sha256"] == app_code_sha256()
    assert versions["content_fingerprint"] == content_fingerprint(content)
    assert versions["model_under_test"] == settings.claude_model and versions["models"] == []
    [approved] = audit_rows(live, "eval.spend_approved")
    assert approved.after["versions_at_start"] == {
        "app_code_sha256": versions["app_code_sha256"],
        "content_fingerprint": versions["content_fingerprint"],
    }
    assert "Model:" not in capsys.readouterr().out


def test_ocr_refuses_the_mock_provider_before_anything_is_shown(imported, live, monkeypatch):
    """Under MOCK_REDUCTO the provider returns empty text: a paid command must refuse, not
    'succeed' with nothing, and it must do so before asking the owner to approve anything."""
    monkeypatch.setattr(settings, "reducto_api_key", "a-key")
    typist = Typist()
    with pytest.raises(ExamRefused, match="MOCK_REDUCTO"):
        asyncio.run(owner_mod.run_ocr(live=live, owner=OWNER, ask=typist, content=ExamContent()))
    monkeypatch.setattr(settings, "mock_reducto", False)
    monkeypatch.setattr(settings, "reducto_api_key", None)
    with pytest.raises(ExamRefused, match="REDUCTO_API_KEY"):
        asyncio.run(owner_mod.run_ocr(live=live, owner=OWNER, ask=typist, content=ExamContent()))
    assert typist.prompts == [] and audit_rows(live, "eval.spend_approved") == []


def test_a_constructed_scan_the_app_would_ocr_is_read_too_but_is_not_a_client_contract(
    imported, live, tmp_path, monkeypatch
):
    """FR-20/FR-35: constructed scans share the OCR step (the plan says any exam file the app
    would OCR), but they are not client data, so only owner files get a real_contract_sent row."""
    docs = tmp_path / "contracts"
    docs.mkdir()
    data = scan_pdf(7)
    (docs / "scan-ae-msa-01.pdf").write_bytes(data)
    monkeypatch.setattr(owner_mod, "DOCS_DIR", docs)
    twin = ExamDocument(
        doc_id="scan-ae-msa-01", category="constructed", dataset=None, file="scan-ae-msa-01.pdf",
        mime_type="application/pdf", file_sha256=hashlib.sha256(data).hexdigest(),
        title="scan-ae-msa-01", twin_of="ae-msa-01", ai_steps=["metadata"],
    )
    fake = FakeOCR()
    run_ocr(live, ask=Typist(ceiling="500"), fill=fake, content=ExamContent(documents=[twin]))
    assert len(fake.calls) == 12 and fake.calls[-1] == "scan-ae-msa-01.pdf"
    assert (owner_mod.OCR_DIR / f"{twin.file_sha256}.json").is_file()
    assert len(sends(live, "sending")) == 11  # the 11 owner files only


# --- the pre-fill (FR-3, FR-16, FR-17, FR-35) -----------------------------------------------


def contract_text(i: int) -> str:
    return (
        f"[[DOC {i:02d}]] SERVICES AGREEMENT\n"
        "This agreement is between Acme Ltd and Beta Pvt Ltd.\n"
        "It is governed by the laws of India.\n"
        "The agreement starts on 1 April 2024 and runs for three years.\n"
        "Either party may end it on ninety days written notice.\n"
    )


def fake_owner_set(*, skip: set[int] | None = None) -> list[OwnerManifestEntry]:
    """An imported and frozen owner set (all 14 as 'text' files), without PDFs: the pre-fill and
    the keys only need the manifest, the file bytes and the frozen texts."""
    entries = []
    for i, name in enumerate(OWNER_FILES, start=1):
        data = f"fake contract bytes {i}".encode()
        sha = hashlib.sha256(data).hexdigest()
        owner_mod._write_atomic(owner_mod.OWNER_DIR / "files" / name, data)
        text = contract_text(i)
        # contract 2 carries a page map, so its 'where' can name a page
        pages = {"1": {"start": 0, "end": 60}, "2": {"start": 61, "end": len(text)}} if i == 2 else None
        frozen_sha = owner_mod._freeze_text(sha, text, pages)
        entries.append(
            OwnerManifestEntry(
                doc_id=f"owner-{i:02d}", name=name, file_sha256=sha, mime_type="application/pdf",
                pages=3, kind="text", imported_at=datetime.now(UTC),
                frozen_text_sha256=None if skip and i in skip else frozen_sha,
                frozen_method="documents_reader",
            )
        )
    owner_mod._write_manifest(entries)
    return entries


STANDARD_ANSWERS = {
    "answers": [
        {"field": "parties", "value": "Acme Ltd | Beta Pvt Ltd",
         "quote": "This agreement is between Acme Ltd and Beta Pvt Ltd."},
        {"field": "effective_date", "value": "1 April 2024",
         "quote": "The agreement starts on 1 April 2024 and runs for three years."},
        # a quote the contract does not contain: FR-17 must catch it
        {"field": "end_date_or_term", "value": "three years",
         "quote": "The agreement ended on the first of April."},
        # spans the line break in the contract
        {"field": "governing_law", "value": "India",
         "quote": "between Acme Ltd and Beta Pvt Ltd. It is governed by the laws of India."},
        {"field": "notice_period", "value": "ninety days",
         "quote": "Either party may end it on ninety days written notice."},
        {"field": "renewal_term", "value": "", "quote": ""},  # the contract does not say
    ]
}


def prefill_answers(monkeypatch, *, raw: dict[int, str] | None = None) -> None:
    """Make the fake Anthropic client answer the pre-fill prompt: the standard answers, or
    ``raw[i]`` verbatim for contract i."""
    raw = raw or {}
    original = exam_fake_anthropic.message_for

    def message_for(payload, usage=None):
        message = original(payload, usage)
        if "output_config" in payload:
            i = int(re.search(r"\[\[DOC (\d+)\]\]", payload["messages"][0]["content"]).group(1))
            answer = raw.get(i, json.dumps(STANDARD_ANSWERS))
            message["content"] = [{"type": "text", "text": answer}]
        return message

    monkeypatch.setattr(exam_fake_anthropic, "message_for", message_for)


def run_prefill(live, *, ask, content=None):
    """The ``prefill`` command, on an empty exam unless told otherwise: the run record's content
    fingerprint must not depend on whatever private files this machine happens to hold."""
    return asyncio.run(
        owner_mod.run_prefill(
            live=live, owner=OWNER, ask=ask,
            content=content if content is not None else ExamContent(),
        )
    )


def sheet_rows(path: Path) -> dict[tuple[str, str], dict]:
    """A facts sheet as the owner sees it: (contract file name, field) -> the cells."""
    by_question = {q: f for f, q in sheets.FACT_QUESTIONS.items()}
    rows = {}
    for r in load_workbook(path)["Facts"].iter_rows(min_row=2, values_only=True):
        rows[(r[1], by_question[r[2]])] = {
            "suggested": r[3], "sentence": r[4], "where": r[5] or "", "mark": r[6]
        }
    return rows


def mark_sheet(path: Path, overrides: dict | None = None) -> None:
    """Mark a facts sheet the way a careful owner would: every supported row correct, rows with
    no supporting sentence left blank, and ``overrides[(contract, field)] = (mark, correction)``."""
    overrides = overrides or {}
    by_question = {q: f for f, q in sheets.FACT_QUESTIONS.items()}
    wb = load_workbook(path)
    for row in wb["Facts"].iter_rows(min_row=2):
        key = (row[1].value, by_question[row[2].value])
        mark, correction = "correct", ""
        if row[4].value == sheets.NO_SUPPORT:
            mark = ""
        mark, correction = overrides.get(key, (mark, correction))
        row[6].value, row[7].value = mark, correction
    wb.save(path)


@pytest.fixture
def frozen_set(private):
    return fake_owner_set()


@pytest.fixture
def answered(frozen_set, live, client_on_fake, monkeypatch):
    """The fake Anthropic client answering the pre-fill prompt with the standard answers."""
    prefill_answers(monkeypatch)
    return client_on_fake


def test_prefill_refuses_when_the_model_under_test_is_the_prefill_model(
    frozen_set, live, client_on_fake, monkeypatch
):
    """AC-16/FR-3: a model must never fill the answers it is graded on. Refused before anything is
    shown, asked or sent -- also when the app's model carries a date suffix."""
    for model_under_test in (PREFILL_MODEL, f"{PREFILL_MODEL}-20260301"):
        monkeypatch.setattr(settings, "claude_model", model_under_test)
        typist = Typist()
        with pytest.raises(ExamRefused, match="model under test"):
            run_prefill(live, ask=typist)
        assert typist.prompts == []
    assert client_on_fake.streamed == []
    assert audit_rows(live, "eval.spend_approved") == [] and audit_rows(live, "eval.spend_refused") == []


def test_prefill_refuses_a_fake_model(frozen_set, live):
    """Under MOCK_CLAUDE the 'answers' are canned: they must never reach a sheet the owner will
    mark as if a real model had read the contracts."""
    assert settings.mock_claude  # the suite's own setting, as CI runs it
    typist = Typist()
    with pytest.raises(ExamRefused, match="MOCK_CLAUDE"):
        run_prefill(live, ask=typist)
    assert typist.prompts == []


def test_prefill_refuses_until_all_14_texts_are_frozen(private, live, client_on_fake):
    """FR-20: keys are tied to a frozen text, so pre-fill needs all 14; the message names the
    ones that are missing (here a scan nobody has OCR-read yet)."""
    fake_owner_set(skip={4, 9})
    typist = Typist()
    with pytest.raises(ExamRefused, match="owner-04, owner-09"):
        run_prefill(live, ask=typist)
    assert typist.prompts == [] and client_on_fake.streamed == []


def test_prefill_sends_each_whole_contract_once_to_the_prefill_model_never_batched(
    answered, live, monkeypatch
):
    """AC-16/FR-16/FR-35/FR-41: 14 synchronous calls on claude-sonnet-5-5 in structured-output
    form (no temperature, no tools), each carrying the whole frozen text; Anthropic is the only
    destination; the ledger rows hold no contract text even with raw-output storage on; and the
    estimate's worst case is exactly what the meter checks, so 'set the ceiling to it' is true."""
    monkeypatch.setattr(settings, "ai_store_raw_outputs", True)
    result = run_prefill(live, ask=Typist(ceiling="25"))
    assert result["outcome"] == "complete" and result["rows"] == 84

    assert len(answered.streamed) == 14 and answered.batch_creates == []
    texts = {i: contract_text(i) for i in range(1, 15)}
    for payload in answered.streamed:
        assert payload["model"] == PREFILL_MODEL
        assert payload["output_config"]["format"]["type"] == "json_schema"
        assert not {"temperature", "tools", "tool_choice"} & payload.keys()
        i = int(re.search(r"\[\[DOC (\d+)\]\]", payload["messages"][0]["content"]).group(1))
        assert texts[i] in payload["messages"][0]["content"]  # the whole frozen text

    rows = ledger_rows(live, result["run_id"])
    assert len(rows) == 14 and {r.model for r in rows} == {PREFILL_MODEL}
    for r in rows:
        assert r.status == "succeeded" and r.batched is False and r.prompt_key == "exam_prefill"
        assert not r.raw_ai_output and r.validated_output is None  # nothing the model said
        assert set(r.input_payload) == {"doc_id", "chars"}  # a summary, never the text
    # each contract is on the record twice: written before it left, then with its tokens and cost
    assert len(sends(live, "sending")) == len(sends(live, "sent")) == 14 and len(sends(live)) == 28
    assert {r.after["destination"] for r in sends(live)} == {f"anthropic:{PREFILL_MODEL}"}
    assert {r.after["purpose"] for r in sends(live)} == {"prefill"}
    assert all(r.after["tokens"] == {"prompt": 100, "completion": 20} for r in sends(live, "sent"))
    assert all(r.after["tokens"] is None for r in sends(live, "sending"))
    [approved] = audit_rows(live, "eval.spend_approved")
    assert approved.after["kind"] == "prefill" and approved.after["delivery"] == "sync"
    assert approved.after["model"] == PREFILL_MODEL
    assert approved.after["estimate"]["worst_case_usd"] == pytest.approx(
        sum(runner.worst_cost(p, batched=False) for p in answered.streamed)
    )
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["outcome"] == "complete" and done.after["actual"]["calls"] == 14
    assert done.after["headline"] == {
        "contracts_answered": 14, "contracts_failed": [], "sheet_written": True
    }
    assert claude.exam_meter is None


def test_the_prefill_run_record_carries_the_models_and_versions(answered, live):
    """FR-16/FR-39: the immutable record names the pre-fill model, the model under test, the
    prompt and the exam content the run started with -- so it can be told later which model made
    the suggestions that the owner confirmed (AC-16)."""
    content = ExamContent()
    run_prefill(live, ask=Typist(), content=content)
    [done] = audit_rows(live, "eval.run_finished")
    v = done.after["versions"]
    assert v["models"] == [PREFILL_MODEL] and v["prefill_model"] == PREFILL_MODEL
    assert v["model_under_test"] == settings.claude_model != PREFILL_MODEL
    assert v["app_code_sha256"] == app_code_sha256()
    assert v["content_fingerprint"] == content_fingerprint(content)
    assert [(p["prompt_key"], p["version"]) for p in v["prompts"]] == [("exam_prefill", "1")]
    assert len(v["prompts"][0]["hash"]) == 64
    [approved] = audit_rows(live, "eval.spend_approved")
    assert approved.after["versions_at_start"] == {
        "app_code_sha256": v["app_code_sha256"],
        "content_fingerprint": v["content_fingerprint"],
        "prefill_model": PREFILL_MODEL,
    }


def test_the_prefill_schema_requires_every_fact_with_a_value_or_null_and_a_quote_or_null(
    answered, live
):
    """FR-16: the structured-output schema makes each answer complete (field, value, quote all
    required; value and quote a string or null, so 'not stated' is an explicit answer)."""
    run_prefill(live, ask=Typist())
    schema = answered.streamed[0]["output_config"]["format"]["schema"]
    item = schema["properties"]["answers"]["items"]
    assert set(item["required"]) == {"field", "value", "quote"}
    assert item["additionalProperties"] is False
    assert item["properties"]["field"]["enum"] == list(owner_mod._PREFILL_FIELDS)
    for name in ("value", "quote"):
        assert item["properties"][name] == {"anyOf": [{"type": "string"}, {"type": "null"}]}
    assert "exactly one answer for each of the six facts" in answered.streamed[0]["system"][0]["text"]


def test_prefill_shows_the_estimate_and_the_model_before_asking(answered, live, capsys):
    """FR-31: the owner sees the pre-fill model (and that it is not the model under test), the
    14 calls and their cost before typing anything."""
    run_prefill(live, ask=Typist())
    out = capsys.readouterr().out
    assert f"Pre-fill model: {PREFILL_MODEL}" in out and "never batched" in out
    assert "14 (0 batch, 14 sync)" in out and "exam_prefill" in out and "Delivery: sync" in out


def test_prefill_without_a_go_ahead_sends_nothing(answered, live):
    """AC-31: no go-ahead, no contract leaves the machine."""
    with pytest.raises(ExamRefused, match="nothing was sent"):
        run_prefill(live, ask=Typist(approve=False))
    assert answered.streamed == [] and sends(live) == []
    assert len(audit_rows(live, "eval.spend_refused")) == 1


def test_prefill_checks_every_quote_and_an_unsupported_row_cannot_be_marked_correct(
    answered, live
):
    """AC-17/FR-17: the sentence shown is the contract's own (a quote spanning a line break is
    found); a quote that is not in the contract reads 'no supporting sentence found' and the
    sheet is rejected whole if that row is marked correct; 'not stated' needs no sentence; a page
    shows when the text has a page map."""
    result = run_prefill(live, ask=Typist())
    path = sheets.SHEETS_DIR / f"facts-{result['sheet_id']}.xlsx"
    rows = sheet_rows(path)
    first = OWNER_FILES[0]
    assert rows[(first, "parties")]["suggested"] == "Acme Ltd | Beta Pvt Ltd"
    law = rows[(first, "governing_law")]
    assert "Beta Pvt Ltd.\nIt is governed by the laws of India." in law["sentence"]
    assert law["where"].startswith("characters ")
    bad = rows[(first, "end_date_or_term")]
    assert bad["suggested"] == "three years" and bad["sentence"] == sheets.NO_SUPPORT
    assert bad["where"] == ""
    open_ended = rows[(first, "renewal_term")]
    assert open_ended["suggested"] == sheets.NOT_STATED
    assert open_ended["sentence"] == sheets.NOTHING_TO_QUOTE
    assert rows[(OWNER_FILES[1], "notice_period")]["where"].startswith("page 2, characters ")
    assert result["unsupported"] == 14

    mark_sheet(path, {(first, "end_date_or_term"): ("correct", "")})
    with pytest.raises(ExamRefused, match="no supporting sentence"):
        owner_mod.load_facts(live, OWNER, path)
    assert owner_mod.read_fact_keys() == []

    mark_sheet(path, {(first, "end_date_or_term"): ("wrong", "three years")})
    owner_mod.load_facts(live, OWNER, path)
    corrected = next(
        k for k in owner_mod.read_fact_keys()
        if (k.doc_id, k.field) == ("owner-01", "end_date_or_term")
    )
    assert corrected.status == "corrected" and corrected.value == "three years"
    assert corrected.quote == "three years" and corrected.suggested_by == PREFILL_MODEL


def test_a_ceiling_that_cannot_cover_a_call_stops_the_prefill_before_anything_is_sent(
    answered, live
):
    """FR-32: the meter refuses a call whose worst case does not fit the ceiling, before the
    request is made -- so no contract is sent, no sheet is written, and the record says why."""
    result = run_prefill(live, ask=Typist(ceiling="0.0001"))
    assert result["outcome"] == "incomplete" and result["sheet_id"] is None
    # the first contract is on the record as about to be sent, then as refused before any request
    assert answered.streamed == [] and sends(live, "sent") == []
    assert [(r.after["doc_id"], r.after["stage"]) for r in sends(live)] == [
        ("owner-01", "sending"), ("owner-01", "not_sent")
    ]
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["incomplete_reason"] == "ceiling_reached"
    assert not sheets.SHEETS_DIR.exists() or not list(sheets.SHEETS_DIR.glob("*.xlsx"))
    assert claude.exam_meter is None


def test_an_unusable_prefill_answer_is_a_failed_call_and_no_sheet_is_written(
    frozen_set, live, client_on_fake, monkeypatch
):
    """FR-40: an answer that is not the asked shape is on the ledger as failed, counted in the run
    record ('contains errors'), and the sheet is not written with a contract missing from it.
    Every contract was attempted, so the run is complete (with errors), not incomplete."""
    prefill_answers(monkeypatch, raw={5: "this is not json"})
    result = run_prefill(live, ask=Typist())
    assert result["outcome"] == "complete" and result["failed"] == ["owner-05"]
    assert result["sheet_id"] is None
    statuses = sorted(r.status for r in ledger_rows(live, result["run_id"]))
    assert statuses == ["succeeded"] * 13 + ["validation_failed"]
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["contains_errors"] is True and done.after["incomplete_reason"] is None
    assert done.after["failed_call_rate"] == pytest.approx(1 / 14)
    assert done.after["headline"] == {
        "contracts_answered": 13, "contracts_failed": ["owner-05"], "sheet_written": False
    }
    assert len(sends(live, "sent")) == 14  # it was sent, so it is audited


def without(field: str) -> str:
    """The standard answer with one fact left out of the reply."""
    return json.dumps({"answers": [a for a in STANDARD_ANSWERS["answers"] if a["field"] != field]})


@pytest.mark.parametrize(
    ("reply", "why"),
    [
        (without("notice_period"), "no answer for notice_period"),
        (
            json.dumps({"answers": [*STANDARD_ANSWERS["answers"], STANDARD_ANSWERS["answers"][0]]}),
            "more than one answer for parties",
        ),
        (
            json.dumps({"answers": [
                *STANDARD_ANSWERS["answers"], {"field": "liability_cap", "value": None, "quote": None}
            ]}),
            "1 answer(s) for facts that were not asked",
        ),
        (
            json.dumps({"answers": [
                {k: v for k, v in a.items() if k != "quote"} if a["field"] == "parties" else a
                for a in STANDARD_ANSWERS["answers"]
            ]}),
            "not the asked JSON",
        ),
    ],
)
def test_a_reply_that_leaves_a_fact_out_is_a_failed_prefill_never_a_silent_not_stated(
    frozen_set, live, client_on_fake, monkeypatch, capsys, reply, why
):
    """FR-16/FR-3: a fact the model forgot must not reach the owner as 'not stated' -- the owner
    could mark that correct and a wrong answer would enter the key. It is a failed pre-fill for
    that contract (reported with the reason, on the ledger as failed, no sheet), and the reason
    names facts only, never what the model wrote."""
    prefill_answers(monkeypatch, raw={3: reply})
    result = run_prefill(live, ask=Typist())
    assert result["failed"] == ["owner-03"] and result["sheet_id"] is None
    shown = capsys.readouterr().out
    assert "owner-03: the pre-fill answer was unusable" in shown and why in shown
    [bad] = [r for r in ledger_rows(live, result["run_id"]) if r.status == "validation_failed"]
    assert why in bad.validation_error and "Acme" not in bad.validation_error
    assert not sheets.SHEETS_DIR.exists() or not list(sheets.SHEETS_DIR.glob("*.xlsx"))


def test_a_stated_null_is_a_real_not_stated_answer(frozen_set, live, client_on_fake, monkeypatch):
    """The other side of the rule above: null (or an empty string) for both value and quote IS the
    model saying 'the contract does not say', and is shown as exactly that."""
    answers = {"answers": [
        {"field": a["field"], "value": None, "quote": None} for a in STANDARD_ANSWERS["answers"]
    ]}
    answers["answers"][5] = {"field": "renewal_term", "value": "", "quote": ""}
    prefill_answers(monkeypatch, raw={i: json.dumps(answers) for i in range(1, 15)})
    result = run_prefill(live, ask=Typist())
    assert result["failed"] == [] and result["rows"] == 84 and result["unsupported"] == 0
    rows = sheet_rows(sheets.SHEETS_DIR / f"facts-{result['sheet_id']}.xlsx")
    assert {r["suggested"] for r in rows.values()} == {sheets.NOT_STATED}
    assert {r["sentence"] for r in rows.values()} == {sheets.NOTHING_TO_QUOTE}


def test_an_interrupted_prefill_leaves_the_send_on_the_record(
    frozen_set, live, client_on_fake, monkeypatch
):
    """FR-35: Ctrl-C during a call is a KeyboardInterrupt. The contract in flight keeps its
    'sending' row and gains an 'interrupted' one; the run is recorded incomplete (interrupted),
    the meter is removed, and no sheet is written."""
    prefill_answers(monkeypatch)
    real_call = claude.claude_client.complete_json

    async def interrupted(**kwargs):
        if "[[DOC 04]]" in kwargs["user_prompt"]:
            raise KeyboardInterrupt
        return await real_call(**kwargs)

    monkeypatch.setattr(claude.claude_client, "complete_json", interrupted)
    with pytest.raises(KeyboardInterrupt):
        run_prefill(live, ask=Typist())
    mine = [r.after["stage"] for r in sends(live) if r.after["doc_id"] == "owner-04"]
    assert mine == ["sending", "interrupted"]
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["outcome"] == "incomplete" and done.after["incomplete_reason"] == "interrupted"
    assert claude.exam_meter is None
    assert not sheets.SHEETS_DIR.exists() or not list(sheets.SHEETS_DIR.glob("*.xlsx"))


def test_each_contract_is_on_the_record_before_it_is_sent_to_the_prefill_model(
    answered, live, monkeypatch
):
    """FR-35: the 'sending' row exists before the provider is asked, for every contract."""
    real_call = claude.claude_client.complete_json
    seen: dict[str, list[str]] = {}

    async def watching(**kwargs):
        doc = re.search(r"\[\[DOC (\d+)\]\]", kwargs["user_prompt"]).group(1)
        seen[doc] = [r.after["stage"] for r in sends(live) if r.after["doc_id"] == f"owner-{doc}"]
        return await real_call(**kwargs)

    monkeypatch.setattr(claude.claude_client, "complete_json", watching)
    run_prefill(live, ask=Typist())
    assert len(seen) == 14 and all(stages == ["sending"] for stages in seen.values())


def test_prefill_without_an_audit_row_sends_nothing(answered, live, monkeypatch):
    """FR-35: if the 'sending' row cannot be written the contract is not sent."""
    approved_audit = env.audit

    def audit_only_the_gate(live_, *, action, **kwargs):
        if action == "eval.real_contract_sent":
            raise RuntimeError("audit database unavailable")
        return approved_audit(live_, action=action, **kwargs)

    monkeypatch.setattr(env, "audit", audit_only_the_gate)
    with pytest.raises(RuntimeError, match="audit database"):
        run_prefill(live, ask=Typist())
    assert answered.streamed == []
    [done] = audit_rows(live, "eval.run_finished")
    assert done.after["incomplete_reason"] == "interrupted" and claude.exam_meter is None


def test_a_failed_prefill_call_is_on_the_ledger_and_the_audit_trail_with_its_error(
    frozen_set, live, client_on_fake, monkeypatch
):
    """FR-40: a provider error is counted and listed, never silently dropped; the contract was
    still handed to the provider, so the send is audited with the error class."""
    prefill_answers(monkeypatch)
    real_call = claude.claude_client.complete_json

    async def flaky(**kwargs):
        if "[[DOC 02]]" in kwargs["user_prompt"]:
            raise ConnectionError("network down")
        return await real_call(**kwargs)

    monkeypatch.setattr(claude.claude_client, "complete_json", flaky)
    result = run_prefill(live, ask=Typist())
    assert result["failed"] == ["owner-02"] and result["sheet_id"] is None
    failed = [r for r in ledger_rows(live, result["run_id"]) if r.status == "failed"]
    assert [r.error_class for r in failed] == ["ConnectionError"]
    errors = [r for r in sends(live) if r.after.get("error")]
    assert [(r.after["doc_id"], r.after["stage"]) for r in errors] == [("owner-02", "failed")]


# --- the keys: load, replace, go out of date (FR-18, FR-19, FR-20) --------------------------


@pytest.fixture
def prefilled(answered, live):
    """A facts sheet pre-filled by the fake model for the 14 frozen contracts."""
    result = run_prefill(live, ask=Typist())
    return sheets.SHEETS_DIR / f"facts-{result['sheet_id']}.xlsx"


def test_load_facts_turns_marks_into_keys_and_audits_every_row(prefilled, live):
    """AC-18/AC-19/FR-3: correct -> the suggestion becomes the key; wrong -> the owner's value;
    blank -> unmarked (kept so it can be counted, never scored). The audit row records the sheet's
    fingerprint and, per row, the suggestion and the final value."""
    first = OWNER_FILES[0]
    mark_sheet(
        prefilled,
        {
            (first, "governing_law"): ("wrong", "Singapore"),
            (first, "notice_period"): ("", ""),
        },
    )
    result = owner_mod.load_facts(live, OWNER, prefilled)
    keys = {(k.doc_id, k.field): k for k in owner_mod.read_fact_keys()}
    assert len(keys) == 84
    assert keys[("owner-01", "parties")].status == "confirmed"
    assert keys[("owner-01", "parties")].value == ["Acme Ltd", "Beta Pvt Ltd"]
    law = keys[("owner-01", "governing_law")]
    assert (law.status, law.value, law.suggested_value) == ("corrected", "Singapore", "India")
    assert law.suggested_quote and law.suggested_by == PREFILL_MODEL
    assert keys[("owner-01", "notice_period")].status == "unmarked"
    assert keys[("owner-01", "renewal_term")].value is None  # "not stated" is a real answer
    # 14 rows have no supporting sentence (left blank), plus the notice period the test blanks
    assert result["counts"] == {"confirmed": 68, "corrected": 1, "unmarked": 15}
    manifest = {e.doc_id: e for e in owner_mod.read_manifest()}
    assert all(  # what the scorer compares a key with, or every key would read as out of date
        (k.file_sha256, k.frozen_text_sha256)
        == (manifest[k.doc_id].file_sha256, manifest[k.doc_id].frozen_text_sha256)
        for k in keys.values()
    )

    [row] = audit_rows(live, "eval.labels_loaded")
    assert row.resource_type == "eval_labels" and row.resource_id == result["sheet_id"]
    assert row.after["counts"] == result["counts"] and len(row.after["sheet_sha256"]) == 64
    by_row = {(r["doc_id"], r["field"]): r for r in row.after["rows"]}
    assert by_row[("owner-01", "governing_law")] | {"row_key": None} == {
        "row_key": None, "doc_id": "owner-01", "field": "governing_law",
        "suggested_value": "India",
        "suggested_quote": "between Acme Ltd and Beta Pvt Ltd.\nIt is governed by the laws of India.",
        "mark": "wrong", "final_value": "Singapore",
    }
    assert by_row[("owner-01", "notice_period")]["mark"] == "unmarked"


def test_a_sheet_that_no_longer_matches_what_was_written_loads_nothing(prefilled, live):
    """AC-18/FR-19: an edited question column rejects the whole sheet, listing the problem; no key
    file is written and nothing is audited as loaded."""
    wb = load_workbook(prefilled)
    wb["Facts"]["D2"].value = "something the AI never suggested"
    wb.save(prefilled)
    with pytest.raises(ExamRefused, match="nothing was loaded") as raised:
        owner_mod.load_facts(live, OWNER, prefilled)
    assert "a question column was changed" in str(raised.value)
    assert owner_mod.read_fact_keys() == []
    assert audit_rows(live, "eval.labels_loaded") == []


def test_load_facts_refuses_a_routing_sheet(private, live):
    """A routing sheet is loaded by its own command; here it must be refused, not half-read."""
    snapshot = {"fingerprint": "f" * 64, "workflows": [{"id": "w1", "name": "NDA", "enabled": True}]}
    row = {"item_id": "rt-01", "channel": "email", "subject": "s", "body": "b",
           "drafted_for_workflow_id": "w1"}
    sheet_id = sheets.write_routing_sheet([row], snapshot=snapshot, path=sheets.SHEETS_DIR)
    with pytest.raises(ExamRefused, match="routing sheet"):
        owner_mod.load_facts(live, OWNER, sheets.SHEETS_DIR / f"routing-{sheet_id}.xlsx")
    assert owner_mod.read_fact_keys() == []


def test_a_changed_file_or_frozen_text_puts_the_key_out_of_date_until_it_is_reconfirmed(
    prefilled, live, client_on_fake, monkeypatch
):
    """AC-20/FR-20: change a frozen text and that contract's keys are out of date (and counted in
    status); `sheet facts --reconfirm` offers only those, suggesting what the owner confirmed, with
    each quote checked again against the new text; loading it replaces just those keys; and a
    replaced contract file is caught the same way."""
    mark_sheet(prefilled)
    owner_mod.load_facts(live, OWNER, prefilled)
    assert owner_mod.stale_keys(owner_mod.read_fact_keys()) == []
    assert owner_mod.owner_status()["keys"]["out_of_date"] == 0

    entry = owner_mod.read_manifest()[2]
    changed = contract_text(3).replace("ninety days", "sixty days")
    owner_mod._freeze_text(entry.file_sha256, changed, None)
    stale = owner_mod.stale_keys(owner_mod.read_fact_keys())
    assert {k.doc_id for k in stale} == {"owner-03"} and len(stale) == 6
    assert owner_mod.owner_status()["keys"]["out_of_date"] == 6

    path = owner_mod.regenerate_facts_sheet(reconfirm=True)
    rows = sheet_rows(path)
    assert len(rows) == 5 and {c for c, _ in rows} == {entry.name}  # the unmarked row has nothing to confirm
    assert rows[(entry.name, "parties")]["suggested"] == "Acme Ltd | Beta Pvt Ltd"
    assert rows[(entry.name, "parties")]["sentence"] != sheets.NO_SUPPORT
    # the confirmed sentence about ninety days is gone from the new text
    assert rows[(entry.name, "notice_period")]["suggested"] == "ninety days"
    assert rows[(entry.name, "notice_period")]["sentence"] == sheets.NO_SUPPORT
    sheet_id = path.stem.removeprefix("facts-")
    expected = json.loads((path.parent / f"{sheet_id}.expected.json").read_text())
    assert expected["rows"] and all(r["suggested_by"] == "previous key" for r in expected["rows"])

    mark_sheet(path, {(entry.name, "notice_period"): ("wrong", "sixty days")})
    owner_mod.load_facts(live, OWNER, path)
    keys = owner_mod.read_fact_keys()
    assert len(keys) == 84  # replaced in place, not appended
    notice = next(k for k in keys if (k.doc_id, k.field) == ("owner-03", "notice_period"))
    assert notice.value == "sixty days" and notice.status == "corrected"
    left = owner_mod.stale_keys(keys)
    assert [(k.doc_id, k.field, k.status) for k in left] == [("owner-03", "end_date_or_term", "unmarked")]

    (owner_mod.OWNER_DIR / "files" / OWNER_FILES[4]).write_bytes(b"a different file")
    assert {k.doc_id for k in owner_mod.stale_keys(keys)} == {"owner-03", "owner-05"}


def test_reconfirm_with_nothing_out_of_date_is_refused(prefilled, live):
    """No sheet is made for keys that are still valid."""
    mark_sheet(prefilled)
    owner_mod.load_facts(live, OWNER, prefilled)
    with pytest.raises(ExamRefused, match="nothing to re-confirm"):
        owner_mod.regenerate_facts_sheet(reconfirm=True)


def test_owner_keys_for_a_document_that_was_never_imported_or_given_twice_are_refused(
    tmp_path, monkeypatch
):
    """FR-20: a key whose document is not in the owner manifest is scored against nothing and its
    contract would read 'without keys' for ever; a duplicate would count twice. Neither fails
    later, so the loader refuses them before a run is estimated. The files `load-facts` writes
    (one key per contract and field, for imported contracts) load."""
    key = {
        "doc_id": "owner-01", "field": "governing_law", "status": "confirmed", "value": "India",
        "quote": "laws of India", "char_start": 1, "char_end": 14, "suggested_value": "India",
        "suggested_quote": "laws of India", "suggested_by": PREFILL_MODEL,
        "file_sha256": "a" * 64, "frozen_text_sha256": "b" * 64, "sheet_sha256": "c" * 64,
        "loaded_at": "2026-10-07T00:00:00+00:00",
    }
    manifest = [{
        "doc_id": "owner-01", "name": OWNER_FILES[0], "file_sha256": "a" * 64,
        "mime_type": "application/pdf", "pages": 3, "kind": "text",
        "frozen_text_sha256": "b" * 64, "frozen_method": "documents_reader",
        "imported_at": "2026-10-07T00:00:00+00:00",
    }]

    def load(keys):
        for rel, value in (("owner/manifest.json", manifest), ("keys/owner_facts.json", keys)):
            path = tmp_path / "private" / rel
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(value), encoding="utf-8")
        monkeypatch.setattr(content_mod, "CONTENT_DIR", tmp_path / "content")
        monkeypatch.setattr(content_mod, "PRIVATE_DIR", tmp_path / "private")
        return content_mod.load_content()

    assert len(load([key]).owner_keys) == 1
    with pytest.raises(ValueError, match="inconsistent") as unknown:
        load([key, {**key, "doc_id": "owner-15"}])
    assert "owner key for owner-15 / governing_law: not an imported owner document" in str(
        unknown.value
    )
    with pytest.raises(ValueError, match="duplicate owner key for owner-01 / governing_law"):
        load([key, key])
