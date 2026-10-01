"""Tracked changes written into a Word file and resolved there.

The file is what a lawyer sends: it has to open in Word with a real redline, a
rejected change has to leave exactly what was there, and nobody's change may be
overwritten by somebody else's. Each test names the way that goes wrong.
"""

import io
import re
import zipfile

import pytest
from docx import Document
from lxml import etree

from app.docstudio.parsing.registry import DOCX_MIME, parser_for
from app.docstudio.redline import Edit, changes, propose, resolve, word_copy


def _docx(*paragraphs, bold: str = "", tab: bool = False) -> bytes:
    document = Document()
    for text in paragraphs:
        p = document.add_paragraph()
        if bold and bold in text:
            head, tail = text.split(bold, 1)
            p.add_run(head)
            p.add_run(bold).bold = True
            p.add_run(tail)
        else:
            p.add_run(text)
        if tab:
            p.runs[-1].add_tab()
            p.runs[-1].add_text("after the tab")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _text(content: bytes) -> str:
    """What our reader — and Word, with changes accepted — takes from the file."""
    return " | ".join(block.text for block in parser_for(DOCX_MIME).parse(content, filename="x.docx").blocks)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _xml(content: bytes) -> str:
    return zipfile.ZipFile(io.BytesIO(content)).read("word/document.xml").decode()


CAP = "Liability is capped at twelve (12) months of fees."


def test_an_edit_is_a_real_tracked_change_of_only_the_words_that_changed():
    """Replacing the whole selection would redline words that did not change;
    Word's compare marks word by word, and so must this."""
    content, applied = propose(_docx(CAP), [Edit("twelve (12) months", "twenty-four (24) months")], author="You")

    [change] = changes(content)
    assert applied.refused == []
    assert (change.kind, change.deleted, change.inserted, change.author) == (
        "replace", "twelve (12)", "twenty-four (24)", "You")
    assert "<w:del " in _xml(content) and "<w:ins " in _xml(content)
    assert "twenty-four (24) months" in _text(content)


def test_rejecting_leaves_exactly_what_was_there_and_accepting_leaves_no_marks():
    source = _docx(CAP)
    content, _ = propose(source, [Edit("twelve (12) months", "twenty-four (24) months")], author="You")

    rejected, _ = resolve(content, None, accept=False)
    accepted, _ = resolve(content, None, accept=True)

    assert _text(rejected) == _text(source)
    assert changes(accepted) == [] and "w:ins" not in _xml(accepted) and "w:del" not in _xml(accepted)
    assert "twenty-four (24) months" in _text(accepted)
    Document(io.BytesIO(accepted))  # still a document Word's own library opens


def test_the_new_words_keep_their_formatting_and_a_tab_in_the_run_survives():
    """Rebuilding an edited run from its text alone drops its tab; the new words
    must look like the ones they replace."""
    source = _docx(CAP, bold="twelve (12) months", tab=True)

    content, _ = propose(source, [Edit("after the tab", "after the tab, amended")], author="You")
    content, _ = propose(content, [Edit("twelve (12)", "twenty-four (24)")], author="You")
    accepted, _ = resolve(content, None, accept=True)

    runs = Document(io.BytesIO(accepted)).paragraphs[0].runs
    assert any(r.bold and "twenty-four" in r.text for r in runs)
    assert "<w:tab/>" in _xml(accepted)


def test_words_inside_someone_elses_change_are_refused_not_overwritten():
    content, _ = propose(_docx(CAP), [Edit("twelve (12)", "twenty-four (24)")], author="Their counsel")

    _, applied = propose(content, [Edit("twenty-four (24)", "eighteen (18)")], author="You")

    [(_, why)] = applied.refused
    assert "Accept or reject that first" in why


def test_an_edit_across_two_paragraphs_is_refused():
    _, applied = propose(_docx("First paragraph ends here.", "Second one starts."),
                         [Edit("ends here. Second one", "x")], author="You")

    assert "two paragraphs" in applied.refused[0][1]


def test_new_change_ids_never_collide_with_those_already_in_the_file():
    content, first = propose(_docx(CAP), [Edit("twelve (12)", "twenty-four (24)")], author="A")
    content, second = propose(content, [Edit("of fees", "of all fees")], author="B")

    ids = re.findall(r'<w:(?:ins|del) w:id="(\d+)"', _xml(content))
    # A replacement and then a pure insertion ("all" added, nothing removed).
    assert len(ids) == len(set(ids)) == 3
    assert not set(first.edits[0][1]) & set(second.edits[0][1])


def test_one_change_is_resolved_and_the_others_stay():
    content, _ = propose(_docx(CAP, "Payment is due within thirty days."),
                         [Edit("twelve (12)", "twenty-four (24)"), Edit("thirty", "forty-five")], author="You")
    first, second = changes(content)

    content, count = resolve(content, first.ids, accept=True)

    assert count == 2  # the deletion and the insertion of one replacement
    assert [c.inserted for c in changes(content)] == [second.inserted]


def test_accepting_a_removed_paragraph_mark_joins_the_paragraphs_as_word_does():
    source = _docx("Liability is capped", "at twelve months.")
    xml = _xml(source).replace(
        "<w:p>", '<w:p><w:pPr><w:rPr><w:del w:id="90" w:author="X" w:date="2026-01-01T00:00:00Z"/></w:rPr></w:pPr>', 1
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as old, zipfile.ZipFile(buffer, "w") as new:
        for item in old.infolist():
            new.writestr(item, xml if item.filename == "word/document.xml" else old.read(item.filename))

    [change] = changes(buffer.getvalue())
    joined, _ = resolve(buffer.getvalue(), change.ids, accept=True)

    assert change.kind == "paragraph"
    assert [p.text for p in Document(io.BytesIO(joined)).paragraphs] == ["Liability is cappedat twelve months."]


def test_a_replacement_is_one_change_in_either_order_the_file_keeps_it():
    """Word writes a replacement as deletion-then-insertion or the reverse; the
    software licence in the corpus has insertion first. Either is one change."""
    content, _ = propose(_docx(CAP), [Edit("twelve (12)", "twenty-four (24)")], author="You")
    xml = _xml(content)
    deletion = re.search(r"<w:del .*?</w:del>", xml, re.DOTALL).group()
    insertion = re.search(r"<w:ins .*?</w:ins>", xml, re.DOTALL).group()
    swapped = xml.replace(deletion + insertion, insertion + deletion)
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as old, zipfile.ZipFile(buffer, "w") as new:
        for item in old.infolist():
            new.writestr(item, swapped if item.filename == "word/document.xml" else old.read(item.filename))

    [change] = changes(buffer.getvalue())

    assert (change.kind, change.deleted, change.inserted) == ("replace", "twelve (12)", "twenty-four (24)")


def test_a_pdfs_words_become_a_word_document_our_reader_reads_back():
    from types import SimpleNamespace

    clauses = [
        SimpleNamespace(text="16.14 No Assignment", clause_type="heading", level=1),
        SimpleNamespace(text="Neither party may assign this Agreement.", clause_type="clause", level=2),
        SimpleNamespace(text="<table><tr><th>No</th><th>Item</th></tr><tr><td>1</td><td>Fees</td></tr></table>",
                        clause_type="table", level=1),
    ]

    content = word_copy(clauses, "Lockton MSA")

    text = _text(content)
    assert "No Assignment" in text and "Neither party may assign this Agreement." in text
    assert len(Document(io.BytesIO(content)).tables) == 1


@pytest.mark.parametrize("not_word", [b"%PDF-1.4 not a zip", b"plain text"])
def test_a_file_that_is_not_word_is_refused_plainly(not_word):
    from app.docstudio.redline import RedlineError

    with pytest.raises(RedlineError, match="not a Word document"):
        changes(not_word)


def test_a_saved_file_keeps_every_namespace_it_declares():
    """mc:Ignorable names prefixes like w14 that nothing else in the file uses.
    Tidied away with the rest, their declarations went, and Word called every
    redlined file unreadable — which python-docx, lenient, never noticed."""
    content, _ = propose(_docx(CAP), [Edit("twelve (12)", "twenty-four (24)")], author="You")

    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(content)).read("word/document.xml"))
    ignorable = root.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable").split()

    assert ignorable and set(ignorable) <= set(root.nsmap)


# --- a typed PDF as a Word document that looks like it ----------------------------

BODY = [
    "Each party shall keep the terms of this Agreement confidential at all times.",
    "The Supplier shall provide the Services with reasonable skill and care.",
    "Payment is due within thirty days of the date of a valid invoice.",
    "This Agreement is governed by the laws of England and Wales.",
]


def _pdf(*draw) -> bytes:
    import fitz

    document = fitz.open()
    page = document.new_page(width=612, height=792)
    y = 90
    for step in draw:
        y = step(page, y) + 14
    return document.tobytes()


def _lines(lines, pitch=14, font="tiro"):
    def step(page, y):
        for line in lines:
            page.insert_text((72, y), line, fontname=font, fontsize=11)
            y += pitch
        return y
    return step


def _reading(content: bytes) -> str:
    from app.docstudio.ocr import _plain

    blocks = parser_for("application/pdf").parse(content, filename="x.pdf").blocks
    return " ".join(f"{b.number_label or ''} {_plain(b.text)}" for b in blocks)


def _converted(content: bytes) -> tuple[list[str], str]:
    from app.docstudio.toword_build import convert

    word, _ = convert(content)
    return [p.text for p in Document(io.BytesIO(word)).paragraphs if p.text.strip()], _xml(word)


def test_a_typed_pdf_keeps_its_fonts_under_the_names_word_knows():
    """A PDF names fonts as PostScript does — "Times-Roman", "ArialMT" — which
    Word does not recognise, and draws in its default font instead."""
    _, xml = _converted(_pdf(_lines(BODY[:2]), _lines(BODY[2:], font="hebo")))

    fonts = set(re.findall(r'w:ascii="([^"]+)"', xml))
    assert fonts == {"Times", "Helvetica"} and "<w:b/>" in xml


def test_a_number_and_its_heading_written_apart_stay_apart():
    """The space between them is drawn as a span of its own, and dropped the
    licence came out with "1.DEFINITIONS" 41 times."""
    def heading(page, y):
        page.insert_text((72, y), "1.", fontname="helv", fontsize=11)
        page.insert_text((84, y), " ", fontname="cour", fontsize=11)
        page.insert_text((96, y), "DEFINITIONS", fontname="hebo", fontsize=11)
        return y + 14

    paragraphs, _ = _converted(_pdf(heading, _lines(BODY)))

    assert paragraphs[0].split() == ["1.", "DEFINITIONS"]


def test_a_sentence_split_line_by_line_is_one_paragraph_again():
    """A PDF prints lines, not paragraphs. Kept a paragraph each, the text
    would not re-wrap when a word in it is changed — the point of Word."""
    indemnity = [
        "The Supplier shall indemnify the Customer against every loss, claim and",
        "expense arising from any breach of this Agreement by the Supplier or",
        "by any of its officers, employees, agents or subcontractors whatsoever.",
    ]

    paragraphs, _ = _converted(_pdf(_lines(BODY), _lines(indemnity, pitch=17), _lines(BODY)))

    assert " ".join(indemnity) in paragraphs


def test_two_columns_are_read_down_each_column_and_not_across():
    """Read across, a two-column page keeps every word and ruins every
    sentence: "the Licensee shall pay  Confidential Information means"."""
    import fitz

    from app.docstudio.toword_build import convert

    left = ["Clause one, the first column, line one of it.", "Clause one, the first column, line two.",
            "Clause one, the first column, line three here."]
    right = ["Clause two, the second column, line one.", "Clause two, the second column, line two of it.",
             "Clause two, the second column, line three."]
    document = fitz.open()
    for _ in range(2):
        page = document.new_page(width=612, height=792)
        for i in range(10):
            page.insert_text((60, 90 + 16 * i), left[i % 3], fontname="tiro", fontsize=10)
            page.insert_text((330, 90 + 16 * i), right[i % 3], fontname="tiro", fontsize=10)

    made, how = convert(document.tobytes())
    said = " ".join(p.text for p in Document(io.BytesIO(made)).paragraphs)

    assert how["lines_whole"] == 1.0
    assert "the first column, line one of it. Clause one, the first column, line two." in said


def test_one_page_set_in_columns_is_read_down_them_like_the_rest():
    """Every word kept and every sentence ruined — "the Supplier shall indemnify
    the Customer  Confidential Information means" — is the failure a word count
    misses. A schedule set in two columns inside a contract that is otherwise
    one has to be found on its own page."""
    import fitz

    from app.docstudio.toword_build import convert

    document = fitz.open()
    for _ in range(5):
        page = document.new_page(width=612, height=792)
        for i, line in enumerate(BODY * 3):
            page.insert_text((72, 90 + 16 * i), line, fontname="tiro", fontsize=10)
    page = document.new_page(width=612, height=792)
    for i in range(20):
        page.insert_text((60, 90 + 20 * i), "The Supplier shall indemnify the Customer for it.",
                         fontname="tiro", fontsize=9)
        page.insert_text((320, 90 + 20 * i), "Confidential Information means all of the data.",
                         fontname="tiro", fontsize=9)

    made, how = convert(document.tobytes())
    said = " ".join(p.text for p in Document(io.BytesIO(made)).paragraphs)

    assert how["words_kept"] == 1.0
    # Down the left column, then down the right: each column's line follows its
    # own, not the line printed beside it.
    assert said.count("for it. The Supplier shall indemnify") >= 15
    assert said.count("of the data. Confidential Information means") >= 15


def test_the_running_header_and_footer_end_up_in_word_s_margins():
    """Left in the text they are read as clauses, repeat on every page and take
    room from each one; and the page number must follow Word's own count."""
    import fitz

    from app.docstudio.toword_build import convert

    document = fitz.open()
    for number in range(1, 4):
        page = document.new_page(width=612, height=792)
        page.insert_text((72, 40), "ACME LIMITED — CONFIDENTIAL", fontname="hebo", fontsize=9)
        for i, line in enumerate(BODY):
            page.insert_text((72, 120 + 14 * i), line, fontname="tiro", fontsize=11)
        page.insert_text((300, 760), str(number), fontname="tiro", fontsize=9)

    made, how = convert(document.tobytes())
    read = Document(io.BytesIO(made))
    body = " ".join(p.text for p in read.paragraphs)

    assert how["header"] and how["footer"]
    assert "CONFIDENTIAL" in read.sections[0].header.paragraphs[0].text
    assert "CONFIDENTIAL" not in body
    assert "PAGE" in _xml(made) or "PAGE" in read.sections[0].footer.paragraphs[0]._p.xml


def test_a_clause_that_runs_into_the_next_column_is_one_paragraph_again():
    """A page breaks a sentence where it runs out of room — at the foot of a
    column, at the foot of a page — and a Word document has no reason to. Left
    split, the clause is two clauses, a citation quoting it matches neither
    (which is how a liability clause came back "not verified"), and changing a
    word in it re-wraps only half. A clause set in capitals carries on in
    capitals, so the join cannot ask for a small letter."""
    from app.docstudio.toword import Doc, Para, Run, _rejoin

    def para(text: str) -> Para:
        return Para(runs=[Run(text=text, font="Times New Roman", size=10.0)])

    doc = Doc(body=[
        para("8.3 IN NO EVENT WILL WIPRO BE LIABLE FOR ANY LOSS OF PROFITS, COST OF COVER OR"),
        para("INDIRECT, SPECIAL OR CONSEQUENTIAL DAMAGES OF ANY KIND."),
        para("9. TERM. This Agreement begins on the Effective Date."),
        para("The Supplier shall provide the Services with reasonable skill and care,"),
        para("and shall pay each invoice within thirty days."),
    ])

    _rejoin(doc)

    assert [p.text[:24] for p in doc.body] == [
        "8.3 IN NO EVENT WILL WIP", "9. TERM. This Agreement ", "The Supplier shall provi"]
    assert "COST OF COVER OR INDIRECT, SPECIAL" in doc.body[0].text   # joined across the column
    assert doc.body[2].text.endswith("within thirty days.")           # joined across the page


def test_a_justified_line_does_not_come_out_with_double_spaces():
    """A PDF spreads a justified line by writing wide gaps, and its fragments
    already end with a space. A second space added for the gap leaves
    "AGGREGATE  LIABILITY": it reads wrong in Word, and a citation searching for
    the clause matches nothing — which is how a verified citation still failed
    to highlight."""
    from app.docstudio.toword import _rows

    def fragment(x0: float, x1: float, text: str) -> dict:
        return {"bbox": (x0, 100.0, x1, 110.0), "block": 1,
                "spans": [{"text": text, "font": "Times", "size": 10.0, "flags": 0, "color": 0,
                           "bbox": (x0, 100.0, x1, 110.0)}]}

    [row] = _rows([fragment(72, 140, "WIPRO'S "), fragment(150, 220, "MAXIMUM "),
                   fragment(230, 300, "AGGREGATE ")])

    said = "".join(span["text"] for span in row["spans"])
    assert "  " not in said
    assert said.strip() == "WIPRO'S MAXIMUM AGGREGATE"
