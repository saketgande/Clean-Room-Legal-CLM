"""New uploads are read with the Documents reader (app/documents/reader).

Guards the ways swapping the reader could quietly damage a contract: elements
whose offsets don't slice back out of the stored text (every citation then
points at the wrong words), a clause tree lost on the way into the database, a
Word tracked insertion dropped, and a file the reader can't handle producing an
empty contract instead of falling back to the old path (and its OCR).
"""

import io

from docx import Document
from lxml import etree

from app.contract_files.structure import read_with_documents_reader

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _nda() -> bytes:
    doc = Document()
    doc.add_paragraph("MUTUAL NON-DISCLOSURE AGREEMENT")
    doc.add_paragraph("1. CONFIDENTIALITY")
    doc.add_paragraph("1.1 The Recipient shall keep the Confidential Information secret and use it only for the Purpose.")
    doc.add_paragraph("1.2 The Recipient may disclose it to its Representatives who need to know it for the Purpose.")
    doc.add_paragraph("2. TERM")
    p = doc.add_paragraph("2.1 This Agreement lasts for ")
    ins = etree.SubElement(p._p, f"{{{W}}}ins", {f"{{{W}}}id": "1", f"{{{W}}}author": "Globex"})
    r = etree.SubElement(ins, f"{{{W}}}r")
    t = etree.SubElement(r, f"{{{W}}}t")
    t.text = "three (3) years"
    p.add_run(" from the Effective Date.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_offsets_slice_back_exactly_and_the_tree_is_kept():
    read = read_with_documents_reader(_nda(), mime_type=DOCX, filename="nda.docx")
    assert read is not None and read["method"].startswith("documents_reader:docx")
    text, els = read["text"], read["elements"]
    for e in els:
        assert text[e["char_start"]:e["char_end"]] == e["text"]
    by_label = {e["number_label"]: e for e in els if e["number_label"]}
    # 1.1 sits under 1: the parent is the clause numbered "1."
    parent = els[by_label["1.1"]["parent_seq"]]
    assert parent["number_label"].rstrip(".") == "1"
    assert by_label["1.1"]["level"] > parent["level"]
    assert by_label["1.1"]["element_type"] == "clause"


def test_a_tracked_insertion_is_part_of_the_clause():
    read = read_with_documents_reader(_nda(), mime_type=DOCX, filename="nda.docx")
    assert "lasts for three (3) years from the Effective Date" in read["text"]


def test_what_the_reader_cannot_read_falls_back_to_the_old_path():
    """None means "use today's extractor and its OCR decision", never "empty contract"."""
    assert read_with_documents_reader(b"not a word file", mime_type=DOCX, filename="x.docx") is None
    assert read_with_documents_reader(b"\x89PNG\r\n", mime_type="image/png", filename="scan.png") is None


def test_the_element_rows_fit_the_table():
    """Every key the reader emits is a column the persist step can store."""
    from app.contract_files.models import ContractDocumentElement

    columns = set(ContractDocumentElement.__table__.columns.keys())
    read = read_with_documents_reader(_nda(), mime_type=DOCX, filename="nda.docx")
    assert set(read["elements"][0]) - {"parent_seq"} <= columns


def test_stored_elements_link_each_clause_to_its_parent():
    """The tree survives into ContractDocumentElement: parent_id is a real id."""
    from types import SimpleNamespace

    from app.contract_files.service import _persist_document_elements

    class FakeDB:
        def __init__(self):
            self.added = []

        def query(self, *_):
            return SimpleNamespace(filter_by=lambda **_: SimpleNamespace(delete=lambda: 0))

        def add(self, row):
            self.added.append(row)

    read = read_with_documents_reader(_nda(), mime_type=DOCX, filename="nda.docx")
    snap = SimpleNamespace(id="s", org_id="o", contract_id="c", contract_version_id="v", text=read["text"],
                           structure_status="flat_only", element_count=0)
    db = FakeDB()
    _persist_document_elements(db, snap, elements=read["elements"])
    assert snap.structure_status == "structured" and snap.element_count == len(db.added)
    by_id = {r.id: r for r in db.added}
    child = next(r for r in db.added if r.number_label == "1.2")
    assert by_id[child.parent_id].number_label.rstrip(".") == "1"
    assert all("parent_seq" in e for e in read["elements"])  # the reader's own list isn't mutated
