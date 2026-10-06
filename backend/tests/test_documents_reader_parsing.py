"""The Documents reader (from docstudio Phase 1): a document goes in, structured clauses come out.

Every failure guarded here is silent in the system this replaces — the document
still ingests, the quality score still looks fine, and the damage only surfaces
when a lawyer cannot find a clause that is plainly in the contract.
"""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.documents.reader.parsing.base import UnsupportedFormat
from app.documents.reader.parsing.registry import DOCX_MIME, parser_for
from app.documents.reader.structure import build, verify_offsets

PDF_MIME = "application/pdf"


# --- fixtures ---------------------------------------------------------------


def _el(tag: str, **attrs):
    element = OxmlElement(tag)
    for name, value in attrs.items():
        element.set(qn(name), str(value))
    return element


def _numbered_docx() -> bytes:
    """A contract that numbers itself the way real client paper does: a
    multilevel definition in numbering.xml, referenced by the paragraphs."""
    document = Document()
    numbering = document.part.numbering_part.element
    abstract = _el("w:abstractNum", **{"w:abstractNumId": "900"})
    for ilvl, (fmt, text) in enumerate(
        [("decimal", "%1."), ("decimal", "%1.%2"), ("lowerLetter", "(%3)")]
    ):
        lvl = _el("w:lvl", **{"w:ilvl": ilvl})
        lvl.append(_el("w:start", **{"w:val": 1}))
        lvl.append(_el("w:numFmt", **{"w:val": fmt}))
        lvl.append(_el("w:lvlText", **{"w:val": text}))
        abstract.append(lvl)
    numbering.insert(0, abstract)
    num = _el("w:num", **{"w:numId": "90"})
    num.append(_el("w:abstractNumId", **{"w:val": "900"}))
    numbering.append(num)

    def numbered(text: str, ilvl: int):
        paragraph = document.add_paragraph(text)
        num_pr = OxmlElement("w:numPr")
        num_pr.append(_el("w:ilvl", **{"w:val": ilvl}))
        num_pr.append(_el("w:numId", **{"w:val": "90"}))
        paragraph._p.get_or_add_pPr().append(num_pr)

    numbered("DEFINITIONS", 0)
    numbered("Confidential Information means non-public information.", 1)
    numbered("Affiliate means a controlled entity.", 1)
    numbered("TERM", 0)
    numbered("This Agreement continues for three years.", 1)
    numbered("or until terminated for convenience.", 2)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_HELVETICA = {"r": 3.996, "e": 6.672, "q": 6.672, "u": 6.672, "i": 2.664, "s": 6.0,
              "a": 6.672, "w": 8.664, "t": 3.336, "h": 6.672, "l": 2.664, "y": 6.0,
              "o": 6.672, "v": 6.0, " ": 3.336, "c": 6.0, "n": 6.672, "d": 6.672}
_TC = 0.42  # character spacing — the operator naive readers overlook


def _fragmented_pdf(fragments: list[str]) -> bytes:
    """A PDF that draws a line as positioned fragments advancing by exactly the
    width already drawn, with character spacing set.

    This is the shape WordPerfect emits, and it is what makes position-guessing
    readers insert spaces inside words.
    """
    ops = ["BT", "/F1 12 Tf", f"{_TC} Tc", "-0.18 Tw", "72 720 Td"]
    for fragment in fragments:
        ops.append(f"({fragment})Tj")
        advance = sum(_HELVETICA.get(c, 6.672) + _TC for c in fragment)
        ops.append(f"{advance:.4f} 0.0000 TD")
    ops.append("ET")
    stream = "\n".join(ops).encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"),
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1, xref)
    return bytes(out)


def _parse(content: bytes, mime: str):
    return build(parser_for(mime).parse(content, filename="contract"))


# --- Word: the number is not in the text ------------------------------------


def test_word_auto_numbering_is_reconstructed():
    """Word stores "item of list N at level L" and counts "2.1" while drawing
    the page, so paragraph.text carries no number at all. Clause boundaries are
    found by that number, so without this an auto-numbered contract arrives as
    one undifferentiated block and no clause is separable."""
    structure = _parse(_numbered_docx(), DOCX_MIME)
    labels = [c.number_label for c in structure.clauses]

    assert labels[:5] == ["1.", "1.1", "1.2", "2.", "2.1"]
    assert structure.clauses[0].text.startswith("1. DEFINITIONS")


def test_a_deeper_level_restarts_under_each_parent():
    """Numbering is counted, not stored. Without resetting deeper levels when a
    parent increments, section 2's first sub-clause is labelled 2.3 — a
    plausible number for a clause that does not exist."""
    labels = [c.number_label for c in _parse(_numbered_docx(), DOCX_MIME).clauses]
    assert "2.1" in labels
    assert "2.3" not in labels


def test_numbering_defined_on_the_style_is_found():
    """Word's own `List Number` defines numbering in styles.xml rather than on
    the paragraph, and house styles inherit it through basedOn. Reading only the
    paragraph returns nothing for every such clause."""
    document = Document()
    document.add_paragraph("Governing law is New York.", style="List Number")
    document.add_paragraph("Notices go to Schedule A.", style="List Number")
    buffer = BytesIO()
    document.save(buffer)

    labels = [c.number_label for c in _parse(buffer.getvalue(), DOCX_MIME).clauses]
    assert labels == ["1.", "2."]


def test_a_table_is_not_dropped():
    """`Document.paragraphs` omits every table. In a contract that is the fee
    schedule, the SLA tiers and the liability cap — clauses that would then
    exist nowhere in the system while everything reported success."""
    document = Document()
    document.add_paragraph("1. FEES.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Final delivery"
    table.cell(1, 1).text = "USD 2,400,000"
    buffer = BytesIO()
    document.save(buffer)

    structure = _parse(buffer.getvalue(), DOCX_MIME)
    assert "USD 2,400,000" in structure.flat_text
    assert any(c.clause_type == "table" for c in structure.clauses)


# --- PDF: words must survive ------------------------------------------------


def test_a_word_drawn_as_fragments_is_not_split():
    """Readers that predict glyph positions from font metrics ignore the `Tc`
    character-spacing operator and insert spaces inside words — "requi res
    lawyers". Nothing downstream recovers a word split in three, and the damage
    is invisible to any quality score because fragments are printable letters."""
    structure = _parse(_fragmented_pdf(["requi", "res l", "aw", "y", "ers"]), PDF_MIME)

    assert "requires lawyers" in structure.flat_text
    assert "requi res" not in structure.flat_text


def test_pdf_clauses_carry_page_and_coordinates():
    """A citation can only be highlighted on the page it came from if the page
    position was kept at parse time. The parser returns these for free; not
    storing them is why the previous system could never do it."""
    structure = _parse(_fragmented_pdf(["Term and termination applies"]), PDF_MIME)

    assert structure.clauses, "no clauses parsed"
    for clause in structure.clauses:
        assert clause.page_number == 1
        assert set(clause.bbox) == {"x0", "y0", "x1", "y1"}
        assert clause.bbox["x1"] > clause.bbox["x0"]


def test_a_scanned_pdf_says_so_rather_than_looking_thin():
    """A scanned contract holds no text. Stored silently it is indistinguishable
    from a contract that says very little, and every downstream job reports
    success on an empty document.

    Asserted on `needs_ocr` rather than the warning's wording: that flag is what
    ingest actually branches on, so it is the promise that must not regress.
    """
    parsed = parser_for(PDF_MIME).parse(_fragmented_pdf([" "]), filename="c.pdf")

    assert parsed.needs_ocr is True
    assert any("needs OCR" in w for w in parsed.warnings)


@pytest.mark.parametrize(
    "label, expected",
    [("1.", 1), ("2.1", 2), ("3.4.5", 3), ("2.1.", 2), ("(a)", 3)],
)
def test_clause_depth_comes_from_the_number_shape(label, expected):
    """The trailing dot is punctuation, not a separator. Counting it made every
    top-level clause look like a sub-clause, which flattens the whole hierarchy
    by one and makes "2." a child of nothing."""
    from app.documents.reader.parsing.labels import level_for

    assert level_for(label) == expected


# --- structure invariants ---------------------------------------------------


def test_offsets_round_trip_for_every_clause():
    """`flat_text[char_start:char_end]` must be the clause's own text. When it
    is not, every citation into this version points at the wrong words — and
    nothing else in the system can detect that."""
    for content, mime in [(_numbered_docx(), DOCX_MIME),
                          (_fragmented_pdf(["requi", "res l", "aw"]), PDF_MIME)]:
        structure = _parse(content, mime)
        assert verify_offsets(structure) == []


def test_a_sub_clause_names_its_parent():
    """Clause 1.1 belongs under 1. Without the parent link the panel cannot show
    a clause in context, and "the indemnity in 4.2" loses the 4."""
    clauses = _parse(_numbered_docx(), DOCX_MIME).clauses
    by_label = {c.number_label: c for c in clauses}

    assert by_label["1.1"].parent_clause_id == by_label["1."].clause_id
    assert by_label["1."].parent_clause_id is None


def test_clause_ids_are_not_content_hashes():
    """A hash changes the moment a clause is reworded, which is exactly when its
    comments most need to follow it. Two clauses with identical text must still
    be separately addressable."""
    document = Document()
    document.add_paragraph("Reserved.")
    document.add_paragraph("Reserved.")
    buffer = BytesIO()
    document.save(buffer)

    clauses = _parse(buffer.getvalue(), DOCX_MIME).clauses
    assert len(clauses) == 2
    assert clauses[0].clause_id != clauses[1].clause_id


# --- the registry -----------------------------------------------------------


def test_an_unreadable_format_is_refused_with_a_reason():
    """.doc is accepted by the wider app and extracts to nothing, producing an
    empty contract with no error. Refusing loudly is the point: "we could not
    read this" must never be stored as "this contract says nothing"."""
    with pytest.raises(UnsupportedFormat) as excinfo:
        parser_for("application/msword")

    assert ".docx" in str(excinfo.value)


def test_an_image_is_refused_pointing_at_ocr():
    with pytest.raises(UnsupportedFormat) as excinfo:
        parser_for("image/png")
    assert "OCR" in str(excinfo.value)


# --- Section / Article numbering --------------------------------------------


@pytest.mark.parametrize(
    "text, label, level",
    [
        ("Section 4 Liability", "Section 4", 1),
        ("Section 4.2 Limitation of liability", "Section 4.2", 2),
        ("SECTION 12.3.1 Notices", "SECTION 12.3.1", 3),
        ("Article IV MISCELLANEOUS", "Article IV", 1),
    ],
)
def test_a_section_reference_is_read_as_a_clause_number(text, label, level):
    """The pattern allowed digits and roman numerals after "Section" but not
    dots, so "Section 4" matched and "Section 4.2" did not — a contract that
    numbers its sub-clauses that way lost every one of them. And reading the
    depth from the first character made them all top-level, so even the ones
    that matched arrived flat."""
    from app.documents.reader.parsing.labels import level_for, split_label

    found, remainder = split_label(text)
    assert found == label
    assert level_for(found) == level
    assert not remainder.startswith(label)


def test_a_sentence_mentioning_sections_is_not_a_clause_number():
    """"Sections 4 and 5 apply" is prose about clauses, not a clause heading."""
    from app.documents.reader.parsing.labels import split_label

    assert split_label("Sections 4 and 5 apply to this Agreement.")[0] is None


# --- labels that looked like clause numbers and were not, and vice versa -----


@pytest.mark.parametrize("text, label", [
    ("(z) the final lettered item", "(z)"),
    ("(aa) the item after (z)", "(aa)"),
    ("(ab) and the one after that", "(ab)"),
])
def test_a_lettered_list_continues_past_z(text, label):
    """A list of exclusions runs a..z and then continues (aa), (ab). Matching a
    single letter dropped the number from every item past the twenty-sixth, so
    the clauses existed but none of them could be cited."""
    from app.documents.reader.parsing.labels import split_label

    assert split_label(text)[0] == label


@pytest.mark.parametrize("text, expected", [
    ("2026. The year was difficult for the industry.", None),
    ("1999. Revenue fell sharply that year.", None),
    ("999. A genuinely deep clause number", "999."),
    ("1002.1 A deep sub-clause", "1002.1"),
    ("12. FEES", "12."),
])
def test_a_year_opening_a_paragraph_is_not_a_clause_number(text, expected):
    """"2026. The year was difficult" was read as clause 2026. Contracts do not
    number top-level clauses in the thousands, and anything genuinely that deep
    carries a dot, which is still allowed through."""
    from app.documents.reader.parsing.labels import split_label

    assert split_label(text)[0] == expected


def test_a_clause_number_split_from_its_clause_is_reunited():
    """A layout can put the number in its own block. Left alone it becomes a
    clause whose whole content is "3." — nothing to read, nothing to cite —
    while the real clause beside it has no number at all."""
    from app.documents.reader.parsing.text import blocks_from_text

    blocks = blocks_from_text("3.\n\nThe Company shall pay all fees within thirty days.")

    assert len(blocks) == 1
    assert blocks[0].number_label == "3."
    assert blocks[0].text.startswith("The Company shall pay")


def test_a_trailing_orphan_number_is_dropped():
    """A bare number at the end of the text is a page number. Storing it makes
    a clause with no content that a reviewer has to read and dismiss."""
    from app.documents.reader.parsing.text import blocks_from_text

    assert [b.text for b in blocks_from_text("Fees are payable.\n\n12")] == ["Fees are payable."]


def test_an_orphan_number_does_not_steal_the_next_clauses_number():
    """"3." followed by "4. FEES" is a stray marker, not clause 3. Overwriting
    the real number would misfile the clause under one that does not exist."""
    from app.documents.reader.parsing.text import blocks_from_text

    blocks = blocks_from_text("3.\n\n4. FEES are payable within thirty days.")

    assert [b.number_label for b in blocks] == ["4."]


# --- page numbers OCR ran into the text -------------------------------------


def test_a_page_number_run_into_the_next_page_is_not_a_clause_number():
    """OCR glues a page footer to the following page's first words:

        "...income or payroll taxes, and health,"
        "2 accident and workers' compensation benefits..."

    Read as clause 2, it mislabels the clause and sits between the two halves
    of a sentence so reflow cannot rejoin them. On one 11-page scan that was
    9 of 86 clauses. Furniture stripping cannot catch it — OCR never made the
    number a block of its own.
    """
    from app.documents.reader.parsing.text import blocks_from_text

    text = "\n\n".join(
        [f"{n} Body text of page {n} continues here." for n in range(1, 11)]
    )
    labels = [b.number_label for b in blocks_from_text(text, page_count=10)]

    assert labels == [None] * 10


def test_bare_clause_numbers_are_not_mistaken_for_page_numbers():
    """The guard that makes the above safe. Plenty of contracts number clauses
    "1 Definitions", "2 Term" with no dot. What separates a page-number run is
    that its highest value IS the page count; clause numbering does not line up
    with pagination."""
    from app.documents.reader.parsing.text import blocks_from_text

    text = "\n\n".join([f"{n} Clause heading number {n}" for n in range(1, 26)])
    labels = [b.number_label for b in blocks_from_text(text, page_count=36)]

    assert all(labels), "genuine clause numbers were stripped"


def test_a_short_document_is_left_alone():
    """Three chunks in a two-page document is not evidence of anything."""
    from app.documents.reader.parsing.text import blocks_from_text

    text = "1 First clause.\n\n2 Second clause."
    assert [b.number_label for b in blocks_from_text(text, page_count=2)] == ["1", "2"]


def test_without_a_page_count_nothing_is_stripped():
    """Plain text files have no pages, so the signal does not exist and the
    numbers must be left as they are."""
    from app.documents.reader.parsing.text import blocks_from_text

    text = "\n\n".join([f"{n} Body text {n}." for n in range(1, 11)])
    assert all(b.number_label for b in blocks_from_text(text))


@pytest.mark.parametrize("text, label", [
    ("Item 5 - Expiry Date", "Item 5"),
    ("ITEM 2 Name and Address of Licensor", "ITEM 2"),
    ("- 12.3.LICENSOR shall have the right to terminate", "12.3."),
    ("5.GENERAL PROVISIONS", "5."),
    ("U.S. Government rights are reserved.", None),
    ("A.The Company", None),
])
def test_schedule_items_and_numbers_ocr_glued_to_their_text_are_read(text, label):
    """A patent licence's schedule numbers its entries "Item 1" to "Item 9";
    unread, the AI grouped Items 2-4 under Item 1's text. And OCR dropped the
    space in "12.3.LICENSOR", so clause 12.3 lost its number — only a digit
    before the dot counts, or "U.S." would become clause "U."."""
    from app.documents.reader.parsing.labels import split_label

    assert split_label(text)[0] == label
