"""The temporary Phase 1 page: what it must refuse as much as what it does.

It takes uploads, spends real OCR and Claude credits, and displays contract
text nobody has vetted. So the tests that matter most are that it does not
exist in production and that a document's own markup never becomes markup on
the page — the second is what a hostile PDF would aim at.
"""

import re
import uuid
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import delete

import app.models
from app.core.database import SessionLocal
from app.docstudio.devui import _SCRIPT, _STYLE, _sha256, render
from app.docstudio.models import DsAnnotation, DsClause, DsDocument, DsEvent, DsVersion
from app.main import app

URL = "/api/v1/docstudio/dev"


def _docx() -> bytes:
    document = Document()
    # Unique per run, so it never deduplicates against a real contract.
    document.add_paragraph(f"1. TERM. Three years. Reference {uuid.uuid4()}.")
    # No number: whether it continues clause 1 is a question only the AI can
    # answer, so a run always has something to ask it.
    document.add_paragraph("Either party may renew it.")
    document.add_paragraph("2. FEES. Net thirty days.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def made():
    """Documents a test created, removed by id — never by scope, which is
    shared with the user's own runs."""
    ids: list[str] = []
    yield ids
    session = SessionLocal()
    try:
        for document_id in ids:
            version_ids = [
                v.id for v in session.query(DsVersion).filter(DsVersion.document_id == document_id)
            ]
            session.execute(delete(DsAnnotation).where(DsAnnotation.document_id == document_id))
            session.execute(delete(DsEvent).where(DsEvent.document_id == document_id))
            session.execute(delete(DsClause).where(DsClause.version_id.in_(version_ids)))
            session.execute(delete(DsVersion).where(DsVersion.document_id == document_id))
            session.execute(delete(DsDocument).where(DsDocument.id == document_id))
        session.commit()
    finally:
        session.close()


# --- a document's markup never becomes markup -------------------------------


def test_a_script_in_the_document_is_shown_as_text():
    """Clause text goes into the report verbatim, and the report is rendered
    into the page. A hostile PDF only needs one of those to execute."""
    html = render("1. TERM <script>alert(1)</script> and <img src=x onerror=alert(2)>")

    assert "<script" not in html
    assert "<img" not in html
    assert "&lt;script&gt;" in html


def test_ocr_markup_is_shown_as_the_tags_it_is():
    """Reducto's <b>, <signature> and <table> in clause text are extraction
    defects the report exists to surface. Rendering them as formatting would
    hide exactly what the page is for."""
    html = render("MINDTREE LIMITED By: <signature> <b>THE FORD FOUNDATION</b>")

    assert "&lt;signature&gt;" in html
    assert "&lt;b&gt;" in html


def test_a_javascript_link_in_the_document_is_not_a_link():
    assert "href" not in render("[click](javascript:alert(1))")


def test_the_reports_own_formatting_still_renders():
    """The safeguard must not cost the report its tables and headings."""
    html = render("## Findings\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n**CUT OFF**")

    assert "<h2>" in html and "<table>" in html and "<strong>" in html


# --- the page ---------------------------------------------------------------


def test_the_page_is_served_with_its_own_locked_down_policy(client):
    """The API's policy is default-src 'none', which blocks the page entirely;
    the page sends its own. Script may come only from this server, the page's
    one inline script by hash, and the document editor's loader: a page holding
    contract text must never run inline script, eval, or code from anywhere
    else — and the editor's origin may frame, but fetch and load nothing else."""
    from app.docstudio.devui import _EDITOR

    res = client.get(URL)
    policy = res.headers["content-security-policy"]
    directives = dict(d.split(" ", 1) for d in policy.split("; "))

    assert res.status_code == 200
    assert directives["script-src"] == f"'self' '{_sha256(_SCRIPT)}' {_EDITOR}"
    assert directives["frame-src"] == _EDITOR
    assert "unsafe-eval" not in policy
    # The editor is the only other origin, and only for its script and its frame.
    assert [d for d in policy.split("; ") if "http" in d] == [f"script-src 'self' '{_sha256(_SCRIPT)}' {_EDITOR}",
                                                               f"frame-src {_EDITOR}"]


def test_the_hashes_match_what_is_actually_served(client):
    """A hash of anything but the exact served text blocks the page silently —
    which is the failure a browser found and no test did."""
    page = client.get(URL).text
    script = re.search(r"<script>(.*)</script>", page, re.DOTALL).group(1)
    style = re.search(r"<style>(.*)</style>", page, re.DOTALL).group(1)

    assert script == _SCRIPT
    assert style == _STYLE


def test_the_page_loads_nothing_from_elsewhere(client):
    page = client.get(URL).text

    assert "<script src" not in page and "<link" not in page
    assert ' style="' not in page  # an inline style attribute is blocked by the policy


# --- the run ----------------------------------------------------------------


def test_a_run_returns_the_report_and_its_rendered_form(client, made):
    res = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())})

    assert res.status_code == 200
    body = res.json()
    made.append(body["document_id"])
    assert body["clauses"] == 3
    assert "## Clauses" in body["report"]
    assert "<h2>Clauses</h2>" in body["report_html"]
    assert body["structure"].startswith("rules only")


def test_a_client_supplied_path_is_reduced_to_a_file_name(client, made):
    """The name is stored and displayed. A path in it is either a mistake or
    an attempt, and neither should reach the database."""
    res = client.post(f"{URL}/run", files={"file": ("../../etc/msa.docx", _docx())})

    body = res.json()
    made.append(body["document_id"])
    assert body["name"] == "msa.docx"


def test_an_unreadable_type_is_refused(client):
    res = client.post(f"{URL}/run", files={"file": ("scan.png", b"\x89PNG\r\n")})

    assert res.status_code == 422


def test_an_empty_file_is_refused(client):
    res = client.post(f"{URL}/run", files={"file": ("empty.pdf", b"")})

    assert res.status_code == 422


def test_a_run_still_succeeds_when_the_ai_cannot_arrange_it(client, made):
    """With no usable AI the page must still work, on the numbering's structure,
    and say that is what it is showing. The suite runs mocked."""
    res = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "true"})

    assert res.status_code == 200
    made.append(res.json()["document_id"])
    assert "mocked" in res.json()["structure"]


def test_the_page_does_not_exist_in_production(monkeypatch):
    """Mounted inside `main.py`'s non-prod block. If it ever moves out of it,
    an unauthenticated page that uploads files and spends API credit ships to
    production."""
    import app.main as main_module
    from app.core.config import settings

    monkeypatch.setattr(settings, "environment", "production")
    # Production's boot checks would reject a dev configuration long before
    # routing; routing is the only thing under test here.
    monkeypatch.setattr(main_module, "validate_runtime_settings", lambda _settings: None)

    paths = {route.path for route in main_module.create_app().routes}

    assert not any("docstudio/dev" in path for path in paths)


# --- Phase 2: notes that follow a new version --------------------------------


def _docx_of(*paragraphs: str) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_notes_follow_a_new_version_and_a_lost_one_can_be_relinked(client, made):
    """The page's whole Phase 2 loop: run a file, add notes, run an edited file
    as a new version of it, and re-link the note whose words are gone."""
    reference = f"Reference {uuid.uuid4()}."
    first = client.post(
        f"{URL}/run",
        files={"file": ("msa.docx", _docx_of(
            f"1. TERM. Three years. {reference}",
            "2. FEES. Payment is due within thirty days.",
            "3. NOTICES. Notices must be sent by registered post to the office.",
        ))},
        data={"ai": "false"},
    ).json()
    made.append(first["document_id"])
    for quote, body in (
        ("Payment is due within thirty days", "Push for 45."),
        ("Notices must be sent by registered post to the office", "Allow email."),
    ):
        added = client.post(
            f"{URL}/notes", data={"document_id": first["document_id"], "quote": quote, "body": body}
        )
        assert added.status_code == 200, added.text

    second = client.post(
        f"{URL}/run",
        files={"file": ("msa-v2.docx", _docx_of(
            f"1. TERM. Three years. {reference}", "2. FEES. Payment is due within thirty days."
        ))},
        data={"ai": "false", "version_of": first["document_id"]},
    ).json()

    assert second["document_id"] == first["document_id"]
    assert [note["body"] for note in second["lost"]] == ["Allow email."]
    relinked = client.post(
        f"{URL}/notes/{second['lost'][0]['id']}/relink", data={"quote": "Three years"}
    ).json()
    assert relinked["lost"] == []
    assert "## Annotations" in relinked["report"]


def test_words_that_are_missing_or_repeated_are_refused_with_a_reason(client, made):
    first = client.post(
        f"{URL}/run",
        files={"file": ("msa.docx", _docx_of(f"1. FEES. Net thirty days. Net thirty days. {uuid.uuid4()}"))},
        data={"ai": "false"},
    ).json()
    made.append(first["document_id"])

    missing = client.post(f"{URL}/notes", data={"document_id": first["document_id"], "quote": "sixty", "body": "x"})
    twice = client.post(
        f"{URL}/notes", data={"document_id": first["document_id"], "quote": "Net thirty days", "body": "x"}
    )

    assert missing.status_code == 422 and "not in this version" in missing.json()["detail"]
    assert twice.status_code == 422 and "appear 2 times" in twice.json()["detail"]


def test_a_version_of_a_document_that_does_not_exist_is_refused(client):
    res = client.post(
        f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false", "version_of": "nope"}
    )

    assert res.status_code == 422


def test_a_note_is_not_added_to_a_version_other_than_the_one_shown(client, made):
    """Re-uploading an older file shows that older version. A note typed
    against it must not land silently on the newer text."""
    reference = f"Reference {uuid.uuid4()}."
    older = _docx_of(f"1. TERM. Three years. {reference}", "2. FEES. Payment is due within thirty days.")
    first = client.post(f"{URL}/run", files={"file": ("msa.docx", older)}, data={"ai": "false"}).json()
    made.append(first["document_id"])
    client.post(
        f"{URL}/run",
        files={"file": ("msa-v2.docx", _docx_of(f"1. TERM. Four years. {reference}"))},
        data={"ai": "false", "version_of": first["document_id"]},
    )

    shown = client.post(f"{URL}/run", files={"file": ("msa.docx", older)}, data={"ai": "false"}).json()
    refused = client.post(
        f"{URL}/notes",
        data={"document_id": shown["document_id"], "version_id": shown["version_id"],
              "quote": "Payment is due", "body": "x"},
    )

    assert shown["deduplicated"] is True
    assert refused.status_code == 422
    assert "notes go on the latest version (2)" in refused.json()["detail"]


# --- the original file -------------------------------------------------------


def test_the_original_file_is_served_for_the_viewer_never_rendered_by_the_browser(client, made):
    """The Original tab fetches these bytes and draws them with pdf.js or
    docx-preview. As an attachment the browser itself never renders an
    uploaded file in this origin — the product route's rule, kept here."""
    content = f"1. TERM. Three years. Reference {uuid.uuid4()}.\n".encode()
    body = client.post(f"{URL}/run", files={"file": ("msa.txt", content)}).json()
    made.append(body["document_id"])

    res = client.get(f"{URL}/versions/{body['version_id']}/file")

    assert body["original"] is True and body["mime"] == "text/plain"
    assert res.status_code == 200
    assert res.content == content
    assert res.headers["content-disposition"].startswith("attachment")
    assert res.headers["x-content-type-options"] == "nosniff"


def test_a_file_outside_this_tools_scope_is_not_served(client, made):
    """This page has no sign-in, so the only thing keeping it off the rest of
    the database is that the version belongs to the tool's own scope."""
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}).json()
    made.append(body["document_id"])
    session = SessionLocal()
    try:
        session.get(DsVersion, body["version_id"]).org_id = "a-real-organisation"
        session.commit()
    finally:
        session.close()

    assert client.get(f"{URL}/versions/{body['version_id']}/file").status_code == 404
    assert client.get(f"{URL}/versions/{uuid.uuid4()}/file").status_code == 404


def test_a_document_already_read_can_be_opened_without_uploading_it_again(client, made):
    """The viewer existed only in the seconds after a run, so looking at a
    document read yesterday meant re-uploading it — and the page looked empty
    to anyone who just loaded it."""
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}).json()
    made.append(body["document_id"])

    res = client.get(f"{URL}/documents/{body['document_id']}")

    opened = res.json()
    assert res.status_code == 200
    assert opened["version_id"] == body["version_id"]
    assert opened["clauses"] == body["clauses"]
    assert opened["original"] is True
    assert "<h2>Clauses</h2>" in opened["report_html"]
    assert client.get(f"{URL}/documents/{uuid.uuid4()}").status_code == 404


def test_the_page_is_never_served_from_the_cache(client):
    """It changes under the developer reading it; a cached copy is a script
    silently missing whatever was just added."""
    assert client.get(URL).headers["cache-control"] == "no-store"


def test_the_viewer_libraries_are_served_from_this_server_and_nothing_else_is(client):
    """The vendor route takes a name from the request; it must only ever
    match the four files it exists for, never walk to another file."""
    res = client.get(f"{URL}/vendor/pdf.min.mjs")

    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/javascript")
    assert b"pdfjs" in res.content[:4000] or b"PDFJS" in res.content[:4000]
    for name in ("README.md", "LICENSE.pdfjs", "..%2Fdevui.py", "..%2F..%2Fmain.py"):
        assert client.get(f"{URL}/vendor/{name}").status_code == 404


# --- asking --------------------------------------------------------------------


def test_asking_without_an_ai_says_why(client, made):
    """Tests run with the Claude client mocked — the same state as a deployment
    with no key. The page has to say so, not fail somewhere obscure."""
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false"}).json()
    made.append(body["document_id"])

    res = client.post(f"{URL}/ask", data={"version_id": body["version_id"], "question": "What are the fees?"})

    assert res.status_code == 422
    assert "Asking needs the AI" in res.json()["detail"]


def test_an_answer_comes_back_with_its_citation_checked_against_the_clause(client, made, monkeypatch):
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false"}).json()
    made.append(body["document_id"])

    class Fake:
        name = "fake"

        def answer(self, document, question):
            # Cites whichever line the document numbered for the fees clause.
            fees = next(line for line in document.splitlines() if "FEES" in line)
            return {
                "answer": "Payment is due net thirty days [1].",
                "citations": [{"clause": int(fees.split()[0].lstrip("#")), "quote": "Net thirty days"}],
            }

    monkeypatch.setattr("app.docstudio.ask.default_answerer", lambda org_id: Fake())

    reply = client.post(f"{URL}/ask", data={"version_id": body["version_id"], "question": "Fees?"}).json()

    [cite] = reply["citations"]
    assert reply["answer"] == "Payment is due net thirty days [1]."
    assert cite["verified"] is True
    assert cite["label"].startswith("2")


def test_a_question_that_is_empty_or_about_another_scope_is_refused(client):
    assert client.post(f"{URL}/ask", data={"version_id": str(uuid.uuid4()), "question": "  "}).status_code == 422
    assert client.post(f"{URL}/ask", data={"version_id": str(uuid.uuid4()), "question": "Fees?"}).status_code == 404


# --- comments, as Word takes them -------------------------------------------------


def test_a_comment_on_selected_words_lands_on_them_and_comes_back_for_the_margin(client, made):
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false"}).json()
    made.append(body["document_id"])

    res = client.post(f"{URL}/comments", data={
        "version_id": body["version_id"], "quote": "Net  thirty\ndays", "prefix": "2. FEES. ", "body": "Too long?",
    })

    [comment] = res.json()["comments"]
    assert res.status_code == 200
    assert (comment["body"], comment["quote"], comment["kind"]) == ("Too long?", "Net thirty days", "comment")
    assert comment["label"].startswith("2")


def test_a_deleted_comment_leaves_a_record(client, made):
    """Deleting is a person's decision, and is kept as one: an annotation that
    simply vanished is what docstudio exists to prevent."""
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false"}).json()
    made.append(body["document_id"])
    added = client.post(f"{URL}/comments", data={
        "version_id": body["version_id"], "quote": "Net thirty days", "body": "x",
    }).json()

    res = client.delete(f"{URL}/comments/{added['comments'][0]['id']}")

    assert res.status_code == 200 and res.json()["comments"] == []
    session = SessionLocal()
    try:
        events = session.query(DsEvent).filter(DsEvent.document_id == body["document_id"]).all()
        assert "annotation.deleted" in {event.event_type for event in events}
    finally:
        session.close()


def test_a_comment_needs_words_that_are_there_and_something_to_say(client, made):
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false"}).json()
    made.append(body["document_id"])
    ask = lambda **fields: client.post(f"{URL}/comments", data={"version_id": body["version_id"], **fields})

    assert ask(quote="Net thirty days", body="  ").status_code == 422
    assert ask(quote="Net ninety days", body="x").status_code == 422
    assert client.post(f"{URL}/comments", data={
        "version_id": str(uuid.uuid4()), "quote": "Net thirty days", "body": "x",
    }).status_code == 404


def test_a_comment_takes_replies_and_can_be_resolved_and_reopened(client, made):
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false"}).json()
    made.append(body["document_id"])
    added = client.post(f"{URL}/comments", data={
        "version_id": body["version_id"], "quote": "Net thirty days", "body": "Too long?",
    }).json()
    thread = added["comments"][0]["id"]

    replied = client.post(f"{URL}/comments/{thread}/replies", data={"body": "Agreed."}).json()
    resolved = client.post(f"{URL}/comments/{thread}/resolve", data={"resolved": "true"}).json()
    reopened = client.post(f"{URL}/comments/{thread}/resolve", data={"resolved": "false"}).json()

    # A reply lives in its thread, never as a comment of its own in the margin.
    assert [len(replied["comments"]), replied["comments"][0]["replies"][0]["body"]] == [1, "Agreed."]
    assert (resolved["comments"][0]["resolved"], reopened["comments"][0]["resolved"]) == (True, False)
    assert client.post(f"{URL}/comments/{thread}/replies", data={"body": " "}).status_code == 422


# --- redlining ---------------------------------------------------------------------


def test_a_suggested_edit_is_the_next_version_with_a_tracked_change_and_comments_carried(client, made):
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false"}).json()
    made.append(body["document_id"])
    client.post(f"{URL}/comments", data={
        "version_id": body["version_id"], "quote": "Either party may renew it", "body": "Why renew?",
    })

    res = client.post(f"{URL}/redline/suggest", data={
        "version_id": body["version_id"], "quote": "Net thirty days", "replacement": "Net forty-five days",
    })
    after = res.json()
    done = client.post(f"{URL}/redline/resolve", data={
        "version_id": after["version_id"], "everything": "true", "accept": "true",
    }).json()

    [change] = after["changes"]
    assert res.status_code == 200 and after["version_id"] != body["version_id"]
    assert (change["kind"], change["deleted"], change["inserted"], change["author"]) == (
        "replace", "thirty", "forty-five", "You")
    assert [c["body"] for c in after["comments"]] == ["Why renew?"]  # followed its words
    assert done["changes"] == [] and "Net forty-five days" in done["report"]


def test_a_file_that_is_not_word_gets_an_editable_word_version_before_a_redline(client, made):
    """A PDF or text file cannot carry tracked changes: the page says so, and
    offers the Word version rather than failing somewhere in the file."""
    content = f"1. TERM. Three years. Reference {uuid.uuid4()}.\n2. FEES. Net thirty days.\n".encode()
    body = client.post(f"{URL}/run", files={"file": ("msa.txt", content)}, data={"ai": "false"}).json()
    made.append(body["document_id"])

    refused = client.post(f"{URL}/redline/suggest", data={
        "version_id": body["version_id"], "quote": "Net thirty days", "replacement": "x",
    })
    editable = client.post(f"{URL}/redline/editable", data={"version_id": body["version_id"]}).json()
    edited = client.post(f"{URL}/redline/suggest", data={
        "version_id": editable["version_id"], "quote": "Net thirty days", "replacement": "Net sixty days",
    }).json()

    assert refused.status_code == 422 and "editable Word version" in refused.json()["detail"]
    assert editable["editable"] is True and editable["name"].endswith(".docx")
    assert editable["converted"] is None  # plain text has no look to set beside it
    assert [c["inserted"] for c in edited["changes"]] == ["sixty"]


def _typed_pdf() -> bytes:
    import fitz

    document = fitz.open()
    page = document.new_page(width=612, height=792)
    for i, line in enumerate([
        f"1. TERM. This Agreement lasts three years. Reference {uuid.uuid4()}.",
        "2. FEES. Payment is due within thirty days of the date of a valid invoice.",
        "3. LAW. This Agreement is governed by the laws of England and Wales.",
    ]):
        page.insert_text((72, 90 + 18 * i), line, fontname="tiro", fontsize=11)
    return document.tobytes()


def test_a_typed_pdf_becomes_a_word_version_with_its_look_set_beside_it(client, made):
    """The Word version is what gets redlined and sent, so it says it was
    converted, and the page can put the PDF it came from next to it."""
    body = client.post(f"{URL}/run", files={"file": ("msa.pdf", _typed_pdf())}, data={"ai": "false"}).json()
    made.append(body["document_id"])

    word = client.post(f"{URL}/redline/editable", data={"version_id": body["version_id"]}).json()
    pdf = client.get(f"{URL}/versions/{word['converted']['version_id']}/file")

    assert word["name"] == "msa (converted from PDF).docx" and word["editable"] is True
    assert word["converted"] == {"version_id": body["version_id"], "version_number": 1, "look": "kept", "note": ""}
    assert pdf.content.startswith(b"%PDF")


def test_a_scan_is_rebuilt_as_word_from_what_was_read_on_it(client, made):
    """A scan holds no text at all, so there is nothing to convert: the Word
    version is rebuilt from OCR's reading of the picture, and must say so —
    presented as the contract's own words it would put a machine's misreadings
    into a signed agreement. With nothing read yet, there is nothing to build."""
    body = client.post(f"{URL}/run", files={"file": ("scan.pdf", _typed_pdf())}, data={"ai": "false"}).json()
    made.append(body["document_id"])
    session = SessionLocal()
    try:  # read as a scan would have been, but nothing stored from the reading
        session.get(DsVersion, body["version_id"]).parser_name = "pdf+ocr:test"
        session.commit()
    finally:
        session.close()

    opened = client.get(f"{URL}/documents/{body['document_id']}").json()
    unread = client.post(f"{URL}/redline/editable", data={"version_id": body["version_id"]})

    assert body["scanned"] is False and opened["scanned"] is True
    assert unread.status_code == 422 and "not been read yet" in unread.json()["detail"]


def test_the_ais_edits_are_checked_and_written_as_its_own_tracked_changes(client, made, monkeypatch):
    """An edit on words the contract does not contain would be a change nobody
    asked for, placed wherever the words happened to match — so it is dropped,
    and said so."""
    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx())}, data={"ai": "false"}).json()
    made.append(body["document_id"])

    class Fake:
        name = "fake"

        def draft(self, document, instruction):
            fees = next(line for line in document.splitlines() if "FEES" in line)
            seq = int(fees.split()[0].lstrip("#"))
            return {"summary": "Longer payment terms.", "clauses": [seq, seq],
                    "old": ["Net thirty days", "Net ninety days"], "new": ["Net forty-five days", "x"],
                    "why": ["Eases the customer's cash flow.", "Invented."]}

    monkeypatch.setattr("app.docstudio.drafting.default_drafter", lambda org_id: Fake())
    reply = client.post(f"{URL}/redline/draft", data={
        "version_id": body["version_id"], "instruction": "Give the customer 45 days to pay.",
    }).json()

    [change] = reply["changes"]
    assert (change["author"], change["inserted"], change["why"]) == (
        "AEGIS AI", "forty-five", "Eases the customer's cash flow.")
    assert len(reply["drafted"]["placed"]) == 1
    assert "not in the contract" in reply["drafted"]["skipped"][0]


# --- the document editor ---------------------------------------------------------------


def test_the_editor_is_handed_a_signed_config_and_never_a_scan(client, made):
    """An unsigned config would let anyone point the editor at any file; a scan
    is the signed contract, changed by amendment, never edited."""
    import jwt

    from app.docstudio.devui import _EDITOR_SECRET, _HERE_INSIDE

    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx_of(f"Payment in thirty days. {uuid.uuid4()}"))},
                       data={"ai": "false"}).json()
    made.append(body["document_id"])
    scan = client.post(f"{URL}/run", files={"file": ("scan.pdf", _typed_pdf())}, data={"ai": "false"}).json()
    made.append(scan["document_id"])
    session = SessionLocal()
    try:
        session.get(DsVersion, scan["version_id"]).parser_name = "pdf+ocr:test"
        session.commit()
    finally:
        session.close()

    opened = client.get(f"{URL}/editor/config", params={"version_id": body["version_id"]}).json()
    refused = client.get(f"{URL}/editor/config", params={"version_id": scan["version_id"]})
    config = opened["config"]
    signed = jwt.decode(config.pop("token"), _EDITOR_SECRET, algorithms=["HS256"])

    assert signed == config and config["document"]["key"] == body["version_id"]
    assert config["document"]["url"].startswith(_HERE_INSIDE)
    assert config["editorConfig"]["customization"]["review"]["trackChanges"] is True
    assert refused.status_code == 422 and "amendment" in refused.json()["detail"]


def test_a_save_is_believed_only_signed_and_fetched_only_from_the_editor(client, made, monkeypatch):
    """A forged save would write any bytes as the next version; a save naming
    another host would make this server fetch from wherever it was told."""
    import httpx
    import jwt

    from app.docstudio import devui

    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx_of(f"Payment in thirty days. {uuid.uuid4()}"))},
                       data={"ai": "false"}).json()
    made.append(body["document_id"])
    edited = _docx_of(f"Payment in forty-five days. {uuid.uuid4()}")
    fetched = []
    monkeypatch.setattr(httpx, "get", lambda url, **_: fetched.append(url) or httpx.Response(
        200, content=edited, request=httpx.Request("GET", url)))
    where = {"document_id": body["document_id"], "version_id": body["version_id"]}
    save = {"status": 6, "key": body["version_id"], "url": "http://attacker.example/cache/files/abc/output.docx?md5=1"}

    forged = client.post(f"{URL}/editor/saved", params=where, json={**save, "token": "not-a-token"})
    ok = client.post(f"{URL}/editor/saved", params=where,
                     json={**save, "token": jwt.encode(save, devui._EDITOR_SECRET, algorithm="HS256")})
    latest = client.get(f"{URL}/documents/{body['document_id']}").json()

    assert forged.status_code == 403
    assert ok.json() == {"error": 0}
    assert fetched == [f"{devui._EDITOR_INSIDE}/cache/files/abc/output.docx?md5=1"]
    assert latest["version_id"] != body["version_id"] and "forty-five" in latest["report"]


def test_the_editor_closing_after_a_save_makes_no_second_version(client, made, monkeypatch):
    """The editor sends the file again when it closes, repackaged but unchanged;
    stored, History would fill with versions that differ in nothing."""
    import io
    import zipfile

    import httpx
    import jwt

    from app.docstudio import devui

    body = client.post(f"{URL}/run", files={"file": ("msa.docx", _docx_of(f"Payment in thirty days. {uuid.uuid4()}"))},
                       data={"ai": "false"}).json()
    made.append(body["document_id"])
    edited = _docx_of(f"Payment in forty-five days. {uuid.uuid4()}")
    repacked = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(edited)) as old, zipfile.ZipFile(repacked, "w") as new:
        for item in old.infolist():  # the same parts, new zip timestamps
            new.writestr(zipfile.ZipInfo(item.filename, date_time=(2030, 1, 1, 0, 0, 0)), old.read(item.filename))
    files = iter([edited, repacked.getvalue()])
    monkeypatch.setattr(httpx, "get", lambda url, **_: httpx.Response(200, content=next(files),
                                                                     request=httpx.Request("GET", url)))
    where = {"document_id": body["document_id"], "version_id": body["version_id"]}
    for status in (6, 2):
        save = {"status": status, "key": body["version_id"], "url": "http://onlyoffice/cache/files/x/output.docx"}
        client.post(f"{URL}/editor/saved", params=where,
                    json={**save, "token": jwt.encode(save, devui._EDITOR_SECRET, algorithm="HS256")})

    versions = client.get(f"{URL}/documents/{body['document_id']}/history").json()["versions"]
    assert [v["number"] for v in versions] == [2, 1]


def test_the_editor_is_asked_to_accept_where_the_person_can_see_it(client):
    """Accepting on the server means closing the editor, changing the file and
    opening it again — which looks like the document reloading for no reason,
    and leaves nobody sure whether it worked. The editor is asked instead, and
    it does them at once."""
    client.post(f"{URL}/editor/find", data={"version_id": "v-acc", "action": "accept", "quote": ""})

    waiting = client.get(f"{URL}/editor/plugin/command", params={"key": "v-acc", "since": 0}).json()
    refused = client.post(f"{URL}/editor/find", data={"version_id": "v-acc", "action": "shred", "quote": ""})

    assert waiting["action"] == "accept"
    assert refused.status_code == 422


def test_words_to_find_are_handed_to_the_editor_holding_that_version_once(client):
    """Two editors can be open at once — two tabs, a document and the one before
    it — and one queue for both sent a citation from one contract to whichever
    editor asked first, which then found nothing. And the plugin asks again
    every second: an answer that repeats the last words drags the cursor back
    for ever."""
    told = client.post(f"{URL}/editor/find",
                       data={"version_id": "v-one", "quote": "  within thirty days  "}).json()
    client.post(f"{URL}/editor/find", data={"version_id": "v-two", "quote": "governed by the laws"})

    waiting = client.get(f"{URL}/editor/plugin/command", params={"key": "v-one", "since": 0}).json()
    other = client.get(f"{URL}/editor/plugin/command", params={"key": "v-two", "since": 0}).json()
    again = client.get(f"{URL}/editor/plugin/command", params={"key": "v-one", "since": told["id"]}).json()
    nobody = client.get(f"{URL}/editor/plugin/command", params={"key": "v-none", "since": 0}).json()

    assert waiting == {"id": told["id"], "text": "within thirty days", "action": "find"}
    assert other["text"] == "governed by the laws"
    assert again == {"id": told["id"], "text": "", "action": "find"}
    assert nobody == {"id": 0, "text": "", "action": "find"}


def test_the_look_alike_word_copy_is_a_download_and_never_made_of_a_scan(client, made, monkeypatch):
    """It lays each line into a box of its own, so kept as a version it would
    cost the document its clause structure; and OCR text in boxes would be a
    scan's words presented as the contract's."""
    from app.docstudio import devui

    body = client.post(f"{URL}/run", files={"file": ("msa.pdf", _typed_pdf())}, data={"ai": "false"}).json()
    made.append(body["document_id"])
    scan = client.post(f"{URL}/run", files={"file": ("scan.pdf", _typed_pdf())}, data={"ai": "false"}).json()
    made.append(scan["document_id"])
    session = SessionLocal()
    try:
        session.get(DsVersion, scan["version_id"]).parser_name = "pdf+ocr:test"
        session.commit()
    finally:
        session.close()
    monkeypatch.setattr(devui, "lookalike_copy", lambda *a, **k: _docx_of("a page in boxes"))

    made_copy = client.get(f"{URL}/versions/{body['version_id']}/word-lookalike")
    refused = client.get(f"{URL}/versions/{scan['version_id']}/word-lookalike")
    after = client.get(f"{URL}/documents/{body['document_id']}/history").json()["versions"]

    assert made_copy.status_code == 200 and "looks the same" in made_copy.headers["content-disposition"]
    assert refused.status_code == 422 and "scan" in refused.json()["detail"]
    assert [v["number"] for v in after] == [1]  # a download, not a version


def test_the_editor_is_told_to_save_before_the_page_changes_the_file(client, monkeypatch):
    """The editor holds the file while it is open. A change made underneath it
    — the AI's suggestions, accepting, rejecting — is overwritten at its next
    save, so the editor writes out what it has first and the change is made on
    that. "Nothing to save" is not a failure: the file on the server is already
    what the editor has."""
    import httpx
    import jwt

    from app.docstudio import devui

    asked = {}

    def answer(url, **kwargs):
        asked["url"] = url
        asked["sent"] = kwargs["json"]
        return httpx.Response(200, json={"error": asked.get("error", 0), "key": "v1"},
                              request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", answer)

    saved = client.post(f"{URL}/editor/save-now", data={"version_id": "v1"}).json()
    asked["error"] = 4
    nothing = client.post(f"{URL}/editor/save-now", data={"version_id": "v1"}).json()
    asked["error"] = 1
    refused = client.post(f"{URL}/editor/save-now", data={"version_id": "v1"})

    assert saved == {"saved": True, "nothing_to_save": False}
    assert nothing == {"saved": False, "nothing_to_save": True}
    assert refused.status_code == 422
    assert asked["url"].startswith(devui._EDITOR_INSIDE) and "CommandService" in asked["url"]
    # Signed, or anyone on the network could make the editor write out a file.
    assert jwt.decode(asked["sent"]["token"], devui._EDITOR_SECRET, algorithms=["HS256"]) == {
        "c": "forcesave", "key": "v1"}


def test_the_page_can_tell_whether_the_document_service_is_answering(client, monkeypatch):
    """When the editor loses its connection it goes on showing the document and
    quietly stops taking changes — typing does nothing and nothing says why.
    The page waits on this before opening the document again."""
    import httpx

    from app.docstudio import devui

    monkeypatch.setattr(httpx, "get", lambda url, **_: httpx.Response(
        200, text="true", request=httpx.Request("GET", url)))
    up = client.get(f"{URL}/editor/health").json()

    def refuse(url, **_):
        raise httpx.ConnectError("no route", request=httpx.Request("GET", url))

    monkeypatch.setattr(httpx, "get", refuse)
    down = client.get(f"{URL}/editor/health").json()

    assert up == {"up": True} and down == {"up": False}
    assert devui._EDITOR_INSIDE.startswith("http")
