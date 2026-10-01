"""A scan written down as markdown, and rebuilt from it as Word.

A scanned contract has no text to convert — every page is a photograph — so the
Word version is built from what OCR read, and the markdown is the record of
that reading: what was read, where it sat and how big it was. Each test names
the way that goes wrong.
"""

import io

from docx import Document

from app.docstudio.scanmd import document, markdown
from app.docstudio.toword import Paper
from app.docstudio.toword_build import build


def _read(page: int, kind: str, text: str, box: tuple[float, float, float, float]) -> dict:
    return {"type": kind, "page": page, "content": text, "confidence": "high",
            "bbox": {"x0": box[0], "y0": box[1], "x1": box[2], "y1": box[3]}}


BODY = ("The Supplier shall provide the Services with reasonable skill and care, and the Customer "
        "shall pay each invoice within thirty days of the date of that invoice.")
PAGES = [
    _read(1, "Header", "ACME LIMITED — CONFIDENTIAL", (0.1, 0.02, 0.5, 0.035)),
    _read(1, "Title", "MASTER SERVICES AGREEMENT", (0.3, 0.10, 0.7, 0.125)),
    _read(1, "Text", BODY, (0.1, 0.20, 0.9, 0.245)),
    _read(1, "Section Header", "1. DEFINITIONS", (0.1, 0.30, 0.35, 0.315)),
    _read(1, "Footer", "Page 1 of 2", (0.45, 0.95, 0.55, 0.965)),
    _read(2, "Header", "ACME LIMITED - CONFIDENTIAL", (0.1, 0.02, 0.5, 0.035)),
    _read(2, "Text", BODY, (0.1, 0.20, 0.9, 0.245)),
    _read(2, "Footer", "Page 2 of 2", (0.45, 0.95, 0.55, 0.965)),
]


def _said(content: bytes) -> Document:
    return Document(io.BytesIO(content))


def test_the_reading_records_where_each_block_sat_and_how_big_it_was():
    """Rebuilt from the words alone, every line of a scan comes out the same
    size in one column: the sizes and places are the only styling a photograph
    gives up, and the file has to carry them."""
    said = markdown(PAGES, title="msa.pdf", provider="reducto", paper=Paper())

    assert "read_by: reducto" in said and "assumed" in said.split("font:")[1]
    assert "<!-- page=1 size=" in said and "at=0.300,0.100" in said
    assert "# MASTER SERVICES AGREEMENT" in said       # a title is a heading, not a line of text
    assert "header" in said.split("ACME")[0].rsplit("<!--", 1)[-1]


def test_the_running_header_and_footer_are_kept_once_however_ocr_read_them():
    """OCR reads the same header differently on different pages — a dash for an
    em dash — and kept each way, Word's header stacks up a line per page."""
    doc = document(markdown(PAGES, title="msa.pdf", provider="reducto", paper=Paper()))

    assert len(doc.header) == 1 and len(doc.footer) == 1
    assert "CONFIDENTIAL" in doc.header[0].text
    assert not [p for p in doc.body if "CONFIDENTIAL" in p.text]


def test_the_word_version_is_built_from_the_file_so_a_correction_rebuilds_it():
    """The reading is a machine's, and wrong here and there. The markdown is
    the one place it can be corrected — if the Word version came from anywhere
    else, the correction would not reach it."""
    said = markdown(PAGES, title="msa.pdf", provider="reducto", paper=Paper())
    corrected = said.replace("MASTER SERVICES AGREEMENT", "MASTER SERVICES AGREEMENT (2019)")

    read = _said(build(document(corrected)))

    assert any("MASTER SERVICES AGREEMENT (2019)" in p.text for p in read.paragraphs)
    assert read.sections[0].header.paragraphs[0].text.strip().startswith("ACME")


def test_a_heading_keeps_its_weight_and_the_body_one_size():
    """A contract sets its body in one size; measured block by block from a
    photograph it wanders by a point or two, which a reader sees as a mess."""
    read = _said(build(document(markdown(PAGES, title="msa.pdf", provider="reducto", paper=Paper()))))

    sizes = {p.text: p.runs[0].font.size.pt for p in read.paragraphs if p.runs}
    body = [size for text, size in sizes.items() if text.startswith("The Supplier")]

    assert len(set(body)) == 1
    assert sizes["MASTER SERVICES AGREEMENT"] > body[0]
    assert next(p.runs[0].bold for p in read.paragraphs if p.text == "1. DEFINITIONS")
