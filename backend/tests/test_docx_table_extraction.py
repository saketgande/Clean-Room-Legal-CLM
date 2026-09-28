"""DOCX text extraction must include tables.

`Document.paragraphs` is top-level body paragraphs only and silently omits
every table. In a contract that is the payment schedule, the fee table, the
SLA tiers, the liability cap and usually the signature block — so the loss is
invisible to the quality score and lands as a clause that exists nowhere in
the system.
"""

from io import BytesIO

from docx import Document

from app.contract_files.text_extraction import extract_text

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx(build) -> bytes:
    document = Document()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _extract(content: bytes):
    return extract_text(content, mime_type=DOCX_MIME, filename="contract.docx")


def test_payment_table_reaches_the_snapshot():
    """Guards the defect this test exists for: a contract's money lives in a
    table, and dropping it meant the liability cap and payment milestones were
    absent from clause extraction, embeddings, obligations and every citation
    — while the quality score still read 0.89 and never triggered OCR."""

    def build(document):
        document.add_paragraph("1. TERM. This Agreement begins on the Effective Date.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Payment milestone"
        table.cell(0, 1).text = "Amount"
        table.cell(1, 0).text = "Final delivery"
        table.cell(1, 1).text = "USD 2,400,000"
        document.add_paragraph("2. TERMINATION. Either party may terminate on 90 days notice.")

    result = _extract(_docx(build))

    assert "USD 2,400,000" in result.text
    assert "Payment milestone" in result.text
    # The surrounding prose is still there, and still in order.
    assert result.text.index("1. TERM") < result.text.index("Payment milestone")
    assert result.text.index("Payment milestone") < result.text.index("2. TERMINATION")


def test_table_free_document_is_unchanged():
    """Guards the blast radius. Every stored offset and citation anchors to the
    snapshot text, so a document with no tables must extract byte-identically
    to the previous implementation — same paragraphs, same order, same "\\n"
    join, empty paragraphs included."""

    def build(document):
        document.add_paragraph("1. DEFINITIONS")
        document.add_paragraph("")
        document.add_paragraph("Confidential Information means any information...")

    content = _docx(build)
    document = Document(BytesIO(content))
    previous = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert _extract(content).text == previous


def test_merged_header_cell_is_not_repeated():
    """Guards duplicated text from merged cells: python-docx returns the same
    cell once per grid position it spans, so a header merged across three
    columns would otherwise be emitted three times and skew both the quality
    score and any embedding of that row."""

    def build(document):
        table = document.add_table(rows=2, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 2))
        merged.text = "SCHEDULE A - FEES"
        table.cell(1, 0).text = "Tier 1"
        table.cell(1, 1).text = "Tier 2"
        table.cell(1, 2).text = "Tier 3"

    text = _extract(_docx(build)).text

    assert text.count("SCHEDULE A - FEES") == 1
    assert "Tier 1\tTier 2\tTier 3" in text


def test_nested_table_inside_a_cell_is_read():
    """Guards a silent gap in the recursion: exhibits and fee schedules are
    routinely built as a table inside a table cell, and a single-level walk
    reads the outer table and drops the inner one."""

    def build(document):
        outer = document.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "Late fee"
        inner.cell(0, 1).text = "1.5% per month"

    text = _extract(_docx(build)).text

    assert "1.5% per month" in text


def test_a_document_that_is_only_a_table_is_no_longer_scored_as_empty():
    """Guards the worst shape: a schedule or rate card with no body prose used
    to extract to an empty string, score 0.0 and be handed to OCR — paying for
    OCR on a document whose text was sitting in the XML all along."""

    def build(document):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Governing law"
        table.cell(0, 1).text = "England and Wales"
        table.cell(1, 0).text = "Liability cap"
        table.cell(1, 1).text = "150% of fees paid in the preceding 12 months"

    result = _extract(_docx(build))

    assert "150% of fees paid" in result.text
    assert result.needs_ocr is False
