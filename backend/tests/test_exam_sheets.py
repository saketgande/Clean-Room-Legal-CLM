"""Exam spreadsheets: the facts sheet and the routing sheet.

specs/005-ai-eval-baseline T010 (FR-3, FR-17, FR-18, FR-19, FR-21). No database, no network,
nothing paid: sheets are built from invented rows in a temp folder standing in for
``eval/private/sheets``, then edited with openpyxl the way the owner edits them in Excel.
"""

import hashlib
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from eval.exam import sheets as s

BACKEND_DIR = Path(__file__).resolve().parent.parent
FILE_SHA = "a" * 64
TEXT_SHA = "b" * 64
# Invented text: the tests must never carry a real contract.
FROZEN_TEXT = (
    "This Agreement is made between Acme Ltd and Beta Pvt Ltd.\n"
    "It is governed by the laws of the State of\n   Delaware. Either side may leave on 30 days."
)


@pytest.fixture(autouse=True)
def sheets_dir(tmp_path, monkeypatch):
    """Point the module at a temp folder: the real private folder must never be touched."""
    folder = tmp_path / "sheets"
    texts = tmp_path / "texts"
    texts.mkdir()
    monkeypatch.setattr(s, "SHEETS_DIR", folder)
    monkeypatch.setattr(s, "TEXTS_DIR", texts)
    return folder


def _fact(doc_id, field, value, quote, **extra):
    row = {
        "doc_id": doc_id,
        "contract": f"{doc_id}.pdf",
        "field": field,
        "value": value,
        "quote": quote,
        "char_start": 10 if quote else None,
        "char_end": 10 + len(quote) if quote else None,
        "page": 2 if quote else None,
        "suggested_by": "claude-sonnet-5-5",
        "file_sha256": FILE_SHA,
        "frozen_text_sha256": TEXT_SHA,
    }
    return row | extra


def facts_rows():
    return [
        _fact("owner-01", "governing_law", "Delaware", "governed by the laws of Delaware"),
        _fact(
            "owner-01",
            "parties",
            ["Acme Ltd", "Beta Pvt Ltd"],
            "between Acme Ltd and Beta Pvt Ltd",
            page=None,
        ),
        _fact("owner-01", "effective_date", "2024-04-01", None),  # quote not found: unsupported
        _fact("owner-01", "notice_period", None, None),  # "not stated"
        _fact("owner-01", "renewal_term", "renews yearly", "renews for successive one year terms"),
    ]


def _key(sheet_id, doc_id, field):
    """The row_key formula from the plan, restated so a drifted formula fails here."""
    return hashlib.sha256((sheet_id + doc_id + field).encode()).hexdigest()[:16]


def _path(sheet_id):
    return s.SHEETS_DIR / f"facts-{sheet_id}.xlsx"


def make_facts(rows=None):
    sheet_id = s.write_facts_sheet(rows if rows is not None else facts_rows(), path=s.SHEETS_DIR)
    return sheet_id, _path(sheet_id)


def edit(path, row_key, *, mark=None, correction=None, set_cells=None):
    """Edit one row of a written sheet the way the owner would, and save it in place."""
    wb = load_workbook(path)
    ws = wb["Facts"] if "Facts" in wb.sheetnames else wb["Routing"]
    mark_col = ws.max_column - 2
    for row in ws.iter_rows(min_row=2):
        if row[0].value == row_key:
            if mark is not None:
                row[mark_col - 1].value = mark
            if correction is not None:
                row[mark_col].value = correction
            for col, value in (set_cells or {}).items():
                row[col].value = value
            break
    else:
        raise AssertionError("row not on the sheet")
    wb.save(path)


def correct_to(path, sheet_id, field, value):
    """The owner marks one owner-01 row wrong and writes the right value."""
    edit(path, _key(sheet_id, "owner-01", field), mark="wrong", correction=value)


def rejected(path):
    with pytest.raises(s.SheetRejected) as caught:
        s.load_sheet(path)
    return caught.value.problems


# --- the sheet the owner receives (FR-17, FR-18) ------------------------------------------


def test_sheet_shows_question_suggestion_sentence_and_where(sheets_dir):
    """FR-18: the owner can only judge a row if each one shows the question, the suggestion,
    the exact sentence and where it sits; a blank 'where' or a lost sentence makes the mark a
    guess."""
    sheet_id, path = make_facts()
    ws = load_workbook(path)["Facts"]
    row = {c.value: i for i, c in enumerate(ws[1])}
    assert list(row) == s.FACTS_HEADERS
    first = [c.value for c in ws[2]]
    assert first[row["contract"]] == "owner-01.pdf"
    assert first[row["question"]] == "Which law governs this contract?"
    assert first[row["suggested_answer"]] == "Delaware"
    assert first[row["supporting_sentence"]] == "governed by the laws of Delaware"
    assert first[row["where"]] == "page 2, characters 10–42"
    assert ws.cell(row=3, column=6).value == "characters 10–43"  # no page: characters alone
    assert ws.cell(row=3, column=4).value == "Acme Ltd | Beta Pvt Ltd"
    assert first[row["row_key"]] == _key(sheet_id, "owner-01", "governing_law")


def test_notice_period_question_uses_the_owner_definition():
    """D-14: notice period means the notice needed to end the contract or stop it renewing;
    a vaguer question would label a different fact."""
    _, path = make_facts()
    ws = load_workbook(path)["Facts"]
    questions = [ws.cell(row=r, column=3).value for r in range(2, 7)]
    assert "How much notice is needed to end this contract or stop it renewing?" in questions


def test_unsupported_row_offers_only_wrong_and_cannot_be_forced_correct():
    """AC-17 / FR-17: a stated answer whose quote is not in the contract reads 'no supporting
    sentence found' and the dropdown offers only 'wrong'; typing 'correct' over it anyway
    (a locked dropdown is not a lock) must reject the sheet."""
    sheet_id, path = make_facts()
    ws = load_workbook(path)["Facts"]
    assert ws.cell(row=4, column=5).value == s.NO_SUPPORT
    assert ws.cell(row=4, column=6).value in (None, "")
    options = {}
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            options[str(rng)] = dv.formula1
    assert options["G4"] == '"wrong"'
    assert options["G2"] == '"correct,wrong"'

    edit(path, _key(sheet_id, "owner-01", "effective_date"), mark="correct")
    assert any("no supporting sentence" in p for p in rejected(path))


def test_not_stated_suggestion_can_be_confirmed():
    """A 'not stated' answer has no sentence to check, so the owner is the check: marking it
    correct must work, or every absent fact would have to be re-typed as 'wrong'."""
    sheet_id, path = make_facts()
    edit(path, _key(sheet_id, "owner-01", "notice_period"), mark="correct")
    _, keys = s.load_sheet(path)
    key = next(k for k in keys if k.field == "notice_period")
    assert (key.status, key.value, key.quote) == ("confirmed", None, None)


def test_sheet_is_protected_with_hidden_row_key_and_meta():
    """The question columns are locked, the three answer columns are not, the row key and the
    meta tab are hidden, and the saved expected rows sit beside the sheet."""
    sheet_id, path = make_facts()
    wb = load_workbook(path)
    ws = wb["Facts"]
    assert ws.protection.sheet is True
    assert ws.column_dimensions["A"].hidden is True
    assert wb["meta"].sheet_state == "hidden"
    meta = {a.value: b.value for a, b in wb["meta"].iter_rows(min_row=1, max_col=2)}
    assert meta["sheet_id"] == sheet_id
    assert meta["kind"] == "facts"
    assert meta["row_count"] == "5"
    assert len(meta["frozen_manifest_sha256"]) == 64
    assert [ws.cell(row=2, column=c).protection.locked for c in range(1, 10)] == [
        True,
        True,
        True,
        True,
        True,
        True,
        False,
        False,
        False,
    ]
    assert (path.parent / f"{sheet_id}.expected.json").is_file()
    assert wb.active.title == "Facts"


def test_text_that_looks_like_a_formula_or_has_control_characters_round_trips():
    """A quote starting with '=' must stay text (not an Excel formula) and OCR form feeds must
    not crash the writer; either would break the owner's sheet or the locked-cell check."""
    rows = [
        _fact("owner-01", "renewal_term", "=SUM(A1)", "=SUM(A1) renews\x0c yearly"),
        _fact("owner-01", "governing_law", "+1 law", "@home law"),
    ]
    sheet_id, path = make_facts(rows)
    ws = load_workbook(path)["Facts"]
    assert ws.cell(row=2, column=5).data_type == "s"
    assert ws.cell(row=2, column=5).value == "=SUM(A1) renews yearly"
    edit(path, _key(sheet_id, "owner-01", "renewal_term"), mark="correct")
    _, keys = s.load_sheet(path)
    assert [k.status for k in keys] == ["confirmed", "unmarked"]


def test_duplicate_fact_rows_are_refused_at_write():
    """Two rows for one (document, field) would share a row_key and one answer would
    silently overwrite the other."""
    rows = [
        _fact("owner-01", "governing_law", "X", "q"),
        _fact("owner-01", "governing_law", "Y", "q"),
    ]
    with pytest.raises(ValueError, match="duplicate"):
        s.write_facts_sheet(rows, path=s.SHEETS_DIR)


def test_a_file_path_is_used_as_given(sheets_dir):
    """The caller may name the file; a folder gets facts-<sheet_id>.xlsx (the id only exists
    once the sheet is written)."""
    target = sheets_dir / "mine.xlsx"
    sheet_id = s.write_facts_sheet(facts_rows(), path=target)
    assert target.is_file()
    assert (sheets_dir / f"{sheet_id}.expected.json").is_file()
    assert s.load_sheet(target)[0] == "facts"


# --- loading: marks become keys (FR-3, FR-18, FR-19) --------------------------------------


def test_marks_become_keys_and_only_the_owner_decides():
    """AC-18 / FR-3: correct -> the suggestion becomes the key; wrong + value -> the written
    value; blank -> unmarked. The suggestion stays on the key for the 'how often did the owner
    change it' figure but is never the scored value of a wrong row."""
    sheet_id, path = make_facts()
    edit(path, _key(sheet_id, "owner-01", "governing_law"), mark="correct")
    correct_to(path, sheet_id, "parties", "Acme Ltd | Gamma LLC")
    correct_to(path, sheet_id, "effective_date", "2024-05-01")
    edit(
        path,
        _key(sheet_id, "owner-01", "notice_period"),
        mark="wrong",
        correction="30 days",
    )
    kind, keys = s.load_sheet(path)
    by = {k.field: k for k in keys}
    assert kind == "facts"
    assert len(keys) == 5

    law = by["governing_law"]
    assert (law.status, law.value, law.quote) == (
        "confirmed",
        "Delaware",
        "governed by the laws of Delaware",
    )
    assert (law.char_start, law.char_end) == (10, 42)

    parties = by["parties"]
    assert parties.status == "corrected"
    assert parties.value == ["Acme Ltd", "Gamma LLC"]
    assert parties.suggested_value == ["Acme Ltd", "Beta Pvt Ltd"]

    date = by["effective_date"]
    assert (date.status, date.value) == ("corrected", "2024-05-01")
    assert date.suggested_value == "2024-04-01"
    assert date.suggested_quote is None
    assert by["notice_period"].value == "30 days"
    assert by["notice_period"].suggested_value is None

    renewal = by["renewal_term"]
    assert (renewal.status, renewal.value, renewal.char_start) == ("unmarked", None, None)
    assert renewal.suggested_value == "renews yearly"  # kept for the report, never scored
    for k in keys:
        assert k.suggested_by == "claude-sonnet-5-5"
        assert (k.file_sha256, k.frozen_text_sha256) == (FILE_SHA, TEXT_SHA)


def test_unmarked_rows_are_listed_with_status_unmarked_and_counted():
    """AC-18 / FR-19: a row the owner cannot judge must not enter the key, and must stay
    countable; dropping it silently would hide how much of the key is missing."""
    _, path = make_facts()
    _, keys = s.load_sheet(path)
    assert len(keys) == 5
    assert {k.status for k in keys} == {"unmarked"}
    assert all(k.value is None and k.quote is None for k in keys)
    assert sum(k.status == "unmarked" for k in keys) == 5


def test_wrong_with_not_stated_is_a_key_that_says_not_stated():
    """FR-19 / edge case: 'not stated in the document' is a legitimate corrected answer and
    must become None (an AI answer of 'not stated' is then right), not the literal phrase."""
    sheet_id, path = make_facts()
    correct_to(path, sheet_id, "governing_law", "Not Stated in the Document")
    _, keys = s.load_sheet(path)
    key = next(k for k in keys if k.field == "governing_law")
    assert (key.status, key.value) == ("corrected", None)


def test_keys_carry_the_sheet_hash():
    """FR-39: the audit row ties each key to the exact file it came from."""
    _, path = make_facts()
    _, keys = s.load_sheet(path)
    assert {k.sheet_sha256 for k in keys} == {hashlib.sha256(path.read_bytes()).hexdigest()}
    assert all(k.loaded_at.tzinfo is not None for k in keys)


def test_a_mark_is_read_without_regard_to_case_and_spacing():
    """Excel autocorrects the first letter of a typed word; 'Correct ' is still the owner's
    'correct', not a reason to reject an otherwise good sheet."""
    sheet_id, path = make_facts()
    edit(path, _key(sheet_id, "owner-01", "governing_law"), mark=" Correct ")
    _, keys = s.load_sheet(path)
    assert next(k for k in keys if k.field == "governing_law").status == "confirmed"


def test_a_date_typed_into_the_correction_cell_comes_back_as_text():
    """Excel turns a typed 2024-05-01 into a date cell; the key must hold the text the owner
    meant, not a datetime that no scorer could compare."""
    sheet_id, path = make_facts()
    edit(
        path,
        _key(sheet_id, "owner-01", "effective_date"),
        mark="wrong",
        correction=datetime.fromisoformat("2024-05-01"),  # Excel cells carry no timezone
    )
    _, keys = s.load_sheet(path)
    assert next(k for k in keys if k.field == "effective_date").value == "2024-05-01"


def test_corrected_value_is_located_in_the_frozen_text_for_the_25k_split():
    """FR-23: the before/after-25k split needs the offset of the owner's own value, found
    word for word with spacing and line breaks ignored; no text, or no match, gives None."""
    (s.TEXTS_DIR / f"{FILE_SHA}.txt").write_text(FROZEN_TEXT, encoding="utf-8")
    sheet_id, path = make_facts()
    correct_to(path, sheet_id, "governing_law", "State of Delaware")
    correct_to(path, sheet_id, "effective_date", "1 January 1999")
    correct_to(path, sheet_id, "parties", "Acme Ltd | Beta Pvt Ltd")
    _, keys = s.load_sheet(path)
    by = {k.field: k for k in keys}
    law = by["governing_law"]
    assert law.quote == "State of\n   Delaware"
    assert FROZEN_TEXT[law.char_start : law.char_end] == law.quote
    assert (by["effective_date"].quote, by["effective_date"].char_start) == (None, None)
    assert by["parties"].quote is None  # a list of parties has no single sentence


def test_corrected_value_without_a_frozen_text_has_no_location():
    """The frozen text may not exist yet; the key is still valid, just unlocated."""
    sheet_id, path = make_facts()
    correct_to(path, sheet_id, "governing_law", "Delaware")
    _, keys = s.load_sheet(path)
    key = next(k for k in keys if k.field == "governing_law")
    assert (key.value, key.quote, key.char_start, key.char_end) == ("Delaware", None, None, None)


# --- loading: a sheet that does not match is rejected whole (FR-19) -----------------------


def _mutations():
    """Each way a returned sheet can stop matching the questions it was made from."""

    def change_question(path, sid):
        edit(path, _key(sid, "owner-01", "governing_law"), set_cells={2: "Which law is best?"})

    def change_suggestion(path, sid):
        edit(path, _key(sid, "owner-01", "governing_law"), set_cells={3: "California"})

    def change_sentence(path, sid):
        edit(path, _key(sid, "owner-01", "governing_law"), set_cells={4: "some other sentence"})

    def change_contract(path, sid):
        edit(path, _key(sid, "owner-01", "governing_law"), set_cells={1: "other.pdf"})

    def delete_row(path, sid):
        wb = load_workbook(path)
        wb["Facts"].delete_rows(3)
        wb.save(path)

    def add_row(path, sid):
        wb = load_workbook(path)
        wb["Facts"].append(["deadbeefdeadbeef", "x.pdf", "q", "a", "s", "", "correct", "", ""])
        wb.save(path)

    def duplicate_row(path, sid):
        wb = load_workbook(path)
        ws = wb["Facts"]
        ws.append([c.value for c in ws[2]])
        wb.save(path)

    def rename_header(path, sid):
        wb = load_workbook(path)
        wb["Facts"]["D1"].value = "my_answer"
        wb.save(path)

    return {
        "question changed": (change_question, "a question column was changed"),
        "suggested answer changed": (change_suggestion, "a question column was changed"),
        "sentence changed": (change_sentence, "a question column was changed"),
        "contract changed": (change_contract, "a question column was changed"),
        "row deleted": (delete_row, "its row is missing"),
        "row added": (add_row, "a row that is not on the sheet that was made"),
        "row duplicated": (duplicate_row, "appears 2 times"),
        "heading renamed": (rename_header, "the column headings were changed"),
    }


@pytest.mark.parametrize("name", list(_mutations()))
def test_a_sheet_that_no_longer_matches_is_rejected_as_a_whole(name):
    """AC-18 / FR-19: even with every other row marked correctly, one changed, missing, extra
    or duplicated row means the rows can no longer be trusted to be the questions asked, so
    nothing at all is loaded."""
    sheet_id, path = make_facts()
    edit(path, _key(sheet_id, "owner-01", "parties"), mark="correct")
    mutate, expect = _mutations()[name]
    mutate(path, sheet_id)
    assert any(expect in p for p in rejected(path))


def test_every_problem_is_listed_not_just_the_first():
    """FR-19: the owner fixes the sheet once, not once per problem."""
    sheet_id, path = make_facts()
    edit(path, _key(sheet_id, "owner-01", "governing_law"), mark="maybe")
    edit(path, _key(sheet_id, "owner-01", "parties"), mark="wrong")
    edit(path, _key(sheet_id, "owner-01", "effective_date"), mark="correct")
    edit(path, _key(sheet_id, "owner-01", "renewal_term"), set_cells={2: "changed"})
    wb = load_workbook(path)
    wb["Facts"].delete_rows(5)  # notice_period
    wb.save(path)
    problems = rejected(path)
    text = " | ".join(problems)
    assert len(problems) >= 5
    assert "your_mark must be correct, wrong or blank" in text
    assert "no correct value was written" in text
    assert "no supporting sentence" in text
    assert "a question column was changed" in text
    assert "owner-01 / notice_period: its row is missing" in text


@pytest.mark.parametrize(
    ("mark", "correction", "expect"),
    [
        ("maybe", None, "must be correct, wrong or blank"),
        ("wrong", None, "no correct value was written"),
        ("correct", "Texas", "the mark is not wrong"),
        (None, "Texas", "the mark is not wrong"),
    ],
)
def test_mark_rules(mark, correction, expect):
    """FR-19: a mark outside correct/wrong/blank, a 'wrong' with no correction, and a
    correction that disagrees with the mark are all refused: the system never guesses what
    the owner meant."""
    sheet_id, path = make_facts()
    edit(path, _key(sheet_id, "owner-01", "governing_law"), mark=mark, correction=correction)
    assert any(expect in p for p in rejected(path))


def test_whitespace_only_difference_in_a_locked_cell_is_not_a_change():
    """Excel normalises line endings and a stray space is not a changed question; rejecting it
    would send the owner back for nothing."""
    sheet_id, path = make_facts()
    edit(path, _key(sheet_id, "owner-01", "governing_law"), set_cells={3: "  Delaware "})
    assert s.load_sheet(path)[0] == "facts"


def test_unknown_sheet_id_is_rejected():
    """A spreadsheet the exam did not make has no questions to be matched against."""
    sheet_id, path = make_facts()
    (path.parent / f"{sheet_id}.expected.json").unlink()
    assert any("unknown sheet id" in p for p in rejected(path))


def test_a_hostile_sheet_id_cannot_reach_other_files(sheets_dir):
    """The id inside the file is untrusted input that becomes a file name; '../' must not walk
    out of the sheets folder."""
    _, path = make_facts()
    wb = load_workbook(path)
    wb["meta"]["B1"].value = "../../etc/passwd"
    wb.save(path)
    assert any("unknown sheet id" in p for p in rejected(path))


@pytest.mark.parametrize("junk", [b"not a spreadsheet", b""])
def test_unreadable_file_is_rejected(sheets_dir, junk):
    """A corrupt or empty file is a rejection with a reason, not a traceback."""
    sheets_dir.mkdir(parents=True)
    path = sheets_dir / "bad.xlsx"
    path.write_bytes(junk)
    assert rejected(path) == ["not a readable .xlsx spreadsheet"]


def test_missing_file_is_rejected(sheets_dir):
    assert rejected(sheets_dir / "gone.xlsx") == ["the spreadsheet could not be read"]


def test_sheets_are_only_written_and_read_inside_the_private_sheets_folder(tmp_path):
    """FR-36: these files quote real contracts; one written or loaded from anywhere else could
    end up in a tracked folder."""
    with pytest.raises(ValueError, match="written to"):
        s.write_facts_sheet(facts_rows(), path=tmp_path / "elsewhere")
    _, inside = make_facts()
    outside = tmp_path / "copy.xlsx"
    outside.write_bytes(inside.read_bytes())
    assert any("read from" in p for p in rejected(outside))
    sneaky = s.SHEETS_DIR / ".." / "copy.xlsx"
    assert any("read from" in p for p in rejected(sneaky))


def test_importing_sheets_does_not_import_the_app():
    """The exam repoints the app's settings before any app module loads; a sheets import that
    pulled the app in would make that too late."""
    code = (
        "import sys, eval.exam.sheets\n"
        "bad = sorted(m for m in sys.modules if m == 'app' or m.startswith('app.')"
        " or m.split('.')[0] == 'sqlalchemy')\n"
        "print(bad)\n"
        "sys.exit(1 if bad else 0)\n"
    )
    done = subprocess.run(
        [sys.executable, "-c", code],
        cwd=BACKEND_DIR,
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert done.returncode == 0, done.stdout + done.stderr


# --- the routing sheet (FR-21) ------------------------------------------------------------

SNAPSHOT = {
    "fingerprint": "f" * 64,
    "taken_at": "2026-10-06T00:00:00Z",
    "workflows": [
        {"id": "wf-nda", "name": "NDA review", "enabled": True},
        {"id": "wf-msa", "name": "MSA review", "enabled": True},
        {"id": "wf-old", "name": "Retired flow", "enabled": False},
    ],
}


def routing_rows():
    return [
        {
            "item_id": "rt-01",
            "channel": "email",
            "subject": "Need an NDA",
            "body": "Please prepare a mutual NDA for a vendor.",
            "drafted_for_workflow_id": "wf-nda",
            "drafting_model": "claude-sonnet-5-5",
        },
        {
            "item_id": "rt-02",
            "channel": "chat",
            "subject": "Services contract",
            "body": "We need to sign a services agreement with a new supplier.",
            "drafted_for_workflow_id": "wf-msa",
            "drafting_model": "claude-sonnet-5-5",
        },
        {
            "item_id": "rt-03",
            "channel": "notice",
            "subject": "Termination notice",
            "body": "A counterparty has sent a termination notice.",
            "drafted_for_workflow_id": "wf-msa",
            "drafting_model": "claude-sonnet-5-5",
        },
    ]


def make_routing(snapshot=None, rows=None):
    sheet_id = s.write_routing_sheet(
        rows if rows is not None else routing_rows(),
        snapshot=snapshot or SNAPSHOT,
        path=s.SHEETS_DIR,
    )
    return sheet_id, s.SHEETS_DIR / f"routing-{sheet_id}.xlsx"


def test_routing_sheet_layout_and_workflow_dropdown():
    """FR-21: each row shows the request and the workflow it was drafted for as a suggestion
    only; the dropdown offers today's enabled workflows plus 'no workflow'. A disabled
    workflow must not be offered, and the list lives on its own tab because an inline Excel
    list cannot hold 32 names."""
    sheet_id, path = make_routing()
    wb = load_workbook(path)
    ws = wb["Routing"]
    assert [c.value for c in ws[1]] == s.ROUTING_HEADERS
    first = [c.value for c in ws[2]]
    assert first[1] == "email"
    assert first[2] == "Need an NDA\n\nPlease prepare a mutual NDA for a vendor."
    assert first[3] == "NDA review"
    assert first[0] == _key(sheet_id, "rt-01", "")
    listed = [c.value for c in wb["workflows"]["A"]]
    assert listed == ["NDA review", "MSA review", s.NO_WORKFLOW]
    assert wb["workflows"].sheet_state == "hidden"
    meta = {a.value: b.value for a, b in wb["meta"].iter_rows(min_row=1, max_col=2)}
    assert meta["workflow_fingerprint"] == SNAPSHOT["fingerprint"]
    assert meta["kind"] == "routing"


def test_routing_marks_become_keys_against_the_snapshot():
    """FR-21 / AC-21: correct -> the drafted-for workflow; wrong -> the workflow the owner
    picked, by name -> id; 'no workflow' -> None; blank -> unmarked. Every key carries the
    fingerprint of the workflow list it was picked against, so a changed list is noticed."""
    sheet_id, path = make_routing()
    edit(path, _key(sheet_id, "rt-01", ""), mark="correct")
    edit(path, _key(sheet_id, "rt-02", ""), mark="wrong", correction="NDA review")
    edit(path, _key(sheet_id, "rt-03", ""), mark="wrong", correction=s.NO_WORKFLOW)
    kind, keys = s.load_sheet(path)
    assert kind == "routing"
    by = {k.item_id: k for k in keys}
    assert (by["rt-01"].status, by["rt-01"].workflow_id) == ("confirmed", "wf-nda")
    assert (by["rt-02"].status, by["rt-02"].workflow_id) == ("corrected", "wf-nda")
    assert (by["rt-03"].status, by["rt-03"].workflow_id) == ("corrected", None)
    assert {k.workflow_fingerprint for k in keys} == {SNAPSHOT["fingerprint"]}
    assert {k.sheet_sha256 for k in keys} == {hashlib.sha256(path.read_bytes()).hexdigest()}


def test_unpicked_routing_requests_are_unmarked_and_counted():
    """FR-21: a request with no owner pick is left out of the key and counted."""
    _, path = make_routing()
    _, keys = s.load_sheet(path)
    assert [k.status for k in keys] == ["unmarked"] * 3
    assert all(k.workflow_id is None for k in keys)


@pytest.mark.parametrize(
    ("mark", "correction", "expect"),
    [
        ("wrong", "No Such Workflow", "not in the list"),
        ("wrong", "Retired flow", "not in the list"),
        ("wrong", None, "no correct value was written"),
        ("correct", "NDA review", "the mark is not wrong"),
    ],
)
def test_routing_correction_must_be_a_snapshot_workflow(mark, correction, expect):
    """FR-21: a pick is a workflow that existed when the sheet was made; a made-up or
    disabled name cannot become a key."""
    sheet_id, path = make_routing()
    edit(path, _key(sheet_id, "rt-01", ""), mark=mark, correction=correction)
    assert any(expect in p for p in rejected(path))


def test_routing_sheet_that_no_longer_matches_is_rejected_whole():
    """FR-19 applies to routing too: an edited request is a different question."""
    sheet_id, path = make_routing()
    edit(path, _key(sheet_id, "rt-02", ""), set_cells={2: "Something else entirely"})
    edit(path, _key(sheet_id, "rt-01", ""), mark="correct")
    assert any("a question column was changed" in p for p in rejected(path))


def test_workflows_sharing_a_name_stay_distinguishable():
    """Two workflows may legitimately share a name; the owner must still be able to name
    exactly one, or a pick would key the wrong workflow."""
    snapshot = SNAPSHOT | {
        "workflows": [
            {"id": "11111111-aaaa", "name": "Review", "enabled": True},
            {"id": "22222222-bbbb", "name": "Review", "enabled": True},
        ]
    }
    rows = [routing_rows()[0] | {"drafted_for_workflow_id": "11111111-aaaa"}]
    sheet_id, path = make_routing(snapshot, rows)
    ws = load_workbook(path)["Routing"]
    assert ws.cell(row=2, column=4).value == "Review (11111111)"
    edit(path, _key(sheet_id, "rt-01", ""), mark="wrong", correction="Review (22222222)")
    _, keys = s.load_sheet(path)
    assert keys[0].workflow_id == "22222222-bbbb"


def test_workflows_whose_names_differ_only_by_case_or_spacing_stay_distinguishable():
    """FR-21: the loader reads a pick without regard to case and spacing, so 'MSA review' and
    'msa  Review' are ONE name to it -- and a pick of the second would silently key the first.
    Both get an id suffix, and a pick of either maps to exactly that workflow."""
    snapshot = SNAPSHOT | {
        "workflows": [
            {"id": "11111111-aaaa", "name": "MSA review", "enabled": True},
            {"id": "22222222-bbbb", "name": "msa  Review", "enabled": True},
            {"id": "33333333-cccc", "name": "NDA review", "enabled": True},
        ]
    }
    rows = [routing_rows()[0] | {"drafted_for_workflow_id": "33333333-cccc"}]
    sheet_id, path = make_routing(snapshot, rows)
    wb = load_workbook(path)
    listed = [c.value for c in wb["workflows"]["A"]]
    assert listed == [
        "MSA review (11111111)", "msa  Review (22222222)", "NDA review", s.NO_WORKFLOW
    ]
    for pick, workflow_id in [
        ("MSA review (11111111)", "11111111-aaaa"),
        ("msa  Review (22222222)", "22222222-bbbb"),
        ("NDA review", "33333333-cccc"),
    ]:
        edit(path, _key(sheet_id, "rt-01", ""), mark="wrong", correction=pick)
        assert s.load_sheet(path)[1][0].workflow_id == workflow_id
    # the bare name is no longer a pick: it could mean either workflow
    edit(path, _key(sheet_id, "rt-01", ""), mark="wrong", correction="msa review")
    assert any("not in the list" in p for p in rejected(path))


def test_a_sheet_whose_dropdown_cannot_tell_two_workflows_apart_rejects_a_pick_of_either():
    """Defence at the load boundary for a sheet whose saved record holds two labels that read the
    same (one written before labels were kept apart): the pick is refused, never guessed."""
    snapshot = SNAPSHOT | {
        "workflows": [
            {"id": "11111111-aaaa", "name": "MSA review", "enabled": True},
            {"id": "22222222-bbbb", "name": "NDA review", "enabled": True},
        ]
    }
    rows = [routing_rows()[0] | {"drafted_for_workflow_id": "22222222-bbbb"}]
    sheet_id, path = make_routing(snapshot, rows)
    saved = path.parent / f"{sheet_id}.expected.json"
    record = json.loads(saved.read_text())
    record["workflows"]["msa  REVIEW"] = "22222222-bbbb"  # the same name as "MSA review"
    saved.write_text(json.dumps(record))
    edit(path, _key(sheet_id, "rt-01", ""), mark="wrong", correction="MSA review")
    assert any("fits more than one" in p for p in rejected(path))
    # confirming the suggestion names no workflow, so nothing is ambiguous
    edit(path, _key(sheet_id, "rt-01", ""), mark="correct", correction="")
    assert s.load_sheet(path)[1][0].workflow_id == "22222222-bbbb"


def test_a_routing_sheet_carries_the_digest_of_the_requests_it_was_made_from():
    """FR-21: the sheet's hidden record and every key hold a digest of the exact request texts, so
    `load-routing` can refuse a sheet from an earlier draft whose ids and workflows are the same.
    The same rows give the same digest; any change to a request's text gives another."""
    _, path = make_routing()
    meta = {a.value: b.value for a, b in load_workbook(path)["meta"].iter_rows(min_row=1, max_col=2)}
    digest = s.routing_requests_sha256(routing_rows())
    assert meta["requests_sha256"] == digest
    _, keys = s.load_sheet(path)
    assert {k.requests_sha256 for k in keys} == {digest}
    changed = routing_rows()
    changed[1]["body"] += " Please hurry."
    assert s.routing_requests_sha256(changed) != digest
    _, other = make_routing(rows=changed)
    assert load_workbook(other)["meta"]["B6"].value != digest


def test_request_drafted_for_a_workflow_outside_the_snapshot_is_refused():
    """A suggestion that points at no workflow in the list the owner picks from would show a
    blank suggestion and could never be confirmed."""
    rows = [routing_rows()[0] | {"drafted_for_workflow_id": "wf-unknown"}]
    with pytest.raises(ValueError, match="not in the snapshot"):
        s.write_routing_sheet(rows, snapshot=SNAPSHOT, path=s.SHEETS_DIR)


def test_two_sheets_have_different_ids_and_row_keys():
    """A facts sheet marked for one round must not be loadable against another round's
    questions: ids and row keys are per sheet."""
    a_id, a_path = make_facts()
    b_id, b_path = make_facts()
    assert a_id != b_id
    ws_a, ws_b = load_workbook(a_path)["Facts"], load_workbook(b_path)["Facts"]
    assert ws_a["A2"].value != ws_b["A2"].value
