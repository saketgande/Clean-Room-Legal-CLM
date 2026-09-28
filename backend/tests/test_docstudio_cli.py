"""`python -m app.docstudio` — the way a person runs Phase 1 on a file.

Mostly glue, so these guard the parts of the glue that are not obvious: that a
second run of the same file is free rather than a second paid OCR call, that
an unreadable file is refused before anything is written, and that the paid
option refuses to run against a mocked client and produce a proposal that
looks real and says nothing.
"""

import uuid
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import delete

import app.models  # noqa: F401  (register every mapper)
from app.core.database import SessionLocal
from app.docstudio.__main__ import main
from app.docstudio.models import DsClause, DsDocument, DsEvent, DsVersion


def _docx(path):
    document = Document()
    # Unique per run: the command-line scope is shared with real runs, and a
    # test must never deduplicate against — or clean up — someone's contract.
    document.add_paragraph(f"1. TERM. Three years. Reference {uuid.uuid4()}.")
    # No number: whether it continues clause 1 is a question only the AI can
    # answer, so a run always has something to ask it.
    document.add_paragraph("Either party may renew it.")
    document.add_paragraph("2. FEES. Net thirty days.")
    buffer = BytesIO()
    document.save(buffer)
    path.write_bytes(buffer.getvalue())
    return path


@pytest.fixture
def made():
    """Documents this test created, deleted by id afterwards — never by scope,
    which would take the user's own command-line runs with them."""
    ids: list[str] = []
    yield ids
    session = SessionLocal()
    try:
        for document_id in ids:
            version_ids = [
                v.id for v in session.query(DsVersion).filter(DsVersion.document_id == document_id)
            ]
            session.execute(delete(DsEvent).where(DsEvent.document_id == document_id))
            session.execute(delete(DsClause).where(DsClause.version_id.in_(version_ids)))
            session.execute(delete(DsVersion).where(DsVersion.document_id == document_id))
            session.execute(delete(DsDocument).where(DsDocument.id == document_id))
        session.commit()
    finally:
        session.close()


def test_a_run_writes_the_report_where_it_says(tmp_path, made, capsys):
    source = _docx(tmp_path / "msa.docx")

    result = main([str(source), "--out", str(tmp_path / "out")])
    made.append(result.document_id)

    report = (tmp_path / "out" / "msa.report.md").read_text()
    assert "## Clauses" in report
    assert "msa.report.md" in capsys.readouterr().out


def test_the_same_file_twice_is_not_read_twice(tmp_path, made, capsys):
    """Each run makes a fresh session with no memory of the last, so without
    finding the earlier ingest by hash a scanned contract would be OCR'd — and
    paid for — every time someone looked at it."""
    source = _docx(tmp_path / "msa.docx")

    first = main([str(source), "--out", str(tmp_path / "out")])
    made.append(first.document_id)
    second = main([str(source), "--out", str(tmp_path / "out")])

    assert second.deduplicated is True
    assert second.version_id == first.version_id
    assert "already ingested" in capsys.readouterr().out


def test_an_unreadable_type_is_refused_before_anything_is_written(tmp_path):
    source = tmp_path / "scan.png"
    source.write_bytes(b"\x89PNG\r\n")

    with pytest.raises(SystemExit):
        main([str(source), "--out", str(tmp_path / "out")])

    assert not (tmp_path / "out" / "scan.report.md").exists()


def test_without_the_ai_the_structure_says_where_it_came_from(tmp_path, made, capsys):
    """The suite runs with MOCK_CLAUDE on, so the undecided clause stays
    undecided — and the output has to say so, or a rules-only tree would be
    mistaken for one the AI settled."""
    source = _docx(tmp_path / "msa.docx")

    result = main([str(source), "--out", str(tmp_path / "out")])
    made.append(result.document_id)

    assert "rules only — the Claude client is mocked; 1 clause left undecided" in (
        capsys.readouterr().out
    )


def test_no_ai_skips_the_arrangement(tmp_path, made, capsys):
    source = _docx(tmp_path / "msa.docx")

    result = main([str(source), "--out", str(tmp_path / "out"), "--no-ai"])
    made.append(result.document_id)

    assert "AI switched off" in capsys.readouterr().out


def test_reading_from_stdin_needs_a_name(tmp_path):
    """Without a name there is no extension, so no way to know how to read it."""
    with pytest.raises(SystemExit):
        main(["-", "--out", str(tmp_path / "out")])
