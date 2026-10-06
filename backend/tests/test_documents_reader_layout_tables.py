"""Contracts put clauses in tables, and a clause in a grid is not a clause.

Two layouts are common: a single-column table drawing a box around a clause,
and a two-column table with the number in the first column. Rendered as a grid
both become one lump of text — so a liability cap in a bordered box has no
number, cannot be cited, cannot be commented on and does not exist as a clause
anywhere in the system.

A fee schedule is also a table and must stay one, so the tests here are narrow
on purpose. Splitting a real data table would scatter the money terms across
rows with no header to read them against, which is worse than the bug.
"""

from io import BytesIO

from docx import Document

from app.documents.reader.parsing.registry import DOCX_MIME, parser_for
from app.documents.reader.structure import build


def _clauses(build_document):
    document = Document()
    build_document(document)
    buffer = BytesIO()
    document.save(buffer)
    return build(parser_for(DOCX_MIME).parse(buffer.getvalue(), filename="c.docx")).clauses


# --- layout tables become clauses -------------------------------------------


def test_a_clause_boxed_in_a_single_column_table_is_a_clause():
    """The defect this file exists for. A one-cell table is a border, not data."""
    def build_document(document):
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].text = "9.1 Liability is capped at fees paid in the prior 12 months."

    clauses = _clauses(build_document)

    assert len(clauses) == 1
    assert clauses[0].number_label == "9.1"
    assert clauses[0].number_path == [9, 1]
    assert clauses[0].clause_type != "table"


def test_a_two_column_number_and_text_layout_becomes_clauses():
    """The number lives in its own cell, so it never appears in the clause's
    text and `split_label` cannot find it. Taken from the column instead."""
    def build_document(document):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "9.1"
        table.cell(0, 1).text = "Liability is capped at fees paid."
        table.cell(1, 0).text = "9.2"
        table.cell(1, 1).text = "Neither party is liable for indirect loss."

    clauses = _clauses(build_document)

    assert [c.number_label for c in clauses] == ["9.1", "9.2"]
    assert clauses[0].text.startswith("9.1 Liability is capped")


def test_a_single_column_table_with_several_paragraphs_yields_several_clauses():
    def build_document(document):
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].text = "9.1 Liability is capped."
        cell.add_paragraph("9.2 Indirect loss is excluded.")

    assert [c.number_label for c in _clauses(build_document)] == ["9.1", "9.2"]


# --- data tables must stay tables -------------------------------------------


def test_a_fee_schedule_stays_a_table():
    """The guard that matters more than the fix. Split into rows, the money
    terms lose the header that says what they are."""
    def build_document(document):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Milestone"
        table.cell(0, 1).text = "Amount"
        table.cell(1, 0).text = "Final delivery"
        table.cell(1, 1).text = "USD 2,400,000"

    clauses = _clauses(build_document)

    assert len(clauses) == 1
    assert clauses[0].clause_type == "table"
    assert "Milestone\tAmount" in clauses[0].text


def test_three_columns_stays_a_table_even_with_a_number_column():
    """"1. | Setup fee | USD 5,000" opens with a clause number and is still a
    price list. Only two columns can be a number-and-clause layout."""
    def build_document(document):
        table = document.add_table(rows=1, cols=3)
        table.cell(0, 0).text = "1."
        table.cell(0, 1).text = "Setup fee"
        table.cell(0, 2).text = "USD 5,000"

    assert _clauses(build_document)[0].clause_type == "table"


def test_two_columns_whose_first_holds_prose_stays_a_table():
    """The first column must be a clause number and nothing else. "Governing
    law | New York" is a two-column table of facts, not clauses."""
    def build_document(document):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Governing law"
        table.cell(0, 1).text = "New York"
        table.cell(1, 0).text = "Venue"
        table.cell(1, 1).text = "Manhattan"

    assert _clauses(build_document)[0].clause_type == "table"


def test_a_nested_table_keeps_the_whole_thing_a_table():
    """A table inside a cell is structure, and flattening the outer one would
    strand the inner one's rows with nothing to read them against."""
    def build_document(document):
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Schedule of Charges"
        inner = cell.add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "LIABILITY CAP"
        inner.cell(0, 1).text = "USD 5,000,000"

    clauses = _clauses(build_document)

    assert clauses[0].clause_type == "table"
    assert "USD 5,000,000" in clauses[0].text
